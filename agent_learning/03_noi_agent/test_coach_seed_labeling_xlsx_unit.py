import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from evals.aichat import export_coach_seed_labeling_workbook_xlsx as xlsx_export


class CoachSeedLabelingXlsxTests(unittest.TestCase):
    def test_export_xlsx_uses_chinese_headers_and_validation_sheet(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "coach.xlsx"
            row_count = xlsx_export.export_xlsx(
                input_csv=Path("docs/research/coach_seed_labeling_workbook_v1.csv"),
                focus_registry=Path("docs/research/focus_registry_v1.json"),
                output_xlsx=output_path,
            )

            workbook = load_workbook(output_path)

        self.assertEqual(50, row_count)
        self.assertIn("标注表", workbook.sheetnames)
        self.assertIn("标签说明", workbook.sheetnames)
        sheet = workbook["标注表"]
        self.assertEqual("学生当前问题（先看）", sheet["C1"].value)
        self.assertEqual("学生当前状态（必选）", sheet["G1"].value)
        self.assertEqual("coach_problem_solving_state", sheet["G2"].value)
        self.assertEqual("cp_bridge_001", sheet["A3"].value)
        self.assertTrue(sheet.freeze_panes)
        self.assertEqual(15, len(sheet.data_validations.dataValidation))
        options = workbook["下拉选项"]
        self.assertEqual("连题意或目标都没看懂（text_comprehension_blocked）", options["A2"].value)
        self.assertEqual("表示/状态桥：不知道状态、对象、变量表示什么（representation_bridge）", options["B2"].value)
        self.assertEqual("轻提示：要上下文/当前尝试/一个观察方向，不补关键桥（L1）", options["H2"].value)
        self.assertEqual("不需要（false）", options["K2"].value)
        self.assertEqual("已标完（labeled）", options["N3"].value)


if __name__ == "__main__":
    unittest.main()
