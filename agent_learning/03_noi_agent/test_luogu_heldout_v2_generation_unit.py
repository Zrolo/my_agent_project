import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from evals.aichat import generate_luogu_heldout_v2_50


def _write_ndjson(path: Path, rows: list[object]) -> None:
    lines = []
    for row in rows:
        if row == "__BAD_JSON__":
            lines.append("{bad json")
        else:
            lines.append(json.dumps(row, ensure_ascii=False, sort_keys=True))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _problem(pid: str, title: str, tags: list[str], description: str | None = None) -> dict:
    return {
        "pid": pid,
        "title": title,
        "difficulty": 3,
        "tags": tags,
        "description": description
        or f"{title}。给定若干输入数据，需要根据题目条件设计正确算法并输出答案。"
        "这是一个适合算法竞赛教学分析的题面片段，包含足够的对象、限制和目标。",
        "inputFormat": "第一行输入若干整数，随后输入题目需要的数据。",
        "outputFormat": "输出题目要求的答案。",
        "hint": "数据范围保证可以使用相应算法。",
        "samples": [{"input": "3\n1 2 3\n", "output": "6\n"}],
    }


class LuoguHeldoutV2GenerationTests(unittest.TestCase):
    def test_iter_luogu_problem_rows_skips_bad_json_and_incomplete_rows(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            ndjson = Path(tmpdir) / "latest.ndjson"
            _write_ndjson(
                ndjson,
                [
                    _problem("P1001", "DP 入门", ["动态规划 DP"]),
                    "__BAD_JSON__",
                    {"pid": "P1002", "title": "缺题面", "tags": ["贪心"]},
                    _problem("P1003", "二分答案", ["二分"]),
                ],
            )

            rows = list(generate_luogu_heldout_v2_50.iter_luogu_problem_rows(ndjson))

        self.assertEqual(["P1001", "P1003"], [row["pid"] for row in rows])
        self.assertIn("problem_statement", rows[0])
        self.assertIn("problem_statement_public_summary", rows[0])

    def test_select_candidates_by_bucket_returns_tag_matched_candidates(self):
        rows = [
            _problem("P2001", "背包状态", ["动态规划 DP"]),
            _problem("P2002", "二分答案", ["二分"]),
            _problem("P2003", "线段树维护", ["线段树"]),
        ]

        candidates = generate_luogu_heldout_v2_50.select_candidates_by_bucket(
            rows,
            quotas={
                "state_representation_semantics": 1,
                "predicate_check_semantics": 1,
                "data_structure_operation_semantics": 1,
            },
            multiplier=1,
        )

        self.assertEqual("P2001", candidates["state_representation_semantics"][0]["pid"])
        self.assertEqual("P2002", candidates["predicate_check_semantics"][0]["pid"])
        self.assertEqual("P2003", candidates["data_structure_operation_semantics"][0]["pid"])

    def test_build_cases_from_candidates_has_source_fields_and_distribution(self):
        candidates = {
            "state_representation_semantics": [
                _problem("P3001", "采药", ["动态规划 DP"]),
                _problem("P3002", "数位 DP", ["动态规划 DP"]),
            ],
            "predicate_check_semantics": [
                _problem("P3003", "木材加工", ["二分"]),
                _problem("P3004", "跳石头", ["二分"]),
            ],
        }

        cases = generate_luogu_heldout_v2_50.build_cases_from_candidates(
            candidates,
            quotas={
                "state_representation_semantics": 2,
                "predicate_check_semantics": 2,
            },
            length_plan=["short", "medium_short", "medium_long", "long"],
            recent_dialogue_plan=["none", "short", "long", "long"],
            min_code_cases=1,
        )
        errors = generate_luogu_heldout_v2_50.validate_generated_cases(
            cases,
            expected_count=4,
            quotas={
                "state_representation_semantics": 2,
                "predicate_check_semantics": 2,
            },
            min_code_cases=1,
            max_none_recent_dialogue=1,
            min_long_recent_dialogue=2,
            length_distribution={
                "short": 1,
                "medium_short": 1,
                "medium_long": 1,
                "long": 1,
            },
        )

        self.assertEqual([], errors)
        self.assertEqual("luogu", cases[0]["problem_source_platform"])
        self.assertTrue(cases[0]["problem_source_url"].startswith("https://www.luogu.com.cn/problem/"))
        self.assertEqual("draft_needs_coach_review", cases[0]["reference_label_status"])
        self.assertIn("synthetic-but-grounded", cases[0]["case_generation_method"])

    def test_full_default_generation_uses_varied_student_message_style(self):
        candidates = {
            bucket: [
                _problem(f"P{idx:04d}", f"{bucket} 样例 {idx}", config["tags"][:2] or ["模拟"])
                for idx in range(start, start + quota)
            ]
            for start, (bucket, quota), config in zip(
                range(6000, 7200, 100),
                generate_luogu_heldout_v2_50.DEFAULT_QUOTAS.items(),
                [
                    generate_luogu_heldout_v2_50.BUCKET_CONFIG[bucket]
                    for bucket in generate_luogu_heldout_v2_50.DEFAULT_QUOTAS
                ],
            )
        }

        cases = generate_luogu_heldout_v2_50.build_cases_from_candidates(candidates)
        messages = [case["student_message"] for case in cases]
        by_bucket = {}
        for case in cases:
            by_bucket.setdefault(case["bridge_bucket"], []).append(case["student_message"])

        self.assertEqual(50, len(messages))
        self.assertEqual(50, len(set(messages)))
        for bucket, bucket_messages in by_bucket.items():
            self.assertEqual(
                len(bucket_messages),
                len(set(bucket_messages)),
                f"duplicate student_message in {bucket}",
            )
        ai_like_fragments = [
            "我不想直接看答案",
            "题面里的哪一条关系",
            "当前关系",
            "可迁移",
            "如果不能直接给",
        ]
        for message in messages:
            self.assertFalse(any(fragment in message for fragment in ai_like_fragments), message)
            clauses = [
                clause.strip()
                for clause in message.replace("？", "。").replace("！", "。").split("。")
                if len(clause.strip()) >= 8
            ]
            self.assertEqual(len(clauses), len(set(clauses)), message)

    def test_recent_dialogues_use_natural_student_ai_context(self):
        candidates = {
            bucket: [
                _problem(f"P{idx:04d}", f"{bucket} 样例 {idx}", config["tags"][:2] or ["模拟"])
                for idx in range(start, start + quota)
            ]
            for start, (bucket, quota), config in zip(
                range(6000, 7200, 100),
                generate_luogu_heldout_v2_50.DEFAULT_QUOTAS.items(),
                [
                    generate_luogu_heldout_v2_50.BUCKET_CONFIG[bucket]
                    for bucket in generate_luogu_heldout_v2_50.DEFAULT_QUOTAS
                ],
            )
        }

        cases = generate_luogu_heldout_v2_50.build_cases_from_candidates(candidates)
        non_empty_dialogues = [
            case["recent_dialogue"]
            for case in cases
            if case["recent_dialogue"].strip().upper() != "N/A"
        ]

        self.assertEqual(40, len(non_empty_dialogues))
        self.assertGreaterEqual(len(set(non_empty_dialogues)), 35)
        ai_like_fragments = [
            "学生知道",
            "缺少",
            "当前最卡住的一步",
            "概念、判断、建模还是实现细节",
            "先不要急着要完整代码",
            "可迁移",
            "当前关系",
        ]
        for case in cases:
            dialogue = case["recent_dialogue"]
            if dialogue.strip().upper() == "N/A":
                continue
            self.assertIn("学生：", dialogue)
            self.assertIn("AI：", dialogue)
            self.assertFalse(any(fragment in dialogue for fragment in ai_like_fragments), dialogue)
            role_lines = [
                line
                for line in dialogue.splitlines()
                if line.startswith(("学生：", "AI："))
            ]
            if case["recent_dialogue_bucket"] == "short":
                self.assertGreaterEqual(len(role_lines), 2, dialogue)
            if case["recent_dialogue_bucket"] == "long":
                self.assertGreaterEqual(len(role_lines), 4, dialogue)
            self.assertTrue(role_lines[-1].startswith("AI："), dialogue)

    def test_validation_rejects_missing_source_and_distribution_shortfall(self):
        case = generate_luogu_heldout_v2_50.build_case(
            problem=_problem("P4001", "背包状态", ["动态规划 DP"]),
            bridge_bucket="state_representation_semantics",
            case_number=1,
            length_bucket="short",
            recent_dialogue_bucket="none",
            include_code_excerpt=False,
        )
        case.pop("problem_source_url")

        errors = generate_luogu_heldout_v2_50.validate_generated_cases(
            [case],
            expected_count=1,
            quotas={"state_representation_semantics": 1},
            min_code_cases=1,
            max_none_recent_dialogue=0,
            min_long_recent_dialogue=1,
            length_distribution={"long": 1},
        )
        error_codes = {error["code"] for error in errors}

        self.assertIn("missing_required_field", error_codes)
        self.assertIn("too_few_code_excerpts", error_codes)
        self.assertIn("too_many_no_recent_dialogue", error_codes)
        self.assertIn("too_few_long_recent_dialogue", error_codes)
        self.assertIn("too_few_student_message_length_bucket", error_codes)

    def test_exports_review_workbook_grouped_by_bridge_bucket(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "review.xlsx"
            cases = [
                generate_luogu_heldout_v2_50.build_case(
                    problem=_problem("P5001", "采药", ["动态规划 DP"]),
                    bridge_bucket="state_representation_semantics",
                    case_number=1,
                    length_bucket="short",
                    recent_dialogue_bucket="none",
                    include_code_excerpt=False,
                ),
                generate_luogu_heldout_v2_50.build_case(
                    problem=_problem("P5002", "木材加工", ["二分"]),
                    bridge_bucket="predicate_check_semantics",
                    case_number=2,
                    length_bucket="medium_short",
                    recent_dialogue_bucket="long",
                    include_code_excerpt=True,
                ),
            ]

            generate_luogu_heldout_v2_50.export_source_and_case_review_workbook(cases, output)
            workbook = load_workbook(output)

        self.assertIn("状态表示", workbook.sheetnames[0])
        self.assertTrue(any("判定条件" in name for name in workbook.sheetnames))
        first_sheet = workbook[workbook.sheetnames[0]]
        headers = [first_sheet.cell(row=1, column=column).value for column in range(1, first_sheet.max_column + 1)]
        self.assertIn("原题链接", headers)
        self.assertIn("必要题面", headers)
        self.assertIn("学生问题", headers)
        self.assertIn("目标 bridge", headers)


if __name__ == "__main__":
    unittest.main()
