#!/usr/bin/env python3
"""Export an EAIT manuscript Markdown draft to LaTeX and Word.

This is a formatting/export utility only. It does not recompute results,
touch experiment data, or modify any online system.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SPECIALS = {
    "&": r"\&",
    "%": r"\%",
    "$": r"\$",
    "#": r"\#",
    "_": r"\_",
    "{": r"\{",
    "}": r"\}",
    "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}",
}


def parse_blocks(text: str) -> list[dict]:
    lines = text.splitlines()
    blocks: list[dict] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            lang = line.strip().strip("`") or ""
            code: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append({"type": "code", "lang": lang, "text": "\n".join(code)})
            continue
        if line.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = []
            for raw in table_lines:
                cells = [c.strip() for c in raw.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells):
                    continue
                rows.append(cells)
            if rows:
                blocks.append({"type": "table", "rows": rows})
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            blocks.append({"type": "heading", "level": len(m.group(1)), "text": m.group(2).strip()})
            i += 1
            continue
        para = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if not nxt.strip() or nxt.startswith("#") or nxt.startswith("|") or nxt.startswith("```"):
                break
            para.append(nxt.strip())
            i += 1
        blocks.append({"type": "para", "text": " ".join(para)})
    return blocks


def protect_cites(text: str) -> tuple[str, list[str]]:
    cites: list[str] = []

    def repl(match: re.Match) -> str:
        cites.append(match.group(0))
        return f"@@CITE{len(cites)-1}@@"

    return re.sub(r"\\cite\{[^}]+\}", repl, text), cites


def latex_inline(text: str) -> str:
    text, cites = protect_cites(text)

    code_tokens: list[str] = []

    def code_repl(match: re.Match) -> str:
        code_tokens.append(match.group(1))
        return f"@@CODE{len(code_tokens)-1}@@"

    text = re.sub(r"`([^`]+)`", code_repl, text)
    out = []
    for ch in text:
        out.append(SPECIALS.get(ch, ch))
    text = "".join(out)
    for idx, code in enumerate(code_tokens):
        escaped = "".join(SPECIALS.get(ch, ch) for ch in code)
        text = text.replace(f"@@CODE{idx}@@", rf"\texttt{{{escaped}}}")
    for idx, cite in enumerate(cites):
        text = text.replace(f"@@CITE{idx}@@", cite)
    return text


def latex_table(rows: list[list[str]]) -> str:
    cols = max(len(r) for r in rows)
    spec = ">{\\raggedright\\arraybackslash}X" * cols
    table_size = r"\scriptsize" if cols >= 5 or len(rows) > 15 else r"\small"
    out = [r"\begingroup", table_size, r"\begin{xltabular}{\textwidth}{" + spec + r"}", r"\toprule"]
    for idx, row in enumerate(rows):
        row = row + [""] * (cols - len(row))
        out.append(" & ".join(latex_inline(cell) for cell in row) + r" \\")
        if idx == 0:
            out.append(r"\midrule")
    out.extend([r"\bottomrule", r"\end{xltabular}", r"\endgroup"])
    return "\n".join(out)


def export_latex(blocks: list[dict], out_path: Path, bib_name: str) -> None:
    title = "CP-MissingBridgeBench: Evaluating Critical-Bridge Leakage in LLM Tutors for Competitive Programming"
    body: list[str] = []
    for block in blocks:
        if block["type"] == "heading" and block["level"] == 1:
            title = block["text"]
            continue
        if block["type"] == "heading":
            level = block["level"]
            text = latex_inline(block["text"])
            if level == 2:
                body.append(rf"\section{{{text}}}")
            elif level == 3:
                body.append(rf"\subsection{{{text}}}")
            else:
                body.append(rf"\subsubsection{{{text}}}")
        elif block["type"] == "para":
            body.append(latex_inline(block["text"]))
        elif block["type"] == "code":
            body.append("\\begin{verbatim}\n" + block["text"] + "\n\\end{verbatim}")
        elif block["type"] == "table":
            body.append(latex_table(block["rows"]))
    preamble = rf"""\documentclass[11pt]{{article}}
\usepackage[margin=1in]{{geometry}}
\usepackage{{fontspec}}
\usepackage{{xeCJK}}
\usepackage{{tabularx}}
\usepackage{{xltabular}}
\usepackage{{booktabs}}
\usepackage{{array}}
\usepackage{{hyperref}}
\usepackage{{natbib}}
\usepackage{{url}}
\setmainfont{{Times New Roman}}
\setCJKmainfont{{Songti SC}}
\sloppy
\title{{{latex_inline(title)}}}
\author{{}}
\date{{}}

\begin{{document}}
\maketitle

"""
    tail = rf"""

\bibliographystyle{{apalike}}
\bibliography{{{bib_name}}}
\end{{document}}
"""
    out_path.write_text(preamble + "\n\n".join(body) + tail, encoding="utf-8")


def docx_citation_text(text: str) -> str:
    return re.sub(r"\\cite\{([^}]+)\}", lambda m: "[" + "; ".join(x.strip() for x in m.group(1).split(",")) + "]", text)


def add_hyperlink_style(document: Document) -> None:
    styles = document.styles
    if "Hyperlink" in styles:
        return
    style = styles.add_style("Hyperlink", 2)
    style.font.color.rgb = RGBColor(5, 99, 193)
    style.font.underline = True


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 8.0) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(docx_citation_text(text.replace("`", "")))
    run.bold = bold
    run.font.size = Pt(size)


def export_docx(blocks: list[dict], references_text: str, out_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    add_hyperlink_style(doc)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = RGBColor(46, 116, 181)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.color.rgb = RGBColor(46, 116, 181)
    styles["Heading 3"].font.size = Pt(12)
    styles["Heading 3"].font.color.rgb = RGBColor(31, 77, 120)

    title_done = False
    for block in blocks:
        if block["type"] == "heading":
            text = block["text"]
            if block["level"] == 1 and not title_done:
                p = doc.add_paragraph(style="Title")
                p.add_run(text)
                title_done = True
            elif block["level"] == 2:
                doc.add_heading(text, level=1)
            elif block["level"] == 3:
                doc.add_heading(text, level=2)
            else:
                doc.add_heading(text, level=3)
        elif block["type"] == "para":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.1
            p.add_run(docx_citation_text(block["text"].replace("`", "")))
        elif block["type"] == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(block["text"])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif block["type"] == "table":
            rows = block["rows"]
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                for c_idx in range(cols):
                    text = row[c_idx] if c_idx < len(row) else ""
                    cell = table.cell(r_idx, c_idx)
                    if r_idx == 0:
                        set_cell_shading(cell, "F2F4F7")
                    set_cell_text(cell, text, bold=(r_idx == 0), size=7.5 if cols >= 5 else 8.5)
            doc.add_paragraph()

    if references_text.strip():
        doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_heading("References", level=1)
        in_refs = False
        for raw in references_text.splitlines():
            line = raw.strip()
            if line == "## Draft Main-Text Reference List":
                in_refs = True
                continue
            if line.startswith("## Excluded From"):
                break
            if not in_refs or not line:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(6)
            p.add_run(line.replace("*", ""))

    doc.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-md", required=True)
    parser.add_argument("--references-md", required=True)
    parser.add_argument("--tex-out", required=True)
    parser.add_argument("--docx-out", required=True)
    parser.add_argument("--bib-name", default="eait_v0_12_references_draft_20260520")
    args = parser.parse_args()

    source = Path(args.input_md).read_text(encoding="utf-8")
    references = Path(args.references_md).read_text(encoding="utf-8")
    blocks = parse_blocks(source)
    export_latex(blocks, Path(args.tex_out), args.bib_name)
    export_docx(blocks, references, Path(args.docx_out))


if __name__ == "__main__":
    main()
