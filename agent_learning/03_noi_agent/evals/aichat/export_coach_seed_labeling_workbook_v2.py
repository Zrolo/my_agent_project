import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from evals.aichat.coach_labeling_schema_v2 import (
    BRIDGE_FAMILIES,
    BRIDGE_SPECIFIC_FORBIDDEN_CONTENT,
    BRIDGE_SUBTYPES,
    CHINESE_HEADERS,
    COLUMN_WIDTHS,
    CONFIDENCE_OPTIONS,
    DIAGNOSIS_UNCERTAINTY,
    EVIDENCE_TYPES,
    FIELD_HELP,
    FOCUS_MATCH_STATUS,
    GENERAL_FORBIDDEN_CONTENT,
    HELP_FORMS,
    HELP_SEEKING_TYPES,
    INPUT_COLUMNS,
    LABEL_COLUMNS,
    LEAKAGE_RISKS,
    MULTI_VALUE_FIELDS,
    NOTES_OPTIONS,
    OPTIONS_BY_FIELD,
    REVIEW_STATUSES,
    SCAFFOLD_LEVELS,
    STUDENT_ALREADY_STATED_BRIDGE,
    STUDENT_ATTEMPT_LEVELS,
    STUDENT_STATES,
    POLICY_RISK_TYPES,
    TURN_TYPES,
    V2_COLUMNS,
    option_label,
    option_labels,
)


DEFAULT_SEED_JSONL = Path("docs/research/bridgebench_cp_seed_v1.jsonl")
DEFAULT_FOCUS_REGISTRY = Path("docs/research/focus_registry_v1.json")
DEFAULT_OUTPUT_XLSX = Path("docs/research/coach_seed_labeling_workbook_v2.zh.xlsx")


def load_seed_rows(path: Path) -> list[dict]:
    rows: list[dict] = []
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def load_focus_options(path: Path) -> list[str]:
    labels = {"unknown": "未知/不确定（unknown）", "not_applicable": "不适用（not_applicable）"}
    if path.exists():
        data = json.loads(path.read_text(encoding="utf-8"))
        focuses = data.get("focuses", data) if isinstance(data, dict) else data
        for item in focuses:
            if isinstance(item, str):
                labels[item] = option_label(item, item)
            elif isinstance(item, dict) and item.get("focus_id"):
                if item.get("status") == "deprecated":
                    continue
                focus_id = str(item["focus_id"])
                aliases = item.get("aliases") or []
                readable = str(aliases[0] if aliases else item.get("description") or focus_id)
                labels[focus_id] = option_label(focus_id, readable)
    return [labels["unknown"], labels["not_applicable"], *[labels[key] for key in sorted(labels) if key not in {"unknown", "not_applicable"}]]


def _normalize_seed_row(row: dict) -> dict:
    return {
        "case_id": row.get("id") or row.get("case_id") or "",
        "problem_ref": row.get("problem_ref") or "",
        "problem_source_platform": row.get("problem_source_platform") or "",
        "problem_source_id": row.get("problem_source_id") or "",
        "problem_source_url": row.get("problem_source_url") or "",
        "problem_statement": row.get("problem_statement") or "",
        "problem_statement_public_summary": row.get("problem_statement_public_summary") or "",
        "problem_statement_rights_note": row.get("problem_statement_rights_note") or "",
        "problem_statement_access_level": row.get("problem_statement_access_level") or "",
        "student_message": row.get("student_message") or "",
        "problem_context": row.get("problem_context") or "",
        "recent_dialogue": row.get("recent_dialogue") or "N/A",
        "student_code_excerpt": row.get("student_code_excerpt") or "N/A",
        "source_case_id": row.get("source_case_id") or "",
        "turn_position": row.get("turn_position") or "",
        "context_type": row.get("context_type") or "",
        "student_scaffold_followability": row.get("student_scaffold_followability") or "",
        "followability_label_confidence": row.get("followability_label_confidence") or "",
        "followability_evidence_quote": row.get("followability_evidence_quote") or "",
        "followability_uncertainty_reason": row.get("followability_uncertainty_reason") or "",
        "prior_ai_scaffold": row.get("prior_ai_scaffold") or "",
        "student_reply_to_prior_scaffold": row.get("student_reply_to_prior_scaffold") or "",
        "expected_tutor_move": row.get("expected_tutor_move") or "",
        "fixed_recent_dialogue_source": row.get("fixed_recent_dialogue_source") or "",
        "review_status": option_label("unlabeled", REVIEW_STATUSES["unlabeled"]),
    }


def _add_options_sheet(workbook: Workbook, focus_options: list[str]) -> dict[str, str]:
    sheet = workbook.create_sheet("下拉选项")
    lists = {
        "turn_type": option_labels(TURN_TYPES),
        "diagnosis_uncertainty": option_labels(DIAGNOSIS_UNCERTAINTY),
        "student_problem_solving_state": option_labels(STUDENT_STATES),
        "student_attempt_level": option_labels(STUDENT_ATTEMPT_LEVELS),
        "student_already_stated_bridge": option_labels(STUDENT_ALREADY_STATED_BRIDGE),
        "policy_risk_type": option_labels(POLICY_RISK_TYPES),
        "bridge_family": option_labels(BRIDGE_FAMILIES),
        "bridge_subtype": option_labels(BRIDGE_SUBTYPES),
        "evidence_type": option_labels(EVIDENCE_TYPES),
        "registered_focus_id": focus_options,
        "focus_match_status": option_labels(FOCUS_MATCH_STATUS),
        "max_scaffold_level": option_labels(SCAFFOLD_LEVELS),
        "help_seeking_type": option_labels(HELP_SEEKING_TYPES),
        "help_forms": option_labels(HELP_FORMS),
        "general_forbidden_content": option_labels(GENERAL_FORBIDDEN_CONTENT),
        "bridge_specific_forbidden_content": option_labels(BRIDGE_SPECIFIC_FORBIDDEN_CONTENT),
        "leakage_risk": option_labels(LEAKAGE_RISKS),
        "coach_confidence": option_labels(CONFIDENCE_OPTIONS),
        "coach_note_tags": option_labels(NOTES_OPTIONS),
        "review_status": option_labels(REVIEW_STATUSES),
    }
    ranges = {}
    for col_idx, (name, values) in enumerate(lists.items(), 1):
        sheet.cell(row=1, column=col_idx, value=name)
        sheet.cell(row=1, column=col_idx).font = Font(bold=True)
        for row_idx, value in enumerate(values, 2):
            sheet.cell(row=row_idx, column=col_idx, value=value)
        col_letter = sheet.cell(row=1, column=col_idx).column_letter
        ranges[name] = f"'下拉选项'!${col_letter}$2:${col_letter}${len(values) + 1}"
        sheet.column_dimensions[col_letter].width = 34
    return ranges


def _add_data_validation(sheet, cell_range: str, formula_range: str) -> None:
    validation = DataValidation(type="list", formula1=f"={formula_range}", allow_blank=True)
    validation.error = "请从下拉列表中选择，或留空。多选字段可用英文分号 ; 连接多个值。"
    validation.errorTitle = "无效选项"
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def _style_labeling_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    field_fill = PatternFill("solid", fgColor="F2F5F7")
    input_fill = PatternFill("solid", fgColor="FFF7D6")
    thin = Side(style="thin", color="D7DEE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="17324D")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    for cell in sheet[2]:
        cell.font = Font(size=9, color="667085")
        cell.fill = field_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    for row in sheet.iter_rows(min_row=3, max_row=sheet.max_row, max_col=len(V2_COLUMNS)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            if cell.column > len(INPUT_COLUMNS):
                cell.fill = input_fill
    for col_letter, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[col_letter].width = width
    sheet.row_dimensions[1].height = 38
    sheet.row_dimensions[2].height = 28
    for row_idx in range(3, sheet.max_row + 1):
        sheet.row_dimensions[row_idx].height = 72
    first_label_column = sheet.cell(row=1, column=len(INPUT_COLUMNS) + 1).column_letter
    sheet.freeze_panes = f"{first_label_column}3"
    sheet.auto_filter.ref = f"A2:{sheet.cell(row=2, column=len(V2_COLUMNS)).coordinate}"


def _add_instruction_sheet(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "开始这里"
    rows = [
        ["标注单位", "只标当前这一轮学生消息。不要把整段聊天、学生整体水平或题目知识点混成一个标签。"],
        ["第一步", "先选“这一轮类型”。不是每一轮都适合强行诊断 missing bridge。"],
        ["第二步", "再判断学生当前解题状态、已有尝试、是否已经说出关键桥。"],
        ["第三步", "选择主要 bridge family、subtype、系统已有焦点 ID，并写证据原话。"],
        ["多选字段", "求助类型、帮助形式、禁止内容可以用英文分号 ; 连接多个选项。"],
        ["L0", "L0 是只澄清/要证据。推荐帮助强度只允许 L0/L1/L2/L3；泄露评价另行标注。"],
        ["不要填 1", "没有近期对话或代码时用 N/A；待标注字段留空，标注状态默认 unlabeled。"],
    ]
    sheet.append(["教练标注 v2 说明", ""])
    for row in rows:
        sheet.append(row)
    sheet["A1"].font = Font(bold=True, size=16, color="17324D")
    sheet["A1"].fill = PatternFill("solid", fgColor="D9EAF7")
    for row in sheet.iter_rows(min_row=2, max_col=2):
        row[0].font = Font(bold=True, color="17324D")
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 110


def _add_guide_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("标签说明")
    sections = [
        ("这一轮类型", TURN_TYPES),
        ("诊断不确定性", DIAGNOSIS_UNCERTAINTY),
        ("学生当前解题状态", STUDENT_STATES),
        ("学生已有尝试程度", STUDENT_ATTEMPT_LEVELS),
        ("学生是否已说出关键桥", STUDENT_ALREADY_STATED_BRIDGE),
        ("策略风险类型", POLICY_RISK_TYPES),
        ("缺失桥梁大类", BRIDGE_FAMILIES),
        ("桥梁细分", BRIDGE_SUBTYPES),
        ("证据类型", EVIDENCE_TYPES),
        ("焦点匹配状态", FOCUS_MATCH_STATUS),
        ("求助类型", HELP_SEEKING_TYPES),
        ("帮助强度", SCAFFOLD_LEVELS),
        ("帮助形式", HELP_FORMS),
        ("通用禁止内容", GENERAL_FORBIDDEN_CONTENT),
        ("桥梁禁止内容", BRIDGE_SPECIFIC_FORBIDDEN_CONTENT),
        ("泄露风险", LEAKAGE_RISKS),
        ("备注标签", NOTES_OPTIONS),
    ]
    row_idx = 1
    for title, options in sections:
        sheet.cell(row=row_idx, column=1, value=title)
        sheet.cell(row=row_idx, column=1).font = Font(bold=True, size=13, color="17324D")
        row_idx += 1
        sheet.cell(row=row_idx, column=1, value="机器 ID")
        sheet.cell(row=row_idx, column=2, value="中文解释")
        sheet.cell(row=row_idx, column=1).font = Font(bold=True)
        sheet.cell(row=row_idx, column=2).font = Font(bold=True)
        row_idx += 1
        for option_id, description in options.items():
            sheet.cell(row=row_idx, column=1, value=option_id)
            sheet.cell(row=row_idx, column=2, value=description)
            row_idx += 1
        row_idx += 2
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 95
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def export_xlsx(
    *,
    seed_jsonl: Path = DEFAULT_SEED_JSONL,
    focus_registry: Path = DEFAULT_FOCUS_REGISTRY,
    output_xlsx: Path = DEFAULT_OUTPUT_XLSX,
    limit: int | None = None,
) -> int:
    seed_rows = load_seed_rows(seed_jsonl)
    if limit is not None:
        if limit < 0:
            raise ValueError("limit must be >= 0")
        seed_rows = seed_rows[:limit]
    if not seed_rows:
        raise ValueError(f"No rows found in {seed_jsonl}")
    workbook = Workbook()
    _add_instruction_sheet(workbook)
    ranges = _add_options_sheet(workbook, load_focus_options(focus_registry))
    _add_guide_sheet(workbook)

    sheet = workbook.create_sheet("标注表", 1)
    sheet.append([CHINESE_HEADERS.get(header, header) for header in V2_COLUMNS])
    sheet.append(V2_COLUMNS)
    for seed_row in seed_rows:
        normalized = _normalize_seed_row(seed_row)
        sheet.append([normalized.get(header, "") for header in V2_COLUMNS])
    _style_labeling_sheet(sheet)

    field_to_col = {field: idx + 1 for idx, field in enumerate(V2_COLUMNS)}
    max_row = sheet.max_row
    validation_ranges = {
        "turn_type": ranges["turn_type"],
        "diagnosis_uncertainty": ranges["diagnosis_uncertainty"],
        "student_problem_solving_state": ranges["student_problem_solving_state"],
        "student_attempt_level": ranges["student_attempt_level"],
        "student_already_stated_bridge": ranges["student_already_stated_bridge"],
        "policy_risk_type": ranges["policy_risk_type"],
        "primary_bridge_family": ranges["bridge_family"],
        "secondary_bridge_family": ranges["bridge_family"],
        "primary_bridge_subtype_id": ranges["bridge_subtype"],
        "secondary_bridge_subtype_id": ranges["bridge_subtype"],
        "evidence_type": ranges["evidence_type"],
        "registered_focus_id": ranges["registered_focus_id"],
        "secondary_registered_focus_id": ranges["registered_focus_id"],
        "focus_match_status": ranges["focus_match_status"],
        "max_scaffold_level": ranges["max_scaffold_level"],
        "help_seeking_type": ranges["help_seeking_type"],
        "help_forms": ranges["help_forms"],
        "general_forbidden_content": ranges["general_forbidden_content"],
        "bridge_specific_forbidden_content": ranges["bridge_specific_forbidden_content"],
        "leakage_risk": ranges["leakage_risk"],
        "coach_confidence": ranges["coach_confidence"],
        "coach_note_tags": ranges["coach_note_tags"],
        "review_status": ranges["review_status"],
    }
    for field, formula_range in validation_ranges.items():
        col_idx = field_to_col[field]
        col = sheet.cell(row=1, column=col_idx).column_letter
        _add_data_validation(sheet, f"{col}3:{col}{max_row}", formula_range)
    for field, help_text in FIELD_HELP.items():
        col_idx = field_to_col.get(field)
        if col_idx:
            sheet.cell(row=1, column=col_idx).comment = Comment(help_text, "Codex")
    for field in MULTI_VALUE_FIELDS:
        col_idx = field_to_col[field]
        sheet.cell(row=1, column=col_idx).comment = Comment(
            "可多选。Excel 原生下拉一次只能选一个；需要多选时用英文分号 ; 手动连接多个机器 ID 或中文选项。",
            "Codex",
        )

    status_col = sheet.cell(row=1, column=field_to_col["review_status"]).column_letter
    labeled_value = option_label("labeled", REVIEW_STATUSES["labeled"])
    needs_discussion_value = option_label("needs_discussion", REVIEW_STATUSES["needs_discussion"])
    full_range = f"A3:{sheet.cell(row=3, column=len(V2_COLUMNS)).column_letter}{max_row}"
    sheet.conditional_formatting.add(
        full_range,
        FormulaRule(formula=[f'${status_col}3="{labeled_value}"'], fill=PatternFill("solid", fgColor="E9F7EF")),
    )
    sheet.conditional_formatting.add(
        full_range,
        FormulaRule(
            formula=[f'${status_col}3="{needs_discussion_value}"'],
            fill=PatternFill("solid", fgColor="FFF1CC"),
        ),
    )

    workbook["下拉选项"].sheet_state = "hidden"
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx)
    return len(seed_rows)


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export Chinese coach seed labeling XLSX workbook v2.")
    parser.add_argument("--seed-jsonl", type=Path, default=DEFAULT_SEED_JSONL)
    parser.add_argument("--focus-registry", type=Path, default=DEFAULT_FOCUS_REGISTRY)
    parser.add_argument("--output-xlsx", type=Path, default=DEFAULT_OUTPUT_XLSX)
    parser.add_argument("--limit", type=int, help="Optional row limit for trial labeling workbooks.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(sys.argv[1:] if argv is None else argv)
    row_count = export_xlsx(
        seed_jsonl=args.seed_jsonl,
        focus_registry=args.focus_registry,
        output_xlsx=args.output_xlsx,
        limit=args.limit,
    )
    print(json.dumps({"output_xlsx": str(args.output_xlsx), "row_count": row_count}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
