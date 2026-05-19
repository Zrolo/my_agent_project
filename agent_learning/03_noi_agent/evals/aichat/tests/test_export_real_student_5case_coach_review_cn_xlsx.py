import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from evals.aichat import export_real_student_5case_coach_review_cn_xlsx as xlsx_export


class RealStudentCoachReviewCnXlsxTest(unittest.TestCase):
    def test_export_xlsx_has_dropdowns_colors_and_chinese_options(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            input_csv = Path(tmpdir) / "coach_review_cn.csv"
            output_xlsx = Path(tmpdir) / "coach_review_cn.xlsx"
            with input_csv.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=xlsx_export.COACH_REVIEW_CN_COLUMNS)
                writer.writeheader()
                row = {field: "" for field in xlsx_export.COACH_REVIEW_CN_COLUMNS}
                row.update(
                    {
                        "干跑案例编号": "dryrun_20260519_01",
                        "候选轮次编号": "rs_screen_20260519_001",
                        "试点案例编号": "rs_online_20260519_001",
                        "30例候选序号": "1",
                        "时间桶": "2026-W19",
                        "题目名称/匿名题号": "匿名题目A",
                        "题目任务摘要": "见本地私有表",
                        "关键约束/输入输出摘要": "见本地私有表",
                        "当前学生状态摘要": "见本地私有表",
                        "是否有完整题目上下文": "不确定",
                        "题目摘要是否足够评分": "不确定",
                        "是否需要查看完整题面才能评分": "不确定",
                        "题目摘要来源": "线上AIChat元数据",
                        "题目/场景摘要": "见本地私有表",
                        "学生问题（已脱敏）": "见本地私有表",
                        "近期对话（已脱敏，可空）": "见本地私有表",
                        "学生代码片段（已脱敏，可空）": "见本地私有表",
                        "当前AIChat回复字段说明": "线上已展示回复，观察项，非实验条件",
                        "当前AIChat回复（已脱敏）": "见本地私有表",
                        "复核状态": "待复核",
                        "隐私复核状态": "待复核",
                        "知情/报告门": "待完成",
                    }
                )
                writer.writerow(row)

            row_count = xlsx_export.export_xlsx(input_csv=input_csv, output_xlsx=output_xlsx)
            workbook = load_workbook(output_xlsx)

        self.assertEqual(row_count, 1)
        self.assertIn("教练复核表", workbook.sheetnames)
        self.assertIn("填写说明", workbook.sheetnames)
        self.assertIn("下拉选项", workbook.sheetnames)
        self.assertIn("50-case字段对照", workbook.sheetnames)

        sheet = workbook["教练复核表"]
        headers = [sheet.cell(row=1, column=col).value for col in range(1, sheet.max_column + 1)]
        self.assertEqual(xlsx_export.COACH_REVIEW_CN_COLUMNS, headers)
        self.assertIn("题目名称/匿名题号", headers)
        self.assertIn("题目任务摘要", headers)
        self.assertIn("题目摘要是否足够评分", headers)
        self.assertIn("当前AIChat回复字段说明", headers)
        self.assertEqual("T2", sheet.freeze_panes)
        self.assertGreaterEqual(len(sheet.data_validations.dataValidation), 28)
        self.assertGreaterEqual(len(sheet.conditional_formatting), 6)

        options_text = "\n".join(
            str(cell.value or "")
            for row in workbook["下拉选项"].iter_rows()
            for cell in row
        )
        self.assertIn("轻微关键桥泄露", options_text)
        self.assertIn("重大关键桥泄露", options_text)
        self.assertIn("答案或代码泄露", options_text)
        self.assertIn("线上已展示回复，观察项，非实验条件", options_text)
        self.assertNotIn("minor_bridge_leakage", options_text)

        guide_text = "\n".join(
            str(cell.value or "")
            for row in workbook["填写说明"].iter_rows()
            for cell in row
        )
        self.assertIn("和 50-case 人审维度一致", guide_text)
        self.assertIn("下拉选择", guide_text)
        self.assertIn("知情/报告门", guide_text)
        self.assertIn("题目任务摘要", guide_text)
        self.assertIn("观察项，非实验条件", guide_text)


if __name__ == "__main__":
    unittest.main()
