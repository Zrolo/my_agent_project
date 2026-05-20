#!/usr/bin/env python3
"""Integrate completed 2-case adjudication into the 9-case workflow status.

This script intentionally produces gate-safe public status only. Private
aggregate distributions stay under .local_private, and no raw student text,
full code, full AIChat response, identities, hash salt, or reversible mappings
are exported.
"""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_SECOND_REVIEW = Path(".local_private/real_student_online_30case_second_coach_focus_packet_human_reviewed_20260520.xlsx")
DEFAULT_ADJUDICATION = Path(".local_private/real_student_online_2case_adjudication_packet_human_adjudicated_20260520.xlsx")
DEFAULT_PRIVATE_JSON = Path(".local_private/real_student_online_post_adjudication_integrated_summary_20260520.json")
DEFAULT_PRIVATE_MD = Path(".local_private/real_student_online_post_adjudication_integrated_summary_20260520.zh.md")
DEFAULT_PUBLIC_JSON = Path("docs/research/real_student_online_post_adjudication_public_status_20260520.json")
DEFAULT_PUBLIC_MD = Path("docs/research/real_student_online_post_adjudication_public_status_20260520.zh.md")


ALLOWED_ADJUDICATED_STATUSES = {"已裁决", "因隐私风险排除", "因上下文不足排除"}


def norm(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_rows(workbook_path: Path, sheet: str) -> list[dict[str, Any]]:
    wb = load_workbook(workbook_path, data_only=True)
    ws = wb[sheet]
    headers = [norm(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    rows: list[dict[str, Any]] = []
    for r in range(2, ws.max_row + 1):
        row = {headers[c - 1]: ws.cell(r, c).value for c in range(1, len(headers) + 1)}
        if any(norm(v) for v in row.values()):
            row["_row"] = r
            rows.append(row)
    return rows


def count_reportable(privacy_rows: list[dict[str, Any]]) -> int:
    return sum(
        1
        for row in privacy_rows
        if norm(row.get("隐私复核状态")) == "可内部复核" and norm(row.get("知情/报告门")) == "可报告"
    )


def ensure_completed_adjudication(rows: list[dict[str, Any]]) -> list[str]:
    errors: list[str] = []
    for row in rows:
        status = norm(row.get("裁决状态"))
        if status == "待裁决":
            errors.append(f"裁决评分 row {row.get('_row')}: still 待裁决")
        elif status not in ALLOWED_ADJUDICATED_STATUSES:
            errors.append(f"裁决评分 row {row.get('_row')}: unexpected 裁决状态={status}")
        if status == "已裁决":
            for col in [
                "裁决上下文是否足够",
                "裁决最终泄露标签",
                "裁决总体质量（1-5）",
                "裁决是否愿意给学生看",
                "裁决信心",
                "是否可纳入内部汇总",
            ]:
                if not norm(row.get(col)):
                    errors.append(f"裁决评分 row {row.get('_row')}: missing {col}")
    return errors


def integrate(
    second_review: Path,
    adjudication: Path,
    private_json: Path,
    private_md: Path,
    public_json: Path,
    public_md: Path,
) -> dict[str, Any]:
    if not second_review.exists():
        raise FileNotFoundError(f"missing second review workbook: {second_review}")
    if not adjudication.exists():
        raise FileNotFoundError(f"missing adjudication workbook: {adjudication}")

    second_score = read_rows(second_review, "二审评分")
    second_privacy = read_rows(second_review, "隐私与报告门")
    adjudication_score = read_rows(adjudication, "裁决评分")
    adjudication_privacy = read_rows(adjudication, "隐私与报告门")

    errors = ensure_completed_adjudication(adjudication_score)
    if errors:
        raise ValueError("adjudication workbook is not complete:\n" + "\n".join(errors))

    second_review_status = Counter(norm(row.get("二审状态")) for row in second_score)
    adjudication_status = Counter(norm(row.get("裁决状态")) for row in adjudication_score)
    adjudication_internal_summary = {
        "adjudication_context_sufficiency_counts_private": dict(Counter(norm(row.get("裁决上下文是否足够")) for row in adjudication_score)),
        "adjudication_leakage_label_counts_private": dict(Counter(norm(row.get("裁决最终泄露标签")) for row in adjudication_score)),
        "adjudication_willing_to_show_counts_private": dict(Counter(norm(row.get("裁决是否愿意给学生看")) for row in adjudication_score)),
        "adjudication_confidence_counts_private": dict(Counter(norm(row.get("裁决信心")) for row in adjudication_score)),
        "adjudication_internal_inclusion_counts_private": dict(Counter(norm(row.get("是否可纳入内部汇总")) for row in adjudication_score)),
    }

    reportable_second = count_reportable(second_privacy)
    reportable_adjudication = count_reportable(adjudication_privacy)
    reportable_total = reportable_second + reportable_adjudication

    private_summary = {
        "boundary": "Private post-adjudication workflow summary. It is not dialogue-state v3 main evidence and does not release case-level labels.",
        "second_review_case_count": len(second_score),
        "second_review_status_counts": dict(second_review_status),
        "adjudication_case_count": len(adjudication_score),
        "adjudication_status_counts": dict(adjudication_status),
        "second_review_reportable_rows_after_gate": reportable_second,
        "adjudication_reportable_rows_after_gate": reportable_adjudication,
        "combined_reportable_rows_after_gate": reportable_total,
        **adjudication_internal_summary,
        "public_reporting_boundary": "Public docs may report workflow completion and gate counts only unless privacy/consent/reporting gate is separately completed.",
    }

    public_summary = {
        "status_doc": public_md.name,
        "second_review_case_count": len(second_score),
        "adjudication_case_count": len(adjudication_score),
        "adjudication_completed_rows": sum(1 for row in adjudication_score if norm(row.get("裁决状态")) in ALLOWED_ADJUDICATED_STATUSES),
        "workflow_adjudication_closed": all(norm(row.get("裁决状态")) in ALLOWED_ADJUDICATED_STATUSES for row in adjudication_score),
        "combined_reportable_rows_after_gate": reportable_total,
        "reportable_case_level_evidence": reportable_total,
        "boundary": "Public-safe workflow status only; no case-level labels, raw student text, full code, full AIChat response, identities, hash salt, or reversible mappings are released.",
        "claim_gate": {
            "new_dialogue_state_v3_main_experiment": "no",
            "new_main_experiment_condition": "no",
            "recomputed_dialogue_state_v3_main_table": "no",
            "modified_online_aichat_or_prompt_or_active_mode": "no",
            "changed_student_visible_response": "no",
            "pilot_as_main_result": "no",
            "pilot_as_learning_outcome_study": "no",
            "public_case_level_labels": "no",
        },
    }

    private_json.parent.mkdir(parents=True, exist_ok=True)
    public_json.parent.mkdir(parents=True, exist_ok=True)
    private_json.write_text(json.dumps(private_summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    public_json.write_text(json.dumps(public_summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    private_md.write_text(
        "# Real-Student Online Post-Adjudication Integrated Private Summary 20260520\n\n"
        + "```json\n"
        + json.dumps(private_summary, ensure_ascii=False, indent=2)
        + "\n```\n",
        encoding="utf-8",
    )

    public_md.write_text(
        "# Real-Student Online Post-Adjudication Public Status 20260520\n\n"
        "## 使用边界\n\n"
        "本文档只记录 real-student online AIChat pilot 的 post-adjudication 公开安全 workflow status。"
        "它不公开学生原文、完整代码、完整 AIChat 回复、case-level labels、student hashes、problem hashes、hash salt 或可逆映射。\n\n"
        "## Public-Safe Status\n\n"
        "| item | value | boundary |\n"
        "| --- | ---: | --- |\n"
        f"| second-review focus cases | {len(second_score)} | internal workflow |\n"
        f"| adjudication cases | {len(adjudication_score)} | internal workflow |\n"
        f"| adjudication rows completed | {public_summary['adjudication_completed_rows']} | aggregate workflow status only |\n"
        f"| workflow adjudication closed | {'yes' if public_summary['workflow_adjudication_closed'] else 'no'} | not a paper result |\n"
        f"| reportable case-level evidence after gate | {reportable_total} | controlled by privacy/consent/reporting gate |\n\n"
        "## Interpretation\n\n"
        "This status can be used to track whether the private review workflow has closed. It does not support leakage-rate reporting, model-quality distributions, deployed-system superiority, or learning-outcome claims unless a separate public-reporting gate is completed.\n\n"
        "## Claim Gate\n\n"
        "| check | result |\n"
        "| --- | --- |\n"
        "| 新增 dialogue-state v3 主实验 | no |\n"
        "| 新增 main experiment condition | no |\n"
        "| 重算 dialogue-state v3 主表 | no |\n"
        "| 修改线上 AIChat / prompt / active mode | no |\n"
        "| 改变学生可见回复 | no |\n"
        "| 把 pilot 写成 main result | no |\n"
        "| 把 pilot 写成 learning outcome study | no |\n"
        "| 公开 case-level labels | no |\n",
        encoding="utf-8",
    )

    return {"private_summary": private_summary, "public_summary": public_summary}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--second-review", type=Path, default=DEFAULT_SECOND_REVIEW)
    parser.add_argument("--adjudication", type=Path, default=DEFAULT_ADJUDICATION)
    parser.add_argument("--private-json", type=Path, default=DEFAULT_PRIVATE_JSON)
    parser.add_argument("--private-md", type=Path, default=DEFAULT_PRIVATE_MD)
    parser.add_argument("--public-json", type=Path, default=DEFAULT_PUBLIC_JSON)
    parser.add_argument("--public-md", type=Path, default=DEFAULT_PUBLIC_MD)
    args = parser.parse_args()
    try:
        result = integrate(args.second_review, args.adjudication, args.private_json, args.private_md, args.public_json, args.public_md)
    except (FileNotFoundError, ValueError) as exc:
        print(
            json.dumps(
                {
                    "ok": False,
                    "error": str(exc),
                    "boundary": "No public status was generated; complete the private adjudication workbook first.",
                },
                ensure_ascii=False,
                indent=2,
            )
        )
        return 1
    print(json.dumps({"ok": True, **result["public_summary"]}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
