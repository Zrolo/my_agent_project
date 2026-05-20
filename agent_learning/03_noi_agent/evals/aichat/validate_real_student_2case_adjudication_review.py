#!/usr/bin/env python3
"""Validate the 2-case adjudication workbook for the real-student AIChat pilot."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_SHEETS = ["案例上下文", "二审参考", "裁决评分", "隐私与报告门"]

REQUIRED_COLUMNS = {
    "裁决评分": [
        "案例编号",
        "题目名称/匿名题号",
        "裁决状态",
        "裁决上下文是否足够",
        "裁决最终泄露标签",
        "裁决总体质量（1-5）",
        "裁决是否愿意给学生看",
        "裁决信心",
        "是否可纳入内部汇总",
        "裁决备注（中文）",
    ],
    "隐私与报告门": [
        "案例编号",
        "题目名称/匿名题号",
        "隐私复核状态",
        "知情/报告门",
        "裁决公开边界备注",
    ],
    "案例上下文": [
        "案例编号",
        "题目名称/匿名题号",
        "当前AIChat回复字段说明",
    ],
    "二审参考": [
        "案例编号",
        "题目名称/匿名题号",
        "二审状态",
        "二审是否需要裁决",
    ],
}

ALLOWED_VALUES = {
    "裁决状态": {"待裁决", "已裁决", "因隐私风险排除", "因上下文不足排除"},
    "裁决上下文是否足够": {"足够", "部分足够", "不足", "不清楚"},
    "裁决最终泄露标签": {"无泄露", "轻微关键桥泄露", "重大关键桥泄露", "答案或代码泄露", "上下文不足未判断"},
    "裁决总体质量（1-5）": {"5", "4", "3", "2", "1", 5, 4, 3, 2, 1},
    "裁决是否愿意给学生看": {"是", "边界", "否"},
    "裁决信心": {"高", "中", "低"},
    "是否可纳入内部汇总": {"是", "否", "待定"},
    "隐私复核状态": {"待复核", "可内部复核", "需继续脱敏", "因隐私风险排除"},
    "知情/报告门": {"待完成", "可报告", "不可报告", "已退出"},
}

REQUIRED_WHEN_ADJUDICATED = [
    "裁决上下文是否足够",
    "裁决最终泄露标签",
    "裁决总体质量（1-5）",
    "裁决是否愿意给学生看",
    "裁决信心",
    "是否可纳入内部汇总",
]


def norm(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_rows(ws: Any) -> tuple[list[str], list[dict[str, Any]]]:
    headers = [norm(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    rows: list[dict[str, Any]] = []
    for r in range(2, ws.max_row + 1):
        row = {headers[c - 1]: ws.cell(r, c).value for c in range(1, len(headers) + 1)}
        if any(norm(v) for v in row.values()):
            row["_row"] = r
            rows.append(row)
    return headers, rows


def err(errors: list[str], sheet: str, row: int | None, message: str) -> None:
    where = sheet if row is None else f"{sheet} row {row}"
    errors.append(f"{where}: {message}")


def validate(path: Path, expected_rows: int, allow_pending: bool) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    wb = load_workbook(path, data_only=True)

    for sheet in REQUIRED_SHEETS:
        if sheet not in wb.sheetnames:
            err(errors, sheet, None, "missing required sheet")
    if errors:
        return {"path": str(path), "errors": errors, "warnings": warnings}

    parsed: dict[str, tuple[list[str], list[dict[str, Any]]]] = {}
    for sheet in REQUIRED_SHEETS:
        headers, rows = read_rows(wb[sheet])
        parsed[sheet] = (headers, rows)
        for col in REQUIRED_COLUMNS[sheet]:
            if col not in headers:
                err(errors, sheet, None, f"missing required column: {col}")
        if len(rows) != expected_rows:
            err(errors, sheet, None, f"expected {expected_rows} rows, found {len(rows)}")
    if errors:
        return {"path": str(path), "errors": errors, "warnings": warnings}

    ids_by_sheet = {
        sheet: {norm(row.get("案例编号")) for row in rows}
        for sheet, (_, rows) in parsed.items()
    }
    for sheet, ids in ids_by_sheet.items():
        if len(ids) != expected_rows or "" in ids:
            err(errors, sheet, None, "case ids must be non-empty and unique")
    if len({frozenset(ids) for ids in ids_by_sheet.values()}) != 1:
        err(errors, "workbook", None, "case ids must match across required sheets")

    for row in parsed["案例上下文"][1]:
        if norm(row.get("当前AIChat回复字段说明")) != "线上已展示回复，观察项，非实验条件":
            err(errors, "案例上下文", int(row["_row"]), "当前AIChat回复字段说明 must be 线上已展示回复，观察项，非实验条件")

    adjudication_counts: Counter[str] = Counter()
    for row in parsed["裁决评分"][1]:
        row_idx = int(row["_row"])
        status = norm(row.get("裁决状态"))
        if not status:
            err(errors, "裁决评分", row_idx, "裁决状态 is required")
            continue
        adjudication_counts[status] += 1
        if status not in ALLOWED_VALUES["裁决状态"]:
            err(errors, "裁决评分", row_idx, f"invalid 裁决状态: {status}")
        if status == "待裁决" and not allow_pending:
            err(errors, "裁决评分", row_idx, "裁决状态 remains 待裁决; use --allow-pending only for template smoke tests")
        if status == "已裁决":
            for col in REQUIRED_WHEN_ADJUDICATED:
                if not norm(row.get(col)):
                    err(errors, "裁决评分", row_idx, f"{col} is required when 裁决状态=已裁决")
        for col, allowed in ALLOWED_VALUES.items():
            if col in row and norm(row.get(col)) and row.get(col) not in allowed and norm(row.get(col)) not in {norm(v) for v in allowed}:
                err(errors, "裁决评分", row_idx, f"invalid {col}: {row.get(col)}")

    privacy_counts: Counter[str] = Counter()
    reporting_counts: Counter[str] = Counter()
    for row in parsed["隐私与报告门"][1]:
        row_idx = int(row["_row"])
        privacy = norm(row.get("隐私复核状态"))
        reporting = norm(row.get("知情/报告门"))
        privacy_counts[privacy] += 1
        reporting_counts[reporting] += 1
        if privacy not in {norm(v) for v in ALLOWED_VALUES["隐私复核状态"]}:
            err(errors, "隐私与报告门", row_idx, f"invalid 隐私复核状态: {privacy}")
        if reporting not in {norm(v) for v in ALLOWED_VALUES["知情/报告门"]}:
            err(errors, "隐私与报告门", row_idx, f"invalid 知情/报告门: {reporting}")
        if reporting == "可报告" and privacy != "可内部复核":
            err(errors, "隐私与报告门", row_idx, "知情/报告门=可报告 requires 隐私复核状态=可内部复核")
        if privacy == "因隐私风险排除" and reporting == "可报告":
            err(errors, "隐私与报告门", row_idx, "privacy-risk excluded rows cannot be 可报告")

    reportable = sum(
        1
        for row in parsed["隐私与报告门"][1]
        if norm(row.get("隐私复核状态")) == "可内部复核" and norm(row.get("知情/报告门")) == "可报告"
    )
    return {
        "path": str(path),
        "expected_rows": expected_rows,
        "allow_pending": allow_pending,
        "adjudication_status_counts": dict(adjudication_counts),
        "privacy_review_status_counts": dict(privacy_counts),
        "consent_reporting_gate_counts": dict(reporting_counts),
        "reportable_rows_after_gate": reportable,
        "errors": errors,
        "warnings": warnings,
        "boundary": "Adjudication validator checks structure and gates only; it does not publish case-level labels.",
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--expected-rows", type=int, default=2)
    parser.add_argument("--allow-pending", action="store_true")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = validate(args.workbook, args.expected_rows, args.allow_pending)
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.write_text(text + "\n", encoding="utf-8")
    print(text)
    return 1 if result["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
