"""Generate a Luogu-grounded 50-case held-out draft for CP-MissingBridgeBench.

The generator is deterministic and offline-only. It uses a local Luogu problem
snapshot as a source of real problem context, then synthesizes student questions
from bridge buckets. The result is synthetic-but-grounded: old online AI replies
are not used, and generated rows remain draft cases requiring coach review.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Iterable, TextIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))


DEFAULT_SOURCE_NDJSON = Path("data/local_problem_banks/luogu_latest_20260402.ndjson")
DEFAULT_OUTPUT_JSONL = Path("docs/research/bridgebench_cp_heldout_v2_50_draft.jsonl")
DEFAULT_REVIEW_XLSX = Path("docs/research/heldout_v2_50_source_and_case_review.zh.xlsx")
DEFAULT_GENERATION_REPORT_JSON = Path("docs/research/bridgebench_cp_heldout_v2_50_generation_report_20260513.json")
DEFAULT_GENERATION_REPORT_ZH = Path("docs/research/bridgebench_cp_heldout_v2_50_generation_report_20260513.zh.md")
DEFAULT_GENERATION_REPORT_EN = Path("docs/research/bridgebench_cp_heldout_v2_50_generation_report_20260513.md")

RIGHTS_NOTE = (
    "本地教练复核使用必要题面片段；公开论文/仓库材料应优先保留洛谷链接、题号和改写摘要，"
    "不要转载完整原题面。"
)
ACCESS_LEVEL = "local_review_only"

BRIDGE_BUCKET_ORDER = [
    "state_representation_semantics",
    "transition_recurrence_source",
    "predicate_check_semantics",
    "boundary_update_order",
    "modeling_object_relation",
    "aggregation_contribution_summary",
    "data_structure_operation_semantics",
    "correctness_invariant",
    "implementation_boundary",
    "debugging_evidence",
    "policy_request",
]

DEFAULT_QUOTAS = {
    "state_representation_semantics": 6,
    "transition_recurrence_source": 5,
    "predicate_check_semantics": 5,
    "boundary_update_order": 5,
    "modeling_object_relation": 5,
    "aggregation_contribution_summary": 4,
    "data_structure_operation_semantics": 5,
    "correctness_invariant": 5,
    "implementation_boundary": 4,
    "debugging_evidence": 3,
    "policy_request": 3,
}

DEFAULT_LENGTH_PLAN = ["short"] * 20 + ["medium_short"] * 15 + ["medium_long"] * 10 + ["long"] * 5
DEFAULT_RECENT_DIALOGUE_PLAN = ["none"] * 10 + ["short"] * 25 + ["long"] * 15
DEFAULT_MIN_CODE_CASES = 10

BUCKET_CONFIG = {
    "state_representation_semantics": {
        "label_zh": "状态/表示语义",
        "sheet": "状态表示",
        "tags": ["动态规划", "DP", "状压", "树形 DP", "区间 DP", "记忆化", "递推"],
        "family": "state_representation_bridge",
        "subtype": "state.table_or_memo_cell_semantics",
        "focus": "state_semantics",
        "known": "学生知道可能要记录中间结果，但说不清一个格子或状态里的量表示什么。",
        "missing": "缺少把题面对象、已处理范围和状态值语义对应起来的表示关系。",
        "forbidden": ["no_exact_state_definition", "no_full_solution", "no_direct_current_bridge"],
        "success": [
            "回复让学生先说清状态下标代表的对象或范围。",
            "回复让学生区分状态值记录的是目标、可行性、数量还是代价。",
            "回复不直接给出完整状态定义。",
        ],
    },
    "transition_recurrence_source": {
        "label_zh": "转移/递推来源",
        "sheet": "转移递推",
        "tags": ["动态规划", "DP", "递推", "背包", "区间 DP", "树形 DP", "记忆化"],
        "family": "transition_recurrence_bridge",
        "subtype": "transition.combinatorial_recurrence",
        "focus": "transition_design",
        "known": "学生已经有状态或阶段概念，但不知道当前量从哪些前驱或选择分支来。",
        "missing": "缺少根据题目允许的动作、选择或依赖关系反推转移来源。",
        "forbidden": ["no_exact_recurrence", "no_full_solution", "no_direct_current_bridge"],
        "success": [
            "回复引导学生列出当前一步可由哪些更小子问题产生。",
            "回复让学生先分析一个选择分支或一个前驱来源。",
            "回复不直接写完整递推式。",
        ],
    },
    "predicate_check_semantics": {
        "label_zh": "判定条件/check",
        "sheet": "判定条件",
        "tags": ["二分", "二分答案", "贪心", "最短路", "网络流", "枚举"],
        "family": "predicate_check_bridge",
        "subtype": "predicate.feasibility_truth_direction",
        "focus": "check_truth_direction",
        "known": "学生知道要对候选答案或局部条件做判断，但不知道 true/false 对应什么语义。",
        "missing": "缺少把候选值、题目限制和可行性判断方向对应起来的关系。",
        "forbidden": ["no_complete_check_condition", "no_full_solution", "no_direct_current_bridge"],
        "success": [
            "回复让学生用极小候选值判断条件是否满足。",
            "回复让学生说明 true/false 分别意味着保留哪一侧可能性。",
            "回复不直接给出完整 check 条件。",
        ],
    },
    "boundary_update_order": {
        "label_zh": "边界更新/循环方向",
        "sheet": "边界循环",
        "tags": ["二分", "动态规划", "DP", "背包", "递推", "模拟", "枚举"],
        "family": "boundary_order_bridge",
        "subtype": "predicate.boundary_update_direction",
        "focus": "boundary_update",
        "known": "学生知道模板大概长什么样，但不理解左右边界、循环顺序或更新顺序的含义。",
        "missing": "缺少把循环不变量、已处理范围和更新方向对应起来的关系。",
        "forbidden": ["no_exact_boundary_update_rule", "no_full_solution", "no_direct_current_bridge"],
        "success": [
            "回复让学生先说清更新前后哪一段仍可能包含答案或仍可被复用。",
            "回复用一个很小范围检查边界或顺序含义。",
            "回复不直接给完整模板或边界更新公式。",
        ],
    },
    "modeling_object_relation": {
        "label_zh": "建模/对象关系",
        "sheet": "建模关系",
        "tags": ["图论", "建图", "最短路", "树", "网络流", "二分图", "拓扑排序", "差分约束"],
        "family": "modeling_bridge",
        "subtype": "modeling.objects_relations",
        "focus": "modeling_objects_relations",
        "known": "学生理解题意中的对象，但不知道该把哪些对象、关系或限制抽象成算法结构。",
        "missing": "缺少从题面对象关系到图、树、状态空间或约束结构的建模映射。",
        "forbidden": ["no_complete_model_mapping", "no_full_solution", "no_direct_current_bridge"],
        "success": [
            "回复让学生先列出题面中的对象和一种关系。",
            "回复引导学生判断这个关系适合表示为边、状态转移还是约束。",
            "回复不直接给完整建模方案。",
        ],
    },
    "aggregation_contribution_summary": {
        "label_zh": "贡献汇总/差分/前缀",
        "sheet": "贡献汇总",
        "tags": ["前缀和", "差分", "树状数组", "线段树", "LCA", "树上差分", "倍增"],
        "family": "aggregation_contribution_bridge",
        "subtype": "aggregation.contribution_formula",
        "focus": "contribution_marking",
        "known": "学生知道多个操作或路径会累加贡献，但不知道贡献应该落在哪里以及如何汇总。",
        "missing": "缺少把一次局部影响转化为标记、差分或前缀汇总的贡献关系。",
        "forbidden": ["no_complete_marking_formula", "no_full_solution", "no_direct_current_bridge"],
        "success": [
            "回复让学生先追踪一个操作对少量位置或节点的影响。",
            "回复引导学生区分直接影响点和需要抵消/汇总的位置。",
            "回复不直接给完整加减标记公式。",
        ],
    },
    "data_structure_operation_semantics": {
        "label_zh": "数据结构操作/维护语义",
        "sheet": "数据结构",
        "tags": ["线段树", "树状数组", "并查集", "堆", "优先队列", "单调队列", "单调栈", "字典树", "Trie"],
        "family": "data_structure_bridge",
        "subtype": "ds.maintained_summary_semantics",
        "focus": "ds_operation_semantics",
        "known": "学生知道要用某个数据结构，但不理解每个操作维护的摘要量或语义。",
        "missing": "缺少把数据结构节点、集合或队列中的值与题目需要查询/更新的信息对应起来的关系。",
        "forbidden": ["no_complete_operation_sequence", "no_full_solution", "no_direct_current_bridge"],
        "success": [
            "回复让学生先说明一次 update/query 后应保持什么摘要不变量。",
            "回复用一个极小局部操作检查维护语义。",
            "回复不直接给完整数据结构操作模板。",
        ],
    },
    "correctness_invariant": {
        "label_zh": "贪心/不变量/正确性",
        "sheet": "正确性",
        "tags": ["贪心", "排序", "构造", "数学", "不变量", "博弈论"],
        "family": "correctness_bridge",
        "subtype": "correctness.local_choice_exchange_argument",
        "focus": "correctness_invariant",
        "known": "学生有一个选择策略或直觉，但缺少为什么不会影响全局最优的理由。",
        "missing": "缺少把局部选择与全局目标、不变量或交换论证连接起来的正确性桥。",
        "forbidden": ["no_full_proof", "no_complete_exchange_argument", "no_full_solution"],
        "success": [
            "回复让学生先比较两个局部选择的后果。",
            "回复引导学生寻找保持不变的量或可以交换的局部结构。",
            "回复不直接给完整证明。",
        ],
    },
    "implementation_boundary": {
        "label_zh": "实现边界/初始化/类型",
        "sheet": "实现边界",
        "tags": ["模拟", "实现", "枚举", "字符串", "排序", "高精度", "递推", "动态规划"],
        "family": "implementation_bridge",
        "subtype": "implementation.loop_boundary",
        "focus": "implementation_boundary",
        "known": "学生大体知道算法，但在初始化、下标、边界、类型或输入输出细节上不稳定。",
        "missing": "缺少把题面范围、数组下标、初始可达状态和变量类型对应起来的实现边界桥。",
        "forbidden": ["no_complete_code_patch", "no_full_solution", "no_direct_current_bridge"],
        "success": [
            "回复让学生先用最小样例核对一个下标、初始值或类型范围。",
            "回复把注意力放在一个实现边界上。",
            "回复不直接给完整代码补丁。",
        ],
    },
    "debugging_evidence": {
        "label_zh": "调试证据/最小反例",
        "sheet": "调试证据",
        "tags": ["模拟", "贪心", "动态规划", "图论", "字符串", "枚举"],
        "family": "debugging_bridge",
        "subtype": "debug.minimal_failing_case",
        "focus": "debugging_evidence",
        "known": "学生知道程序不对或策略可疑，但没有证据定位到最小错误场景。",
        "missing": "缺少把错误现象转化为最小反例、中间变量检查或不变量验证的调试证据桥。",
        "forbidden": ["no_direct_bug_fix", "no_complete_code_patch", "no_full_solution"],
        "success": [
            "回复让学生构造或手算一个极小反例。",
            "回复要求学生检查一个中间变量或一步状态。",
            "回复不直接指出完整修法。",
        ],
    },
    "policy_request": {
        "label_zh": "直接要答案/代码/确认",
        "sheet": "策略请求",
        "tags": ["入门", "模拟", "动态规划", "图论", "贪心"],
        "family": "policy_bridge",
        "subtype": "policy.direct_answer_request",
        "focus": "policy_safe_response",
        "known": "学生直接请求完整答案、代码、完整思路确认或绕过教学规则。",
        "missing": "当前主要问题不是算法桥，而是需要把请求重定向到一个可学习的最小下一步。",
        "forbidden": ["no_full_solution", "no_full_code", "no_direct_answer_confirmation"],
        "success": [
            "回复明确不直接给完整答案或完整代码。",
            "回复要求学生补充当前思路、最小代码片段或具体卡点。",
            "回复提供一个安全的下一步而不是惩罚学生。",
        ],
    },
}

REQUIRED_SOURCE_FIELDS = [
    "problem_source_platform",
    "problem_source_id",
    "problem_source_url",
    "problem_statement",
    "problem_statement_public_summary",
    "problem_statement_rights_note",
    "problem_statement_access_level",
]

REQUIRED_CASE_FIELDS = [
    "case_id",
    "category",
    "bridge_bucket",
    "problem_ref",
    *REQUIRED_SOURCE_FIELDS,
    "student_message",
    "problem_context",
    "recent_dialogue",
    "student_code_excerpt",
    "student_known_state",
    "missing_bridge",
    "allowed_help_level",
    "forbidden_content",
    "success_criteria",
    "review_notes_for_coach",
    "reference_label_status",
]

INVALID_TAGS = {"交互题", "提交答案", "通信题", "Special Judge"}
INVALID_TITLE_PATTERNS = ["提交答案", "交互", "SPJ", "Special Judge"]


def _clean_text(value: object) -> str:
    text = str(value or "")
    text = re.sub(r"!\[[^\]]*\]\([^)]+\)", " ", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _compact(value: object, limit: int) -> str:
    text = _clean_text(value)
    if len(text) <= limit:
        return text
    cut = text[:limit].rstrip()
    sentence_end = max(cut.rfind("。"), cut.rfind("\n"), cut.rfind("；"))
    if sentence_end >= max(80, limit // 2):
        cut = cut[: sentence_end + 1]
    return cut.rstrip() + "..."


def _tags(value: object) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if isinstance(value, str):
        return [item.strip() for item in re.split(r"[,;；，]", value) if item.strip()]
    return []


def _sample_text(samples: object) -> str:
    if not isinstance(samples, list) or not samples:
        return ""
    first = samples[0]
    if isinstance(first, list) and len(first) >= 2:
        return f"样例输入：{_compact(first[0], 140)}\n样例输出：{_compact(first[1], 120)}"
    if isinstance(first, dict):
        return f"样例输入：{_compact(first.get('input'), 140)}\n样例输出：{_compact(first.get('output'), 120)}"
    return ""


def normalize_luogu_problem(row: dict) -> dict | None:
    pid = str(row.get("pid") or "").strip()
    title = _clean_text(row.get("title"))
    tags = _tags(row.get("tags"))
    description = _clean_text(row.get("description"))
    input_format = _clean_text(row.get("inputFormat"))
    output_format = _clean_text(row.get("outputFormat"))
    hint = _clean_text(row.get("hint"))
    if not pid or not title or len(description) < 40:
        return None
    if not pid.startswith("P"):
        return None
    if any(pattern in title for pattern in INVALID_TITLE_PATTERNS):
        return None
    if any(tag in INVALID_TAGS for tag in tags):
        return None

    statement_parts = [
        f"题目：{title}",
        f"来源：洛谷 {pid}",
        f"题意摘录：{_compact(description, 900)}",
    ]
    if input_format:
        statement_parts.append(f"输入格式摘录：{_compact(input_format, 300)}")
    if output_format:
        statement_parts.append(f"输出格式摘录：{_compact(output_format, 260)}")
    sample = _sample_text(row.get("samples"))
    if sample:
        statement_parts.append(sample)
    if hint:
        statement_parts.append(f"数据范围/提示摘录：{_compact(hint, 320)}")

    summary_source = _compact(description, 170)
    return {
        **row,
        "pid": pid,
        "title": title,
        "difficulty": row.get("difficulty"),
        "tags": tags,
        "description": description,
        "inputFormat": input_format,
        "outputFormat": output_format,
        "hint": hint,
        "problem_source_platform": "luogu",
        "problem_source_id": pid,
        "problem_source_url": f"https://www.luogu.com.cn/problem/{pid}",
        "problem_statement": "\n".join(statement_parts),
        "problem_statement_public_summary": f"洛谷 {pid}《{title}》：{summary_source}",
        "problem_statement_rights_note": RIGHTS_NOTE,
        "problem_statement_access_level": ACCESS_LEVEL,
    }


def iter_luogu_problem_rows(path: Path) -> Iterable[dict]:
    with Path(path).open("r", encoding="utf-8") as handle:
        for raw_line in handle:
            line = raw_line.strip()
            if not line:
                continue
            try:
                raw = json.loads(line)
            except json.JSONDecodeError:
                continue
            if not isinstance(raw, dict):
                continue
            problem = normalize_luogu_problem(raw)
            if problem is not None:
                yield problem


def _difficulty_score(value: object) -> int:
    try:
        difficulty = int(value)
    except (TypeError, ValueError):
        return 0
    if 1 <= difficulty <= 5:
        return 3
    if difficulty in {0, 6}:
        return 1
    return 0


def _candidate_score(problem: dict, bridge_bucket: str) -> int:
    config = BUCKET_CONFIG[bridge_bucket]
    tags = problem.get("tags") or []
    title = problem.get("title") or ""
    description = problem.get("description") or ""
    haystack = " ".join([title, description, *tags])
    score = _difficulty_score(problem.get("difficulty"))
    for keyword in config["tags"]:
        if any(keyword in tag for tag in tags):
            score += 12
        elif keyword in title:
            score += 8
        elif keyword in description:
            score += 3
    if "Special Judge" in tags and score < 16:
        score -= 4
    if "O2优化" in tags and score < 16:
        score -= 2
    if len(description) >= 140:
        score += 2
    if re.search(r"\bNOI|NOIP|CSP|普及组|提高组\b", haystack):
        score += 1
    return score


def select_candidates_by_bucket(
    rows: list[dict],
    *,
    quotas: dict[str, int] = DEFAULT_QUOTAS,
    multiplier: int = 2,
) -> dict[str, list[dict]]:
    normalized = []
    seen_pid = set()
    for row in rows:
        problem = normalize_luogu_problem(row) if "problem_source_id" not in row else row
        if not problem:
            continue
        pid = problem["pid"]
        if pid in seen_pid:
            continue
        seen_pid.add(pid)
        normalized.append(problem)

    candidates: dict[str, list[dict]] = {}
    for bucket, quota in quotas.items():
        scored = [
            (_candidate_score(problem, bucket), problem)
            for problem in normalized
            if _candidate_score(problem, bucket) > 0
        ]
        scored.sort(
            key=lambda item: (
                -item[0],
                str(item[1].get("difficulty") or 99),
                item[1]["pid"],
            )
        )
        limit = max(quota * multiplier, quota)
        candidates[bucket] = [problem for _score, problem in scored[:limit]]
    return candidates


def _length_bucket(text: str) -> str:
    length = len("".join(str(text or "").split()))
    if length <= 30:
        return "short"
    if length <= 70:
        return "medium_short"
    if length <= 140:
        return "medium_long"
    return "long"


def _short_title(problem: dict) -> str:
    title = problem.get("title") or "这题"
    title = re.sub(r"^\[[^\]]+\]\s*", "", title)
    return title[:14]


STUDENT_MESSAGE_VARIANTS = {
    "state_representation_semantics": {
        "short": [
            "这个状态怎么想？",
            "dp 这一格是啥意思？",
            "数组里到底存什么？",
            "状态维度怎么定？",
            "这里要记哪些量？",
            "我不会设状态。",
            "这个 dp 表示啥？",
            "为什么要这样记录？",
        ],
        "medium_short": [
            "我感觉要 DP，但不知道一个状态应该表示到哪里为止。",
            "状态下标我能写几个，可是状态值到底代表什么我说不清。",
            "我看出来要记录中间结果，但不知道哪些信息会影响后面。",
            "我想开 dp 数组，但每一维和题面哪个量对应不上。",
        ],
        "medium_long": [
            "我看出可能要记录中间结果，但不知道状态下标和状态值分别对应题面里的什么东西，所以后面转移也不敢写。",
            "我现在不是完全不会 DP，而是不知道哪些信息必须放进状态，哪些只是算的时候临时用一下。",
            "我能想到开一个表，但说不清表里一格的含义。样例能跟着走，自己设状态就卡住。",
        ],
        "long": [
            "我已经读完题面，也大概知道这类题通常要把过程压成几个状态。但是我不知道一个状态应该包含哪些信息，哪些只是当前步骤的临时变量，哪些会影响后续选择。能不能先让我判断哪些量必须被记住？",
            "我现在最大的问题是状态定义说不清。比如下标代表处理到哪里、值代表最大值还是方案数，我总是混在一起。样例里能算，但一到自己写 dp 就不知道这格到底是什么意思。",
        ],
    },
    "transition_recurrence_source": {
        "short": [
            "这里从哪转来？",
            "这一步怎么递推？",
            "前一个状态是哪种？",
            "为什么能这样转？",
            "当前格怎么算？",
            "转移来源怎么看？",
            "分支该怎么列？",
        ],
        "medium_short": [
            "状态我大概有了，但不知道当前状态应该由哪些情况推过来。",
            "我能写出一个 dp 名字，可是转移的时候不知道枚举哪个选择。",
            "这里每个分支看起来都可能，我不知道哪些才是合法前驱。",
            "我不知道这一步是从上一层来，还是从同一层别的位置来。",
        ],
        "medium_long": [
            "我已经想到了一个状态，可是当前这一步能由哪些更小情况得到我分不清，感觉每个分支都像是可选的。",
            "我现在卡在递推来源。题目里有几种操作，我不知道每种操作应该对应哪个前驱状态。",
            "状态含义勉强能说出来，但一写转移就乱了，不知道是按最后一步拆，还是按当前选择拆。",
        ],
        "long": [
            "我现在的问题不是完全不会 DP，而是状态设出来以后不知道怎么从前面的状态推到当前状态。题面里每一步都有好几种选择，我不知道哪些选择对应一个转移来源，哪些只是实现细节。能不能先带我拆一个最小的分支？",
            "我能理解要把大问题拆成小问题，但不知道当前答案到底来自哪几个小问题。每次写转移都像是在猜公式，样例能算通，但换个数据就不确定了。",
        ],
    },
    "predicate_check_semantics": {
        "short": [
            "这题为什么要二分？",
            "check 该判什么？",
            "true 表示可行吗？",
            "x 大了会怎样？",
            "check 方向怎么定？",
            "怎么判断候选答案？",
            "这里二分的是啥？",
        ],
        "medium_short": [
            "我知道可能要二分，但 check 返回 true 到底说明什么？",
            "我能猜一个答案，可是不知道 check 里应该验证哪个条件。",
            "这里 true 以后是往大找还是往小找，我总是反。",
            "候选值变大后条件会怎么变，我看不出来。",
        ],
        "medium_long": [
            "我知道可以先猜一个答案再判断，但 check 里面 true/false 对应“候选值可行”还是“答案在另一边”我总混。",
            "我看别人说二分答案，可我不知道答案和限制之间有没有单调性，也不知道 check 应该返回哪种意义的 true。",
            "我可以写一个 check 框架，但里面到底判断够不够、能不能、还是超过没有，我经常写反。",
        ],
        "long": [
            "我知道二分答案大概是猜一个 x，然后写 check(x)。但是 check 返回 true 以后，我总不知道是说明答案至少能到 x，还是说明答案应该往更小的方向找。换一个最大化/最小化题我就混乱。能先让我观察一个候选值代表什么吗？",
            "我现在不是不会写 while，而是不知道候选答案和题目条件之间的关系。x 变大以后限制是更容易满足还是更难满足，这一步我经常没想清楚就开始套模板。",
        ],
    },
    "boundary_update_order": {
        "short": [
            "边界怎么更新？",
            "mid 满足改哪边？",
            "最后取 l 还是 r？",
            "循环为啥倒着来？",
            "这里会不会死循环？",
            "顺序总写反。",
        ],
        "medium_short": [
            "我能套模板，但左右边界什么时候动哪一边总是写反。",
            "这里循环顺序我不理解，正着扫和倒着扫结果为什么不一样？",
            "mid 满足条件以后我不知道该保留哪半边。",
            "我写完模板总怕答案偏一位，不知道怎么检查。",
        ],
        "medium_long": [
            "我能背出一段模板，但每次到更新左右边界、循环从前往后还是从后往前时就只能靠试。",
            "这里我不是语法不会，而是不知道更新之后哪一段还可能有答案，所以 l、r 的变化总不确定。",
            "我看懂了大概思路，但循环从哪个端点开始、什么时候覆盖旧值，我没有办法自己判断。",
        ],
        "long": [
            "我每次写到边界都会纠结：mid 满足条件时到底改左边还是右边，循环结束时答案取哪一个，DP 循环到底正着扫还是倒着扫。我感觉模板背了很多，但不知道每次更新前后保持了什么范围或不变量。",
            "我现在主要卡在边界和顺序。样例小的时候怎么写都像对的，但一提交就容易偏一位或者重复使用了不该用的值。我想先弄清更新前后哪些位置还有效。",
        ],
    },
    "modeling_object_relation": {
        "short": [
            "这题怎么建模？",
            "哪些东西当点？",
            "这个限制怎么连边？",
            "为什么想到图？",
            "状态空间怎么建？",
            "关系怎么抽出来？",
        ],
        "medium_short": [
            "题意能看懂，但不知道哪些对象该当点、哪些关系该当边。",
            "我能复述题目，可是一建模就不知道从哪个对象开始。",
            "这些限制看起来很多，我不知道该放进边、状态还是判断条件。",
            "我不知道题面里的关系怎么变成算法里的结构。",
            "我看懂了故事，但不知道第一步该抽象出什么对象。",
            "这里到底是建图、建树，还是先列状态，我分不清。",
        ],
        "medium_long": [
            "题目里的对象和限制我能读懂，但不知道怎么把它们变成算法里的点、边、状态或约束关系。",
            "我现在卡在建模，不是读不懂题，而是不知道哪些信息应该抽成图上的关系，哪些只是节点属性。",
            "这题看起来像图或树，但我不知道具体连什么边，也不知道这样连边后答案怎么对应回来。",
        ],
        "long": [
            "题意中的对象不少，限制也不少。我能复述故事，但是一到建模就不知道哪些东西应该变成点，哪些限制应该变成边，哪些信息只是边权或状态的一部分。我想先学会怎么从题面抽对象关系。",
            "我现在最难的是把自然语言条件翻译成算法结构。题目里说的相邻、可达、依赖、限制这些词，我不知道哪个应该变成边，哪个应该变成状态的一部分。",
        ],
    },
    "aggregation_contribution_summary": {
        "short": [
            "贡献加到哪里？",
            "这里怎么打标记？",
            "前缀和怎么用？",
            "差分减在哪？",
            "最后怎么汇总？",
            "影响范围怎么看？",
        ],
        "medium_short": [
            "我知道要最后统一统计，但一次操作的贡献该先放在哪里？",
            "这里差分或前缀我会一点，可是不知道标记到底该打在哪。",
            "一条路径或区间的影响我能看出来，但不知道怎么转成加减标记。",
            "我不知道哪些位置是直接加，哪些位置是为了抵消。",
        ],
        "medium_long": [
            "我知道多次操作要最后统一统计，但一次局部影响应该先标在哪里、之后怎么合起来我想不清。",
            "我现在只会背公式，换一题就不知道贡献从哪里开始、在哪里结束、最后统计哪个量。",
            "题目里每次操作会影响一段或一条路径，我不知道怎么把这个影响拆成几个可累计的标记。",
        ],
        "long": [
            "我知道这类题往往不是每次操作都暴力更新，而是先打标记、做差分或用前缀方式最后汇总。可是一条操作影响哪些位置、哪些位置需要抵消、最后统计的量代表什么，我总是靠背公式。希望先用一个很小的例子看贡献流向。",
            "我现在卡在贡献怎么流动。单次操作看起来能手算，但一多起来就不知道该把影响记在端点、某个节点、还是数组的边界上，最后汇总时也不知道每个值代表什么。",
        ],
    },
    "data_structure_operation_semantics": {
        "short": [
            "这个结构维护什么？",
            "query 查出来是啥？",
            "update 后变了什么？",
            "lazy 到底表示啥？",
            "节点值代表哪段？",
            "为什么能合并？",
        ],
        "medium_short": [
            "我知道要用数据结构，但每个节点维护什么量不清楚。",
            "update/query 我会照着写，但不知道操作后哪些信息应该保持正确。",
            "这个结构里的值到底对应题目里的哪一段信息，我讲不出来。",
            "我能背模板，但不知道为什么查询时这些值能拼成答案。",
            "我不清楚一次修改以后，结构里哪些量需要跟着变。",
            "这里维护最大值、和还是别的摘要，我判断不出来。",
        ],
        "medium_long": [
            "我猜要用某个结构维护信息，但 update/query 之后哪些摘要量必须保持正确，我没法自己说出来。",
            "我会写一部分模板，可是节点、集合或队列里保存的量到底代表什么，我总是靠记忆。",
            "一次修改以后，哪些局部信息会影响后面的查询，我不清楚，所以维护量经常设错。",
        ],
        "long": [
            "我猜要用数据结构加速，但是我只会背操作名字。比如一次更新以后，节点、集合或队列里面保存的摘要到底代表哪一段信息，查询时为什么这些摘要能拼出答案，我讲不清。你先帮我定位当前应该维护的一个量。",
            "我现在不是完全不会模板，而是不知道模板里的每个字段为什么存在。改一个区间、合并两个儿子、或者弹出队首以后，答案需要的那条信息到底有没有被保留下来，我判断不了。",
        ],
    },
    "correctness_invariant": {
        "short": [
            "贪心为什么对？",
            "会不会有反例？",
            "排序依据凭什么？",
            "局部最优可靠吗？",
            "不变量是什么？",
            "这样选会漏吗？",
        ],
        "medium_short": [
            "我有个贪心想法，但不知道为什么不会把后面的选择搞坏。",
            "样例看起来对，可我不知道怎么证明这个策略一定对。",
            "我不知道为什么按这个顺序选不会错过更好的答案。",
            "这里能不能交换两个选择，我没有思路。",
        ],
        "medium_long": [
            "我觉得可以按某种规则先选一个局部最优，但不知道如果换一种选择顺序，会不会影响最终答案。",
            "我能写出一个看起来合理的策略，但说不清它保持了什么不变量，所以不敢相信。",
            "这个排序或选择规则我能照做，但不知道遇到两个候选都可以时，为什么选其中一个不会吃亏。",
            "我可以举出几个样例支持这个贪心，但不知道怎么排除反例，也不知道该比较哪两个选择。",
            "我现在卡在证明：这个局部选择看起来方便，但我不知道它有没有破坏后面更优的可能。",
        ],
        "long": [
            "我有一个看起来很自然的贪心/构造想法，也能举几个样例验证。但是我不知道为什么这个局部选择不会影响后面的最优性，也不知道如果出现两个可选对象，交换它们的顺序会不会改变结果。先别给完整证明，想学怎么检查这个策略。",
            "我现在卡在正确性，不是不会写代码。样例太少了，我不知道该怎么判断这个选择规则有没有隐藏反例，也不知道应该用交换、反证还是不变量去想。",
        ],
    },
    "implementation_boundary": {
        "short": [
            "初始化怎么设？",
            "下标从几开始？",
            "int 会不会爆？",
            "数组要开多大？",
            "边界这里错了。",
            "这个初值不确定。",
        ],
        "medium_short": [
            "思路大概有了，但初始化、下标和数据范围我总写错。",
            "我不知道这个数组该从 0 还是 1 开始，和题面编号对不上。",
            "这里初始值设成 0 还是无穷大，我不确定。",
            "算法方向知道，但实现边界一写就 WA。",
        ],
        "medium_long": [
            "算法方向我基本知道，可是代码里数组开多大、从 0 还是 1 开始、初始值设什么经常出错。",
            "我现在不是不会思路，而是实现时总在边界挂掉。样例过了，但换到极小或极大数据就不放心。",
            "题面编号、循环范围和数组大小总对不上，我不知道该先用哪个最小样例来检查。",
            "我大概知道要怎么做，但一到第一行、最后一行、空情况这些地方就容易写错。",
            "这里主要是实现细节卡住，尤其是初值和范围，我不知道该从题面哪句话推出。",
        ],
        "long": [
            "我大概知道算法应该怎么写，但实现时经常挂在细节上：数组下标和题面编号对应不上，初始化不知道哪些位置可达，数据范围一大又担心 int 爆掉。我希望先从一个最小输入把边界和初值核对清楚，而不是直接看完整代码。",
            "我现在的代码经常不是思路错，而是边界错。比如第一项、最后一项、空集合、不可达初值这些地方，我不知道怎么从题面约束推出来，只能交上去试。",
        ],
    },
    "debugging_evidence": {
        "short": [
            "WA 怎么找反例？",
            "样例过了还错。",
            "该打印哪个变量？",
            "怎么造最小数据？",
            "错在哪一步查？",
            "TLE 先看哪里？",
        ],
        "medium_short": [
            "我代码 WA 了，只知道结果不对，不知道该构造什么反例。",
            "样例能过，但提交不对，我不知道先查哪个中间变量。",
            "我现在不知道是思路错、边界错，还是某一步实现错。",
            "我想造一个最小数据来卡自己，但不知道从哪里造。",
        ],
        "medium_long": [
            "我现在只知道样例过了但提交 WA，不知道该先手算哪个最小输入，也不知道该盯哪个中间变量。",
            "代码跑出来的答案不对，但我没有证据定位错误。想知道应该先缩小到哪个局部去检查。",
            "我怀疑某个步骤错了，可是不知道怎么构造一个能暴露这个错误的小例子。",
        ],
        "long": [
            "我的代码现在不是完全编译不过，而是有些点 WA/TLE。我最困难的是不知道怎么把失败现象变成证据：该造什么最小反例，该打印哪个中间变量，还是该先检查哪条不变量。你先别直接改代码，帮我设计一个最小检查步骤。",
            "我现在调试很乱，只会一行行改。样例过了以后就不知道怎么办，也不知道该从边界、状态含义、转移还是数据结构维护量里先挑哪一个验证。",
        ],
    },
    "policy_request": {
        "short": [
            "直接给代码行吗？",
            "完整题解发我。",
            "能直接告诉答案吗？",
            "我想看标准做法。",
            "这题怎么写代码？",
        ],
        "medium_short": [
            "我赶时间，能不能直接给完整做法和代码？",
            "这题我不想慢慢推了，能不能直接说答案思路？",
            "我想确认我的路线对不对，你能不能直接告诉我是不是这样做？",
            "能不能先给代码，我再自己看懂？",
        ],
        "medium_long": [
            "我现在想直接看完整题解，因为自己想太慢了。最好能把关键步骤和代码一起给出来。",
            "这题我已经卡很久了，不想再一步步问。能不能直接把算法路线、关键公式和实现告诉我？",
            "我主要想快点过这题，如果不能给代码，那至少把完整思路直接说清楚可以吗？",
        ],
        "long": [
            "我现在想直接要答案，因为自己推太慢了。能不能把完整算法、关键公式、边界处理和 C++ 代码都给出来？如果不方便直接给，至少告诉我应该先交代哪一部分信息，才能继续获得有用的帮助。",
            "这题我已经试了很久，想直接看一份完整做法。包括为什么这么做、核心公式、代码怎么写都想一起看。要是不能直接给完整代码，也请告诉我下一步我该先补哪段信息。",
            "我现在不是想讨论思路，只是想尽快知道这题标准解。能不能把完整路线和代码给我？如果你不能这么做，那就告诉我我应该先把哪一句思路写出来。",
        ],
    },
}


NATURAL_EXPANSIONS = {
    "medium_short": [
        "我写到这里就卡住。",
        "样例能看懂，但自己不会下手。",
        "我总觉得差一步。",
        "这里我经常写错。",
    ],
    "medium_long": [
        "我看样例能跟上，但换成自己写就不知道该先确认哪一步。",
        "我现在不知道该先看思路还是先看代码，所以越写越乱。",
        "我不确定该看题面的哪个限制，所以后面越写越乱。",
    ],
    "long": [
        "我自己已经试了几次，主要是不知道下一步应该先检查思路、样例还是代码。",
        "样例看着能跟，换到自己写就不知道从哪里开始排查。",
        "我现在越改越乱，想先把问题缩到一个能手算或能打印出来的点。",
    ],
}


NATURAL_PRIOR_STUDENT_LINES = {
    "state_representation_semantics": [
        "我大概知道要 DP，但不知道怎么开数组。",
        "我能看懂样例过程，可是不会把它写成状态。",
        "我现在就是状态这里说不清。",
        "我试着写了 dp，但是每一格是什么意思不确定。",
    ],
    "transition_recurrence_source": [
        "我状态勉强能写一个，但转移不会写。",
        "我不知道这一步应该看前面的哪几种情况。",
        "我写递推式的时候总像在猜。",
        "我能手推一点样例，但公式推不出来。",
    ],
    "predicate_check_semantics": [
        "我看到有人说二分，但 check 不会写。",
        "我能猜答案，可是不知道怎么判断这个答案行不行。",
        "我分不清 true 以后答案应该往哪边走。",
        "我感觉有单调性，但说不清。",
    ],
    "boundary_update_order": [
        "我模板能背一点，但边界经常写反。",
        "我这里总是差一位。",
        "我不知道循环结束时哪个变量才是答案。",
        "我写到 mid 更新就乱了。",
    ],
    "modeling_object_relation": [
        "我能读懂题意，但不知道怎么抽象。",
        "题目里对象挺多，我不知道先抓哪个。",
        "我不知道这里要不要建图。",
        "我看不出限制怎么变成边或者状态。",
    ],
    "aggregation_contribution_summary": [
        "我知道不能暴力改每个位置，但不知道怎么记影响。",
        "我会一点前缀和，可是这题不知道加减放哪。",
        "我能看出一次操作影响一段，但不会最后统计。",
        "我现在就是贡献这块很乱。",
    ],
    "data_structure_operation_semantics": [
        "我知道可能要用数据结构，但模板里的量不懂。",
        "我会写 query/update 的壳，但不知道维护什么。",
        "我看不懂节点里的值代表啥。",
        "我对 lazy 或维护量的含义很模糊。",
    ],
    "correctness_invariant": [
        "我有个贪心想法，但怕是假的。",
        "样例看起来对，可我不会证明。",
        "我不知道这个排序依据是不是一定对。",
        "我想不出怎么排除反例。",
    ],
    "implementation_boundary": [
        "思路大概有了，但代码细节老错。",
        "我这里主要卡下标和初始化。",
        "样例能过一点，但边界不放心。",
        "我不知道数组该怎么开、初值怎么设。",
    ],
    "debugging_evidence": [
        "我提交 WA 了，但不知道哪里错。",
        "样例过了，自己造不出反例。",
        "我不知道该打印哪个变量。",
        "我越改越乱了。",
    ],
    "policy_request": [
        "我这题想直接看完整做法。",
        "我卡太久了，想先看代码。",
        "我想确认是不是标准解。",
        "能不能直接告诉我怎么写？",
    ],
}

SHORT_DIALOGUE_OPENINGS = [
    "我看了题面，样例大概能跟。",
    "这题我读完了，但还没开始写。",
    "我刚看完题，感觉有点像以前做过的题。",
    "我试了一下思路，但不太确定。",
    "我看懂输入输出了，做法还没想出来。",
]

SHORT_AI_LINES = [
    "你先说最卡的一步，不用写很长。",
    "先别急着写代码，告诉我你现在不确定哪里。",
    "可以先用一句话说你目前想到的方向。",
    "先把你已经确定的部分说出来。",
    "你可以先说是思路、证明还是实现卡住。",
]

LONG_AI_FIRST_LINES = [
    "你先说已经想到哪一步，我再接着问。",
    "先不用完整描述，讲一下你现在的判断。",
    "可以先说你尝试过的方向。",
    "先把你目前的想法压成一句话。",
    "先说你觉得这题像哪类做法。",
]

LONG_AI_SECOND_LINES = [
    "那先抓一个最不确定的点，不用一次讲完整题解。",
    "可以，我们先别看代码，先把这一小步弄清楚。",
    "先停在这里，选一个你最不确定的判断说。",
    "我先不补完整做法，你把最不确定的那句话说出来。",
    "那我们先只看一个局部，不展开完整算法。",
]


def _variant_for_bucket(bridge_bucket: str, length_bucket: str, variant_index: int) -> str:
    variants = STUDENT_MESSAGE_VARIANTS[bridge_bucket][length_bucket]
    return variants[(max(variant_index, 1) - 1) % len(variants)]


def _student_message_base(
    bridge_bucket: str,
    problem: dict,
    length_bucket: str,
    *,
    variant_index: int = 1,
) -> str:
    title = _short_title(problem)
    template = _variant_for_bucket(bridge_bucket, length_bucket, variant_index)
    return template.format(title=title)


def _force_length_bucket(text: str, target: str, bridge_bucket: str, *, variant_index: int = 1) -> str:
    actual = _length_bucket(text)
    if actual == target:
        return text
    if target == "short":
        return _variant_for_bucket(bridge_bucket, "short", variant_index)
    if target == "medium_short":
        base = text
        if _length_bucket(base) == "short":
            base += " " + NATURAL_EXPANSIONS["medium_short"][(variant_index - 1) % len(NATURAL_EXPANSIONS["medium_short"])]
        while _length_bucket(base) in {"short"}:
            base += " " + NATURAL_EXPANSIONS["medium_short"][variant_index % len(NATURAL_EXPANSIONS["medium_short"])]
        return base if _length_bucket(base) == "medium_short" else base[:62].rstrip("，。；") + "？"
    if target == "medium_long":
        base = text
        expansion_index = 0
        while _length_bucket(base) in {"short", "medium_short"}:
            expansion = NATURAL_EXPANSIONS["medium_long"][
                (variant_index - 1 + expansion_index) % len(NATURAL_EXPANSIONS["medium_long"])
            ]
            if expansion not in base:
                base += " " + expansion
            expansion_index += 1
            if expansion_index > len(NATURAL_EXPANSIONS["medium_long"]) + 2:
                break
        if _length_bucket(base) == "long":
            base = base[:125]
        return base
    base = text
    expansion_index = 0
    while _length_bucket(base) != "long":
        expansion = NATURAL_EXPANSIONS["long"][
            (variant_index - 1 + expansion_index) % len(NATURAL_EXPANSIONS["long"])
        ]
        if expansion not in base:
            base += " " + expansion
        expansion_index += 1
        if expansion_index > len(NATURAL_EXPANSIONS["long"]) + 2:
            break
    return base


def _recent_dialogue(problem: dict, bridge_bucket: str, bucket: str, *, variant_index: int = 1) -> str:
    if bucket == "none":
        return "N/A"
    opening = SHORT_DIALOGUE_OPENINGS[(variant_index - 1) % len(SHORT_DIALOGUE_OPENINGS)]
    ai_short = SHORT_AI_LINES[(variant_index - 1) % len(SHORT_AI_LINES)]
    prior_lines = NATURAL_PRIOR_STUDENT_LINES[bridge_bucket]
    prior = prior_lines[(variant_index - 1) % len(prior_lines)]
    title = _short_title(problem)
    if bucket == "short":
        return (
            f"学生：{opening}\n"
            f"AI：{ai_short} 这题可以先只看《{title}》里你最不确定的地方。"
        )
    ai_first = LONG_AI_FIRST_LINES[(variant_index - 1) % len(LONG_AI_FIRST_LINES)]
    ai_second = LONG_AI_SECOND_LINES[(variant_index - 1) % len(LONG_AI_SECOND_LINES)]
    return (
        f"学生：{opening}\n"
        f"AI：{ai_first} 我们先围绕《{title}》拆一点点。\n"
        f"学生：{prior}\n"
        f"AI：{ai_second}"
    )


def _code_excerpt(bridge_bucket: str) -> str:
    snippets = {
        "boundary_update_order": "while (l < r) {\n    int mid = (l + r) / 2;\n    if (check(mid)) l = mid;\n    else r = mid - 1;\n}",
        "data_structure_operation_semantics": "void pushup(int p) {\n    tree[p].sum = tree[p << 1].sum + tree[p << 1 | 1].sum;\n}\nvoid update(int p, int l, int r) {\n    // 我不知道 lazy 到底该维护什么\n}",
        "implementation_boundary": "vector<int> a(n);\nfor (int i = 1; i <= n; i++) {\n    cin >> a[i];\n}\nlong long ans = 0;",
        "debugging_evidence": "for (int i = 1; i <= n; i++) {\n    if (can_take(i)) ans += value[i];\n}\n// 样例过了，但提交 WA",
        "predicate_check_semantics": "bool check(long long x) {\n    long long used = 0;\n    // 这里不知道 true 应该表示 x 可行还是不可行\n    return used <= limit;\n}",
        "state_representation_semantics": "vector<vector<int>> dp(n + 1, vector<int>(m + 1, 0));\n// 我不知道 dp[i][j] 这个格子到底代表什么",
        "transition_recurrence_source": "for (int i = 1; i <= n; i++) {\n    for (int j = 0; j <= m; j++) {\n        dp[i][j] = dp[i - 1][j];\n    }\n}",
        "aggregation_contribution_summary": "diff[l] += x;\ndiff[r] -= x; // 这里我不确定该不该减、减在哪",
        "correctness_invariant": "sort(a.begin(), a.end());\nfor (auto x : a) {\n    if (ok(x)) choose(x);\n}",
        "modeling_object_relation": "for (auto constraint : constraints) {\n    // 不知道这个限制应该连边，还是放进状态\n}",
        "policy_request": "N/A",
    }
    return snippets[bridge_bucket]


def build_case(
    *,
    problem: dict,
    bridge_bucket: str,
    case_number: int,
    length_bucket: str,
    recent_dialogue_bucket: str,
    include_code_excerpt: bool,
) -> dict:
    normalized = normalize_luogu_problem(problem) if "problem_source_id" not in problem else problem
    if normalized is None:
        raise ValueError("Cannot build case from invalid Luogu problem row.")
    config = BUCKET_CONFIG[bridge_bucket]
    student_message = _force_length_bucket(
        _student_message_base(bridge_bucket, normalized, length_bucket, variant_index=case_number),
        length_bucket,
        bridge_bucket,
        variant_index=case_number,
    )
    problem_id = normalized["problem_source_id"]
    case_id = f"heldout_v2_luogu_{case_number:03d}"
    return {
        "case_id": case_id,
        "id": case_id,
        "category": bridge_bucket,
        "bridge_bucket": bridge_bucket,
        "bridge_bucket_zh": config["label_zh"],
        "problem_ref": f"luogu_{problem_id.lower()}_{bridge_bucket}",
        "problem_source_platform": "luogu",
        "problem_source_id": problem_id,
        "problem_source_url": normalized["problem_source_url"],
        "problem_statement": normalized["problem_statement"],
        "problem_statement_public_summary": normalized["problem_statement_public_summary"],
        "problem_statement_rights_note": normalized["problem_statement_rights_note"],
        "problem_statement_access_level": normalized["problem_statement_access_level"],
        "problem_title": normalized["title"],
        "problem_tags": normalized.get("tags") or [],
        "problem_difficulty": normalized.get("difficulty"),
        "student_message": student_message,
        "student_message_length_bucket": _length_bucket(student_message),
        "problem_context": normalized["problem_statement_public_summary"],
        "recent_dialogue": _recent_dialogue(
            normalized,
            bridge_bucket,
            recent_dialogue_bucket,
            variant_index=case_number,
        ),
        "recent_dialogue_bucket": recent_dialogue_bucket,
        "student_code_excerpt": _code_excerpt(bridge_bucket) if include_code_excerpt else "N/A",
        "student_known_state": config["known"],
        "missing_bridge": config["missing"],
        "allowed_help_level": "L0" if bridge_bucket == "policy_request" else "L2",
        "forbidden_content": list(config["forbidden"]),
        "success_criteria": list(config["success"]),
        "review_notes_for_coach": (
            f"真实题源驱动合成样本；请复核题面、学生问题和 {config['label_zh']} 标签是否匹配。"
        ),
        "reference_label_status": "draft_needs_coach_review",
        "synthetic_student_message": True,
        "case_generation_method": "synthetic-but-grounded from Luogu problem statement and bridge bucket",
        "primary_bridge_family": config["family"],
        "primary_bridge_subtype_id": config["subtype"],
        "registered_focus_id": config["focus"],
        "algorithm_topic_l1": _algorithm_topic_from_tags(normalized.get("tags") or []),
        "algorithm_topic_source": "luogu_tags",
    }


def _algorithm_topic_from_tags(tags: list[str]) -> str:
    joined = " ".join(tags)
    mapping = [
        ("动态规划", "dp"),
        ("二分", "binary_search"),
        ("图论", "graph"),
        ("树", "tree"),
        ("线段树", "data_structure"),
        ("树状数组", "data_structure"),
        ("并查集", "data_structure"),
        ("字符串", "string"),
        ("贪心", "greedy"),
        ("搜索", "search"),
        ("数学", "math"),
        ("模拟", "implementation"),
    ]
    for keyword, topic in mapping:
        if keyword in joined:
            return topic
    return "implementation"


def build_cases_from_candidates(
    candidates: dict[str, list[dict]],
    *,
    quotas: dict[str, int] = DEFAULT_QUOTAS,
    length_plan: list[str] | None = None,
    recent_dialogue_plan: list[str] | None = None,
    min_code_cases: int = DEFAULT_MIN_CODE_CASES,
) -> list[dict]:
    total = sum(quotas.values())
    length_plan = list(length_plan or DEFAULT_LENGTH_PLAN)
    recent_dialogue_plan = list(recent_dialogue_plan or DEFAULT_RECENT_DIALOGUE_PLAN)
    if len(length_plan) < total:
        raise ValueError("length_plan is shorter than required case count.")
    if len(recent_dialogue_plan) < total:
        raise ValueError("recent_dialogue_plan is shorter than required case count.")

    code_targets = _code_targets_for_quotas(quotas, min_code_cases)
    cases = []
    used_pids = set()
    code_count = 0
    case_number = 1
    for bucket in BRIDGE_BUCKET_ORDER:
        quota = quotas.get(bucket, 0)
        bucket_candidates = candidates.get(bucket) or []
        selected = []
        for problem in bucket_candidates:
            normalized = normalize_luogu_problem(problem) if "problem_source_id" not in problem else problem
            if not normalized:
                continue
            pid = normalized["problem_source_id"]
            if pid in used_pids:
                continue
            selected.append(normalized)
            used_pids.add(pid)
            if len(selected) >= quota:
                break
        if len(selected) < quota:
            raise ValueError(f"Not enough candidates for {bucket}: needed {quota}, got {len(selected)}")
        for problem in selected:
            target_for_bucket = code_targets.get(bucket, 0)
            already_for_bucket = sum(
                1
                for existing in cases
                if existing.get("bridge_bucket") == bucket
                and _code_bucket(existing.get("student_code_excerpt") or "") == "present"
            )
            include_code = already_for_bucket < target_for_bucket
            if include_code:
                code_count += 1
            cases.append(
                build_case(
                    problem=problem,
                    bridge_bucket=bucket,
                    case_number=case_number,
                    length_bucket=length_plan[case_number - 1],
                    recent_dialogue_bucket=recent_dialogue_plan[case_number - 1],
                    include_code_excerpt=include_code,
                )
            )
            case_number += 1
    return cases


def _code_targets_for_quotas(quotas: dict[str, int], min_code_cases: int) -> dict[str, int]:
    """Allocate code-present cases to implementation/debug/DS first, then fallback.

    The held-out v2 plan wants code/error-code scenes mainly in implementation,
    debugging, and data-structure cases. Small unit-test quota slices may not
    include those buckets, so this function falls back deterministically.
    """

    preferred_order = [
        "implementation_boundary",
        "debugging_evidence",
        "data_structure_operation_semantics",
        "boundary_update_order",
        "predicate_check_semantics",
        "aggregation_contribution_summary",
        "state_representation_semantics",
        "transition_recurrence_source",
        "modeling_object_relation",
        "correctness_invariant",
        "policy_request",
    ]
    target_caps = {
        "implementation_boundary": 4,
        "debugging_evidence": 3,
        "data_structure_operation_semantics": 3,
    }
    targets: dict[str, int] = {}
    remaining = min_code_cases
    for bucket in preferred_order:
        if remaining <= 0:
            break
        quota = quotas.get(bucket, 0)
        if quota <= 0:
            continue
        cap = target_caps.get(bucket, quota)
        assign = min(quota, cap, remaining)
        targets[bucket] = assign
        remaining -= assign
    return targets


def _is_blank(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    if isinstance(value, list):
        return not value
    return False


def _recent_bucket(value: str) -> str:
    text = (value or "").strip()
    if not text or text.upper() == "N/A":
        return "none"
    role_lines = [
        line
        for line in text.splitlines()
        if line.strip().lower().startswith(("student:", "assistant:", "学生：", "ai：", "教练："))
    ]
    return "long" if len(role_lines) >= 4 else "short"


def _dialogue_role_lines(value: str) -> list[tuple[str, str]]:
    lines = []
    for raw_line in str(value or "").splitlines():
        line = raw_line.strip()
        if not line or "：" not in line:
            continue
        role, content = line.split("：", 1)
        if role.strip() == "学生":
            lines.append(("student", content.strip()))
        elif role.strip().lower() in {"ai", "assistant", "教练", "老师"}:
            lines.append(("assistant", content.strip()))
    return lines


def _recent_dialogue_alignment(case: dict) -> str:
    recent_dialogue = str(case.get("recent_dialogue") or "").strip()
    if not recent_dialogue or recent_dialogue.upper() == "N/A":
        return "no_recent_dialogue"
    role_lines = _dialogue_role_lines(recent_dialogue)
    if not role_lines:
        return "unparseable_recent_dialogue"
    last_role, _ = role_lines[-1]
    if last_role == "assistant":
        return "aligned_prior_context_ends_with_assistant"
    return "recent_dialogue_last_student_mismatch"


def _code_bucket(value: str) -> str:
    text = (value or "").strip()
    if not text or text.upper() in {"N/A", "NA", "NONE"}:
        return "none"
    return "present"


def validate_generated_cases(
    cases: list[dict],
    *,
    expected_count: int = 50,
    quotas: dict[str, int] = DEFAULT_QUOTAS,
    min_code_cases: int = DEFAULT_MIN_CODE_CASES,
    max_none_recent_dialogue: int = 10,
    min_long_recent_dialogue: int = 15,
    length_distribution: dict[str, int] | None = None,
) -> list[dict]:
    errors = []
    if len(cases) != expected_count:
        errors.append({"code": "unexpected_row_count", "expected": expected_count, "actual": len(cases)})

    case_ids = Counter(str(case.get("case_id") or "") for case in cases)
    problem_ids = Counter(str(case.get("problem_source_id") or "") for case in cases)
    bridge_counts = Counter(str(case.get("bridge_bucket") or case.get("category") or "") for case in cases)
    length_counts = Counter(_length_bucket(str(case.get("student_message") or "")) for case in cases)
    recent_counts = Counter(_recent_bucket(str(case.get("recent_dialogue") or "")) for case in cases)
    code_counts = Counter(_code_bucket(str(case.get("student_code_excerpt") or "")) for case in cases)
    source_platform_counts = Counter(str(case.get("problem_source_platform") or "") for case in cases)

    for idx, case in enumerate(cases, 1):
        case_id = case.get("case_id")
        for field in REQUIRED_CASE_FIELDS:
            if field not in case or _is_blank(case.get(field)):
                errors.append(
                    {
                        "code": "missing_required_field",
                        "row": idx,
                        "case_id": case_id,
                        "field": field,
                    }
                )
        url = str(case.get("problem_source_url") or "")
        if url and not url.startswith("https://www.luogu.com.cn/problem/"):
            errors.append({"code": "invalid_problem_source_url", "case_id": case_id, "url": url})
        if case.get("reference_label_status") != "draft_needs_coach_review":
            errors.append(
                {
                    "code": "invalid_reference_label_status",
                    "case_id": case_id,
                    "reference_label_status": case.get("reference_label_status"),
                }
            )
        alignment = _recent_dialogue_alignment(case)
        if alignment in {"recent_dialogue_last_student_mismatch", "unparseable_recent_dialogue"}:
            errors.append(
                {
                    "code": alignment,
                    "case_id": case_id,
                }
            )

    for case_id, count in case_ids.items():
        if case_id and count > 1:
            errors.append({"code": "duplicate_case_id", "case_id": case_id})
    for problem_id, count in problem_ids.items():
        if problem_id and count > 1:
            errors.append({"code": "duplicate_problem_source_id", "problem_source_id": problem_id})

    for bucket, expected in quotas.items():
        actual = bridge_counts.get(bucket, 0)
        if actual != expected:
            errors.append({"code": "unexpected_bridge_bucket_count", "bucket": bucket, "expected": expected, "actual": actual})

    if source_platform_counts != {"luogu": len(cases)}:
        errors.append(
            {
                "code": "unexpected_problem_source_platform_counts",
                "actual": dict(source_platform_counts),
            }
        )
    if code_counts.get("present", 0) < min_code_cases:
        errors.append(
            {
                "code": "too_few_code_excerpts",
                "min_required": min_code_cases,
                "actual_count": code_counts.get("present", 0),
            }
        )
    if recent_counts.get("none", 0) > max_none_recent_dialogue:
        errors.append(
            {
                "code": "too_many_no_recent_dialogue",
                "max_allowed": max_none_recent_dialogue,
                "actual_count": recent_counts.get("none", 0),
            }
        )
    if recent_counts.get("long", 0) < min_long_recent_dialogue:
        errors.append(
            {
                "code": "too_few_long_recent_dialogue",
                "min_required": min_long_recent_dialogue,
                "actual_count": recent_counts.get("long", 0),
            }
        )
    for bucket, expected in (length_distribution or {}).items():
        actual = length_counts.get(bucket, 0)
        if actual < expected:
            errors.append(
                {
                    "code": "too_few_student_message_length_bucket",
                    "bucket": bucket,
                    "min_required": expected,
                    "actual_count": actual,
                }
            )
        if actual > expected:
            errors.append(
                {
                    "code": "too_many_student_message_length_bucket",
                    "bucket": bucket,
                    "max_allowed": expected,
                    "actual_count": actual,
                }
            )
    return errors


REVIEW_COLUMNS = [
    ("case_id", "样本编号"),
    ("bridge_bucket_zh", "桥梁 bucket"),
    ("problem_source_id", "洛谷题号"),
    ("problem_title", "题目标题"),
    ("problem_source_url", "原题链接"),
    ("problem_tags", "洛谷标签"),
    ("problem_statement", "必要题面"),
    ("student_message", "学生问题"),
    ("student_message_length_bucket", "学生问题长度"),
    ("recent_dialogue", "近期对话"),
    ("student_code_excerpt", "学生代码片段"),
    ("student_known_state", "学生已知状态"),
    ("missing_bridge", "目标 bridge"),
    ("forbidden_content", "禁止内容"),
    ("success_criteria", "成功标准"),
    ("review_notes_for_coach", "复核提示"),
]


def _cell_value(value: object) -> str:
    if isinstance(value, list):
        return "\n".join(str(item) for item in value)
    return str(value or "")


def export_source_and_case_review_workbook(cases: list[dict], output_xlsx: Path) -> int:
    workbook = Workbook()
    first_sheet = workbook.active
    workbook.remove(first_sheet)
    grouped = {bucket: [case for case in cases if case.get("bridge_bucket") == bucket] for bucket in BRIDGE_BUCKET_ORDER}
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for bucket in BRIDGE_BUCKET_ORDER:
        rows = grouped.get(bucket) or []
        if not rows:
            continue
        sheet_name = BUCKET_CONFIG[bucket]["sheet"][:31]
        sheet = workbook.create_sheet(sheet_name)
        sheet.append([label for _key, label in REVIEW_COLUMNS])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="17324D")
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for row in rows:
            sheet.append([_cell_value(row.get(key)) for key, _label in REVIEW_COLUMNS])
        for row_cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row_cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        widths = {
            "A": 20,
            "B": 18,
            "C": 14,
            "D": 28,
            "E": 42,
            "F": 28,
            "G": 72,
            "H": 42,
            "I": 14,
            "J": 52,
            "K": 42,
            "L": 42,
            "M": 48,
            "N": 36,
            "O": 44,
            "P": 44,
        }
        for col, width in widths.items():
            sheet.column_dimensions[col].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx)
    return len(cases)


def write_jsonl(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "\n".join(json.dumps(row, ensure_ascii=False, sort_keys=True) for row in rows) + "\n",
        encoding="utf-8",
    )


def build_generation_payload(cases: list[dict], errors: list[dict], source_path: Path, output_jsonl: Path, review_xlsx: Path) -> dict:
    return {
        "source_path": str(source_path),
        "output_jsonl": str(output_jsonl),
        "review_xlsx": str(review_xlsx),
        "row_count": len(cases),
        "ok": not errors,
        "error_count": len(errors),
        "errors": errors,
        "bridge_bucket_counts": dict(Counter(case["bridge_bucket"] for case in cases)),
        "student_message_length_distribution": dict(Counter(case["student_message_length_bucket"] for case in cases)),
        "recent_dialogue_distribution": dict(Counter(case["recent_dialogue_bucket"] for case in cases)),
        "student_code_excerpt_distribution": dict(
            Counter(_code_bucket(case.get("student_code_excerpt") or "") for case in cases)
        ),
        "problem_source_platform_counts": dict(Counter(case["problem_source_platform"] for case in cases)),
        "reference_label_status_counts": dict(Counter(case["reference_label_status"] for case in cases)),
        "selected_problem_ids": [case["problem_source_id"] for case in cases],
    }


def render_generation_report_zh(payload: dict) -> str:
    lines = [
        "# Luogu-grounded Held-out 50 Generation Report",
        "",
        f"本报告记录 `{payload['output_jsonl']}` 的生成结果。该数据集是 draft，等待教练复核；学生问题为 synthetic-but-grounded，不来自旧线上 AI 回复。",
        "",
        f"- source_path: `{payload['source_path']}`",
        f"- output_jsonl: `{payload['output_jsonl']}`",
        f"- review_xlsx: `{payload['review_xlsx']}`",
        f"- row_count: `{payload['row_count']}`",
        f"- ok: `{payload['ok']}`",
        "",
        "## 分布",
        "",
        f"- bridge_bucket_counts: `{payload['bridge_bucket_counts']}`",
        f"- student_message_length_distribution: `{payload['student_message_length_distribution']}`",
        f"- recent_dialogue_distribution: `{payload['recent_dialogue_distribution']}`",
        f"- student_code_excerpt_distribution: `{payload['student_code_excerpt_distribution']}`",
        "",
        "## 使用边界",
        "",
        "- 原始洛谷快照不进入 GitHub；本地路径由 snapshot 文档记录。",
        "- JSONL 中包含必要题面摘录，仅供本地教练复核；公开材料优先使用题号、链接和改写摘要。",
        "- 洛谷标签保留为 metadata / workbook 独立列，不写入 `problem_statement`，避免后续生成回复时误把算法标签当作题面泄露给 tutor。",
        f"- `{payload['review_xlsx']}` 是题源/case 复核表，不是 AI 回复盲评表；它故意不包含待评分 AI 回复。",
        "- 学生问题为 synthetic-but-grounded，并使用真实线上学生短问风格的多变体模板；后续仍需教练确认是否自然、是否匹配题面。",
        "- 本版本不称为 gold，只能作为 `draft_needs_coach_review`。",
    ]
    if payload["errors"]:
        lines.extend(["", "## Errors", "", json.dumps(payload["errors"], ensure_ascii=False, indent=2)])
    return "\n".join(lines) + "\n"


def render_generation_report_en(payload: dict) -> str:
    lines = [
        "# Luogu-grounded Held-out 50 Generation Report",
        "",
        f"This report records the generation of `{payload['output_jsonl']}`. The dataset is a draft pending coach review; student messages are synthetic-but-grounded and do not use prior online AI replies.",
        "",
        f"- source_path: `{payload['source_path']}`",
        f"- output_jsonl: `{payload['output_jsonl']}`",
        f"- review_xlsx: `{payload['review_xlsx']}`",
        f"- row_count: `{payload['row_count']}`",
        f"- ok: `{payload['ok']}`",
        "",
        "## Distributions",
        "",
        f"- bridge_bucket_counts: `{payload['bridge_bucket_counts']}`",
        f"- student_message_length_distribution: `{payload['student_message_length_distribution']}`",
        f"- recent_dialogue_distribution: `{payload['recent_dialogue_distribution']}`",
        f"- student_code_excerpt_distribution: `{payload['student_code_excerpt_distribution']}`",
        "",
        "## Boundaries",
        "",
        "- The raw Luogu snapshot is not committed to GitHub; the local path is recorded in the snapshot document.",
        "- The JSONL contains necessary statement excerpts for local coach review; public materials should prefer problem IDs, links, and rewritten summaries.",
        "- Luogu tags are kept as metadata / workbook columns and are not embedded in `problem_statement`, so response-generation prompts can avoid leaking algorithm labels as statement text.",
        f"- `{payload['review_xlsx']}` is a source/case review workbook, not an AI response blind-review workbook; it intentionally has no target AI response.",
        "- Student messages are synthetic-but-grounded and use varied templates inspired by real online short student questions; coaches still need to confirm naturalness and statement fit.",
        "- This version is not gold; every row remains `draft_needs_coach_review`.",
    ]
    if payload["errors"]:
        lines.extend(["", "## Errors", "", json.dumps(payload["errors"], ensure_ascii=False, indent=2)])
    return "\n".join(lines) + "\n"


def generate_dataset(
    *,
    source_ndjson: Path = DEFAULT_SOURCE_NDJSON,
    output_jsonl: Path = DEFAULT_OUTPUT_JSONL,
    review_xlsx: Path = DEFAULT_REVIEW_XLSX,
    generation_report_json: Path = DEFAULT_GENERATION_REPORT_JSON,
    generation_report_zh: Path = DEFAULT_GENERATION_REPORT_ZH,
    generation_report_en: Path = DEFAULT_GENERATION_REPORT_EN,
) -> dict:
    problems = list(iter_luogu_problem_rows(source_ndjson))
    candidates = select_candidates_by_bucket(problems, quotas=DEFAULT_QUOTAS, multiplier=6)
    cases = build_cases_from_candidates(
        candidates,
        quotas=DEFAULT_QUOTAS,
        length_plan=DEFAULT_LENGTH_PLAN,
        recent_dialogue_plan=DEFAULT_RECENT_DIALOGUE_PLAN,
        min_code_cases=DEFAULT_MIN_CODE_CASES,
    )
    errors = validate_generated_cases(
        cases,
        expected_count=sum(DEFAULT_QUOTAS.values()),
        quotas=DEFAULT_QUOTAS,
        min_code_cases=DEFAULT_MIN_CODE_CASES,
        max_none_recent_dialogue=10,
        min_long_recent_dialogue=15,
        length_distribution={
            "short": 20,
            "medium_short": 15,
            "medium_long": 10,
            "long": 5,
        },
    )
    write_jsonl(cases, output_jsonl)
    export_source_and_case_review_workbook(cases, review_xlsx)
    payload = build_generation_payload(cases, errors, source_ndjson, output_jsonl, review_xlsx)
    generation_report_json.parent.mkdir(parents=True, exist_ok=True)
    generation_report_json.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    generation_report_zh.write_text(render_generation_report_zh(payload), encoding="utf-8")
    generation_report_en.write_text(render_generation_report_en(payload), encoding="utf-8")
    return payload


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate Luogu-grounded held-out v2 50 draft.")
    parser.add_argument("--source-ndjson", type=Path, default=DEFAULT_SOURCE_NDJSON)
    parser.add_argument("--output-jsonl", type=Path, default=DEFAULT_OUTPUT_JSONL)
    parser.add_argument("--review-xlsx", type=Path, default=DEFAULT_REVIEW_XLSX)
    parser.add_argument("--generation-report-json", type=Path, default=DEFAULT_GENERATION_REPORT_JSON)
    parser.add_argument("--generation-report-zh", type=Path, default=DEFAULT_GENERATION_REPORT_ZH)
    parser.add_argument("--generation-report-en", type=Path, default=DEFAULT_GENERATION_REPORT_EN)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None, *, stdout: TextIO | None = None) -> int:
    args = _parse_args(sys.argv[1:] if argv is None else argv)
    payload = generate_dataset(
        source_ndjson=args.source_ndjson,
        output_jsonl=args.output_jsonl,
        review_xlsx=args.review_xlsx,
        generation_report_json=args.generation_report_json,
        generation_report_zh=args.generation_report_zh,
        generation_report_en=args.generation_report_en,
    )
    output = stdout or sys.stdout
    print(json.dumps(payload, ensure_ascii=False), file=output)
    return 0 if payload["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
