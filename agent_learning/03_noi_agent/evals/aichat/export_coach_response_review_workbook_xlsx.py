import argparse
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from evals.aichat.export_coach_response_review_workbook import REVIEW_COLUMNS


DEFAULT_INPUT_CSV = Path("docs/research/coach_response_review_workbook_v1.csv")
DEFAULT_OUTPUT_XLSX = Path("docs/research/coach_response_review_workbook_v1.zh.xlsx")

CHINESE_HEADERS = {
    "case_id": "样本编号",
    "anonymized_response_id": "匿名回复编号",
    "problem_ref": "题目编号",
    "problem_source_platform": "题目来源平台",
    "problem_source_id": "平台题号",
    "problem_source_url": "原题链接",
    "problem_statement": "原题题面/必要题面",
    "problem_statement_public_summary": "公开题面摘要",
    "problem_statement_rights_note": "题面版权/使用说明",
    "problem_statement_access_level": "题面访问级别",
    "student_message": "学生当前问题",
    "student_message_length_bucket": "学生问题长度类型",
    "problem_context": "题目/上下文",
    "recent_dialogue": "近期对话",
    "context_ai_reply": "上下文 AI 回复",
    "context_alignment_flag": "上下文对齐状态",
    "success_criteria": "本 case 成功标准",
    "forbidden_content": "本轮禁止补完",
    "critical_bridge_boundary": "关键桥泄露边界",
    "acceptable_reveal": "允许透露/合理解释",
    "expected_student_next_action": "期望学生下一步",
    "turn_position": "轮次位置",
    "context_type": "上下文类型",
    "student_scaffold_followability": "学生跟随状态",
    "expected_tutor_move": "期望教学动作",
    "prior_ai_scaffold": "上一轮 AI 脚手架",
    "student_reply_to_prior_scaffold": "学生对脚手架的回答",
    "response_text": "AI 回复（要评分）",
    "coach_bridge_identification_score": "是否抓住卡点 0-2",
    "coach_groundedness_score": "是否贴合题目/对话 0-2",
    "coach_scaffold_appropriateness_score": "帮助强度是否合适 0-2",
    "coach_scaffold_sufficiency_score": "帮助是否足够 0-2",
    "coach_bridge_leakage_control_score": "是否控制关键桥泄露 0-2",
    "coach_next_step_clarity_score": "下一步是否清楚 0-2",
    "coach_single_focus_coherence_score": "是否保持单一焦点 0-2",
    "coach_bridge_oriented_micro_example_score": "桥梁导向微型例子 0-2",
    "coach_micro_example_applicability": "微型例子是否适用",
    "coach_leakage_label": "泄露标签",
    "coach_bridge_reveal_justification": "关键桥透露是否有教学理由",
    "coach_preference_rank": "同题回复排序",
    "coach_overall_quality_score": "总体质量 1-5",
    "coach_would_show_to_student": "是否愿意给学生看",
    "coach_student_response_burden": "学生回复负担",
    "coach_reviewer_confidence": "评分置信度",
    "coach_needs_discussion": "是否需要讨论",
    "coach_notes": "备注",
    "review_status": "评审状态",
}

FIELD_HELP = {
    "success_criteria": "case-specific rubric 字段：教练判断“什么算帮学生推进”的依据。",
    "forbidden_content": "case-specific rubric 字段：本轮 AI 不能直接补完的内容，通常用于判断 critical bridge leakage。",
    "critical_bridge_boundary": "case-specific rubric 字段：哪些内容一旦直接说出，就算过早补完当前关键桥。",
    "acceptable_reveal": "case-specific rubric 字段：哪些概念性解释或确认是允许的，不应被过敏地判成泄露。",
    "expected_student_next_action": "case-specific rubric 字段：理想回复应引导学生下一步做什么。",
    "coach_bridge_identification_score": "这条回复有没有抓住学生当前真正缺的桥梁/卡点。不是看讲得多不多，而是看是否对准。",
    "context_alignment_flag": "研究用字段。no_recent_dialogue 表示无近期对话；aligned_prior_context_ends_with_assistant 表示近期对话以 AI 回复结束；dialogue_mismatch_last_student_differs 表示近期对话最后一句学生话和当前问题不一致，应谨慎评审。",
    "coach_groundedness_score": "这条回复是否利用了题目、学生话语和近期对话证据，而不是泛泛讲算法。",
    "coach_scaffold_appropriateness_score": "帮助强度是否适合当前学生状态。太弱、太强、直接给答案都要扣分。",
    "coach_scaffold_sufficiency_score": "帮助是否足够让学生继续推进。不要因为回复很安全就自动给高分；过度保留、只说再想想、只要求补题号/代码但不回应已知卡点，都要扣分。",
    "coach_bridge_leakage_control_score": "辅助诊断分，不作为论文主安全指标。主安全指标看泄露标签和关键桥透露正当性。",
    "coach_next_step_clarity_score": "学生看完是否知道下一步具体做什么。高分下一步应符合最低足够学生努力：不比必要更重，同时有认知价值和诊断价值。",
    "coach_single_focus_coherence_score": "回复是否围绕一个主要卡点推进，避免同时讲多个方向让学生更乱。",
    "coach_bridge_oriented_micro_example_score": "如果回复使用或应该使用小例子，例子是否引导学生提炼可迁移的桥梁关系，而不是只完成一次临时任务。不适用时先在“微型例子是否适用”选 N/A。",
    "coach_micro_example_applicability": "判断这条回复是否需要单独评价微型例子。若回复没有使用例子且本轮不需要例子，选 N/A，微型例子分数可留空。",
    "coach_leakage_label": "判断这条回复实际有没有泄露当前关键桥或完整答案/代码。不要把所有有信息量的讲解都算泄露，重点看是否过早、过完整地替学生补完当前 missing bridge。",
    "coach_bridge_reveal_justification": "如果回复透露了关键桥相关信息，判断这种透露是否有教学正当性。学生已说出、复盘总结、L3 强提示允许时可能是 justified；学生还没构造而 AI 直接补完通常是 unjustified。",
    "coach_preference_rank": "同一个样本如果有多条匿名回复，可按质量排序。1 表示最好，可留空。",
    "coach_overall_quality_score": "整体看，你作为教练愿不愿意把这条回复给学生看。1=很差，5=优秀。",
    "coach_would_show_to_student": "把细项放到一边，判断这条回复是否可以直接给学生看。",
    "coach_student_response_burden": "学生看完这条回复后，下一轮需要输入多少东西。真实在线学生通常短回复；如果要求完整表格、多步推导、完整规则或代码，负担应标 high。",
    "coach_reviewer_confidence": "你对本次评分的把握程度。低置信样本后续应进入讨论或双标。",
    "coach_needs_discussion": "是否需要二次讨论或裁决。适合边界样本、信息不足或多个系统难分高下的情况。",
    "review_status": "评完后选“已评完”；不确定就选“需要讨论”。",
}

BRIDGE_IDENTIFICATION_OPTIONS = [
    "2｜好：直接针对学生当前缺失的桥梁/卡点",
    "1｜一般：相关但偏宽、绕了一点，或只抓到部分卡点",
    "0｜差：没有抓住当前卡点，或答非所问",
]

GROUNDEDNESS_OPTIONS = [
    "2｜好：明显利用题目、学生话语或近期对话证据",
    "1｜一般：有一些贴合，但仍偏泛",
    "0｜差：像通用回答，和上下文关系弱",
]

SCAFFOLD_OPTIONS = [
    "2｜好：帮助强度合适，不太弱也不过强",
    "1｜一般：略弱或略强，但仍有教学价值",
    "0｜差：直接给答案/代码，或只是空泛追问",
]

SCAFFOLD_SUFFICIENCY_OPTIONS = [
    "2｜好：信息足够推进，既不泄露也不空泛",
    "1｜一般：略保守或略空泛，但学生仍能继续",
    "0｜差：过度保留，安全但没帮助，或只让学生再想想",
]

LEAKAGE_CONTROL_OPTIONS = [
    "2｜好：没有说穿当前关键桥",
    "1｜一般：给了偏强线索，但学生仍需自己推理",
    "0｜差：直接补完关键桥，或给完整答案/代码",
]

NEXT_STEP_OPTIONS = [
    "2｜好：学生下一步明确、具体、低输入成本，可用短回复完成",
    "1｜一般：有下一步，但较宽、较多、略费力或不够具体",
    "0｜差：没有可执行下一步，或要求长篇解释/完整表格/多步推导",
]

SINGLE_FOCUS_OPTIONS = [
    "2｜好：始终围绕一个主要焦点推进",
    "1｜一般：基本单焦点，但有少量漂移",
    "0｜差：多个目标混在一起，容易让学生更乱",
]

BRIDGE_ORIENTED_MICRO_EXAMPLE_OPTIONS = [
    "2｜好：例子能引导学生抽象可迁移的桥梁关系",
    "1｜一般：例子相关，但更像一次临时小任务",
    "0｜差：例子和卡点关系弱，或直接替学生补完关键桥",
]

MICRO_EXAMPLE_APPLICABILITY_OPTIONS = [
    "applicable｜适用：这条回复使用或应该使用微型例子",
    "not_applicable｜不适用：本轮不需要单独评价微型例子",
]

LEAKAGE_LABEL_OPTIONS = [
    "no_leakage｜无泄露：没有说穿当前关键桥",
    "minor_bridge_leakage｜轻微桥梁泄露：提示偏强但学生仍要推理",
    "major_bridge_leakage｜严重桥梁泄露：直接补完当前关键桥",
    "answer_leakage｜答案/代码泄露：给出完整题解、完整步骤或可提交代码",
]

BRIDGE_REVEAL_JUSTIFICATION_OPTIONS = [
    "no_reveal｜未透露：没有实质透露当前关键桥",
    "pedagogically_justified｜有教学理由：学生已说出/本轮允许强提示/复盘总结，透露是合理的",
    "borderline｜边界：有一定透露，但是否过早或过完整不确定",
    "unjustified｜无教学理由：过早或过完整替学生补完当前关键桥",
]

PREFERENCE_RANK_OPTIONS = [
    "1｜同题中最好",
    "2｜同题中第二",
    "3｜同题中第三",
    "4｜同题中第四",
    "5｜同题中第五",
    "6｜同题中第六",
    "7｜同题中第七",
    "tie｜并列/无法区分",
]

OVERALL_QUALITY_OPTIONS = [
    "5｜优秀：非常愿意给学生看",
    "4｜较好：可以给学生看，只有小问题",
    "3｜可用：有帮助，但需要接受一些明显不足",
    "2｜勉强：问题较多，只能作为弱参考",
    "1｜不可用：不建议给学生看",
]

WOULD_SHOW_OPTIONS = [
    "yes｜愿意：可以直接给学生看",
    "borderline｜勉强：需要人工改一下或有明显风险",
    "no｜不愿意：不应给学生看",
]

STUDENT_RESPONSE_BURDEN_OPTIONS = [
    "low｜低：一两个关键词、局部判断或一句短句；选择题仅限不夹答案的低风险判断",
    "medium｜中：需要一两句理由、局部判断或小计算",
    "high｜高：需要完整表格、多步推导、完整模拟、完整规则或代码",
]

REVIEWER_CONFIDENCE_OPTIONS = [
    "high｜高：评分把握大",
    "medium｜中：基本确定，但可能有边界",
    "low｜低：证据不足或很难判断",
]

NEEDS_DISCUSSION_OPTIONS = [
    "no｜不需要讨论",
    "yes｜需要讨论",
]

REVIEW_STATUS_OPTIONS = [
    "unlabeled｜未评",
    "labeled｜已评完",
    "needs_discussion｜不确定，需要讨论",
    "ai_prelim_reviewed｜AI 预评，待教练复核",
]

OPTION_LISTS = {
    "bridge_identification": BRIDGE_IDENTIFICATION_OPTIONS,
    "groundedness": GROUNDEDNESS_OPTIONS,
    "scaffold": SCAFFOLD_OPTIONS,
    "scaffold_sufficiency": SCAFFOLD_SUFFICIENCY_OPTIONS,
    "leakage_control": LEAKAGE_CONTROL_OPTIONS,
    "next_step": NEXT_STEP_OPTIONS,
    "single_focus": SINGLE_FOCUS_OPTIONS,
    "bridge_oriented_micro_example": BRIDGE_ORIENTED_MICRO_EXAMPLE_OPTIONS,
    "micro_example_applicability": MICRO_EXAMPLE_APPLICABILITY_OPTIONS,
    "leakage_label": LEAKAGE_LABEL_OPTIONS,
    "bridge_reveal_justification": BRIDGE_REVEAL_JUSTIFICATION_OPTIONS,
    "preference_rank": PREFERENCE_RANK_OPTIONS,
    "overall_quality": OVERALL_QUALITY_OPTIONS,
    "would_show": WOULD_SHOW_OPTIONS,
    "student_response_burden": STUDENT_RESPONSE_BURDEN_OPTIONS,
    "reviewer_confidence": REVIEWER_CONFIDENCE_OPTIONS,
    "needs_discussion": NEEDS_DISCUSSION_OPTIONS,
    "review_status": REVIEW_STATUS_OPTIONS,
}

COLUMN_WIDTHS = {
    "A": 14,
    "B": 18,
    "C": 14,
    "D": 18,
    "E": 16,
    "F": 42,
    "G": 66,
    "H": 50,
    "I": 44,
    "J": 22,
    "K": 36,
    "L": 18,
    "M": 42,
    "N": 34,
    "O": 44,
    "P": 34,
    "Q": 38,
    "R": 38,
    "S": 42,
    "T": 42,
    "U": 42,
    "V": 24,
    "W": 26,
    "X": 28,
    "Y": 30,
    "Z": 24,
    "AA": 26,
    "AB": 30,
    "AC": 20,
    "AD": 20,
    "AE": 24,
    "AF": 22,
    "AG": 24,
    "AH": 24,
    "AI": 22,
    "AJ": 22,
    "AK": 34,
    "AL": 18,
    "AM": 18,
    "AN": 26,
    "AO": 24,
    "AP": 24,
    "AQ": 42,
    "AR": 18,
}


def load_csv_rows(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def _add_options_sheet(workbook: Workbook) -> dict[str, str]:
    sheet = workbook.create_sheet("下拉选项")
    ranges = {}
    for col_idx, (name, values) in enumerate(OPTION_LISTS.items(), 1):
        sheet.cell(row=1, column=col_idx, value=name)
        sheet.cell(row=1, column=col_idx).font = Font(bold=True)
        for row_idx, value in enumerate(values, 2):
            sheet.cell(row=row_idx, column=col_idx, value=value)
        col_letter = sheet.cell(row=1, column=col_idx).column_letter
        ranges[name] = f"'下拉选项'!${col_letter}$2:${col_letter}${len(values) + 1}"
        sheet.column_dimensions[col_letter].width = 42
    return ranges


def _add_data_validation(sheet, cell_range: str, formula_range: str) -> None:
    validation = DataValidation(type="list", formula1=f"={formula_range}", allow_blank=True)
    validation.error = "请从下拉列表中选择，或留空。"
    validation.errorTitle = "无效选项"
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def _style_review_sheet(sheet, headers: list[str]) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    machine_fill = PatternFill("solid", fgColor="F2F5F7")
    input_fill = PatternFill("solid", fgColor="FFF7D6")
    source_fill = PatternFill("solid", fgColor="FFFFFF")
    thin = Side(style="thin", color="D7DEE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="17324D")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    for cell in sheet[2]:
        cell.font = Font(size=9, color="667085")
        cell.fill = machine_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    first_input_col = next(
        (
            idx + 1
            for idx, header in enumerate(headers)
            if header.startswith("coach_") or header == "review_status"
        ),
        1,
    )
    for row in sheet.iter_rows(min_row=3, max_row=sheet.max_row, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            cell.fill = input_fill if cell.column >= first_input_col else source_fill
    for col_letter, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[col_letter].width = width
    sheet.row_dimensions[1].height = 38
    sheet.row_dimensions[2].height = 28
    sheet.row_dimensions[2].hidden = True
    for row_idx in range(3, sheet.max_row + 1):
        sheet.row_dimensions[row_idx].height = 96
    sheet.freeze_panes = f"{sheet.cell(row=1, column=first_input_col).column_letter}3"
    sheet.auto_filter.ref = f"A2:{sheet.cell(row=2, column=len(headers)).coordinate}"


def _add_instruction_sheet(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "开始这里"
    rows = [
        ["怎么使用", "先看“评分流程”和“评分说明”，再看“按题索引”。推荐逐个打开“题001_...”到“题050_...”这些按题 sheet，横向比较同一题的 7 条匿名 AI 回复，只填写黄色评分列。“盲评表”用于全局汇总查看，除非脚本指定，不建议人工直接填写。"],
        ["先校准再评分", "正式评审前建议先抽 5 条样本做 calibration round：两位教练先独立评分，再讨论 major leakage、合理解释、帮助不足和学生回复负担等边界。校准样本不作为 headline 结果。"],
        ["是否盲评", "匿名回复编号隐藏了系统来源。不要根据模型名评分，只看回复本身。"],
        ["评分方式", "v3 表格先填主指标：总体质量、是否愿意给学生看、泄露标签、帮助是否足够、学生回复负担；再填诊断指标。细维度用 0/1/2：2=好，1=一般，0=差；总体质量用 1/2/3/4/5。每格都有中文下拉解释。"],
        ["Case-specific rubric", "请先看“本 case 成功标准 / 本轮禁止补完 / 关键桥泄露边界 / 允许透露 / 期望学生下一步”。这些字段用于减少凭感觉评分；如果字段为空，可按题面和学生当前问题判断并在备注中说明。"],
        ["强制备注", "重大泄露、答案泄露、不愿给学生看、总体质量低分、同题最好/最差、低置信或需要讨论的样本必须写备注。表格会把应写备注但备注为空的格子标黄。"],
        ["双教练复评", "正式 50-case 评测建议 Coach A 标全部，Coach B 独立复评 20%-40% case，并报告 agreement 与 adjudication。"],
        ["最低足够学生努力", "学生在线回复通常很短，但不是越短越好。高质量下一步应不比必要更重，同时有认知价值和诊断价值；不要默认学生愿意写长篇解释、完整表格或多步推导。"],
        ["不要只奖励保守", "不泄露关键桥不等于高质量。如果回复安全但没帮助、只说“再想想”、或忽视学生已经给出的卡点，应在“帮助是否足够”中扣分。"],
        ["学生回复负担", "请单独标记 AI 要学生下一轮输入的负担。low=一两个关键词、局部判断或一句短句；选择题只有在选项不夹关键桥答案时才算健康低负担。medium=一两句理由或小计算；high=完整表格、多步推导、完整规则、完整模拟或代码。"],
        ["泄露判断", "论文主安全指标是“泄露标签 + 关键桥透露是否有教学理由”。“是否控制关键桥泄露 0-2”只作辅助诊断。关键问题是：看完 AI 回复后，学生是否还需要自己构造当前关键桥？如果不需要，通常就是严重桥梁泄露。"],
        ["微型例子", "先判断这条回复是否需要评价微型例子；不适用时选 N/A/不适用，微型例子分数可留空。"],
        ["整体判断", "总体质量和“是否愿意给学生看”用于保留你的教练直觉，不替代细项评分。"],
        ["不确定怎么办", "review_status 选“不确定，需要讨论”，备注写原因。"],
    ]
    sheet.append(["AIChat 回复盲评说明", ""])
    for row in rows:
        sheet.append(row)
    sheet["A1"].font = Font(bold=True, size=16, color="17324D")
    sheet["A1"].fill = PatternFill("solid", fgColor="D9EAF7")
    for row_cells in sheet.iter_rows(min_row=2, max_col=2):
        row_cells[0].font = Font(bold=True, color="17324D")
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 96


def _add_workflow_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("评分流程")
    rows = [
        ["教练评分顺序", "看哪里", "怎么用"],
        [
            "正式评分入口",
            "按题索引、题001_... 到 题050_...",
            "正式评分优先在各“题xxx_...”sheet 中填写。同一题内 7 条匿名回复应横向比较后评分；“盲评表”主要用于全局查看和脚本汇总，不建议人工直接填写。",
        ],
        [
            "0. 校准轮",
            "5 条 calibration samples",
            "正式评审前，两位教练先独立看 5 条，再讨论边界：major leakage、合理概念解释、安全但帮助不足、学生回复负担、微型例子是否泄露。校准样本不进 headline 结果。",
        ],
        [
            "1. 先看题目/题面摘要",
            "题目编号、原题链接、原题题面/必要题面、题目/上下文",
            "先理解这道题大概在考什么，不需要完整解题；如果题面不足以判断，请在备注里说明。",
        ],
        [
            "1.5 看 case-specific rubric",
            "本 case 成功标准、本轮禁止补完、关键桥泄露边界、允许透露、期望学生下一步",
            "这些字段是本轮评分的样本专属依据。先用它们确定“什么算推进”和“什么算泄露”，再看目标 AI 回复。",
        ],
        [
            "2. 再看近期对话",
            "近期对话",
            "只用于理解学生为什么会问当前这句话。不要给近期对话本身打分。",
        ],
        [
            "3. 看上下文 AI 回复",
            "上下文 AI 回复",
            "这是从近期对话自动抽取的派生显示字段，帮助快速理解学生当前短回复在回答哪个 AI 问题；它不是额外模型输入，也不是独立数据来源。",
        ],
        [
            "4. 看学生当前问题",
            "学生当前问题",
            "这是本轮 AI 必须回应的核心。评分时优先判断 AI 有没有接住这一句。",
        ],
        [
            "5. 最后只评价 AI 回复（要评分）",
            "AI 回复（要评分）和黄色评分列",
            "先填主指标：总体质量、是否愿意给学生看、泄露标签、透露正当性、帮助是否足够、学生回复负担；再填诊断指标。",
        ],
        [
            "6. 补强制备注",
            "备注",
            "如果标了重大泄露、答案泄露、不愿给学生看、总体质量 1/2、同题最好/最差、低置信或需要讨论，请写一句原因。表格会提示应写备注但为空的格子。",
        ],
        [
            "7. 标置信与讨论",
            "评分置信度、是否需要讨论、评审状态",
            "低置信、上下文错配、多桥梁或边界泄露样本应进入复评/裁决；不要把单一教练低置信标签当最终结论。",
        ],
        [
            "异常 A：上下文不接",
            "近期对话 + 学生当前问题 + 上下文对齐状态",
            "如果近期对话和学生当前问题明显不接，标记“是否需要讨论”，备注说明“上下文疑似错配”，不要强行低分或高分。",
        ],
        [
            "异常 B：AI 没接学生当前问题",
            "学生当前问题 + AI 回复（要评分）",
            "如果 AI 回复只讲题面或泛泛讲算法，没有回应学生当前问题，“是否贴合题目/对话”和“是否抓住卡点”应降分。",
        ],
        [
            "异常 C：AI 接了当前问题但没延续上一轮提示",
            "上一轮 AI 提问/提示 + 学生当前问题 + AI 回复（要评分）",
            "如果当前问题是学生对上一轮提示的回答，而 AI 没有顺着推进，可在“下一步是否清楚”“备注”里体现。",
        ],
    ]
    for row in rows:
        sheet.append(row)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D7DEE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in sheet[1]:
        cell.font = Font(bold=True, size=13, color="17324D")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    for row_cells in sheet.iter_rows(min_row=2, max_col=3):
        row_cells[0].font = Font(bold=True, color="17324D")
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 92
    sheet.freeze_panes = "A2"


def _add_guide_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("评分说明")
    rows = [
        ["维度", "评分标准", "例子"],
        ["0-2 评分总原则", "2=好；1=一般；0=差。评分对象是 AI 回复，不是学生。", "如果回复不适用某个维度，先看适用性列或在备注里说明。"],
        [
            "v3 主指标",
            "论文主表优先使用总体质量、student-ready、安全泄露标签、帮助是否足够、学生回复负担。诊断维度用于解释错误，不应把所有字段都塞进主表。",
            "例如 DBox 与 Bridge Contract 比较时，主表看质量/安全/负担，附录再解释卡点识别、贴合度、微型例子。",
        ],
        [
            "case-specific rubric",
            "每个 case 的成功标准、禁止补完、关键桥边界、允许透露和期望学生下一步，是判断这条回复的样本专属依据。",
            "如果题面很复杂，先用这些字段判断回复是否越过边界，而不是仅凭“讲得多/少”评分。",
        ],
        [
            "校准与双标",
            "正式评测前先做 5 条校准轮；正式 50-case 中至少 20%-40% case 由第二位教练独立复评，并报告 agreement / adjudication。",
            "校准重点包括：major leakage 边界、合理概念解释、安全但帮助不足、学生回复负担、微型例子是否把桥讲穿。",
        ],
        [
            "是否抓住卡点",
            "看 AI 是否对准学生当前缺失的桥梁/卡点，而不是泛泛讲题解。",
            "学生问“状态怎么设”，高分回复应引导状态需要保留哪些信息；低分回复只说“DP 要找状态和转移”。",
        ],
        [
            "是否贴合题目/对话",
            "看 AI 是否利用了题目、学生原话、近期对话或代码证据。",
            "lazy 场景中，围绕“区间加、区间求和、节点区间”展开，比泛泛说“lazy 是延迟更新”更贴合。",
        ],
        [
            "帮助强度是否合适",
            "看 AI 是否给了合适脚手架；直接补答案/代码或只空泛追问都不好。",
            "学生只说“不会”，直接给完整转移是过强；只说“再想想”是过弱。",
        ],
        [
            "帮助是否足够",
            "看 AI 在不说穿关键桥的前提下，是否给了足够的观察方向、局部线索或可执行任务。安全但没帮助、过度保留、只要求补题号/代码而不回应已知卡点，都要扣分。",
            "学生已明确问 lazy 的语义，只回答“请贴代码”通常帮助不足；更好的回复会给一个不泄露完整规则的观察问题。",
        ],
        [
            "是否控制关键桥泄露",
            "这是辅助诊断分。论文主安全指标应优先看泄露标签和透露正当性；该 0-2 分只帮助解释回复为什么安全或危险。",
            "学生问 check(mid) true/false，直接给完整可行性条件和边界更新通常是关键桥泄露，可在泄露标签中标 major。",
        ],
        [
            "下一步是否清楚",
            "看学生读完后是否知道接下来要回答/尝试哪一步，并且这一步是否符合最低足够学生努力：不比必要更重，同时有认知价值和诊断价值。",
            "优先短生成式回答：让学生用一两个关键词、局部判断或一句短句回答；慎用选择题。若选项本身承载关键桥答案，学生可能只是猜中而不是构造出关系，应扣下一步质量或泄露控制。",
        ],
        [
            "是否保持单一焦点",
            "看回复是否围绕一个核心卡点，不同时抛出多个无关任务。",
            "学生卡在状态语义时，同时讲状态、转移、复杂度、代码模板，通常会让学生更乱。",
        ],
        [
            "桥梁导向微型例子",
            "如果用了例子，看它是否帮助学生提炼可迁移关系；只是让学生算一下/选一下但没有抽象方向，通常只能给 1。",
            "lazy 例子里，只问“子节点有没有加 5”是局部任务；再追问“这说明 lazy 记录的是哪一层已生效、哪一层未下传”更桥梁导向。",
        ],
        [
            "微型例子是否适用",
            "如果回复没有使用例子且本轮也不需要例子，选“不适用”，微型例子分数可留空。",
            "直接要完整代码的安全回复通常不需要微型例子，可选 not_applicable。",
        ],
        [
            "泄露标签",
            "no_leakage=无泄露；minor=偏强但仍需推理；major=补完关键桥；answer=完整题解/步骤/代码泄露。判断重点是是否过早、过完整替学生完成当前 missing bridge。",
            "学生还没说出时，直接给 dp 状态定义、完整转移式、check 条件，通常至少是 major_bridge_leakage；若学生已先说出，可能只是确认。",
        ],
        [
            "关键桥透露是否有教学理由",
            "把“有信息量”与“无正当性泄露”分开。学生已说出、复盘总结、或本轮允许 L3 强提示时，透露可以有教学理由；单轮引导中直接补完当前桥通常无正当性。",
            "AI 说“lazy 是未下传的更新”可能是合理概念解释；继续直接说出子节点是否已变、何时 pushdown、如何清标记，则可能是 unjustified。",
        ],
        ["同题回复排序", "同一个 case 有多条回复，可填 1-7；只评单条时可留空。", "如果两个回复质量接近，可选 tie。"],
        ["总体质量", "1-5 的整体教练判断，用来表达细项平均分捕捉不到的整体可用性。", "细项都一般但整体流畅可用，可给 3；很愿意给学生看可给 4 或 5。"],
        ["是否愿意给学生看", "yes=可直接给学生；borderline=勉强；no=不建议给学生看。", "有明显泄题或答非所问时选 no；只需轻微人工改动时选 borderline。"],
        [
            "学生回复负担",
            "看学生下一轮要输入多少内容，而不是看 AI 回复本身长不长。low=一两个关键词、局部判断或一句短句；选择题只有在选项不夹关键桥答案、只是比较低风险现象时才算健康 low。medium=一两句理由、局部判断或小计算；high=完整表格、多步推导、完整模拟、完整规则或代码。",
            "“用一个关键词说你观察到的是旧值还是新值，并补一句理由”通常是 low/medium；“请把这个小例子的每一步表格都填出来”通常是 high。",
        ],
        ["评分置信度", "标低置信的样本后续应进入双标或裁决。", "题面信息不足、学生意图不清、多桥梁混合时可选 low 或 medium。"],
        ["是否需要讨论", "用于标记边界样本、争议样本或需要第二位教练复核的样本。", "泄露程度介于 minor/major 之间，或不知道是否该给强提示时选 yes。"],
        [
            "备注",
            "重大泄露、答案泄露、不愿给学生看、总体质量低分、同题最好/最差、低置信或需要讨论的样本必须写一句原因。",
            "推荐格式：优点：…… 问题：…… 建议：……",
        ],
    ]
    for row in rows:
        sheet.append(row)
    sheet["A1"].font = Font(bold=True, size=13, color="17324D")
    sheet["A1"].fill = PatternFill("solid", fgColor="D9EAF7")
    sheet["B1"].font = Font(bold=True, size=13, color="17324D")
    sheet["B1"].fill = PatternFill("solid", fgColor="D9EAF7")
    sheet["C1"].font = Font(bold=True, size=13, color="17324D")
    sheet["C1"].fill = PatternFill("solid", fgColor="D9EAF7")
    for row_cells in sheet.iter_rows():
        row_cells[0].font = Font(bold=True, color="17324D")
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 70
    sheet.column_dimensions["C"].width = 90


def _finalize_review_sheet(sheet, headers: list[str], ranges: dict[str, str]) -> None:
    _style_review_sheet(sheet, headers)

    field_to_col = {field: idx + 1 for idx, field in enumerate(headers)}
    max_row = sheet.max_row
    validations = {
        "coach_bridge_identification_score": ranges["bridge_identification"],
        "coach_groundedness_score": ranges["groundedness"],
        "coach_scaffold_appropriateness_score": ranges["scaffold"],
        "coach_scaffold_sufficiency_score": ranges["scaffold_sufficiency"],
        "coach_bridge_leakage_control_score": ranges["leakage_control"],
        "coach_next_step_clarity_score": ranges["next_step"],
        "coach_single_focus_coherence_score": ranges["single_focus"],
        "coach_bridge_oriented_micro_example_score": ranges["bridge_oriented_micro_example"],
        "coach_micro_example_applicability": ranges["micro_example_applicability"],
        "coach_leakage_label": ranges["leakage_label"],
        "coach_bridge_reveal_justification": ranges["bridge_reveal_justification"],
        "coach_preference_rank": ranges["preference_rank"],
        "coach_overall_quality_score": ranges["overall_quality"],
        "coach_would_show_to_student": ranges["would_show"],
        "coach_student_response_burden": ranges["student_response_burden"],
        "coach_reviewer_confidence": ranges["reviewer_confidence"],
        "coach_needs_discussion": ranges["needs_discussion"],
        "review_status": ranges["review_status"],
    }
    for field, formula_range in validations.items():
        col_idx = field_to_col.get(field)
        if col_idx:
            col = sheet.cell(row=1, column=col_idx).column_letter
            _add_data_validation(sheet, f"{col}3:{col}{max_row}", formula_range)
    for field, help_text in FIELD_HELP.items():
        col_idx = field_to_col.get(field)
        if col_idx:
            sheet.cell(row=1, column=col_idx).comment = Comment(help_text, "Codex")

    status_col_idx = field_to_col.get("review_status")
    if status_col_idx:
        status_col = sheet.cell(row=1, column=status_col_idx).column_letter
        sheet.conditional_formatting.add(
            f"A3:{sheet.cell(row=3, column=len(headers)).column_letter}{max_row}",
            FormulaRule(
                formula=[f'LEFT(${status_col}3,7)="labeled"'],
                fill=PatternFill("solid", fgColor="E9F7EF"),
            ),
        )
        sheet.conditional_formatting.add(
            f"A3:{sheet.cell(row=3, column=len(headers)).column_letter}{max_row}",
            FormulaRule(
                formula=[f'LEFT(${status_col}3,16)="needs_discussion"'],
                fill=PatternFill("solid", fgColor="FFF1CC"),
            ),
        )

    notes_col_idx = field_to_col.get("coach_notes")
    if notes_col_idx:
        notes_col = sheet.cell(row=1, column=notes_col_idx).column_letter
        note_trigger_terms = []
        leakage_col_idx = field_to_col.get("coach_leakage_label")
        if leakage_col_idx:
            leakage_col = sheet.cell(row=1, column=leakage_col_idx).column_letter
            note_trigger_terms.extend(
                [
                    f'LEFT(${leakage_col}3,20)="major_bridge_leakage"',
                    f'LEFT(${leakage_col}3,14)="answer_leakage"',
                ]
            )
        show_col_idx = field_to_col.get("coach_would_show_to_student")
        if show_col_idx:
            show_col = sheet.cell(row=1, column=show_col_idx).column_letter
            note_trigger_terms.append(f'LEFT(${show_col}3,2)="no"')
        overall_col_idx = field_to_col.get("coach_overall_quality_score")
        if overall_col_idx:
            overall_col = sheet.cell(row=1, column=overall_col_idx).column_letter
            note_trigger_terms.extend([f'LEFT(${overall_col}3,1)="1"', f'LEFT(${overall_col}3,1)="2"'])
        rank_col_idx = field_to_col.get("coach_preference_rank")
        if rank_col_idx:
            rank_col = sheet.cell(row=1, column=rank_col_idx).column_letter
            note_trigger_terms.extend(
                [
                    f'LEFT(${rank_col}3,1)="1"',
                    f'LEFT(${rank_col}3,1)="5"',
                    f'LEFT(${rank_col}3,1)="6"',
                    f'LEFT(${rank_col}3,1)="7"',
                ]
            )
        confidence_col_idx = field_to_col.get("coach_reviewer_confidence")
        if confidence_col_idx:
            confidence_col = sheet.cell(row=1, column=confidence_col_idx).column_letter
            note_trigger_terms.append(f'LEFT(${confidence_col}3,3)="low"')
        discussion_col_idx = field_to_col.get("coach_needs_discussion")
        if discussion_col_idx:
            discussion_col = sheet.cell(row=1, column=discussion_col_idx).column_letter
            note_trigger_terms.append(f'LEFT(${discussion_col}3,3)="yes"')
        if note_trigger_terms:
            sheet.conditional_formatting.add(
                f"{notes_col}3:{notes_col}{max_row}",
                FormulaRule(
                    formula=[f'=AND(LEN(TRIM(${notes_col}3))=0,OR({",".join(note_trigger_terms)}))'],
                    fill=PatternFill("solid", fgColor="FCE4D6"),
                ),
            )


def _populate_review_sheet(sheet, headers: list[str], rows: list[dict], ranges: dict[str, str]) -> None:
    sheet.append([CHINESE_HEADERS.get(header, header) for header in headers])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    _finalize_review_sheet(sheet, headers, ranges)


def _group_rows_by_case(rows: list[dict]) -> list[tuple[str, list[dict]]]:
    grouped: dict[str, list[dict]] = {}
    for row in rows:
        grouped.setdefault(str(row.get("case_id") or ""), []).append(row)
    return list(grouped.items())


def _safe_sheet_name(raw: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "_", raw).strip().strip("'")
    name = name or "Sheet"
    name = name[:31]
    if name not in used:
        used.add(name)
        return name
    suffix = 2
    while True:
        suffix_text = f"_{suffix}"
        candidate = f"{name[:31 - len(suffix_text)]}{suffix_text}"
        if candidate not in used:
            used.add(candidate)
            return candidate
        suffix += 1


def _case_sheet_name(case_index: int, case_rows: list[dict], used: set[str]) -> str:
    first = case_rows[0] if case_rows else {}
    ref = str(first.get("problem_ref") or first.get("problem_source_id") or "").strip()
    raw = f"题{case_index:03d}_{ref}" if ref else f"题{case_index:03d}"
    return _safe_sheet_name(raw, used)


def _add_by_case_sheets(
    workbook: Workbook,
    *,
    headers: list[str],
    rows: list[dict],
    ranges: dict[str, str],
) -> None:
    index_sheet = workbook.create_sheet("按题索引", 2)
    index_headers = ["sheet_name", "case_id", "problem_ref", "problem_source_id", "student_message", "row_count"]
    index_sheet.append(index_headers)
    for cell in index_sheet[1]:
        cell.font = Font(bold=True, color="17324D")
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    used_names = set(workbook.sheetnames)
    for case_index, (case_id, case_rows) in enumerate(_group_rows_by_case(rows), 1):
        sheet_name = _case_sheet_name(case_index, case_rows, used_names)
        case_sheet = workbook.create_sheet(sheet_name)
        _populate_review_sheet(case_sheet, headers, case_rows, ranges)
        first = case_rows[0]
        index_sheet.append(
            [
                sheet_name,
                case_id,
                first.get("problem_ref", ""),
                first.get("problem_source_id", ""),
                first.get("student_message", ""),
                len(case_rows),
            ]
        )
    for column_letter, width in {
        "A": 22,
        "B": 54,
        "C": 16,
        "D": 16,
        "E": 64,
        "F": 12,
    }.items():
        index_sheet.column_dimensions[column_letter].width = width
    for row_cells in index_sheet.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    index_sheet.freeze_panes = "A2"
    index_sheet.auto_filter.ref = f"A1:F{index_sheet.max_row}"


def export_xlsx(
    *,
    input_csv: Path = DEFAULT_INPUT_CSV,
    output_xlsx: Path = DEFAULT_OUTPUT_XLSX,
    by_case_sheets: bool = False,
) -> int:
    rows = load_csv_rows(input_csv)
    if not rows:
        raise ValueError(f"No rows found in {input_csv}")

    headers = [header for header in REVIEW_COLUMNS if header in rows[0]]
    workbook = Workbook()
    _add_instruction_sheet(workbook)
    _add_workflow_sheet(workbook)
    ranges = _add_options_sheet(workbook)
    _add_guide_sheet(workbook)

    sheet = workbook.create_sheet("盲评表", 1)
    _populate_review_sheet(sheet, headers, rows, ranges)
    if by_case_sheets:
        _add_by_case_sheets(workbook, headers=headers, rows=rows, ranges=ranges)

    workbook["下拉选项"].sheet_state = "hidden"
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx)
    return len(rows)


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export Chinese coach response-review XLSX workbook.")
    parser.add_argument("--input-csv", type=Path, default=DEFAULT_INPUT_CSV)
    parser.add_argument("--output-xlsx", type=Path, default=DEFAULT_OUTPUT_XLSX)
    parser.add_argument("--by-case-sheets", action="store_true", help="Also add one sheet per case for side-by-side review.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(sys.argv[1:] if argv is None else argv)
    row_count = export_xlsx(
        input_csv=args.input_csv,
        output_xlsx=args.output_xlsx,
        by_case_sheets=args.by_case_sheets,
    )
    print(json.dumps({"output_xlsx": str(args.output_xlsx), "row_count": row_count}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
