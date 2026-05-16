"""Export a source-completion CSV for the 50-case held-out draft.

This helper keeps real problem-source metadata out of ad-hoc JSON editing. It
exports the current draft context plus blank source fields, so researchers or
coaches can fill platform ids, original links, and statement visibility notes in
a spreadsheet before merging them back into JSONL.
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import TextIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from evals.aichat.export_coach_response_review_workbook import _student_message_length_bucket


DEFAULT_INPUT_JSONL = Path("docs/research/bridgebench_cp_heldout_v1_50_ai_reference_draft_20260513.jsonl")
DEFAULT_OUTPUT_CSV = Path("docs/research/heldout_50_source_completion_workbook_20260513.csv")
DEFAULT_OUTPUT_XLSX = Path("docs/research/heldout_50_source_completion_workbook_20260513.zh.xlsx")

SOURCE_METADATA_FIELDS = [
    "problem_source_platform",
    "problem_source_id",
    "problem_source_url",
    "problem_statement",
    "problem_statement_public_summary",
    "problem_statement_rights_note",
    "problem_statement_access_level",
    "source_completion_status",
    "source_notes",
]

SOURCE_COMPLETION_COLUMNS = [
    "case_id",
    "category",
    "problem_ref",
    "student_message",
    "student_message_length_bucket",
    "problem_context",
    "recent_dialogue",
    "student_code_excerpt",
    "missing_bridge",
    *SOURCE_METADATA_FIELDS,
]

CHINESE_HEADERS = {
    "case_id": "样本编号",
    "category": "类别",
    "problem_ref": "内部题目编号",
    "student_message": "学生当前问题",
    "student_message_length_bucket": "学生问题长度类型",
    "problem_context": "题目/上下文",
    "recent_dialogue": "近期对话",
    "student_code_excerpt": "学生代码片段",
    "missing_bridge": "缺失桥梁",
    "problem_source_platform": "题目来源平台",
    "problem_source_id": "平台题号",
    "problem_source_url": "原题链接",
    "problem_statement": "原题题面/必要题面",
    "problem_statement_public_summary": "公开题面摘要",
    "problem_statement_rights_note": "题面版权/使用说明",
    "problem_statement_access_level": "题面访问级别",
    "source_completion_status": "题源补全状态",
    "source_notes": "题源备注",
}

ACCESS_LEVEL_OPTIONS = [
    "local_review_only",
    "public_summary_only",
    "open_license",
    "original_link_only",
]


def load_jsonl(path: Path) -> list[dict]:
    rows = []
    for line_number, raw_line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        line = raw_line.strip()
        if not line:
            continue
        try:
            rows.append(json.loads(line))
        except json.JSONDecodeError as exc:
            raise ValueError(f"Invalid JSONL at {path}:{line_number}: {exc}") from exc
    return rows


def build_source_completion_rows(rows: list[dict]) -> list[dict]:
    output = []
    for row in rows:
        student_message = str(row.get("student_message") or "")
        output.append(
            {
                "case_id": row.get("case_id") or row.get("id") or "",
                "category": row.get("category") or "",
                "problem_ref": row.get("problem_ref") or "",
                "student_message": student_message,
                "student_message_length_bucket": _student_message_length_bucket(student_message),
                "problem_context": row.get("problem_context") or "",
                "recent_dialogue": row.get("recent_dialogue") or "",
                "student_code_excerpt": row.get("student_code_excerpt") or "",
                "missing_bridge": row.get("missing_bridge") or row.get("gold_missing_link") or "",
                **{field: row.get(field) or "" for field in SOURCE_METADATA_FIELDS},
            }
        )
    return output


def write_source_completion_csv(output: TextIO, rows: list[dict]) -> None:
    writer = csv.DictWriter(output, fieldnames=SOURCE_COMPLETION_COLUMNS, extrasaction="ignore", lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)


def export_source_completion_csv(
    *,
    input_jsonl: Path = DEFAULT_INPUT_JSONL,
    output_csv: Path = DEFAULT_OUTPUT_CSV,
) -> int:
    rows = build_source_completion_rows(load_jsonl(input_jsonl))
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", encoding="utf-8-sig", newline="") as handle:
        write_source_completion_csv(handle, rows)
    return len(rows)


def export_source_completion_xlsx(
    *,
    input_jsonl: Path = DEFAULT_INPUT_JSONL,
    output_xlsx: Path = DEFAULT_OUTPUT_XLSX,
) -> int:
    rows = build_source_completion_rows(load_jsonl(input_jsonl))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "题源补全表"
    sheet.append([CHINESE_HEADERS.get(header, header) for header in SOURCE_COMPLETION_COLUMNS])
    sheet.append(SOURCE_COMPLETION_COLUMNS)
    for row in rows:
        sheet.append([row.get(header, "") for header in SOURCE_COMPLETION_COLUMNS])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    machine_fill = PatternFill("solid", fgColor="F2F5F7")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="17324D")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for cell in sheet[2]:
        cell.font = Font(size=9, color="667085")
        cell.fill = machine_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for row_cells in sheet.iter_rows(min_row=3, max_row=sheet.max_row):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {
        "A": 16,
        "B": 22,
        "C": 28,
        "D": 36,
        "E": 18,
        "F": 44,
        "G": 40,
        "H": 32,
        "I": 44,
        "J": 18,
        "K": 18,
        "L": 44,
        "M": 60,
        "N": 48,
        "O": 44,
        "P": 22,
        "Q": 20,
        "R": 36,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "J3"
    sheet.auto_filter.ref = f"A2:{sheet.cell(row=2, column=len(SOURCE_COMPLETION_COLUMNS)).coordinate}"

    options = workbook.create_sheet("访问级别选项")
    options.append(["problem_statement_access_level"])
    for value in ACCESS_LEVEL_OPTIONS:
        options.append([value])
    options.sheet_state = "hidden"
    access_col_idx = SOURCE_COMPLETION_COLUMNS.index("problem_statement_access_level") + 1
    access_col = sheet.cell(row=1, column=access_col_idx).column_letter
    validation = DataValidation(
        type="list",
        formula1=f"='访问级别选项'!$A$2:$A${len(ACCESS_LEVEL_OPTIONS) + 1}",
        allow_blank=True,
    )
    sheet.add_data_validation(validation)
    validation.add(f"{access_col}3:{access_col}{sheet.max_row}")

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx)
    return len(rows)


def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export held-out source-completion CSV.")
    parser.add_argument("--input-jsonl", type=Path, default=DEFAULT_INPUT_JSONL)
    parser.add_argument("--output-csv", type=Path, default=DEFAULT_OUTPUT_CSV)
    parser.add_argument("--output-xlsx", type=Path, default=None)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(argv)
    row_count = export_source_completion_csv(input_jsonl=args.input_jsonl, output_csv=args.output_csv)
    payload = {"output_csv": str(args.output_csv), "row_count": row_count}
    if args.output_xlsx:
        export_source_completion_xlsx(input_jsonl=args.input_jsonl, output_xlsx=args.output_xlsx)
        payload["output_xlsx"] = str(args.output_xlsx)
    print(json.dumps(payload, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
