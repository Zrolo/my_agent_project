"""Export a side-by-side adjudication workbook for Coach A/B disagreements."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from evals.aichat.summarize_coach_label_agreement import load_jsonl_by_case_id


DEFAULT_AGREEMENT_JSON = Path("docs/research/coach_agreement_heldout_v1_50_overlap20_20260512.json")
DEFAULT_ANNOTATOR_A_JSONL = Path("docs/research/coach_reference_heldout_v1_50_a.jsonl")
DEFAULT_ANNOTATOR_B_JSONL = Path("docs/research/coach_reference_heldout_v1_50_b_overlap_20.jsonl")
DEFAULT_DRAFT_JSONL = Path("docs/research/bridgebench_cp_heldout_v1_50_draft.jsonl")
DEFAULT_OUTPUT_XLSX = Path("docs/research/coach_adjudication_workbook_heldout_v1_50_overlap20.zh.xlsx")

ADJUDICATION_COLUMNS = [
    "case_id",
    "student_message",
    "problem_context",
    "recent_dialogue",
    "student_code_excerpt",
    "a_student_problem_solving_state",
    "b_student_problem_solving_state",
    "adjudicated_student_problem_solving_state",
    "a_primary_bridge_family",
    "b_primary_bridge_family",
    "adjudicated_primary_bridge_family",
    "a_primary_bridge_subtype_id",
    "b_primary_bridge_subtype_id",
    "adjudicated_primary_bridge_subtype_id",
    "a_registered_focus_id",
    "b_registered_focus_id",
    "adjudicated_registered_focus_id",
    "a_max_scaffold_level",
    "b_max_scaffold_level",
    "adjudicated_max_scaffold_level",
    "a_leakage_risk",
    "b_leakage_risk",
    "adjudicated_leakage_risk",
    "a_forbidden_content",
    "b_forbidden_content",
    "adjudicated_forbidden_content",
    "a_evidence_quote",
    "b_evidence_quote",
    "adjudication_status",
    "adjudication_notes",
]

CHINESE_HEADERS = {
    "case_id": "样本编号",
    "student_message": "学生当前问题",
    "problem_context": "题目/上下文",
    "recent_dialogue": "近期对话",
    "student_code_excerpt": "学生代码片段",
    "a_student_problem_solving_state": "Coach A 学生状态",
    "b_student_problem_solving_state": "Coach B 学生状态",
    "adjudicated_student_problem_solving_state": "裁决学生状态",
    "a_primary_bridge_family": "Coach A 主要桥梁大类",
    "b_primary_bridge_family": "Coach B 主要桥梁大类",
    "adjudicated_primary_bridge_family": "裁决主要桥梁大类",
    "a_primary_bridge_subtype_id": "Coach A 桥梁细分",
    "b_primary_bridge_subtype_id": "Coach B 桥梁细分",
    "adjudicated_primary_bridge_subtype_id": "裁决桥梁细分",
    "a_registered_focus_id": "Coach A focus",
    "b_registered_focus_id": "Coach B focus",
    "adjudicated_registered_focus_id": "裁决 focus",
    "a_max_scaffold_level": "Coach A 最多帮助强度",
    "b_max_scaffold_level": "Coach B 最多帮助强度",
    "adjudicated_max_scaffold_level": "裁决最多帮助强度",
    "a_leakage_risk": "Coach A 泄露风险",
    "b_leakage_risk": "Coach B 泄露风险",
    "adjudicated_leakage_risk": "裁决泄露风险",
    "a_forbidden_content": "Coach A 禁止补完",
    "b_forbidden_content": "Coach B 禁止补完",
    "adjudicated_forbidden_content": "裁决禁止补完",
    "a_evidence_quote": "Coach A 证据",
    "b_evidence_quote": "Coach B 证据",
    "adjudication_status": "裁决状态",
    "adjudication_notes": "裁决备注",
}


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _read_jsonl_by_case_id(path: Path) -> dict[str, dict]:
    rows = {}
    for line_number, raw_line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        line = raw_line.strip()
        if not line:
            continue
        row = json.loads(line)
        case_id = row.get("case_id") or row.get("id")
        if not case_id:
            raise ValueError(f"Missing case_id at {path}:{line_number}")
        rows[str(case_id)] = row
    return rows


def _join(value: object) -> str:
    if isinstance(value, list):
        return ";".join(str(item).strip() for item in value if str(item).strip())
    return str(value or "").strip()


def _forbidden_content(row: dict) -> str:
    values = []
    for key in ("general_forbidden_content", "bridge_specific_forbidden_content", "forbidden_content"):
        value = row.get(key)
        if isinstance(value, list):
            values.extend(str(item).strip() for item in value if str(item).strip())
        elif isinstance(value, str) and value.strip():
            values.append(value.strip())
    return ";".join(values)


def _case_ids_for_export(
    *,
    agreement: dict,
    annotator_a: dict[str, dict],
    annotator_b: dict[str, dict],
    include_all_paired: bool,
) -> list[str]:
    if include_all_paired:
        return sorted(set(annotator_a) & set(annotator_b))
    return list(agreement.get("needs_adjudication_case_ids") or [])


def build_adjudication_rows(
    *,
    agreement: dict,
    annotator_a: dict[str, dict],
    annotator_b: dict[str, dict],
    draft_rows: dict[str, dict],
    include_all_paired: bool = False,
) -> list[dict]:
    rows = []
    for case_id in _case_ids_for_export(
        agreement=agreement,
        annotator_a=annotator_a,
        annotator_b=annotator_b,
        include_all_paired=include_all_paired,
    ):
        draft = draft_rows.get(case_id, {})
        a = annotator_a.get(case_id, {})
        b = annotator_b.get(case_id, {})
        rows.append(
            {
                "case_id": case_id,
                "student_message": draft.get("student_message", ""),
                "problem_context": draft.get("problem_context", ""),
                "recent_dialogue": draft.get("recent_dialogue", ""),
                "student_code_excerpt": draft.get("student_code_excerpt", ""),
                "a_student_problem_solving_state": a.get("student_problem_solving_state", ""),
                "b_student_problem_solving_state": b.get("student_problem_solving_state", ""),
                "adjudicated_student_problem_solving_state": "",
                "a_primary_bridge_family": a.get("primary_bridge_family", ""),
                "b_primary_bridge_family": b.get("primary_bridge_family", ""),
                "adjudicated_primary_bridge_family": "",
                "a_primary_bridge_subtype_id": a.get("primary_bridge_subtype_id", ""),
                "b_primary_bridge_subtype_id": b.get("primary_bridge_subtype_id", ""),
                "adjudicated_primary_bridge_subtype_id": "",
                "a_registered_focus_id": a.get("registered_focus_id", ""),
                "b_registered_focus_id": b.get("registered_focus_id", ""),
                "adjudicated_registered_focus_id": "",
                "a_max_scaffold_level": a.get("max_scaffold_level", ""),
                "b_max_scaffold_level": b.get("max_scaffold_level", ""),
                "adjudicated_max_scaffold_level": "",
                "a_leakage_risk": a.get("leakage_risk", ""),
                "b_leakage_risk": b.get("leakage_risk", ""),
                "adjudicated_leakage_risk": "",
                "a_forbidden_content": _forbidden_content(a),
                "b_forbidden_content": _forbidden_content(b),
                "adjudicated_forbidden_content": "",
                "a_evidence_quote": _join(a.get("evidence_quote", "")),
                "b_evidence_quote": _join(b.get("evidence_quote", "")),
                "adjudication_status": "",
                "adjudication_notes": "",
            }
        )
    return rows


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    machine_fill = PatternFill("solid", fgColor="F2F5F7")
    adjudication_fill = PatternFill("solid", fgColor="FFF7D6")
    thin = Side(style="thin", color="D7DEE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="17324D")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    for cell in sheet[2]:
        cell.font = Font(size=9, color="667085")
        cell.fill = machine_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    for row in sheet.iter_rows(min_row=3, max_row=sheet.max_row, max_col=len(ADJUDICATION_COLUMNS)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            header = sheet.cell(row=2, column=cell.column).value
            if str(header or "").startswith("adjudicated_") or header in {"adjudication_status", "adjudication_notes"}:
                cell.fill = adjudication_fill
    widths = {
        "A": 18,
        "B": 34,
        "C": 42,
        "D": 44,
        "E": 32,
        "F": 24,
        "G": 24,
        "H": 26,
        "I": 28,
        "J": 28,
        "K": 30,
        "L": 32,
        "M": 32,
        "N": 34,
        "O": 28,
        "P": 28,
        "Q": 30,
        "R": 18,
        "S": 18,
        "T": 20,
        "U": 18,
        "V": 18,
        "W": 20,
        "X": 36,
        "Y": 36,
        "Z": 38,
        "AA": 32,
        "AB": 32,
        "AC": 18,
        "AD": 40,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    sheet.row_dimensions[1].height = 38
    sheet.row_dimensions[2].height = 28
    for row_idx in range(3, sheet.max_row + 1):
        sheet.row_dimensions[row_idx].height = 96
    sheet.freeze_panes = "F3"
    sheet.auto_filter.ref = f"A2:{sheet.cell(row=2, column=len(ADJUDICATION_COLUMNS)).coordinate}"


def export_xlsx(
    *,
    agreement_json: Path = DEFAULT_AGREEMENT_JSON,
    annotator_a_jsonl: Path = DEFAULT_ANNOTATOR_A_JSONL,
    annotator_b_jsonl: Path = DEFAULT_ANNOTATOR_B_JSONL,
    draft_jsonl: Path = DEFAULT_DRAFT_JSONL,
    output_xlsx: Path = DEFAULT_OUTPUT_XLSX,
    include_all_paired: bool = False,
) -> int:
    agreement = _read_json(Path(agreement_json))
    annotator_a = load_jsonl_by_case_id(Path(annotator_a_jsonl))
    annotator_b = load_jsonl_by_case_id(Path(annotator_b_jsonl))
    draft_rows = _read_jsonl_by_case_id(Path(draft_jsonl))
    rows = build_adjudication_rows(
        agreement=agreement,
        annotator_a=annotator_a,
        annotator_b=annotator_b,
        draft_rows=draft_rows,
        include_all_paired=include_all_paired,
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "裁决表"
    sheet.append([CHINESE_HEADERS[column] for column in ADJUDICATION_COLUMNS])
    sheet.append(ADJUDICATION_COLUMNS)
    for row in rows:
        sheet.append([row.get(column, "") for column in ADJUDICATION_COLUMNS])
    _style_sheet(sheet)

    guide = workbook.create_sheet("说明")
    guide.append(["用途", "本表用于裁决 Coach A / Coach B 的 held-out overlap 分歧，不是 AI 回复盲评表。"])
    guide.append(["黄色列", "请填写裁决后的 reference label；必要时在 adjudication_notes 说明原因。"])
    guide.append(["裁决状态", "建议填写 resolved / needs_discussion / exclude_from_headline。"])
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 90
    for row_cells in guide.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    output_xlsx = Path(output_xlsx)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx)
    return len(rows)


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export a Coach A/B adjudication workbook for held-out overlap cases.")
    parser.add_argument("--agreement-json", type=Path, default=DEFAULT_AGREEMENT_JSON)
    parser.add_argument("--annotator-a-jsonl", type=Path, default=DEFAULT_ANNOTATOR_A_JSONL)
    parser.add_argument("--annotator-b-jsonl", type=Path, default=DEFAULT_ANNOTATOR_B_JSONL)
    parser.add_argument("--draft-jsonl", type=Path, default=DEFAULT_DRAFT_JSONL)
    parser.add_argument("--output-xlsx", type=Path, default=DEFAULT_OUTPUT_XLSX)
    parser.add_argument(
        "--include-all-paired",
        action="store_true",
        help="Export every paired overlap row instead of only needs_adjudication_case_ids.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(sys.argv[1:] if argv is None else argv)
    row_count = export_xlsx(
        agreement_json=args.agreement_json,
        annotator_a_jsonl=args.annotator_a_jsonl,
        annotator_b_jsonl=args.annotator_b_jsonl,
        draft_jsonl=args.draft_jsonl,
        output_xlsx=args.output_xlsx,
        include_all_paired=args.include_all_paired,
    )
    print(json.dumps({"output_xlsx": str(args.output_xlsx), "row_count": row_count}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
