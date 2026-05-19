#!/usr/bin/env python3
"""Export Chinese coach-facing real-student pilot review forms to XLSX.

The XLSX form mirrors the coach-review workbook style used in the 50-case
human review: dropdowns, color cues, frozen headers, and a visible option sheet.
It is only a pilot/dry-run review artifact and does not modify dialogue-state
v3 main experiment data or online AIChat behavior.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


DEFAULT_INPUT_CSV = Path("docs/research/real_student_online_5case_coach_review_form_cn_v2.csv")
DEFAULT_OUTPUT_XLSX = Path("docs/research/real_student_online_5case_coach_review_form_cn_v2.xlsx")

COACH_REVIEW_CN_COLUMNS = [
    "干跑案例编号",
    "候选轮次编号",
    "试点案例编号",
    "30例候选序号",
    "时间桶",
    "题目/场景摘要",
    "学生问题（已脱敏）",
    "近期对话（已脱敏，可空）",
    "学生代码片段（已脱敏，可空）",
    "当前AIChat回复（已脱敏）",
    "复核状态",
    "隐私复核状态",
    "上下文是否足够",
    "识别当前缺失桥是否准确（0/1/2）",
    "回复是否基于当前材料（0/1/2）",
    "脚手架是否合适（0/1/2）",
    "是否控制关键桥泄露（0/1/2）",
    "下一步是否清楚（0/1/2）",
    "是否单焦点连贯（0/1/2）",
    "微例是否围绕桥梁（0/1/2/不适用）",
    "微例适用性",
    "泄露标签",
    "总体质量（1-5）",
    "是否愿意给学生看",
    "学生回答负担",
    "复核信心",
    "是否需要讨论",
    "缺失桥家族（中文）",
    "当前缺失桥实例（中文简述）",
    "禁止直接说出的内容（中文）",
    "可以提示到什么程度（中文）",
    "期望学生下一步（中文）",
    "是否匹配现有taxonomy",
    "新桥候选（如无填“无”）",
    "观察到的下一轮进展",
    "是否同意AI预标注",
    "是否需要裁决",
    "知情/报告门",
    "教练备注（中文）",
]

OPTION_LISTS = {
    "复核状态": ["待复核", "已复核", "需第二教练", "因隐私风险排除", "因上下文不足排除"],
    "隐私复核状态": ["待复核", "可内部复核", "需继续脱敏", "因隐私风险排除"],
    "上下文是否足够": ["足够", "部分足够", "不足", "不清楚"],
    "识别当前缺失桥是否准确（0/1/2）": ["2", "1", "0"],
    "回复是否基于当前材料（0/1/2）": ["2", "1", "0"],
    "脚手架是否合适（0/1/2）": ["2", "1", "0"],
    "是否控制关键桥泄露（0/1/2）": ["2", "1", "0"],
    "下一步是否清楚（0/1/2）": ["2", "1", "0"],
    "是否单焦点连贯（0/1/2）": ["2", "1", "0"],
    "微例是否围绕桥梁（0/1/2/不适用）": ["2", "1", "0", "不适用"],
    "微例适用性": ["适用", "不适用", "不清楚"],
    "泄露标签": ["无泄露", "轻微关键桥泄露", "重大关键桥泄露", "答案或代码泄露", "上下文不足未判断"],
    "总体质量（1-5）": ["5", "4", "3", "2", "1"],
    "是否愿意给学生看": ["是", "边界", "否"],
    "学生回答负担": ["低", "中", "高"],
    "复核信心": ["高", "中", "低"],
    "是否需要讨论": ["否", "是"],
    "缺失桥家族（中文）": [
        "二分检查语义桥",
        "DP状态/转移桥",
        "贡献/差分汇总桥",
        "数据结构语义桥",
        "图/树遍历桥",
        "调试定位桥",
        "实现边界桥",
        "上下文不足/澄清优先",
        "其他/不确定",
    ],
    "是否匹配现有taxonomy": ["是", "否", "不确定"],
    "新桥候选（如无填“无”）": ["无", "有：见教练备注", "不确定：见教练备注"],
    "观察到的下一轮进展": ["有进展", "部分进展", "无进展", "不清楚", "不可用"],
    "是否同意AI预标注": ["同意", "部分同意", "不同意", "不适用"],
    "是否需要裁决": ["否", "是"],
    "知情/报告门": ["待完成", "可报告", "不可报告", "已退出"],
}

FIELD_COMMENTS = {
    "识别当前缺失桥是否准确（0/1/2）": "对齐 50-case: coach_bridge_identification_score。2=准确抓住；1=部分抓住；0=没抓住。",
    "回复是否基于当前材料（0/1/2）": "对齐 50-case: coach_groundedness_score。看回复是否基于题目、学生话语和近期对话。",
    "脚手架是否合适（0/1/2）": "对齐 50-case: coach_scaffold_appropriateness_score。太弱、太强、直接给答案都扣分。",
    "是否控制关键桥泄露（0/1/2）": "对齐 50-case: coach_bridge_leakage_control_score。0=严重泄露，1=边界/轻微，2=控制良好。",
    "泄露标签": "对齐 50-case: coach_leakage_label。判断当前 AIChat 是否提前补完 critical bridge。",
    "总体质量（1-5）": "对齐 50-case: coach_overall_quality_score。1=不可用，5=很适合给学生。",
    "是否愿意给学生看": "对齐 50-case: coach_would_show_to_student。",
    "学生回答负担": "对齐 50-case: student_response_burden。低=短回复即可，高=完整推导/表格/代码。",
    "知情/报告门": "未完成前保持“待完成”。不是“可报告”时，不能写成公开 deep-pilot evidence。",
}

ALIGNMENT_ROWS = [
    ("中文列名", "50-case 对应字段", "说明"),
    ("识别当前缺失桥是否准确（0/1/2）", "coach_bridge_identification_score", "是否抓住学生当前缺失桥/卡点"),
    ("回复是否基于当前材料（0/1/2）", "coach_groundedness_score", "是否基于题目、学生话语和近期对话"),
    ("脚手架是否合适（0/1/2）", "coach_scaffold_appropriateness_score", "帮助强度是否适合当前学生状态"),
    ("是否控制关键桥泄露（0/1/2）", "coach_bridge_leakage_control_score", "是否避免提前补完 critical bridge"),
    ("下一步是否清楚（0/1/2）", "coach_next_step_clarity_score", "学生看完是否知道下一步做什么"),
    ("是否单焦点连贯（0/1/2）", "coach_single_focus_coherence_score", "是否围绕一个主要卡点推进"),
    ("微例是否围绕桥梁（0/1/2/不适用）", "coach_bridge_oriented_micro_example_score", "微例是否导向可迁移桥梁"),
    ("微例适用性", "coach_micro_example_applicability", "是否需要评价微例"),
    ("泄露标签", "coach_leakage_label", "no/minor/major/answer leakage 的中文口径"),
    ("总体质量（1-5）", "coach_overall_quality_score", "整体教学质量"),
    ("是否愿意给学生看", "coach_would_show_to_student", "是否可直接给学生"),
    ("学生回答负担", "student_response_burden", "学生下一轮输入负担"),
    ("复核信心", "coach_reviewer_confidence", "教练评分把握"),
    ("是否需要讨论", "coach_needs_discussion", "是否需二次讨论/裁决"),
    ("教练备注（中文）", "coach_notes", "中文简要说明判断理由"),
]


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _add_data_validation(sheet, cell_range: str, formula_range: str) -> None:
    validation = DataValidation(type="list", formula1=f"={formula_range}", allow_blank=True)
    validation.error = "请从下拉列表中选择中文选项，或留空。"
    validation.errorTitle = "无效选项"
    validation.prompt = "请用下拉选择，保证能被脚本统计。"
    validation.promptTitle = "结构化复核"
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def _build_option_sheet(workbook: Workbook) -> dict[str, str]:
    sheet = workbook.create_sheet("下拉选项")
    ranges: dict[str, str] = {}
    for col_idx, (field, values) in enumerate(OPTION_LISTS.items(), start=1):
        col = get_column_letter(col_idx)
        sheet.cell(row=1, column=col_idx, value=field)
        sheet.cell(row=1, column=col_idx).font = Font(bold=True, color="FFFFFF")
        sheet.cell(row=1, column=col_idx).fill = PatternFill("solid", fgColor="305496")
        for row_idx, value in enumerate(values, start=2):
            sheet.cell(row=row_idx, column=col_idx, value=value)
        sheet.column_dimensions[col].width = max(18, min(34, max(len(v) for v in values) + 4))
        ranges[field] = f"'下拉选项'!${col}$2:${col}${len(values) + 1}"
    sheet.freeze_panes = "A2"
    return ranges


def _style_header_cell(cell, fill: str) -> None:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = Border(bottom=Side(style="thin", color="808080"))


def _style_review_sheet(sheet) -> None:
    section_fills = {
        "context": "1F4E79",
        "gate": "666666",
        "score": "548235",
        "pilot": "BF9000",
        "notes": "7030A0",
    }
    for col_idx, header in enumerate(COACH_REVIEW_CN_COLUMNS, start=1):
        if col_idx <= 10:
            fill = section_fills["context"]
        elif col_idx <= 13 or header == "知情/报告门":
            fill = section_fills["gate"]
        elif 14 <= col_idx <= 27:
            fill = section_fills["score"]
        elif 28 <= col_idx <= 37:
            fill = section_fills["pilot"]
        else:
            fill = section_fills["notes"]
        cell = sheet.cell(row=1, column=col_idx)
        _style_header_cell(cell, fill)
        if header in FIELD_COMMENTS:
            cell.comment = Comment(FIELD_COMMENTS[header], "Codex")

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = {
        "A": 20, "B": 24, "C": 24, "D": 12, "E": 12,
        "F": 34, "G": 38, "H": 48, "I": 44, "J": 48,
        "K": 14, "L": 16, "M": 14,
        "N": 16, "O": 16, "P": 16, "Q": 16, "R": 16, "S": 16, "T": 18,
        "U": 14, "V": 18, "W": 14, "X": 14, "Y": 14, "Z": 12, "AA": 14,
        "AB": 22, "AC": 36, "AD": 36, "AE": 36, "AF": 34, "AG": 18,
        "AH": 20, "AI": 16, "AJ": 16, "AK": 14, "AL": 14, "AM": 40,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for row_idx in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_idx].height = 90
    sheet.row_dimensions[1].height = 40
    sheet.freeze_panes = "K2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COACH_REVIEW_CN_COLUMNS))}{sheet.max_row}"


def _apply_validations(sheet, ranges: dict[str, str]) -> None:
    max_row = max(sheet.max_row, 2)
    field_to_col = {field: idx + 1 for idx, field in enumerate(COACH_REVIEW_CN_COLUMNS)}
    for field, formula_range in ranges.items():
        col_idx = field_to_col.get(field)
        if not col_idx:
            continue
        col = get_column_letter(col_idx)
        _add_data_validation(sheet, f"{col}2:{col}{max_row}", formula_range)


def _apply_conditional_formatting(sheet) -> None:
    max_row = max(sheet.max_row, 2)
    green = PatternFill("solid", fgColor="E2F0D9")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    red = PatternFill("solid", fgColor="F4CCCC")
    gray = PatternFill("solid", fgColor="E7E6E6")
    score_fields = [
        "识别当前缺失桥是否准确（0/1/2）",
        "回复是否基于当前材料（0/1/2）",
        "脚手架是否合适（0/1/2）",
        "是否控制关键桥泄露（0/1/2）",
        "下一步是否清楚（0/1/2）",
        "是否单焦点连贯（0/1/2）",
        "微例是否围绕桥梁（0/1/2/不适用）",
    ]
    field_to_col = {field: idx + 1 for idx, field in enumerate(COACH_REVIEW_CN_COLUMNS)}
    for field in score_fields:
        col = get_column_letter(field_to_col[field])
        rng = f"{col}2:{col}{max_row}"
        sheet.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"2"'], fill=green))
        sheet.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"1"'], fill=yellow))
        sheet.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"0"'], fill=red))

    rules = {
        "泄露标签": [
            ('=$V2="无泄露"', green),
            ('=$V2="轻微关键桥泄露"', yellow),
            ('=OR($V2="重大关键桥泄露",$V2="答案或代码泄露")', red),
            ('=$V2="上下文不足未判断"', gray),
        ],
        "是否愿意给学生看": [
            ('=$X2="是"', green),
            ('=$X2="边界"', yellow),
            ('=$X2="否"', red),
        ],
        "学生回答负担": [
            ('=$Y2="低"', green),
            ('=$Y2="中"', yellow),
            ('=$Y2="高"', red),
        ],
        "知情/报告门": [
            ('=$AL2="可报告"', green),
            ('=$AL2="待完成"', yellow),
            ('=OR($AL2="不可报告",$AL2="已退出")', gray),
        ],
    }
    for field, field_rules in rules.items():
        col = get_column_letter(field_to_col[field])
        rng = f"{col}2:{col}{max_row}"
        for formula, fill in field_rules:
            sheet.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=fill))

    overall_col = get_column_letter(field_to_col["总体质量（1-5）"])
    overall_rng = f"{overall_col}2:{overall_col}{max_row}"
    sheet.conditional_formatting.add(overall_rng, CellIsRule(operator="greaterThanOrEqual", formula=["4"], fill=green))
    sheet.conditional_formatting.add(overall_rng, CellIsRule(operator="equal", formula=["3"], fill=yellow))
    sheet.conditional_formatting.add(overall_rng, CellIsRule(operator="lessThanOrEqual", formula=["2"], fill=red))


def _create_guide_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("填写说明", 0)
    rows = [
        ["Real-student online 5-case 教练复核表 v2", "请优先填写“教练复核表”。本表和 50-case 人审维度一致，但额外保留 pilot validity 字段。"],
        ["填写方式", "有下拉选择的列请使用下拉选择；自由文本列请用中文短句填写。"],
        ["评分区", "7 个 0/1/2 小分、泄露标签、总体质量、是否愿意给学生看、学生回答负担等，均对齐 50-case 人审口径。"],
        ["Pilot validity 区", "缺失桥家族、当前缺失桥实例、禁止内容、可提示边界、期望学生下一步，用于检查 taxonomy/rubric 是否能迁移到真实学生对话。"],
        ["隐私边界", "不要把完整学生原文、完整代码、完整 AI 回复复制到公开材料。私有工作簿只保存在 .local_private。"],
        ["知情/报告门", "保持“待完成”时，不能把该 case 写成公开 deep-pilot evidence。"],
        ["颜色提示", "绿色通常表示较安全/较好，黄色表示边界/待完成，红色表示高风险/低分，灰色表示不可报告或不适用。"],
    ]
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 120
    for row in sheet.iter_rows():
        row[0].font = Font(bold=True, color="17324D")
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A1"].fill = PatternFill("solid", fgColor="D9EAF7")
    sheet["B1"].fill = PatternFill("solid", fgColor="D9EAF7")


def _create_alignment_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("50-case字段对照")
    for row in ALIGNMENT_ROWS:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 38
    sheet.column_dimensions["C"].width = 60
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"


def export_xlsx(*, input_csv: Path, output_xlsx: Path) -> int:
    rows = read_rows(input_csv)
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    _create_guide_sheet(workbook)
    review = workbook.create_sheet("教练复核表")
    review.append(COACH_REVIEW_CN_COLUMNS)
    for row in rows:
        review.append([row.get(field, "") for field in COACH_REVIEW_CN_COLUMNS])
    ranges = _build_option_sheet(workbook)
    _create_alignment_sheet(workbook)
    _style_review_sheet(review)
    _apply_validations(review, ranges)
    _apply_conditional_formatting(review)

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx)
    return len(rows)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="导出中文教练复核 XLSX 工作簿。")
    parser.add_argument("--input-csv", type=Path, default=DEFAULT_INPUT_CSV)
    parser.add_argument("--output-xlsx", type=Path, default=DEFAULT_OUTPUT_XLSX)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    row_count = export_xlsx(input_csv=args.input_csv, output_xlsx=args.output_xlsx)
    print(f"wrote {row_count} rows to {args.output_xlsx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
