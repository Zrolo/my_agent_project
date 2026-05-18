"""Build a blinded same-candidate Repair before/after review workbook.

The workbook randomizes the original candidate and repaired response into
Response A / Response B. The mapping is written only to a separate key CSV.

This is an offline research artifact builder. It does not call models, change
prompts, modify online AIChat behavior, or alter main experiment labels.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


DEFAULT_CANDIDATES_CSV = Path(
    "evals/aichat/ad_hoc_runs/dialogue_state_v3_main_20260516_merged/human_review/"
    "repair_same_candidate_stress_candidate_list_20260517.csv"
)
DEFAULT_REFERENCE_LABELS_JSONL = Path(
    "evals/aichat/ad_hoc_runs/dialogue_state_v3_main_20260516_merged/human_review/"
    "dialogue_state_v3_main_priority60_adjudicated_plus_coachA_labels_20260517.jsonl"
)
DEFAULT_OUTPUT_XLSX = Path(
    "evals/aichat/ad_hoc_runs/dialogue_state_v3_main_20260516_merged/human_review/"
    "repair_same_candidate_stress_blind_review_workbook_20260517.zh.xlsx"
)
DEFAULT_KEY_CSV = Path(
    "evals/aichat/ad_hoc_runs/dialogue_state_v3_main_20260516_merged/human_review/"
    "repair_same_candidate_stress_blind_review_key_20260517.csv"
)

REVIEW_HEADERS = [
    "stress_id",
    "case_id",
    "slice",
    "bridge_bucket",
    "problem_source_url",
    "problem_statement_or_summary",
    "student_message",
    "recent_dialogue",
    "context_ai_reply",
    "success_criteria",
    "forbidden_content",
    "critical_bridge_boundary",
    "acceptable_reveal",
    "expected_student_next_action",
    "response_A_text",
    "A_overall_quality_score",
    "A_would_show_to_student",
    "A_leakage_label",
    "A_student_response_burden",
    "A_too_vague",
    "response_B_text",
    "B_overall_quality_score",
    "B_would_show_to_student",
    "B_leakage_label",
    "B_student_response_burden",
    "B_too_vague",
    "pair_preference",
    "reviewer_confidence",
    "review_notes",
]

KEY_HEADERS = [
    "stress_id",
    "case_id",
    "slice",
    "bridge_bucket",
    "source_condition",
    "source_run",
    "response_A_source",
    "response_B_source",
    "response_A_sha256",
    "response_B_sha256",
]


def _read_csv(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _read_jsonl(path: Path) -> list[dict]:
    rows = []
    for line_number, raw_line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        line = raw_line.strip()
        if not line:
            continue
        try:
            rows.append(json.loads(line))
        except json.JSONDecodeError as exc:
            raise ValueError(f"Invalid JSONL at {path}:{line_number}: {exc}") from exc
    return rows


def _case_reference_map(labels_path: Path) -> dict[str, dict]:
    by_case: dict[str, dict] = {}
    for row in _read_jsonl(labels_path):
        case_id = str(row.get("case_id") or "")
        if case_id and case_id not in by_case:
            by_case[case_id] = row
    return by_case


def _hash_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]


def _shorten(text: object, limit: int = 3500) -> str:
    value = str(text or "").strip()
    if len(value) <= limit:
        return value
    return value[: limit - 24].rstrip() + "\n...[truncated for review]"


def _make_validation(values: list[str]) -> DataValidation:
    quoted = ",".join(values)
    return DataValidation(type="list", formula1=f'"{quoted}"', allow_blank=True)


def build_rows(
    candidate_rows: list[dict],
    *,
    references: dict[str, dict],
    seed: int,
) -> tuple[list[dict], list[dict]]:
    rng = random.Random(seed)
    review_rows: list[dict] = []
    key_rows: list[dict] = []

    for row in candidate_rows:
        before_text = str(row.get("original_candidate") or "").strip()
        after_text = str(row.get("repair_output") or "").strip()
        if not before_text or not after_text:
            continue

        before_first = rng.random() < 0.5
        response_a = before_text if before_first else after_text
        response_b = after_text if before_first else before_text
        source_a = "before_repair" if before_first else "after_repair"
        source_b = "after_repair" if before_first else "before_repair"

        case_id = str(row.get("case_id") or "")
        ref = references.get(case_id, {})
        problem_text = ref.get("problem_statement") or ref.get("problem_context") or ""

        review_item = {
            "stress_id": row.get("stress_id") or "",
            "case_id": case_id,
            "slice": row.get("slice") or "",
            "bridge_bucket": row.get("bridge_bucket") or "",
            "problem_source_url": ref.get("problem_source_url") or "",
            "problem_statement_or_summary": _shorten(problem_text),
            "student_message": row.get("student_message") or ref.get("student_message") or "",
            "recent_dialogue": _shorten(ref.get("recent_dialogue"), 1500),
            "context_ai_reply": _shorten(ref.get("context_ai_reply"), 1500),
            "success_criteria": ref.get("success_criteria") or "",
            "forbidden_content": ref.get("forbidden_content") or "",
            "critical_bridge_boundary": ref.get("critical_bridge_boundary") or "",
            "acceptable_reveal": ref.get("acceptable_reveal") or "",
            "expected_student_next_action": ref.get("expected_student_next_action") or "",
            "response_A_text": response_a,
            "A_overall_quality_score": "",
            "A_would_show_to_student": "",
            "A_leakage_label": "",
            "A_student_response_burden": "",
            "A_too_vague": "",
            "response_B_text": response_b,
            "B_overall_quality_score": "",
            "B_would_show_to_student": "",
            "B_leakage_label": "",
            "B_student_response_burden": "",
            "B_too_vague": "",
            "pair_preference": "",
            "reviewer_confidence": "",
            "review_notes": "",
        }
        review_rows.append(review_item)

        key_rows.append(
            {
                "stress_id": row.get("stress_id") or "",
                "case_id": case_id,
                "slice": row.get("slice") or "",
                "bridge_bucket": row.get("bridge_bucket") or "",
                "source_condition": row.get("source_condition") or "",
                "source_run": row.get("source_run") or "",
                "response_A_source": source_a,
                "response_B_source": source_b,
                "response_A_sha256": _hash_text(response_a),
                "response_B_sha256": _hash_text(response_b),
            }
        )
    return review_rows, key_rows


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = {
        "A": 18,
        "B": 34,
        "C": 22,
        "D": 30,
        "E": 34,
        "F": 58,
        "G": 36,
        "H": 42,
        "I": 42,
        "J": 42,
        "K": 42,
        "L": 48,
        "M": 42,
        "N": 42,
        "O": 70,
        "U": 70,
        "AC": 48,
    }
    for col_idx in range(1, len(REVIEW_HEADERS) + 1):
        col = get_column_letter(col_idx)
        sheet.column_dimensions[col].width = widths.get(col, 18)
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(REVIEW_HEADERS))}{sheet.max_row}"


def _add_validations(sheet, n_rows: int) -> None:
    if n_rows <= 1:
        return
    validation_by_header = {
        "A_overall_quality_score": ["1", "2", "3", "4", "5"],
        "B_overall_quality_score": ["1", "2", "3", "4", "5"],
        "A_would_show_to_student": ["yes", "borderline", "no"],
        "B_would_show_to_student": ["yes", "borderline", "no"],
        "A_leakage_label": [
            "no_leakage",
            "minor_bridge_leakage",
            "major_bridge_leakage",
            "answer_leakage",
        ],
        "B_leakage_label": [
            "no_leakage",
            "minor_bridge_leakage",
            "major_bridge_leakage",
            "answer_leakage",
        ],
        "A_student_response_burden": ["low", "medium", "high"],
        "B_student_response_burden": ["low", "medium", "high"],
        "A_too_vague": ["yes", "no", "unclear"],
        "B_too_vague": ["yes", "no", "unclear"],
        "pair_preference": ["A", "B", "tie", "unclear"],
        "reviewer_confidence": ["high", "medium", "low"],
    }
    for header, values in validation_by_header.items():
        col_idx = REVIEW_HEADERS.index(header) + 1
        col = get_column_letter(col_idx)
        dv = _make_validation(values)
        sheet.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n_rows}")


def write_workbook(rows: list[dict], path: Path) -> None:
    workbook = Workbook()
    info = workbook.active
    info.title = "说明"
    info_rows = [
        ["Repair same-candidate 盲评说明"],
        ["目标", "比较同一 candidate 的 before_repair 与 after_repair，但评审者只看到 Response A / Response B。"],
        ["请勿推断", "不要根据措辞猜测哪条是 Repair；逐条按 case-specific rubric 独立评分。"],
        ["评分字段", "overall 1-5；would_show yes/borderline/no；leakage label；student burden；too_vague；pair preference。"],
        ["关键口径", "本实验用于估计 Repair same-candidate effect；主实验 condition 均值不能替代该因果测试。"],
    ]
    for item in info_rows:
        info.append(item)
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 110
    for row in info.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    info["A1"].font = Font(bold=True, size=14)

    sheet = workbook.create_sheet("盲评表")
    sheet.append(REVIEW_HEADERS)
    for row in rows:
        sheet.append([row.get(header, "") for header in REVIEW_HEADERS])
    _style_sheet(sheet)
    _add_validations(sheet, sheet.max_row)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_key(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=KEY_HEADERS)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: row.get(header, "") for header in KEY_HEADERS})


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--candidates-csv", type=Path, default=DEFAULT_CANDIDATES_CSV)
    parser.add_argument("--reference-labels-jsonl", type=Path, default=DEFAULT_REFERENCE_LABELS_JSONL)
    parser.add_argument("--output-xlsx", type=Path, default=DEFAULT_OUTPUT_XLSX)
    parser.add_argument("--key-csv", type=Path, default=DEFAULT_KEY_CSV)
    parser.add_argument("--seed", type=int, default=20260517)
    args = parser.parse_args(argv)

    candidate_rows = _read_csv(args.candidates_csv)
    references = _case_reference_map(args.reference_labels_jsonl)
    review_rows, key_rows = build_rows(candidate_rows, references=references, seed=args.seed)
    write_workbook(review_rows, args.output_xlsx)
    write_key(key_rows, args.key_csv)
    print(f"Wrote blind review workbook with {len(review_rows)} pairs: {args.output_xlsx}")
    print(f"Wrote blind key: {args.key_csv}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
