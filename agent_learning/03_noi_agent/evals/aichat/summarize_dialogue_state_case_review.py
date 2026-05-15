"""Summarize filled dialogue-state v3 case/source review workbooks."""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))


DEFAULT_INPUT_XLSX = Path("docs/research/dialogue_state_v3_50_source_and_case_review.zh.xlsx")
DEFAULT_OUTPUT_JSON = Path("docs/research/dialogue_state_v3_case_review_summary_20260515.json")
DEFAULT_OUTPUT_ZH = Path("docs/research/dialogue_state_v3_case_review_summary_20260515.zh.md")
DEFAULT_OUTPUT_EN = Path("docs/research/dialogue_state_v3_case_review_summary_20260515.md")

STRUCTURED_FIELDS = [
    "source_ok",
    "context_coherent",
    "student_message_realistic",
    "followability_ok",
    "missing_bridge_ok",
    "forbidden_content_ok",
    "success_criteria_ok",
    "leakage_boundary_ok",
    "case_decision",
    "issue_type",
    "reviewer_confidence",
]

CHOICE_NORMALIZATION = {
    "是": "yes",
    "部分": "partial",
    "否": "no",
    "需修改": "revise",
    "不适用": "not_applicable",
    "过严": "too_strict",
    "过松": "too_loose",
    "不清楚": "unclear",
    "接受": "accept",
    "修改": "revise",
    "丢弃": "drop",
    "讨论": "discuss",
    "无": "none",
    "题源问题": "source_issue",
    "上下文不一致": "context_mismatch",
    "学生话术不像真实学生": "student_language_artificial",
    "跟随状态问题": "followability_issue",
    "桥梁标签问题": "bridge_label_issue",
    "禁止内容问题": "forbidden_content_issue",
    "成功标准问题": "success_criteria_issue",
    "泄露边界问题": "leakage_boundary_issue",
    "其他": "other",
    "高": "high",
    "中": "medium",
    "低": "low",
}


def _norm(value) -> str:
    text = str(value or "").strip()
    if not text:
        return "blank"
    return CHOICE_NORMALIZATION.get(text, text)


def _sheet_for_workbook(workbook, requested: str | None):
    if requested:
        return workbook[requested]
    if "总表" in workbook.sheetnames:
        return workbook["总表"]
    if "dialogue_state_review" in workbook.sheetnames:
        return workbook["dialogue_state_review"]
    return workbook[workbook.sheetnames[0]]


def summarize_workbook(path: Path, *, sheet_name: str | None = None) -> dict:
    workbook = load_workbook(path)
    sheet = _sheet_for_workbook(workbook, sheet_name)
    machine_headers = [sheet.cell(row=2, column=col).value for col in range(1, sheet.max_column + 1)]
    header_to_col = {header: index + 1 for index, header in enumerate(machine_headers) if header}
    missing_fields = [field for field in STRUCTURED_FIELDS if field not in header_to_col]
    if missing_fields:
        raise ValueError(f"Workbook missing structured review fields: {missing_fields}")

    rows = []
    for row_index in range(3, sheet.max_row + 1):
        case_id = sheet.cell(row=row_index, column=header_to_col.get("case_id", 1)).value
        if not str(case_id or "").strip():
            continue
        rows.append(
            {
                field: _norm(sheet.cell(row=row_index, column=header_to_col[field]).value)
                for field in STRUCTURED_FIELDS
            }
        )

    summary: dict[str, object] = {
        "input_xlsx": str(path),
        "sheet_name": sheet.title,
        "row_count": len(rows),
        "reviewed_count": sum(1 for row in rows if row["case_decision"] != "blank"),
        "case_decision_counts": dict(Counter(row["case_decision"] for row in rows)),
        "issue_type_counts": dict(Counter(row["issue_type"] for row in rows)),
        "reviewer_confidence_counts": dict(Counter(row["reviewer_confidence"] for row in rows)),
        "needs_followup_count": sum(1 for row in rows if row["case_decision"] in {"revise", "drop", "discuss"}),
    }
    for field in STRUCTURED_FIELDS:
        if field in {"case_decision", "issue_type", "reviewer_confidence"}:
            continue
        summary[f"{field}_counts"] = dict(Counter(row[field] for row in rows))
    return summary


def _write_markdown(summary: dict, path: Path, *, language: str) -> None:
    is_zh = language == "zh"
    title = "Dialogue-State v3 Case/Source Review Summary" if not is_zh else "Dialogue-State v3 Case/Source 审核汇总"
    lines = [
        f"# {title}",
        "",
        (
            "本报告汇总教练在 case/source 审核表中填写的结构化字段；它不评价 AI 回复质量。"
            if is_zh
            else "This report summarizes structured fields from the coach case/source review workbook. It does not evaluate AI response quality."
        ),
        "",
        f"- input_xlsx: `{summary['input_xlsx']}`",
        f"- sheet_name: `{summary['sheet_name']}`",
        f"- row_count: {summary['row_count']}",
        f"- reviewed_count: {summary['reviewed_count']}",
        f"- needs_followup_count: {summary['needs_followup_count']}",
        "",
        "## Case Decisions" if not is_zh else "## 样本处理决定",
        "",
        f"`{summary['case_decision_counts']}`",
        "",
        "## Issue Types" if not is_zh else "## 问题类型",
        "",
        f"`{summary['issue_type_counts']}`",
        "",
        "## Reviewer Confidence" if not is_zh else "## 审核置信度",
        "",
        f"`{summary['reviewer_confidence_counts']}`",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Summarize dialogue-state v3 case/source review workbook.")
    parser.add_argument("--input-xlsx", type=Path, default=DEFAULT_INPUT_XLSX)
    parser.add_argument("--sheet-name", default=None)
    parser.add_argument("--output-json", type=Path, default=DEFAULT_OUTPUT_JSON)
    parser.add_argument("--output-zh", type=Path, default=DEFAULT_OUTPUT_ZH)
    parser.add_argument("--output-en", type=Path, default=DEFAULT_OUTPUT_EN)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(sys.argv[1:] if argv is None else argv)
    summary = summarize_workbook(args.input_xlsx, sheet_name=args.sheet_name)
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    _write_markdown(summary, args.output_zh, language="zh")
    _write_markdown(summary, args.output_en, language="en")
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
