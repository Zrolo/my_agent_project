import argparse
import json
import sys
from pathlib import Path

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - import-time dependency guard
    raise RuntimeError("openpyxl is required to read real-log review workbooks") from exc


DEFAULT_SHEET_NAMES = ("student_only候选_v3", "候选样本_v2")
DEFAULT_ACCEPTED_STATUS = "保留主评测"

HEADER_MAP = {
    "审查状态": "review_status",
    "可公开展示": "public_display_ok",
    "隐私风险": "privacy_risk",
    "旧AI上下文污染确认": "legacy_ai_context_contamination",
    "旧AI上下文污染初筛": "legacy_ai_context_contamination",
    "是否自包含": "self_contained",
    "是否需要补充题面/代码上下文": "needs_more_context",
    "是否近重复样本": "near_duplicate",
    "是否代表真实卡点": "real_bottleneck",
    "是否适合主评测": "suitable_for_main_eval",
    "删除旧AI后是否可理解": "student_only_eval_usable",
    "建议split": "recommended_split",
    "需重点复判": "needs_priority_recheck",
    "候选ID": "case_id",
    "类别": "category",
    "题号/URL": "problem_ref",
    "题目标题": "problem_title",
    "创建时间": "created_at",
    "含题面": "has_problem_context",
    "含代码标记": "has_student_code",
    "会话消息数": "session_message_count",
    "完整近期对话（目标turn前，供污染/自包含判断）": "recent_dialogue",
    "完整近期对话（只供审计，不给模型）": "recent_dialogue",
    "目标turn前最近AI回复（上下文，不评分）": "context_ai_reply",
    "目标turn前旧AI上下文（只供审计，不给模型）": "context_ai_reply",
    "学生当前问题": "student_message",
    "学生当前问题（主消融输入）": "student_message",
    "目标turn后旧系统回复（baseline观察，非上下文，非gold）": "observed_current_system_response",
    "missing_bridge（教练填）": "missing_bridge",
    "forbidden_content（教练填）": "forbidden_content",
    "success_criteria（教练填）": "success_criteria",
    "教练备注": "coach_notes",
}

REQUIRED_COACH_FIELDS = ("missing_bridge", "forbidden_content", "success_criteria")
YES_VALUES = {"true", "1", "yes", "y", "是"}
NO_VALUES = {"false", "0", "no", "n", "否"}
ALLOWED_LEGACY_CONTEXT_CONTAMINATION = {"none", "low"}


def _clean(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_bool(value: object) -> bool | None:
    text = _clean(value).lower()
    if text in YES_VALUES:
        return True
    if text in NO_VALUES:
        return False
    return None


def _parse_int(value: object) -> int:
    try:
        return int(float(_clean(value)))
    except (TypeError, ValueError):
        return 0


def _choice(value: object) -> str:
    return _clean(value).lower()


def _is_yes(value: object) -> bool:
    return _choice(value) in YES_VALUES


def _is_no(value: object) -> bool:
    return _choice(value) in NO_VALUES


def _select_sheet_name(workbook, requested_sheet_name: str | None) -> str:
    if requested_sheet_name:
        if requested_sheet_name not in workbook.sheetnames:
            raise ValueError(f"sheet not found: {requested_sheet_name}")
        return requested_sheet_name
    for candidate in DEFAULT_SHEET_NAMES:
        if candidate in workbook.sheetnames:
            return candidate
    raise ValueError(
        "sheet not found: expected one of "
        + ", ".join(DEFAULT_SHEET_NAMES)
    )


def load_review_rows(path: Path, sheet_name: str | None = None) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[_select_sheet_name(workbook, sheet_name)]
    rows_iter = sheet.iter_rows(values_only=True)
    headers = [_clean(cell) for cell in next(rows_iter)]
    normalized_headers = [HEADER_MAP.get(header, header) for header in headers]
    rows: list[dict] = []
    for raw_values in rows_iter:
        row = {
            normalized_headers[index]: _clean(value)
            for index, value in enumerate(raw_values)
            if index < len(normalized_headers)
        }
        if any(row.values()):
            rows.append(row)
    return rows


def _validation_reasons(row: dict) -> list[str]:
    reasons: list[str] = []
    missing = [field for field in REQUIRED_COACH_FIELDS if not _clean(row.get(field))]
    if missing:
        reasons.append(f"missing required coach fields: {', '.join(missing)}")
    if _choice(row.get("privacy_risk")) != "none":
        reasons.append("privacy_risk must be none")
    if _choice(row.get("recommended_split")) in {"", "undecided"}:
        reasons.append("recommended_split must be decided before export")
    if not _is_yes(row.get("student_only_eval_usable")):
        reasons.append("student_only_eval_usable must be yes before export")
    if not _is_yes(row.get("self_contained")):
        reasons.append("self_contained must be yes")
    if not _is_no(row.get("needs_more_context")):
        reasons.append("needs_more_context must be no")
    if not _is_no(row.get("near_duplicate")):
        reasons.append("near_duplicate must be no")
    if not _is_yes(row.get("real_bottleneck")):
        reasons.append("real_bottleneck must be yes")
    if not _is_yes(row.get("suitable_for_main_eval")):
        reasons.append("suitable_for_main_eval must be yes")
    if not _is_no(row.get("needs_priority_recheck")):
        reasons.append("needs_priority_recheck must be no")
    if _choice(row.get("legacy_ai_context_contamination")) not in ALLOWED_LEGACY_CONTEXT_CONTAMINATION:
        reasons.append("legacy_ai_context_contamination must be none or low")
    if not _clean(row.get("student_message")):
        reasons.append("student_message must be present")
    if not (_clean(row.get("problem_ref")) or _clean(row.get("problem_title"))):
        reasons.append("problem_ref or problem_title must be present")
    if not _clean(row.get("category")):
        reasons.append("category must be present")
    return reasons


def _validate_main_row(row: dict) -> None:
    reasons = _validation_reasons(row)
    if reasons:
        case_id = row.get("case_id") or "<unknown>"
        raise ValueError(f"{case_id}: " + "; ".join(reasons))


def build_candidate_row(row: dict, *, require_complete: bool = True) -> dict:
    if require_complete:
        _validate_main_row(row)
    return {
        "schema_version": "real_aichat_candidate_v1",
        "case_id": row.get("case_id", ""),
        "source": "online_real_aichat_log",
        "research_use_status": "coach_reviewed_real_log_candidate",
        "recommended_split": row.get("recommended_split", ""),
        "category": row.get("category", ""),
        "problem_ref": row.get("problem_ref", ""),
        "problem_title": row.get("problem_title", ""),
        "created_at": row.get("created_at", ""),
        "has_problem_context": _parse_bool(row.get("has_problem_context")),
        "has_student_code": _parse_bool(row.get("has_student_code")),
        "session_message_count": _parse_int(row.get("session_message_count")),
        "legacy_ai_context_contamination": row.get("legacy_ai_context_contamination", ""),
        "legacy_context_removed": True,
        "legacy_context_policy": "old_ai_before_target_turn_removed_from_generation_input_audit_only",
        "recent_dialogue": "",
        "context_ai_reply": "",
        "student_message": row.get("student_message", ""),
        "generation_input": {
            "student_message": row.get("student_message", ""),
            "problem_ref": row.get("problem_ref", ""),
            "problem_title": row.get("problem_title", ""),
            "recent_dialogue": "",
            "context_ai_reply": "",
            "student_code_excerpt": "",
            "legacy_context_removed": True,
        },
        "audit_context": {
            "recent_dialogue_before_target_turn": row.get("recent_dialogue", ""),
            "context_ai_reply_before_target_turn": row.get("context_ai_reply", ""),
        },
        "observed_current_system_response": row.get("observed_current_system_response", ""),
        "old_assistant_reply_use": "observed_current_system_response_only_not_gold",
        "coach_label": {
            "missing_bridge": row.get("missing_bridge", ""),
            "forbidden_content": row.get("forbidden_content", ""),
            "success_criteria": row.get("success_criteria", ""),
            "notes": row.get("coach_notes", ""),
        },
        "review_flags": {
            "review_status": row.get("review_status", ""),
            "public_display_ok": row.get("public_display_ok", ""),
            "privacy_risk": row.get("privacy_risk", ""),
            "self_contained": row.get("self_contained", ""),
            "needs_more_context": row.get("needs_more_context", ""),
            "near_duplicate": row.get("near_duplicate", ""),
            "real_bottleneck": row.get("real_bottleneck", ""),
            "suitable_for_main_eval": row.get("suitable_for_main_eval", ""),
            "student_only_eval_usable": row.get("student_only_eval_usable", ""),
            "needs_priority_recheck": row.get("needs_priority_recheck", ""),
        },
    }


def export_review_workbook_to_jsonl(
    input_path: Path,
    output_path: Path,
    *,
    accepted_status: str = DEFAULT_ACCEPTED_STATUS,
    require_complete: bool = True,
    sheet_name: str | None = None,
) -> int:
    rows = load_review_rows(input_path, sheet_name=sheet_name)
    accepted_rows = [row for row in rows if row.get("review_status") == accepted_status]
    candidate_rows = [
        build_candidate_row(row, require_complete=require_complete)
        for row in accepted_rows
    ]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        "\n".join(json.dumps(row, ensure_ascii=False) for row in candidate_rows)
        + ("\n" if candidate_rows else ""),
        encoding="utf-8",
    )
    return len(candidate_rows)


def screen_review_workbook(
    input_path: Path,
    *,
    accepted_status: str = DEFAULT_ACCEPTED_STATUS,
    sheet_name: str | None = None,
) -> dict:
    rows = load_review_rows(input_path, sheet_name=sheet_name)
    accepted_rows = [row for row in rows if row.get("review_status") == accepted_status]
    screened_rows = []
    rejected_by_reason: dict[str, int] = {}
    for row in accepted_rows:
        reasons = _validation_reasons(row)
        for reason in reasons:
            rejected_by_reason[reason] = rejected_by_reason.get(reason, 0) + 1
        screened_rows.append(
            {
                "case_id": row.get("case_id", ""),
                "review_status": row.get("review_status", ""),
                "recommended_split": row.get("recommended_split", ""),
                "category": row.get("category", ""),
                "exportable": not reasons,
                "reasons": reasons,
            }
        )
    return {
        "input_path": str(input_path),
        "total_rows": len(rows),
        "accepted_status": accepted_status,
        "accepted_rows": len(accepted_rows),
        "exportable_rows": sum(1 for row in screened_rows if row["exportable"]),
        "rejected_rows": sum(1 for row in screened_rows if not row["exportable"]),
        "rejected_by_reason": rejected_by_reason,
        "rows": screened_rows,
    }


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export coach-reviewed real AIChat log candidates to JSONL.")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output-jsonl", type=Path, required=True)
    parser.add_argument(
        "--sheet-name",
        default=None,
        help="Optional worksheet name. Defaults to student-only v3, then v2 review sheet.",
    )
    parser.add_argument("--accepted-status", default=DEFAULT_ACCEPTED_STATUS)
    parser.add_argument(
        "--allow-incomplete",
        action="store_true",
        help="Write accepted rows even if coach fields, split, or privacy review are incomplete.",
    )
    parser.add_argument(
        "--screen-report-json",
        type=Path,
        default=None,
        help="Optional path for a strict screening report with per-row rejection reasons.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(sys.argv[1:] if argv is None else argv)
    row_count = export_review_workbook_to_jsonl(
        args.input,
        args.output_jsonl,
        accepted_status=args.accepted_status,
        require_complete=not args.allow_incomplete,
        sheet_name=args.sheet_name,
    )
    if args.screen_report_json:
        report = screen_review_workbook(
            args.input,
            accepted_status=args.accepted_status,
            sheet_name=args.sheet_name,
        )
        args.screen_report_json.parent.mkdir(parents=True, exist_ok=True)
        args.screen_report_json.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    print(json.dumps({"output_jsonl": str(args.output_jsonl), "row_count": row_count}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
