import argparse
import csv
import json
import sys
from pathlib import Path
from typing import TextIO

from openpyxl import load_workbook

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from evals.aichat.coach_labeling_schema_v2 import (
    FOCUS_FIELDS,
    MULTI_VALUE_FIELDS,
    OPTIONS_BY_FIELD,
    REVIEW_STATUSES,
    extract_option_id,
    split_option_ids,
)


DEFAULT_WORKBOOK_PATH = Path("docs/research/coach_seed_labeling_workbook_v2.zh.xlsx")
DEFAULT_FOCUS_REGISTRY_PATH = Path("docs/research/focus_registry_v1.json")
DEFAULT_SUBTYPE_REGISTRY_PATH = Path("docs/research/bridge_subtype_registry_v2.json")
PLACEHOLDER_VALUE = "1"
OPTIONAL_FOCUS_VALUES = {"unknown", "not_applicable", ""}
SUBTYPE_UNKNOWN_VALUES = {"", "unknown"}
TURN_TYPE_EXPECTED_RISK = {
    "complete_solution_request": "complete_answer_risk",
    "complete_code_request": "complete_code_risk",
    "critical_bridge_request": "critical_bridge_completion_risk",
    "algorithm_confirmation_request": "algorithm_confirmation_risk",
    "local_completion_request": "local_code_completion_risk",
}


def load_focus_ids(path: Path = DEFAULT_FOCUS_REGISTRY_PATH) -> set[str]:
    if not path.exists():
        return set()
    data = json.loads(path.read_text(encoding="utf-8"))
    focuses = data.get("focuses", data) if isinstance(data, dict) else data
    if not isinstance(focuses, list):
        raise ValueError(f"Focus registry must be a list or contain focuses: {path}")
    focus_ids = set()
    for item in focuses:
        if isinstance(item, str):
            focus_ids.add(item)
        elif isinstance(item, dict) and item.get("focus_id"):
            focus_ids.add(str(item["focus_id"]))
    return focus_ids


def load_subtype_registry(path: Path = DEFAULT_SUBTYPE_REGISTRY_PATH) -> dict[str, str]:
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    subtypes = data.get("subtypes", data) if isinstance(data, dict) else data
    if not isinstance(subtypes, list):
        raise ValueError(f"Subtype registry must be a list or contain subtypes: {path}")
    registry: dict[str, str] = {}
    for item in subtypes:
        if isinstance(item, dict) and item.get("subtype_id") and item.get("family"):
            registry[str(item["subtype_id"])] = str(item["family"])
    return registry


def load_xlsx_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=True)
    if "标注表" not in workbook.sheetnames:
        raise ValueError(f"Workbook must contain 标注表 sheet: {path}")
    sheet = workbook["标注表"]
    headers = [cell.value for cell in sheet[2]]
    rows: list[dict] = []
    for row_cells in sheet.iter_rows(min_row=3, max_row=sheet.max_row, values_only=True):
        if not any(value not in (None, "") for value in row_cells):
            continue
        rows.append({str(header): value for header, value in zip(headers, row_cells) if header})
    return rows


def load_csv_rows(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def _add_error(errors: list[dict], row_number: int, case_id: str, field: str, value: object, message: str) -> None:
    errors.append(
        {
            "row_number": row_number,
            "case_id": case_id,
            "field": field,
            "value": "" if value is None else str(value),
            "message": message,
        }
    )


def _is_labeled(row: dict) -> bool:
    return extract_option_id(row.get("review_status")) == "labeled"


def _validate_single_enum(
    errors: list[dict],
    *,
    row_number: int,
    case_id: str,
    row: dict,
    field: str,
    allowed: set[str],
) -> None:
    raw_value = row.get(field)
    value = extract_option_id(raw_value)
    if value and value not in allowed:
        _add_error(errors, row_number, case_id, field, raw_value, f"Must be one of {sorted(allowed)}.")


def _validate_multi_enum(
    errors: list[dict],
    *,
    row_number: int,
    case_id: str,
    row: dict,
    field: str,
    allowed: set[str],
) -> None:
    raw_value = row.get(field)
    for value in split_option_ids(raw_value):
        if value not in allowed:
            _add_error(errors, row_number, case_id, field, raw_value, f"Each value must be one of {sorted(allowed)}.")


def _validate_focus(
    errors: list[dict],
    *,
    row_number: int,
    case_id: str,
    row: dict,
    field: str,
    focus_ids: set[str],
) -> None:
    raw_value = row.get(field)
    value = extract_option_id(raw_value)
    if value not in OPTIONAL_FOCUS_VALUES and value not in focus_ids:
        _add_error(
            errors,
            row_number,
            case_id,
            field,
            raw_value,
            "Must be unknown, not_applicable, or a focus_id from focus_registry_v1.json.",
        )


def _validate_subtype_family(
    errors: list[dict],
    *,
    row_number: int,
    case_id: str,
    row: dict,
    family_field: str,
    subtype_field: str,
    subtype_registry: dict[str, str],
) -> None:
    family = extract_option_id(row.get(family_field))
    subtype = extract_option_id(row.get(subtype_field))
    if subtype in SUBTYPE_UNKNOWN_VALUES or not family or not subtype_registry:
        return
    expected_family = subtype_registry.get(subtype)
    if not expected_family:
        _add_error(
            errors,
            row_number,
            case_id,
            subtype_field,
            row.get(subtype_field),
            "Subtype is not registered in bridge_subtype_registry_v2.json.",
        )
        return
    if family != expected_family:
        _add_error(
            errors,
            row_number,
            case_id,
            subtype_field,
            row.get(subtype_field),
            f"Subtype belongs to {expected_family}, but {family_field} is {family}.",
        )


def _validate_no_placeholder_ones(errors: list[dict], *, row_number: int, case_id: str, row: dict) -> None:
    skip_fields = {"coach_confidence"}
    for field, raw_value in row.items():
        if field in skip_fields:
            continue
        if str(raw_value).strip() == PLACEHOLDER_VALUE:
            _add_error(errors, row_number, case_id, field, raw_value, "Do not use placeholder value 1; use N/A or leave blank.")


def validate_rows(
    rows: list[dict],
    *,
    focus_ids: set[str],
    subtype_registry: dict[str, str] | None = None,
) -> list[dict]:
    if subtype_registry is None:
        subtype_registry = load_subtype_registry()
    errors: list[dict] = []
    for index, row in enumerate(rows, 3):
        case_id = str(row.get("case_id") or "")
        for field, options in OPTIONS_BY_FIELD.items():
            if field in FOCUS_FIELDS:
                continue
            if field in MULTI_VALUE_FIELDS:
                _validate_multi_enum(
                    errors,
                    row_number=index,
                    case_id=case_id,
                    row=row,
                    field=field,
                    allowed=set(options),
                )
            else:
                _validate_single_enum(
                    errors,
                    row_number=index,
                    case_id=case_id,
                    row=row,
                    field=field,
                    allowed=set(options),
                )
        for field in FOCUS_FIELDS:
            _validate_focus(errors, row_number=index, case_id=case_id, row=row, field=field, focus_ids=focus_ids)
        _validate_subtype_family(
            errors,
            row_number=index,
            case_id=case_id,
            row=row,
            family_field="primary_bridge_family",
            subtype_field="primary_bridge_subtype_id",
            subtype_registry=subtype_registry,
        )
        _validate_subtype_family(
            errors,
            row_number=index,
            case_id=case_id,
            row=row,
            family_field="secondary_bridge_family",
            subtype_field="secondary_bridge_subtype_id",
            subtype_registry=subtype_registry,
        )

        if _is_labeled(row):
            _validate_no_placeholder_ones(errors, row_number=index, case_id=case_id, row=row)
            if not str(row.get("evidence_quote") or "").strip():
                _add_error(
                    errors,
                    index,
                    case_id,
                    "evidence_quote",
                    row.get("evidence_quote"),
                    "Labeled rows must include the quoted evidence used for the label.",
                )
            review_status = extract_option_id(row.get("review_status"))
            if review_status not in REVIEW_STATUSES:
                _add_error(errors, index, case_id, "review_status", row.get("review_status"), "Invalid review status.")
    return errors


def validate_warnings(rows: list[dict]) -> list[dict]:
    warnings: list[dict] = []
    for index, row in enumerate(rows, 3):
        case_id = str(row.get("case_id") or "")
        turn_type = extract_option_id(row.get("turn_type"))
        policy_risk_type = extract_option_id(row.get("policy_risk_type"))
        expected_risk = TURN_TYPE_EXPECTED_RISK.get(turn_type)
        if expected_risk and policy_risk_type not in {expected_risk, "prompt_injection_risk"}:
            _add_error(
                warnings,
                index,
                case_id,
                "policy_risk_type",
                row.get("policy_risk_type"),
                f"{turn_type} usually pairs with {expected_risk}. Keep current value only if this is intentional.",
            )
    return warnings


def validate_workbook(input_path: Path, focus_registry: Path) -> dict:
    if input_path.suffix.lower() == ".xlsx":
        rows = load_xlsx_rows(input_path)
    else:
        rows = load_csv_rows(input_path)
    errors = validate_rows(rows, focus_ids=load_focus_ids(focus_registry), subtype_registry=load_subtype_registry())
    warnings = validate_warnings(rows)
    return {
        "input_path": str(input_path),
        "row_count": len(rows),
        "error_count": len(errors),
        "warning_count": len(warnings),
        "errors": errors,
        "warnings": warnings,
    }


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Validate coach seed labeling workbook v2.")
    parser.add_argument("--input", type=Path, default=DEFAULT_WORKBOOK_PATH)
    parser.add_argument("--focus-registry", type=Path, default=DEFAULT_FOCUS_REGISTRY_PATH)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None, *, stdout: TextIO | None = None) -> int:
    args = _parse_args(sys.argv[1:] if argv is None else argv)
    report = validate_workbook(args.input, args.focus_registry)
    output = stdout or sys.stdout
    output.write(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    return 0 if report["error_count"] == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
