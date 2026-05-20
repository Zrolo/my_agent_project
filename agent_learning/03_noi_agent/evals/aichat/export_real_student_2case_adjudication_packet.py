#!/usr/bin/env python3
"""Export a 2-case adjudication packet from the reviewed 9-case focus workbook.

The generated workbook stays under .local_private and is intended for human
adjudication only. It copies already-redacted context from the private packet
and separates prior references from the adjudication fields.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


DEFAULT_INPUT = Path(".local_private/real_student_online_30case_second_coach_focus_packet_human_reviewed_20260520.xlsx")
DEFAULT_OUTPUT = Path(".local_private/real_student_online_2case_adjudication_packet_20260520.xlsx")
DEFAULT_SUMMARY = Path(".local_private/real_student_online_2case_adjudication_packet_summary_20260520.json")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
GRAY_FILL = PatternFill("solid", fgColor="E7E6E6")
EDIT_FILL = PatternFill("solid", fgColor="FFF2CC")
REFERENCE_FILL = PatternFill("solid", fgColor="D9EAF7")
GATE_FILL = PatternFill("solid", fgColor="E2F0D9")

OPTION_VALUES = {
    "裁决状态": ["待裁决", "已裁决", "因隐私风险排除", "因上下文不足排除"],
    "裁决上下文是否足够": ["足够", "部分足够", "不足", "不清楚"],
    "裁决最终泄露标签": ["无泄露", "轻微关键桥泄露", "重大关键桥泄露", "答案或代码泄露", "上下文不足未判断"],
    "裁决总体质量（1-5）": ["5", "4", "3", "2", "1"],
    "裁决是否愿意给学生看": ["是", "边界", "否"],
    "裁决信心": ["高", "中", "低"],
    "是否可纳入内部汇总": ["是", "否", "待定"],
    "隐私复核状态": ["待复核", "可内部复核", "需继续脱敏", "因隐私风险排除"],
    "知情/报告门": ["待完成", "可报告", "不可报告", "已退出"],
}

ADJUDICATION_HEADERS = [
    "案例编号",
    "题目名称/匿名题号",
    "裁决状态",
    "裁决上下文是否足够",
    "裁决最终泄露标签",
    "裁决总体质量（1-5）",
    "裁决是否愿意给学生看",
    "裁决信心",
    "是否可纳入内部汇总",
    "裁决备注（中文）",
]

PRIVACY_HEADERS = [
    "案例编号",
    "题目名称/匿名题号",
    "隐私复核状态",
    "知情/报告门",
    "裁决公开边界备注",
]


def normalize(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_sheet(wb: Any, sheet: str) -> list[dict[str, Any]]:
    ws = wb[sheet]
    headers = [normalize(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    rows: list[dict[str, Any]] = []
    for r in range(2, ws.max_row + 1):
        row = {headers[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
        if any(normalize(v) for v in row.values()):
            rows.append(row)
    return rows


def make_lookup(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {normalize(row.get("案例编号")): row for row in rows if normalize(row.get("案例编号"))}


def style_header(ws: Any, fill: PatternFill = HEADER_FILL) -> None:
    for cell in ws[1]:
        cell.fill = fill
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def set_widths(ws: Any) -> None:
    for col in ws.columns:
        letter = col[0].column_letter
        header = normalize(col[0].value)
        if "备注" in header or "摘要" in header or "问题" in header or "回复" in header or "代码" in header:
            ws.column_dimensions[letter].width = 42
        elif "案例编号" in header:
            ws.column_dimensions[letter].width = 24
        elif "题目" in header:
            ws.column_dimensions[letter].width = 28
        else:
            ws.column_dimensions[letter].width = 18
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_sheet_rows(wb_out: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]], fill: PatternFill = HEADER_FILL) -> Any:
    ws = wb_out.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    style_header(ws, fill)
    set_widths(ws)
    return ws


def add_dropdown(ws: Any, header: str, values: list[str], start_row: int, end_row: int) -> None:
    headers = [normalize(cell.value) for cell in ws[1]]
    if header not in headers:
        return
    col_idx = headers.index(header) + 1
    col_letter = ws.cell(1, col_idx).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{start_row}:{col_letter}{end_row}")


def export_packet(input_path: Path, output_path: Path, summary_path: Path) -> dict[str, Any]:
    wb_in = load_workbook(input_path, data_only=True)
    scoring_rows = read_sheet(wb_in, "二审评分")
    context_by_id = make_lookup(read_sheet(wb_in, "案例上下文"))
    ai_by_id = make_lookup(read_sheet(wb_in, "AI预复核参考"))
    privacy_by_id = make_lookup(read_sheet(wb_in, "隐私与报告门"))

    selected = [
        row
        for row in scoring_rows
        if normalize(row.get("二审状态")) == "仍需裁决" or normalize(row.get("二审是否需要裁决")) == "是"
    ]
    selected_ids = [normalize(row["案例编号"]) for row in selected]

    wb = Workbook()
    wb.remove(wb.active)

    ws_intro = wb.create_sheet("填写说明")
    intro_rows = [
        ["用途", "真实教练最终裁决包：仅包含 9-case 二审焦点包中仍需裁决的 rows。"],
        ["优先填写", "请填写“裁决评分”和“隐私与报告门”。"],
        ["参考表", "“二审参考”和“AI预复核参考”只供裁决者了解争议来源，不是 final gold。"],
        ["当前AIChat回复", "线上已展示回复，观察项，非实验条件。不是 baseline、condition、control 或 Repair output。"],
        ["公开边界", "不公开学生原文、完整代码、完整 AIChat 回复、case-level labels、hash salt 或可逆映射。"],
    ]
    for row in intro_rows:
        ws_intro.append(row)
    set_widths(ws_intro)

    ws_options = wb.create_sheet("下拉选项")
    for col_idx, (name, values) in enumerate(OPTION_VALUES.items(), 1):
        ws_options.cell(1, col_idx, name)
        for row_idx, value in enumerate(values, 2):
            ws_options.cell(row_idx, col_idx, value)
    style_header(ws_options, GRAY_FILL)
    set_widths(ws_options)

    context_rows = [context_by_id[case_id] for case_id in selected_ids if case_id in context_by_id]
    context_headers = list(context_rows[0].keys()) if context_rows else []
    add_sheet_rows(wb, "案例上下文", context_headers, context_rows, REFERENCE_FILL)

    second_headers = list(selected[0].keys()) if selected else []
    add_sheet_rows(wb, "二审参考", second_headers, selected, REFERENCE_FILL)

    ai_rows = [ai_by_id[case_id] for case_id in selected_ids if case_id in ai_by_id]
    ai_headers = list(ai_rows[0].keys()) if ai_rows else []
    add_sheet_rows(wb, "AI预复核参考", ai_headers, ai_rows, REFERENCE_FILL)

    adjudication_rows = []
    for row in selected:
        adjudication_rows.append(
            {
                "案例编号": row.get("案例编号"),
                "题目名称/匿名题号": row.get("题目名称/匿名题号"),
                "裁决状态": "待裁决",
                "裁决上下文是否足够": "",
                "裁决最终泄露标签": "",
                "裁决总体质量（1-5）": "",
                "裁决是否愿意给学生看": "",
                "裁决信心": "",
                "是否可纳入内部汇总": "待定",
                "裁决备注（中文）": "",
            }
        )
    ws_adj = add_sheet_rows(wb, "裁决评分", ADJUDICATION_HEADERS, adjudication_rows, EDIT_FILL)
    for header, values in OPTION_VALUES.items():
        add_dropdown(ws_adj, header, values, 2, 1 + len(adjudication_rows))
    ws_adj.conditional_formatting.add(
        f"C2:C{1 + len(adjudication_rows)}",
        FormulaRule(formula=[f'$C2="待裁决"'], fill=PatternFill("solid", fgColor="F4CCCC")),
    )

    privacy_rows = []
    for row in selected:
        case_id = normalize(row.get("案例编号"))
        prev = privacy_by_id.get(case_id, {})
        privacy_rows.append(
            {
                "案例编号": row.get("案例编号"),
                "题目名称/匿名题号": row.get("题目名称/匿名题号"),
                "隐私复核状态": prev.get("隐私复核状态") or "待复核",
                "知情/报告门": prev.get("知情/报告门") or "待完成",
                "裁决公开边界备注": "默认不可公开 case-level labels；如需报告需另过 consent/reporting gate。",
            }
        )
    ws_priv = add_sheet_rows(wb, "隐私与报告门", PRIVACY_HEADERS, privacy_rows, GATE_FILL)
    for header in ["隐私复核状态", "知情/报告门"]:
        add_dropdown(ws_priv, header, OPTION_VALUES[header], 2, 1 + len(privacy_rows))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    summary = {
        "source_workbook": str(input_path),
        "output_workbook": str(output_path),
        "adjudication_case_count": len(selected),
        "case_ids_private": selected_ids,
        "boundary": "Private adjudication packet only; not public case-level evidence; not a new experiment or condition.",
        "reportable_case_level_evidence": 0,
    }
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--summary", type=Path, default=DEFAULT_SUMMARY)
    args = parser.parse_args()
    summary = export_packet(args.input, args.output, args.summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
