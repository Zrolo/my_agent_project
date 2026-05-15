"""Generate dialogue-state v3 held-out draft cases.

The v3 dataset extends Luogu-grounded v2 cases with fixed follow-up tutoring
contexts. It does not overwrite v2 and does not use historical online AI
replies as generation inputs.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill


if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))


DEFAULT_SOURCE_JSONL = Path("docs/research/bridgebench_cp_heldout_v2_50_draft.jsonl")
DEFAULT_OUTPUT_JSONL = Path("docs/research/bridgebench_cp_dialogue_state_v3_50_draft.jsonl")
DEFAULT_REVIEW_XLSX = Path("docs/research/dialogue_state_v3_50_source_and_case_review.zh.xlsx")
DEFAULT_REVIEW_XLSX_EN = Path("docs/research/dialogue_state_v3_50_source_and_case_review.en.xlsx")
DEFAULT_REPORT_JSON = Path("docs/research/dialogue_state_v3_generation_report_20260513.json")
DEFAULT_REPORT_ZH = Path("docs/research/dialogue_state_v3_generation_report_20260513.zh.md")
DEFAULT_REPORT_EN = Path("docs/research/dialogue_state_v3_generation_report_20260513.md")

CONTEXT_PLAN = (
    ["initial_question"] * 10
    + ["followup_after_correct_short_answer"] * 7
    + ["followup_after_partial_answer"] * 8
    + ["followup_after_wrong_answer"] * 7
    + ["followup_after_code_attempt"] * 10
    + ["followup_after_prerequisite_gap"] * 5
    + ["policy_direct_answer_special"] * 3
)

FOLLOWABILITY_BY_CONTEXT = {
    "initial_question": "NA",
    "followup_after_correct_short_answer": "F1",
    "followup_after_partial_answer": "F2",
    "followup_after_wrong_answer": "F3",
    "followup_after_code_attempt": "F2",
    "followup_after_prerequisite_gap": "F4",
    "policy_direct_answer_special": "F3",
}

EXPECTED_MOVE_BY_CONTEXT = {
    "initial_question": "micro_step",
    "followup_after_correct_short_answer": "advance",
    "followup_after_partial_answer": "clarify",
    "followup_after_wrong_answer": "micro_step",
    "followup_after_code_attempt": "micro_step",
    "followup_after_prerequisite_gap": "prerequisite_repair",
    "policy_direct_answer_special": "safe_redirect",
}

CONFIDENCE_BY_CONTEXT = {
    "initial_question": "high",
    "followup_after_correct_short_answer": "high",
    "followup_after_partial_answer": "medium",
    "followup_after_wrong_answer": "medium",
    "followup_after_code_attempt": "medium",
    "followup_after_prerequisite_gap": "high",
    "policy_direct_answer_special": "high",
}

CONTEXT_LABEL_ZH = {
    "initial_question": "初始提问",
    "followup_after_correct_short_answer": "能跟上：短答正确后续轮",
    "followup_after_partial_answer": "部分跟上：回答不完整后续轮",
    "followup_after_wrong_answer": "跟得吃力：回答错误后续轮",
    "followup_after_code_attempt": "部分跟上：代码尝试后续轮",
    "followup_after_prerequisite_gap": "基础断层：前置概念缺口后续轮",
    "policy_direct_answer_special": "策略请求：直接要答案/代码",
}

FOLLOWABILITY_LABEL_ZH = {
    "NA": "不适用",
    "F1": "F1 能一直跟上",
    "F2": "F2 勉强/部分跟上",
    "F3": "F3 比较难跟上",
    "F4": "F4 基础断层",
}

EXPECTED_MOVE_ZH = {
    "advance": "推进到下一小步",
    "clarify": "澄清当前回答",
    "micro_step": "拆成更小一步",
    "prerequisite_repair": "补前置概念",
    "safe_redirect": "安全重定向",
}

TURN_POSITION_ZH = {
    "initial": "初始提问",
    "followup": "后续辅导轮",
}

CONFIDENCE_ZH = {
    "high": "高",
    "medium": "中",
    "low": "低",
}

REVIEW_COLUMNS = [
    "case_id",
    "source_case_id",
    "turn_position",
    "context_type",
    "student_scaffold_followability",
    "followability_label_confidence",
    "expected_tutor_move",
    "problem_source_url",
    "problem_ref",
    "problem_statement",
    "problem_statement_public_summary",
    "student_message",
    "recent_dialogue",
    "context_ai_reply",
    "prior_ai_scaffold",
    "student_reply_to_prior_scaffold",
    "followability_evidence_quote",
    "followability_uncertainty_reason",
    "missing_bridge",
    "forbidden_content",
    "success_criteria",
    "review_notes_for_coach",
    "source_ok",
    "context_coherent",
    "student_message_realistic",
    "followability_ok",
    "missing_bridge_ok",
    "forbidden_content_ok",
    "success_criteria_ok",
    "leakage_boundary_ok",
    "case_decision",
    "issue_type",
    "coach_fix_suggestion",
    "reviewer_confidence",
    "reviewer_id",
    "review_round",
]

REVIEW_HEADERS_ZH = {
    "case_id": "样本编号",
    "source_case_id": "来源 v2 样本",
    "turn_position": "轮次位置",
    "context_type": "上下文类型",
    "student_scaffold_followability": "学生跟随状态",
    "followability_label_confidence": "跟随状态置信度",
    "expected_tutor_move": "期望下一步教学动作",
    "problem_source_url": "原题链接",
    "problem_ref": "题目编号",
    "problem_statement": "原题题面/必要题面",
    "problem_statement_public_summary": "公开题面摘要",
    "student_message": "学生当前问题/回复",
    "recent_dialogue": "近期对话",
    "context_ai_reply": "上下文 AI 回复",
    "prior_ai_scaffold": "上一轮 AI 脚手架",
    "student_reply_to_prior_scaffold": "学生对脚手架的回答",
    "followability_evidence_quote": "跟随状态证据",
    "followability_uncertainty_reason": "不确定原因",
    "missing_bridge": "目标缺失桥梁",
    "forbidden_content": "桥梁禁止内容",
    "success_criteria": "成功标准",
    "review_notes_for_coach": "给教练的复核备注",
    "source_ok": "题源是否可用",
    "context_coherent": "上下文是否连贯",
    "student_message_realistic": "学生话术是否真实",
    "followability_ok": "跟随状态是否合理",
    "missing_bridge_ok": "缺失桥梁是否合理",
    "forbidden_content_ok": "禁止内容是否合理",
    "success_criteria_ok": "成功标准是否合理",
    "leakage_boundary_ok": "泄露边界是否清楚",
    "case_decision": "样本处理决定",
    "issue_type": "问题类型",
    "coach_fix_suggestion": "教练修改建议",
    "reviewer_confidence": "审核置信度",
    "reviewer_id": "审核人",
    "review_round": "审核轮次",
}

REVIEW_HEADERS_EN = {
    "case_id": "Case ID",
    "source_case_id": "Source v2 Case",
    "turn_position": "Turn Position",
    "context_type": "Context Type",
    "student_scaffold_followability": "Student Followability",
    "followability_label_confidence": "Followability Label Confidence",
    "expected_tutor_move": "Expected Tutor Move",
    "problem_source_url": "Original Problem URL",
    "problem_ref": "Problem Reference",
    "problem_statement": "Problem Statement / Necessary Statement",
    "problem_statement_public_summary": "Public Problem Summary",
    "student_message": "Current Student Message / Reply",
    "recent_dialogue": "Recent Dialogue",
    "context_ai_reply": "Context AI Reply",
    "prior_ai_scaffold": "Prior AI Scaffold",
    "student_reply_to_prior_scaffold": "Student Reply To Prior Scaffold",
    "followability_evidence_quote": "Followability Evidence Quote",
    "followability_uncertainty_reason": "Followability Uncertainty Reason",
    "missing_bridge": "Target Missing Bridge",
    "forbidden_content": "Forbidden Content",
    "success_criteria": "Success Criteria",
    "review_notes_for_coach": "Coach Review Notes",
    "source_ok": "Source Usability",
    "context_coherent": "Context Coherence",
    "student_message_realistic": "Student Wording Realism",
    "followability_ok": "Followability Label Validity",
    "missing_bridge_ok": "Missing Bridge Validity",
    "forbidden_content_ok": "Forbidden Content Validity",
    "success_criteria_ok": "Success Criteria Validity",
    "leakage_boundary_ok": "Leakage Boundary Clarity",
    "case_decision": "Case Decision",
    "issue_type": "Issue Type",
    "coach_fix_suggestion": "Coach Fix Suggestion",
    "reviewer_confidence": "Reviewer Confidence",
    "reviewer_id": "Reviewer ID",
    "review_round": "Review Round",
}

REVIEW_VALIDATION_CHOICES = {
    "source_ok": ["yes", "partial", "no"],
    "context_coherent": ["yes", "partial", "no"],
    "student_message_realistic": ["yes", "partial", "no"],
    "followability_ok": ["yes", "revise", "no", "not_applicable"],
    "missing_bridge_ok": ["yes", "revise", "no"],
    "forbidden_content_ok": ["yes", "too_strict", "too_loose", "unclear"],
    "success_criteria_ok": ["yes", "revise", "no"],
    "leakage_boundary_ok": ["yes", "revise", "no", "unclear"],
    "case_decision": ["accept", "revise", "drop", "discuss"],
    "issue_type": [
        "none",
        "source_issue",
        "context_mismatch",
        "student_language_artificial",
        "followability_issue",
        "bridge_label_issue",
        "forbidden_content_issue",
        "success_criteria_issue",
        "leakage_boundary_issue",
        "other",
    ],
    "reviewer_confidence": ["high", "medium", "low"],
}

REVIEW_VALIDATION_CHOICES_ZH = {
    "source_ok": ["是", "部分", "否"],
    "context_coherent": ["是", "部分", "否"],
    "student_message_realistic": ["是", "部分", "否"],
    "followability_ok": ["是", "需修改", "否", "不适用"],
    "missing_bridge_ok": ["是", "需修改", "否"],
    "forbidden_content_ok": ["是", "过严", "过松", "不清楚"],
    "success_criteria_ok": ["是", "需修改", "否"],
    "leakage_boundary_ok": ["是", "需修改", "否", "不清楚"],
    "case_decision": ["接受", "修改", "丢弃", "讨论"],
    "issue_type": [
        "无",
        "题源问题",
        "上下文不一致",
        "学生话术不像真实学生",
        "跟随状态问题",
        "桥梁标签问题",
        "禁止内容问题",
        "成功标准问题",
        "泄露边界问题",
        "其他",
    ],
    "reviewer_confidence": ["高", "中", "低"],
}

BUCKET_SHEET_NAMES = {
    "state_representation_semantics": "bucket_state_representation",
    "transition_recurrence_source": "bucket_transition",
    "predicate_check_semantics": "bucket_predicate_check",
    "boundary_update_order": "bucket_boundary_order",
    "modeling_object_relation": "bucket_modeling",
    "aggregation_contribution_summary": "bucket_aggregation",
    "data_structure_operation_semantics": "bucket_data_structure",
    "correctness_invariant": "bucket_correctness",
    "implementation_boundary": "bucket_implementation",
    "debugging_evidence": "bucket_debugging",
    "policy_request": "bucket_policy",
}

BUCKET_SHEET_NAMES_ZH = {
    "state_representation_semantics": "状态表示",
    "transition_recurrence_source": "转移递推",
    "predicate_check_semantics": "判定条件",
    "boundary_update_order": "边界顺序",
    "modeling_object_relation": "建模关系",
    "aggregation_contribution_summary": "贡献汇总",
    "data_structure_operation_semantics": "数据结构",
    "correctness_invariant": "正确性不变量",
    "implementation_boundary": "实现边界",
    "debugging_evidence": "调试证据",
    "policy_request": "策略请求",
}

INSTRUCTION_ROWS = [
    ["Dialogue-State v3 50-case 复核流程"],
    ["1. 先看原题题面/必要题面和原题链接，确认题意、对象、限制和输入输出。"],
    ["2. 再看近期对话和上下文 AI 回复，判断学生当前回复是在接哪一个脚手架问题。"],
    ["3. 再看学生当前问题/回复，确认 F1-F4 跟随状态是否合理。"],
    ["4. 再看目标缺失桥梁、禁止内容和成功标准，判断它们是否覆盖当前卡点。"],
    ["5. 如题面、学生回复、近期对话、bridge 标签或 forbidden content 不一致，请在复核备注中标记。"],
    ["6. 请优先在各桥梁桶 sheet 中填写审核列；总表用于全局查看。审核结束后由脚本汇总各 bucket sheet 的审核结果。"],
    ["7. 结构化审核列使用固定中文选项，统计脚本会映射为英文 canonical values。例如：接受=accept，修改=revise，丢弃=drop，讨论=discuss。"],
    ["8. 本表是 case/source 复核，不是 AI 回复盲评；此时不评价任何 condition 的回复质量。"],
    ["9. context_ai_reply / prior_ai_scaffold 只用于判断学生当前回复是否承接上一轮脚手架，不用于评价该 AI 回复本身好坏。"],
]

INSTRUCTION_ROWS_EN = [
    ["Dialogue-State v3 50-case Review Workflow"],
    ["1. Read the problem statement first, including the original link and necessary statement."],
    ["2. Then read the recent dialogue and context AI reply to understand what scaffold the current student reply is answering."],
    ["3. Then read the current student message/reply and check whether the F1-F4 followability label is reasonable."],
    ["4. Then review the target missing bridge, forbidden content, and success criteria."],
    ["5. If the problem statement, student reply, recent dialogue, bridge label, or forbidden content is inconsistent, mark it in the review notes."],
    ["6. Prefer filling the bucket sheets; the main sheet is for global browsing. The summary script will aggregate structured review fields from bucket sheets."],
    ["7. Structured review columns use fixed English options for aggregation: case_decision is accept / revise / drop / discuss."],
    ["8. This workbook reviews case/source quality only. It is not an AI-response blind review and should not be used to score condition quality."],
    ["9. context_ai_reply / prior_ai_scaffold are only for checking whether the student reply follows the previous scaffold, not for judging that AI reply's quality."],
]


def _load_jsonl(path: Path) -> list[dict]:
    rows = []
    for line_number, raw_line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        line = raw_line.strip()
        if not line:
            continue
        try:
            row = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ValueError(f"Invalid JSONL at {path}:{line_number}: {exc}") from exc
        if not isinstance(row, dict):
            raise ValueError(f"JSONL row must be an object at {path}:{line_number}")
        rows.append(row)
    return rows


def _write_jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "\n".join(json.dumps(row, ensure_ascii=False, sort_keys=True) for row in rows) + "\n",
        encoding="utf-8",
    )


def _as_list(value) -> list[str]:
    if isinstance(value, list):
        return [str(item) for item in value]
    if value in {None, ""}:
        return []
    return [str(value)]


def _problem_short_name(row: dict) -> str:
    title = row.get("problem_title") or row.get("problem_ref") or row.get("problem_source_id") or "这题"
    return str(title).replace("luogu_", "").replace("_", " ")[:24]


def _current_small_question(row: dict) -> str:
    bucket = row.get("bridge_bucket") or row.get("category") or ""
    if "implementation" in bucket:
        return "你先拿最小样例核对一个边界、初值或变量范围。"
    if "debugging" in bucket:
        return "你先构造一个最小反例，或者指出一个中间变量应该是多少。"
    if "state" in bucket or "representation" in bucket:
        return "你先用一句话说：这个状态的一格应该记录目标值、可行性、数量，还是代价？"
    if "transition" in bucket:
        return "你先只看当前这一步，说它可能从哪一种更小情况转过来。"
    if "predicate" in bucket or "check" in bucket:
        return "你先说说 check 返回 true 时，代表候选值是可行还是不可行。"
    if "boundary" in bucket or "order" in bucket:
        return "你先判断更新后哪一边还可能包含答案，或者哪个旧值不能被覆盖。"
    if "modeling" in bucket:
        return "你先列一个题面对象，再说它和另一个对象之间有什么关系。"
    if "aggregation" in bucket:
        return "你先追踪一次操作会影响哪几个最小位置或节点。"
    if "data_structure" in bucket:
        return "你先说一次 update/query 之后，结构里应该保持什么摘要信息。"
    if "correctness" in bucket:
        return "你先比较两个相邻选择互换以后，目标值会不会变差。"
    return "你先把当前卡住的最小一步说出来，不需要写完整解法。"


def _initial_question(row: dict) -> str:
    return str(row.get("student_message") or "这题下一步怎么想？").strip()


def _compact_len(value: str) -> int:
    return len("".join(str(value or "").split()))


def _student_message_length_bucket(value: str) -> str:
    length = _compact_len(value)
    if length <= 30:
        return "short"
    if length <= 70:
        return "medium_short"
    if length <= 140:
        return "medium_long"
    return "long"


def _target_length_bucket(row: dict) -> str:
    value = str(row.get("student_message_length_bucket") or "").strip()
    if value in {"short", "medium_short", "medium_long", "long"}:
        return value
    return _student_message_length_bucket(_initial_question(row))


def _variant(options: list[str], ordinal: int) -> str:
    return options[(ordinal - 1) % len(options)]


def _bucket_contains(row: dict, fragment: str) -> bool:
    return fragment in str(row.get("bridge_bucket") or row.get("category") or "")


def _expand_reply_for_bucket(row: dict, context_type: str, reply: str, ordinal: int) -> str:
    """Keep follow-up replies realistic while preserving held-out length quotas."""

    bucket = _target_length_bucket(row)
    if bucket == "short":
        if _compact_len(reply) > 30:
            short_replies = {}
            if context_type == "followup_after_partial_answer":
                if _bucket_contains(row, "boundary"):
                    short_replies[context_type] = [
                        "我怕旧值被覆盖。",
                        "顺序一变就不对。",
                        "还分不清新旧值。",
                    ]
                elif _bucket_contains(row, "modeling"):
                    short_replies[context_type] = [
                        "关系我还没列清。",
                        "对象之间怎么连？",
                        "覆盖关系我说不准。",
                    ]
                elif _bucket_contains(row, "implementation"):
                    short_replies[context_type] = [
                        "我还不知道先查哪格。",
                        "下标和初值没对上。",
                        "最小样例不会挑。",
                    ]
                else:
                    short_replies[context_type] = [
                        "另一种情况也要算吗？",
                        "我只想到一种来源。",
                        "分支是不是漏了？",
                    ]
            elif context_type == "followup_after_wrong_answer":
                if _bucket_contains(row, "modeling"):
                    short_replies[context_type] = [
                        "对象关系可能反了。",
                        "覆盖关系我搞反了。",
                        "我像在猜怎么连。",
                    ]
                elif _bucket_contains(row, "transition"):
                    short_replies[context_type] = [
                        "我可能少看了一个来源。",
                        "是不是只接一种情况？",
                        "来源方向好像反了。",
                    ]
                elif _bucket_contains(row, "implementation"):
                    short_replies[context_type] = [
                        "下标可能反了。",
                        "输入范围好像没对上。",
                        "空格处理我猜错了。",
                    ]
                else:
                    short_replies[context_type] = [
                        "我是不是理解反了？",
                        "true 是不是该排除？",
                        "这一边要不要丢掉？",
                    ]
            elif context_type == "followup_after_prerequisite_gap":
                if _bucket_contains(row, "implementation"):
                    short_replies[context_type] = [
                        "能先看一个字符吗？",
                        "下标从哪开始？",
                        "输入怎么读没对上。",
                    ]
                elif _bucket_contains(row, "modeling"):
                    short_replies[context_type] = [
                        "能先解释对象关系吗？",
                        "谁和谁有关？",
                        "题面关系没搭起来。",
                    ]
                else:
                    short_replies[context_type] = [
                        "能先用题里的对象解释吗？",
                        "后面的转移我接不上。",
                        "可能是前面概念没搭起来。",
                    ]
            return _variant(short_replies.get(context_type, [reply]), ordinal)
        return reply

    title = _problem_short_name(row)
    if context_type == "followup_after_correct_short_answer":
        short_additions = [
            "但我还不确定能不能继续用。",
            "我怕这只是套词。",
            "下一步我不知道怎么接。",
        ]
        additions = [
            f"我还不确定这个理解能不能直接放回《{title}》的下一步里。",
            "但我怕自己只是套了一个词，没真正对应到题面里的对象。",
            "如果要继续，我想先确认这一格和题目里的哪个阶段对应。",
        ]
    elif context_type == "followup_after_partial_answer":
        if _bucket_contains(row, "boundary"):
            short_additions = [
                "我怕旧值被覆盖。",
                "顺序一变就不对。",
                "还分不清新旧值。",
            ]
            additions = [
                "我能感觉顺序和旧值有关，但不知道怎样判断旧值会不会被刚更新的值覆盖。",
                f"放到《{title}》这里，我知道要看更新前后，但不知道该保护哪一个旧状态。",
                "我怕循环方向一写反，就把本轮刚算出的值又拿来用了。",
            ]
        elif _bucket_contains(row, "modeling"):
            short_additions = [
                "关系我还没列清。",
                "对象之间怎么连不确定。",
                "覆盖关系我说不准。",
            ]
            additions = [
                "我能列出题面对象，但不知道它们之间应该看覆盖、相邻还是依赖关系。",
                f"放到《{title}》这里，我知道有哪些对象，但不知道哪一种关系才是建模核心。",
                "我怕把题面里的对象抽错，后面图或状态就全都偏了。",
            ]
        elif _bucket_contains(row, "implementation"):
            short_additions = [
                "我还不知道先查哪一格。",
                "下标和初值没对上。",
                "最小样例不会挑。",
            ]
            additions = [
                "我知道现在该查实现细节，但不知道先检查下标、初值还是输入范围。",
                f"放到《{title}》这里，我能看懂题意，但代码里哪一个边界最容易错还没抓住。",
                "我怕不是算法问题，而是一个很小的输入或下标细节一直没对上。",
            ]
        else:
            short_additions = [
                "我怕漏掉另一种情况。",
                "另一个来源我不敢确定。",
                "我还不会把分支合起来。",
            ]
            additions = [
                "我现在只能说出一个来源，另一个来源是不是也要算进去就有点拿不准。",
                f"放到《{title}》这里，我能看出前一步有关系，但不知道该怎么分清两类情况。",
                "我怕少考虑一种情况，所以不敢直接把它合成公式。",
            ]
    elif context_type == "followup_after_wrong_answer":
        if _bucket_contains(row, "modeling"):
            short_additions = [
                "对象关系可能反了。",
                "覆盖关系我搞反了。",
                "我像在猜怎么连。",
            ]
            additions = [
                "我感觉自己把对象和关系放反了，但说不出到底应该让谁覆盖谁。",
                f"在《{title}》这题里，我能复述题面，但不知道该把哪个限制当成关系。",
                "我现在更像是在猜建模方式，不确定题面里的关系该怎么抽出来。",
            ]
        elif _bucket_contains(row, "transition"):
            short_additions = [
                "我可能少看了一个来源。",
                "是不是只接一种情况？",
                "来源方向好像反了。",
            ]
            additions = [
                "我感觉自己只看了一个来源，但另一个更小情况是不是也要接进来还没想清。",
                f"在《{title}》这题里，我知道要从更小情况来，但不确定是哪一类前驱。",
                "我现在更像在猜转移方向，不确定当前量到底从哪里来。",
            ]
        elif _bucket_contains(row, "implementation"):
            short_additions = [
                "下标可能反了。",
                "输入范围好像没对上。",
                "空格处理我猜错了。",
            ]
            additions = [
                "我感觉自己把下标或输入边界理解反了，但说不出先该查哪一个最小样例。",
                f"在《{title}》这题里，我能跑样例，但一换输入就不知道哪个细节错了。",
                "我现在更像是在猜实现细节，不确定是下标、初值还是类型出了问题。",
            ]
        else:
            short_additions = [
                "我感觉方向反了。",
                "边界一换我就乱了。",
                "我现在更像在猜模板。",
            ]
            additions = [
                "我感觉自己把 true/false 或左右方向理解反了，但说不出到底反在哪里。",
                f"在《{title}》这题里，我能跟样例，但换一个边界就不知道该保留哪边。",
                "我现在更像是在猜模板，不确定这个判断背后的含义。",
            ]
    elif context_type == "followup_after_code_attempt":
        short_additions = [
            "我怀疑不是单纯少写一行。",
            "更新顺序一变结果就不同。",
            "我想先知道该打印哪个量。",
        ]
        additions = [
            "我改过一次循环范围，样例还是不稳，所以怀疑不是单纯少写一行。",
            f"我对《{title}》的整体思路大概能跟，但代码里这个更新顺序一变结果就不一样。",
            "我想先知道应该打印哪个中间量，别一上来就整段重写。",
        ]
    elif context_type == "followup_after_prerequisite_gap":
        if _bucket_contains(row, "implementation"):
            short_additions = [
                "能先看一个字符吗？",
                "下标从哪开始我没懂。",
                "输入怎么读我没对上。",
            ]
            additions = [
                "如果可以的话，能不能先用这题里的一个字符解释该检查哪个下标或输入边界？",
                f"我现在看《{title}》时连字符、空格和计数范围都没对上，所以代码只能靠猜。",
                "我可能不是算法不会，而是输入、下标或初值这个基础细节没有搭起来。",
            ]
        elif _bucket_contains(row, "modeling"):
            short_additions = [
                "能先解释对象关系吗？",
                "我连谁和谁有关都没懂。",
                "题面关系没搭起来。",
            ]
            additions = [
                "如果可以的话，能不能先用题里的两个对象解释它们到底是什么关系？",
                f"我现在看《{title}》时连对象之间的基础关系都没对上，所以后面建模接不上。",
                "我可能不是不会写算法，而是题面里的对象和约束还没有搭起来。",
            ]
        else:
            short_additions = [
                "能先用题里的对象解释吗？",
                "后面的转移我接不上。",
                "可能是前面概念没搭起来。",
            ]
            additions = [
                "如果可以的话，能不能先用这题里的一个小对象解释这个词，而不是直接讲完整做法？",
                f"我现在看《{title}》时连这个基础关系都没对上，所以后面的转移/维护都接不上。",
                "我可能不是这一步不会写，而是前面那个概念就没有搭起来。",
            ]
    elif context_type == "policy_direct_answer_special":
        short_additions = [
            "不行的话先告诉我补哪步。",
            "我不知道先排查哪里。",
            "给我一个能动手的小检查点。",
        ]
        additions = [
            "如果不能给完整代码，那请先把我应该补充的最小信息说清楚。",
            f"我不是想跳过学习，只是《{title}》卡太久了，不知道该先排查思路还是实现。",
            "你可以先不给答案，但最好给我一个能马上动手验证的小检查点。",
        ]
    else:
        short_additions = ["我只想先把这一步对上。"]
        additions = ["我不需要完整题解，只想先把当前这一步对上。"]

    result = f"{reply} {_variant(short_additions, ordinal)}" if bucket == "medium_short" else f"{reply} {_variant(additions, ordinal)}"
    if bucket == "medium_short":
        return result

    medium_extra = _variant(
        [
            "我自己试着顺了一遍样例，但只能看到表面过程，看不出为什么下一步一定要这样接。",
            "我希望先确认这个小关系，因为一旦这里错了，后面写代码就会一直靠猜。",
            "现在我最怕的是把一个模板套上去，结果换成别的数据就又不知道该怎么判断。",
        ],
        ordinal + 1,
    )
    result = f"{result} {medium_extra}"
    if bucket == "medium_long":
        return result

    long_extra = _variant(
        [
            "我已经试过按题面手算一个小例子，但算到中间就不知道该记录哪个量，代码里也不知道该检查哪一行。你先别直接给完整公式，帮我把这个小关系拆清楚就行。",
            "如果你只问我一个问题，我比较希望这个问题能让我判断自己到底是对象关系没建好，还是更新/转移方向没搞清楚。这样我后面才知道要改思路还是改代码。",
            "我现在愿意自己继续推，但需要一个不会直接暴露答案的提示。最好是让我对着一个最小样例说出下一步，而不是直接告诉我完整模板。",
        ],
        ordinal + 2,
    )
    return f"{result} {long_extra}"


def _evidence_quote(reply: str, candidates: list[str]) -> str:
    for candidate in candidates:
        if candidate and candidate in reply:
            return candidate
    compact = "".join(reply.split())
    return compact[: min(8, len(compact))]


def _prior_ai_scaffold(row: dict) -> str:
    question = _current_small_question(row)
    return (
        f"先别急着要完整做法。围绕《{_problem_short_name(row)}》，我们只拆当前一小步："
        f"{question}"
    )


def _reply_for_context(row: dict, context_type: str, ordinal: int) -> tuple[str, str, str]:
    if context_type == "followup_after_correct_short_answer":
        if _bucket_contains(row, "predicate"):
            reply = _variant(
                [
                    "true 就是这个候选值能满足限制。",
                    "我理解 true 应该是当前候选能过限制。",
                    "是不是 true 代表这个候选还可行？",
                ],
                ordinal,
            )
            evidence = "能满足限制"
        elif _bucket_contains(row, "transition"):
            reply = _variant(
                [
                    "我能说出一个来源：它应该从前一个更小阶段的结果接过来。",
                    "一个来源应该是已经处理完前面部分后的结果。",
                    "我觉得当前量至少有一个来源，是去掉当前选择后的更小情况。",
                ],
                ordinal,
            )
            evidence = "一个来源"
        elif _bucket_contains(row, "boundary"):
            reply = _variant(
                [
                    "我理解这里要保护旧值，不能让刚更新的值又被本轮用到。",
                    "应该是看哪一边还是旧状态，避免被这轮更新覆盖。",
                    "我知道顺序和旧值有关，不能把新值当成旧值再用。",
                ],
                ordinal,
            )
            evidence = "旧值"
        elif _bucket_contains(row, "modeling"):
            reply = _variant(
                [
                    "我能列出对象了：时间段和能覆盖它的选择之间有关系。",
                    "对象应该是题面里的实体，关系是它们满足或覆盖某个限制。",
                    "我知道要先找两个对象，再看它们之间是不是覆盖或依赖。",
                ],
                ordinal,
            )
            evidence = "对象"
        elif _bucket_contains(row, "implementation"):
            reply = _variant(
                [
                    "我知道先拿最小输入检查一个下标或初值。",
                    "我应该先看输入里的一个字符或边界位置。",
                    "先确认数组范围和初始值，再看后面的循环。",
                ],
                ordinal,
            )
            evidence = "下标"
        else:
            reply = _variant(
                [
                    "我觉得是记录已经处理到这里时的最优/可行情况。",
                    "是不是先记录到当前阶段为止还能成立的情况？",
                    "我理解这一格是在记当前阶段的一个结果。",
                ],
                ordinal,
            )
            evidence = "最优/可行情况"
        final_reply = _expand_reply_for_bucket(row, context_type, reply, ordinal)
        if _bucket_contains(row, "predicate"):
            evidence = _evidence_quote(final_reply, ["true", "候选", "可行", "限制"])
        elif _bucket_contains(row, "transition"):
            evidence = _evidence_quote(final_reply, ["来源", "前一个", "更小", "选择"])
        elif _bucket_contains(row, "boundary"):
            evidence = _evidence_quote(final_reply, ["旧值", "顺序", "覆盖", "新值"])
        elif _bucket_contains(row, "modeling"):
            evidence = _evidence_quote(final_reply, ["对象", "关系", "覆盖", "依赖"])
        elif _bucket_contains(row, "implementation"):
            evidence = _evidence_quote(final_reply, ["下标", "输入", "字符", "初值"])
        else:
            evidence = _evidence_quote(final_reply, ["当前阶段", "这一格", "最优", "结果"])
        return final_reply, evidence, ""
    if context_type == "followup_after_partial_answer":
        if _bucket_contains(row, "boundary"):
            reply = _variant(
                [
                    "我知道要保护旧值，但不知道顺序怎么避免覆盖。",
                    "我能看出要看更新前后，可是不知道哪一个值会被覆盖。",
                    "好像和旧值顺序有关，但从大到小还是从小到大我还没想明白。",
                ],
                ordinal,
            )
        elif _bucket_contains(row, "modeling"):
            reply = _variant(
                [
                    "我能列出对象，但它们之间到底是覆盖、相邻还是依赖，我不确定。",
                    "对象我大概知道，可关系怎么连还没想明白。",
                    "我知道题面里有哪些东西，但不知道该抽成哪种关系。",
                ],
                ordinal,
            )
        elif _bucket_contains(row, "implementation"):
            reply = _variant(
                [
                    "我知道要查代码细节，但不知道先查下标、初值还是输入范围。",
                    "我能跑样例，可是不知道该先打印哪一个中间量。",
                    "看起来像边界错，但我还没想明白该拿哪个最小样例查。",
                ],
                ordinal,
            )
        else:
            reply = _variant(
                [
                    "大概是看前面能不能推过来，但我不知道要不要把另一种情况也算进去。",
                    "我能想到一个来源，但另一种情况是不是也要单独算，我不确定。",
                    "这一步好像能从前面接过来，可是分支怎么拆我还没想明白。",
                ],
                ordinal,
            )
        final_reply = _expand_reply_for_bucket(row, context_type, reply, ordinal)
        return (
            final_reply,
            _evidence_quote(final_reply, ["不知道", "不确定", "没想明白", "另一种", "分支"]),
            "学生方向接近，但没有说清分支或条件。",
        )
    if context_type == "followup_after_wrong_answer":
        if _bucket_contains(row, "modeling"):
            reply = _variant(
                [
                    "我是不是把对象和关系放反了？应该看谁覆盖谁吗？",
                    "对象是时间段还是选择本身？我感觉关系有点反。",
                    "我是不是该把限制当成点，把对象当成边？感觉弄反了。",
                ],
                ordinal,
            )
        elif _bucket_contains(row, "transition"):
            reply = _variant(
                [
                    "是不是只看一个来源就够了？另一个情况不用接进来吗？",
                    "我是不是把来源方向反了，从当前往前推还是从前面推当前？",
                    "当前这一步是不是直接继承就行？我感觉少了一种情况。",
                ],
                ordinal,
            )
        elif _bucket_contains(row, "implementation"):
            reply = _variant(
                [
                    "是不是下标从 1 开始就没事？我感觉样例里空格处理也不对。",
                    "我是不是把输入长度和数组范围弄反了？",
                    "是不是循环多跑一位也没关系？我感觉这里猜错了。",
                ],
                ordinal,
            )
        else:
            reply = _variant(
                [
                    "应该是反过来吧？如果 true 就说明这一边不用管了？",
                    "我是不是该把满足的那一边排除掉？感觉有点反。",
                    "是不是只要当前判断成立，就往另一边继续找？我有点乱。",
                ],
                ordinal,
            )
        final_reply = _expand_reply_for_bucket(row, context_type, reply, ordinal)
        return (
            final_reply,
            _evidence_quote(final_reply, ["反", "另一边", "乱", "猜模板", "丢掉"]),
            "学生给出方向性判断，但明显不稳定。",
        )
    if context_type == "followup_after_code_attempt":
        reply = _variant(
            [
                "我按这个想法写了一点，但样例第二个过不去，感觉边界或者更新顺序有问题。",
                "我把这一步写进代码了，不过有一组样例不对，可能是初始化或者下标错了。",
                "我写了个小版本，输出和手算差一点，不知道该先查哪一个中间量。",
            ],
            ordinal,
        )
        final_reply = _expand_reply_for_bucket(row, context_type, reply, ordinal)
        return (
            final_reply,
            _evidence_quote(final_reply, ["样例", "初始化", "中间量", "打印", "代码"]),
            "学生有尝试和错误证据，但还没有定位到最小桥。",
        )
    if context_type == "followup_after_prerequisite_gap":
        if _bucket_contains(row, "implementation"):
            reply = _variant(
                [
                    "我有点懵，这题里空格和字母要怎么对应到次数？输入范围我也没对上。",
                    "我好像连一个字符该怎么处理都没懂，后面的循环就更接不上。",
                    "你说的下标和初值我没太理解，能不能先用题里的一个字符说明？",
                ],
                ordinal,
            )
        elif _bucket_contains(row, "modeling"):
            reply = _variant(
                [
                    "我有点懵，题面里的对象和关系分别是什么？这个词我没太懂。",
                    "我好像连谁和谁有关系都没懂，后面的建模就更接不上。",
                    "你说的关系我没太理解，能不能先说它在题面里对应哪两个对象？",
                ],
                ordinal,
            )
        else:
            reply = _variant(
                [
                    "我有点懵，什么叫可行性/状态语义？这个词我没太懂。",
                    "我好像连这一格记录的“含义”都没懂，后面的公式就更接不上。",
                    "你说的维护量/不变量我没太理解，能不能先说它在题面里对应什么？",
                ],
                ordinal,
            )
        final_reply = _expand_reply_for_bucket(row, context_type, reply, ordinal)
        return (
            final_reply,
            _evidence_quote(final_reply, ["没懂", "没太理解", "接不上", "基础关系", "前置概念"]),
            "学生暴露前置概念断层，需要先补基础概念。",
        )
    if context_type == "policy_direct_answer_special":
        reply = _variant(
            [
                "能不能直接告诉我完整思路或者代码，我赶时间。",
                "我现在有点想直接看完整题解和代码，不然不知道从哪里继续。",
                "可以直接给我完整做法吗？如果不行，就告诉我最该先补哪一步。",
            ],
            ordinal,
        )
        final_reply = _expand_reply_for_bucket(row, context_type, reply, ordinal)
        return (
            final_reply,
            _evidence_quote(final_reply, ["完整", "直接", "代码", "题解", "答案"]),
            "学生请求完整答案，需要安全重定向。",
        )
    return _initial_question(row), "", ""


def _recent_dialogue_bucket(value: str) -> str:
    text = str(value or "").strip()
    if not text or text.upper() == "N/A":
        return "none"
    role_lines = [
        line
        for line in text.splitlines()
        if line.strip().lower().startswith(("student:", "assistant:", "学生：", "ai：", "assistant："))
    ]
    if len(role_lines) >= 4:
        return "long"
    return "short"


def _long_recent_dialogue(row: dict, prior_ai: str, source_index: int) -> tuple[str, str]:
    bridge_question = _current_small_question(row)
    middle_student = _variant(
        [
            "我能说出一点，但还是不知道它怎么落到这题里。",
            "我试着顺了样例，可是到中间那一步就断了。",
            "我感觉自己在套模板，不确定这个判断是不是对应题意。",
        ],
        source_index,
    )
    last_ai = _variant(
        [
            f"那先别扩大范围。继续只看这一小步：{bridge_question}",
            f"好，我们就停在这个局部。你先回答一个更小的问题：{bridge_question}",
            f"先不用写完整式子。现在只补这个关系：{bridge_question}",
        ],
        source_index + 1,
    )
    recent_dialogue = "\n".join(
        [
            f"学生：{_initial_question(row)}",
            f"AI：{prior_ai}",
            f"学生：{middle_student}",
            f"AI：{last_ai}",
        ]
    )
    return recent_dialogue, last_ai


def _code_excerpt_for_context(row: dict, context_type: str, source_index: int) -> str:
    if context_type != "followup_after_code_attempt":
        return str(row.get("student_code_excerpt") or "N/A")
    return "\n".join(
        [
            "for (int i = 1; i <= n; ++i) {",
            "    // 这里按我的理解更新，但样例有一组不对",
            "    update(i);",
            "}",
            f"// case trace id: dialogue_v3_code_{source_index:03d}",
        ]
    )


def _base_output_row(row: dict, source_index: int, context_type: str) -> dict:
    followability = FOLLOWABILITY_BY_CONTEXT[context_type]
    expected_move = EXPECTED_MOVE_BY_CONTEXT[context_type]
    prior_ai = _prior_ai_scaffold(row)
    reply, evidence, uncertainty = _reply_for_context(row, context_type, source_index)
    turn_position = "initial" if context_type == "initial_question" else "followup"
    if turn_position == "initial":
        student_message = _initial_question(row)
        recent_dialogue = str(row.get("recent_dialogue") or "N/A")
        prior_ai = ""
        reply = ""
        evidence = ""
        uncertainty = ""
        context_ai_reply = ""
    else:
        student_message = reply
        if str(row.get("recent_dialogue_bucket") or "") == "long":
            recent_dialogue, context_ai_reply = _long_recent_dialogue(row, prior_ai, source_index)
            prior_ai = context_ai_reply
        else:
            recent_dialogue = f"学生：{_initial_question(row)}\nAI：{prior_ai}"
            context_ai_reply = prior_ai

    source_case_id = str(row.get("case_id") or row.get("id") or f"source_{source_index:03d}")
    bridge_bucket = row.get("bridge_bucket") or row.get("category") or "unknown"
    output = dict(row)
    output.update(
        {
            "case_id": f"dialogue_v3_{source_index:03d}_{bridge_bucket}",
            "id": f"dialogue_v3_{source_index:03d}_{bridge_bucket}",
            "source_case_id": source_case_id,
            "source_dataset": "bridgebench_cp_heldout_v2_50_draft",
            "turn_position": turn_position,
            "context_type": context_type,
            "context_type_zh": CONTEXT_LABEL_ZH[context_type],
            "student_scaffold_followability": followability,
            "student_scaffold_followability_zh": FOLLOWABILITY_LABEL_ZH[followability],
            "followability_label_confidence": CONFIDENCE_BY_CONTEXT[context_type],
            "followability_evidence_quote": evidence,
            "followability_uncertainty_reason": uncertainty,
            "prior_ai_scaffold": prior_ai,
            "student_reply_to_prior_scaffold": reply,
            "expected_tutor_move": expected_move,
            "expected_tutor_move_zh": EXPECTED_MOVE_ZH[expected_move],
            "fixed_recent_dialogue_source": "synthetic_dialogue_state_v3",
            "student_message": student_message,
            "student_message_length_bucket": _student_message_length_bucket(student_message),
            "recent_dialogue": recent_dialogue,
            "recent_dialogue_bucket": _recent_dialogue_bucket(recent_dialogue),
            "context_ai_reply": context_ai_reply,
            "student_code_excerpt": _code_excerpt_for_context(row, context_type, source_index),
            "reference_label_status": "draft_needs_coach_review",
            "review_notes_for_coach": (
                f"Dialogue-state v3 draft. 请复核 followability={followability}、"
                f"expected_tutor_move={expected_move} 与目标 missing bridge 是否一致。"
            ),
        }
    )
    return output


def build_dialogue_state_cases(v2_rows: list[dict]) -> list[dict]:
    if len(v2_rows) < 50:
        raise ValueError(f"Need at least 50 v2 rows, got {len(v2_rows)}")
    rows = []
    for source_index, (row, context_type) in enumerate(zip(v2_rows[:50], CONTEXT_PLAN, strict=True), 1):
        rows.append(_base_output_row(row, source_index, context_type))
    return rows


def _join_cell(value) -> str:
    if isinstance(value, list):
        return "\n".join(str(item) for item in value)
    return str(value or "")


def _style_review_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    machine_fill = PatternFill("solid", fgColor="F2F5F7")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="17324D")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for cell in sheet[2]:
        cell.font = Font(size=9, color="667085")
        cell.fill = machine_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].hidden = True
    for row_cells in sheet.iter_rows(min_row=3):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {
        "A": 26,
        "B": 24,
        "C": 14,
        "D": 28,
        "E": 18,
        "F": 18,
        "G": 22,
        "H": 40,
        "I": 22,
        "J": 72,
        "K": 52,
        "L": 42,
        "M": 58,
        "N": 54,
        "O": 44,
        "P": 28,
        "Q": 34,
        "R": 48,
        "S": 38,
        "T": 42,
        "U": 46,
        "V": 34,
        "W": 18,
        "X": 18,
        "Y": 20,
        "Z": 20,
        "AA": 20,
        "AB": 22,
        "AC": 22,
        "AD": 22,
        "AE": 18,
        "AF": 26,
        "AG": 44,
        "AH": 18,
        "AI": 18,
        "AJ": 18,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "H3"
    sheet.auto_filter.ref = f"A2:{sheet.cell(row=2, column=len(REVIEW_COLUMNS)).coordinate}"


def _apply_review_validations(sheet, *, language: str) -> None:
    if sheet.max_row < 3:
        return
    validation_choices = REVIEW_VALIDATION_CHOICES if language == "en" else REVIEW_VALIDATION_CHOICES_ZH
    for column_name, choices in validation_choices.items():
        column_index = REVIEW_COLUMNS.index(column_name) + 1
        column_letter = get_column_letter(column_index)
        formula = '"' + ",".join(choices) + '"'
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.error = "Please choose one of the allowed review codes."
        validation.errorTitle = "Invalid review code"
        validation.prompt = "Use the dropdown value so the review can be summarized."
        validation.promptTitle = "Structured review"
        sheet.add_data_validation(validation)
        validation.add(f"{column_letter}3:{column_letter}{sheet.max_row}")


def _headers_for_language(language: str) -> dict[str, str]:
    return REVIEW_HEADERS_EN if language == "en" else REVIEW_HEADERS_ZH


def _display_cell(row: dict, column: str, *, language: str) -> str:
    if language == "zh":
        if column == "turn_position":
            return TURN_POSITION_ZH.get(str(row.get(column) or ""), _join_cell(row.get(column)))
        if column == "context_type":
            return _join_cell(row.get("context_type_zh") or row.get(column))
        if column == "student_scaffold_followability":
            return _join_cell(row.get("student_scaffold_followability_zh") or row.get(column))
        if column == "followability_label_confidence":
            return CONFIDENCE_ZH.get(str(row.get(column) or ""), _join_cell(row.get(column)))
        if column == "expected_tutor_move":
            return _join_cell(row.get("expected_tutor_move_zh") or row.get(column))
    return _join_cell(row.get(column))


def _append_review_rows(sheet, rows: list[dict], *, language: str) -> None:
    headers = _headers_for_language(language)
    sheet.append([headers[column] for column in REVIEW_COLUMNS])
    sheet.append(REVIEW_COLUMNS)
    for row in rows:
        sheet.append([_display_cell(row, column, language=language) for column in REVIEW_COLUMNS])
    _style_review_sheet(sheet)
    _apply_review_validations(sheet, language=language)


def _create_instruction_sheet(workbook: Workbook, *, language: str) -> None:
    sheet = workbook.create_sheet("Instructions" if language == "en" else "评审说明", 0)
    rows = INSTRUCTION_ROWS_EN if language == "en" else INSTRUCTION_ROWS
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 110
    title_fill = PatternFill("solid", fgColor="17324D")
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    sheet["A1"].fill = title_fill
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def export_dialogue_state_review_workbook(rows: list[dict], output_path: Path, *, language: str = "zh") -> None:
    if language not in {"zh", "en"}:
        raise ValueError(f"Unsupported workbook language: {language}")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "总表" if language == "zh" else "dialogue_state_review"
    _append_review_rows(sheet, rows, language=language)
    _create_instruction_sheet(workbook, language=language)

    sheet_names = BUCKET_SHEET_NAMES_ZH if language == "zh" else BUCKET_SHEET_NAMES
    for bucket, sheet_name in sheet_names.items():
        bucket_rows = [row for row in rows if (row.get("bridge_bucket") or row.get("category")) == bucket]
        if not bucket_rows:
            continue
        bucket_sheet = workbook.create_sheet(sheet_name)
        _append_review_rows(bucket_sheet, bucket_rows, language=language)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def build_generation_report(rows: list[dict], *, source_path: Path, output_path: Path) -> dict:
    return {
        "source_jsonl": str(source_path),
        "output_jsonl": str(output_path),
        "row_count": len(rows),
        "turn_position_counts": dict(Counter(row["turn_position"] for row in rows)),
        "context_type_counts": dict(Counter(row["context_type"] for row in rows)),
        "followability_counts": dict(Counter(row["student_scaffold_followability"] for row in rows)),
        "expected_tutor_move_counts": dict(Counter(row["expected_tutor_move"] for row in rows)),
        "low_confidence_count": sum(1 for row in rows if row["followability_label_confidence"] == "low"),
        "reference_label_status": "draft_needs_coach_review",
        "ok": len(rows) == 50,
    }


def _write_report_markdown(report: dict, path: Path, *, language: str) -> None:
    is_zh = language == "zh"
    title = "Dialogue-State v3 50-Case 生成报告" if is_zh else "Dialogue-State v3 50-Case Generation Report"
    note = (
        "本文件是开发阶段草稿报告。v3 case 增加固定后续辅导上下文和 F1-F4 学生跟随状态，仍需教练复核；不作为 gold。"
        if is_zh
        else "This is a development draft report. The v3 cases add fixed follow-up tutoring contexts and F1-F4 scaffold-followability labels; coach review is still required and this is not gold data."
    )
    lines = [
        f"# {title}",
        "",
        note,
        "",
        f"- row_count: {report['row_count']}",
        f"- source_jsonl: `{report['source_jsonl']}`",
        f"- output_jsonl: `{report['output_jsonl']}`",
        f"- reference_label_status: `{report['reference_label_status']}`",
        "",
        "## Distributions" if not is_zh else "## 分布",
        "",
        f"- turn_position_counts: `{report['turn_position_counts']}`",
        f"- context_type_counts: `{report['context_type_counts']}`",
        f"- followability_counts: `{report['followability_counts']}`",
        f"- expected_tutor_move_counts: `{report['expected_tutor_move_counts']}`",
        f"- low_confidence_count: {report['low_confidence_count']}",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate dialogue-state v3 held-out draft cases.")
    parser.add_argument("--source-jsonl", type=Path, default=DEFAULT_SOURCE_JSONL)
    parser.add_argument("--output-jsonl", type=Path, default=DEFAULT_OUTPUT_JSONL)
    parser.add_argument("--review-xlsx", type=Path, default=DEFAULT_REVIEW_XLSX)
    parser.add_argument("--review-xlsx-en", type=Path, default=DEFAULT_REVIEW_XLSX_EN)
    parser.add_argument("--report-json", type=Path, default=DEFAULT_REPORT_JSON)
    parser.add_argument("--report-zh", type=Path, default=DEFAULT_REPORT_ZH)
    parser.add_argument("--report-en", type=Path, default=DEFAULT_REPORT_EN)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(sys.argv[1:] if argv is None else argv)
    source_rows = _load_jsonl(args.source_jsonl)
    rows = build_dialogue_state_cases(source_rows)
    _write_jsonl(args.output_jsonl, rows)
    export_dialogue_state_review_workbook(rows, args.review_xlsx, language="zh")
    export_dialogue_state_review_workbook(rows, args.review_xlsx_en, language="en")
    report = build_generation_report(rows, source_path=args.source_jsonl, output_path=args.output_jsonl)
    args.report_json.parent.mkdir(parents=True, exist_ok=True)
    args.report_json.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    _write_report_markdown(report, args.report_zh, language="zh")
    _write_report_markdown(report, args.report_en, language="en")
    print(json.dumps(report, ensure_ascii=False))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
