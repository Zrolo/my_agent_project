"""Analyze dev-ablation blind-review labels.

The dev-ablation workbook is a repeated-measures review: each case has several
anonymous responses from different tutor/pipeline conditions. This analyzer
merges the Chinese review workbook with the hidden key file, normalizes system
condition names, and reports system summaries plus paired case-level deltas.
"""

from __future__ import annotations

import argparse
import csv
import json
import random
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable


DEFAULT_KEY_CSV = Path(
    "evals/aichat/ad_hoc_runs/dev_ablation_20260512_safe_scaffold_limit10/"
    "coach_response_review_workbook_dev_ablation.key.csv"
)
DEFAULT_OUTPUT_LABELS_JSONL = Path(
    "docs/research/coach_response_review_labels_dev_ablation_safe_scaffold_20260512.jsonl"
)
DEFAULT_OUTPUT_JSON = Path(
    "docs/research/dev_ablation_safe_scaffold_blind_review_analysis_20260512.summary.json"
)
DEFAULT_OUTPUT_MD_ZH = Path(
    "docs/research/dev_ablation_safe_scaffold_blind_review_analysis_20260512.zh.md"
)
DEFAULT_OUTPUT_MD = Path(
    "docs/research/dev_ablation_safe_scaffold_blind_review_analysis_20260512.md"
)

CASE_SPECIFIC_RUBRIC_FIELDS = [
    "success_criteria",
    "forbidden_content",
    "critical_bridge_boundary",
    "acceptable_reveal",
    "expected_student_next_action",
]

REVIEW_FIELD_NAMES = [
    "case_id",
    "anonymized_response_id",
    "problem_ref",
    "problem_source_platform",
    "problem_source_id",
    "problem_source_url",
    "problem_statement",
    "problem_statement_public_summary",
    "problem_statement_rights_note",
    "problem_statement_access_level",
    "student_message",
    "student_message_length_bucket",
    "problem_context",
    "recent_dialogue",
    "context_ai_reply",
    "context_alignment_flag",
    "success_criteria",
    "forbidden_content",
    "critical_bridge_boundary",
    "acceptable_reveal",
    "expected_student_next_action",
    "response_text",
    "coach_overall_quality_score",
    "coach_would_show_to_student",
    "coach_leakage_label",
    "coach_bridge_reveal_justification",
    "coach_scaffold_sufficiency_score",
    "coach_student_response_burden",
    "coach_bridge_identification_score",
    "coach_groundedness_score",
    "coach_scaffold_appropriateness_score",
    "coach_next_step_clarity_score",
    "coach_single_focus_coherence_score",
    "coach_micro_example_applicability",
    "coach_bridge_oriented_micro_example_score",
    "coach_bridge_leakage_control_score",
    "coach_preference_rank",
    "coach_reviewer_confidence",
    "coach_needs_discussion",
    "coach_notes",
    "review_status",
]

SCORE_COLUMN_MAP = {
    "bridge_identification": "coach_bridge_identification_score",
    "groundedness": "coach_groundedness_score",
    "scaffold_appropriateness": "coach_scaffold_appropriateness_score",
    "scaffold_sufficiency": "coach_scaffold_sufficiency_score",
    "bridge_leakage_control": "coach_bridge_leakage_control_score",
    "next_step_clarity": "coach_next_step_clarity_score",
    "single_focus_coherence": "coach_single_focus_coherence_score",
    "bridge_oriented_micro_example": "coach_bridge_oriented_micro_example_score",
}

CORE_SCORE_KEYS = [
    "bridge_identification",
    "groundedness",
    "scaffold_appropriateness",
    "bridge_leakage_control",
    "next_step_clarity",
    "single_focus_coherence",
]

DEFAULT_SYSTEMS = [
    "current_system_deployment",
    "enhanced_prompt_only_clean",
    "socratic_no_answer_clean",
    "codehelp_codeaid_clean",
    "dbox_inspired_clean",
    "dbox_inspired_guard",
    "edf_inspired_clean",
    "edf_inspired_guard",
    "bridge_inspired_expert_decision_clean",
    "single_llm_structured_clean",
    "single_llm_structured_guard",
    "bridge_contract_clean",
    "bridge_contract_guard",
    "bridge_contract_guard_repair",
    "bridge_contract_safe_scaffold",
]

EDF_CORE_SYSTEMS = [
    "enhanced_prompt_only_clean",
    "dbox_inspired_guard",
    "edf_inspired_clean",
    "edf_inspired_guard",
    "bridge_contract_guard",
    "bridge_contract_guard_repair",
]

PROMPT_COMPRESSION_SYSTEMS = [
    "dbox_inspired_guard",
    "bridge_contract_guard",
    "bridge_contract_compact_guard",
    "bridge_contract_minimal_guard",
]

DBOX_BRIDGE_HYBRID_SYSTEMS = [
    "enhanced_prompt_only_clean",
    "dbox_inspired_clean",
    "dbox_inspired_guard",
    "bridge_contract_compact_guard",
    "bridge_guided_dbox_style_guard",
]

DBOX_BRIDGE_HYBRID_FULL_FAIRNESS_SYSTEMS = [
    "enhanced_prompt_only_clean",
    "enhanced_prompt_only_guard",
    "enhanced_prompt_only_guard_repair",
    "dbox_inspired_clean",
    "dbox_inspired_guard",
    "dbox_inspired_guard_repair",
    "bridge_contract_compact_clean",
    "bridge_contract_compact_guard",
    "bridge_contract_compact_guard_repair",
    "bridge_guided_dbox_style_guard",
]

IMPORTANT_COMPARISONS = [
    ("bridge_contract_safe_scaffold", "bridge_contract_clean"),
    ("bridge_contract_safe_scaffold", "bridge_contract_guard"),
    ("bridge_contract_safe_scaffold", "enhanced_prompt_only_clean"),
    ("bridge_contract_safe_scaffold", "dbox_inspired_guard"),
    ("bridge_contract_guard", "dbox_inspired_guard"),
    ("edf_inspired_guard", "edf_inspired_clean"),
    ("edf_inspired_guard", "dbox_inspired_guard"),
    ("bridge_contract_guard", "edf_inspired_guard"),
    ("bridge_contract_guard_repair", "edf_inspired_guard"),
    ("bridge_contract_guard_repair", "bridge_contract_guard"),
    ("bridge_contract_clean", "enhanced_prompt_only_clean"),
    ("dbox_inspired_clean", "bridge_contract_clean"),
    ("dbox_inspired_guard", "dbox_inspired_clean"),
    ("enhanced_prompt_only_guard", "enhanced_prompt_only_clean"),
    ("enhanced_prompt_only_guard_repair", "enhanced_prompt_only_guard"),
    ("dbox_inspired_guard_repair", "dbox_inspired_guard"),
    ("bridge_contract_compact_guard", "bridge_contract_compact_clean"),
    ("bridge_contract_compact_guard_repair", "bridge_contract_compact_guard"),
    ("bridge_contract_compact_guard", "dbox_inspired_guard"),
    ("bridge_contract_compact_guard_repair", "dbox_inspired_guard_repair"),
    ("bridge_guided_dbox_style_guard", "dbox_inspired_guard"),
    ("dbox_inspired_guard", "bridge_guided_dbox_style_guard"),
    ("single_llm_structured_guard", "single_llm_structured_clean"),
]


def systems_for_set(system_set: str) -> list[str]:
    if system_set == "default":
        return list(DEFAULT_SYSTEMS)
    if system_set == "edf_core":
        return list(EDF_CORE_SYSTEMS)
    if system_set == "prompt_compression":
        return list(PROMPT_COMPRESSION_SYSTEMS)
    if system_set == "dbox_bridge_hybrid":
        return list(DBOX_BRIDGE_HYBRID_SYSTEMS)
    if system_set == "dbox_bridge_hybrid_full_fairness":
        return list(DBOX_BRIDGE_HYBRID_FULL_FAIRNESS_SYSTEMS)
    raise ValueError(f"Unsupported system_set: {system_set}")


def _choice_key(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if "｜" in text:
        return text.split("｜", 1)[0].strip()
    return text


def _bool_key(value) -> str:
    return _choice_key(value).lower()


def _float(value, default=0.0) -> float:
    key = _choice_key(value)
    try:
        return float(key)
    except (TypeError, ValueError):
        return default


def _optional_float(value) -> float | None:
    key = _choice_key(value)
    if key == "":
        return None
    try:
        return float(key)
    except (TypeError, ValueError):
        return None


def _mean(values: Iterable[float]) -> float:
    values = list(values)
    return round(sum(values) / len(values), 4) if values else 0.0


def _core_scores(row: dict) -> list[float]:
    scores = row.get("scores") or {}
    return [_float(scores.get(key)) for key in CORE_SCORE_KEYS]


def _micro_score(row: dict) -> float:
    scores = row.get("scores") or {}
    return _float(scores.get("bridge_oriented_micro_example"))


def _micro_score_optional(row: dict) -> float | None:
    scores = row.get("scores") or {}
    value = scores.get("bridge_oriented_micro_example")
    return value if isinstance(value, (float, int)) else None


def _micro_applicable(row: dict) -> bool:
    return row.get("micro_example_applicability") == "applicable"


def _row_with_micro_score(row: dict) -> float:
    values = list(_core_scores(row))
    if _micro_applicable(row):
        micro = _micro_score_optional(row)
        if micro is not None:
            values.append(float(micro))
    return _mean(values)


def student_ready_pass(row: dict) -> bool:
    if _float(row.get("overall_quality_score")) < 4:
        return False
    if row.get("would_show_to_student") != "yes":
        return False
    if row.get("leakage_label") in {"major_bridge_leakage", "answer_leakage"}:
        return False
    return all(score > 0 for score in _core_scores(row))


def student_ready_safe_pass(row: dict) -> bool:
    return student_ready_pass(row) and row.get("leakage_label") == "no_leakage"


def rubric_eval_score_v1(row: dict) -> float:
    """Development-only composite score; never use as the sole paper metric."""

    scores = row.get("scores") or {}
    value = 0.0
    value += min(_float(row.get("overall_quality_score")), 5.0) / 5.0 * 30.0
    value += 20.0 if student_ready_pass(row) else 0.0
    value += min(_float(scores.get("scaffold_sufficiency")), 2.0) / 2.0 * 15.0
    value += min(_float(scores.get("next_step_clarity")), 2.0) / 2.0 * 10.0
    value += min(_float(scores.get("bridge_identification")), 2.0) / 2.0 * 5.0

    leakage_label = row.get("leakage_label")
    if leakage_label == "minor_bridge_leakage":
        value -= 5.0
    elif leakage_label == "major_bridge_leakage":
        value -= 25.0
    elif leakage_label == "answer_leakage":
        value -= 35.0

    show = row.get("would_show_to_student")
    if show == "borderline":
        value -= 5.0
    elif show == "no":
        value -= 15.0

    burden = row.get("student_response_burden")
    if burden == "medium":
        value -= 2.0
    elif burden == "high":
        value -= 8.0

    return round(max(0.0, min(100.0, value)), 4)


def condition_id_from_key(key_row: dict) -> str:
    tutor_mode = str(key_row.get("tutor_mode") or "")
    pipeline_mode = str(key_row.get("pipeline_mode") or "")

    if tutor_mode == "current_system":
        return "current_system_deployment"
    if tutor_mode == "enhanced_prompt_only":
        if pipeline_mode == "tutor_plus_guard":
            return "enhanced_prompt_only_guard"
        if pipeline_mode == "tutor_plus_guard_plus_repair":
            return "enhanced_prompt_only_guard_repair"
        return "enhanced_prompt_only_clean"
    if tutor_mode == "socratic_no_answer_tutor":
        return "socratic_no_answer_clean"
    if tutor_mode == "codehelp_codeaid_no_direct_solution_tutor":
        return "codehelp_codeaid_clean"
    if tutor_mode == "dbox_inspired_decomposition_tutor":
        if pipeline_mode == "tutor_plus_guard":
            return "dbox_inspired_guard"
        if pipeline_mode == "tutor_plus_guard_plus_repair":
            return "dbox_inspired_guard_repair"
        return "dbox_inspired_clean"
    if tutor_mode == "bridge_guided_dbox_style_tutor":
        if pipeline_mode == "tutor_plus_guard":
            return "bridge_guided_dbox_style_guard"
        if pipeline_mode == "tutor_plus_guard_plus_repair":
            return "bridge_guided_dbox_style_guard_repair"
        return "bridge_guided_dbox_style_clean"
    if tutor_mode == "edf_inspired_adaptive_scaffolding_tutor":
        if pipeline_mode == "tutor_plus_guard":
            return "edf_inspired_guard"
        if pipeline_mode == "tutor_plus_guard_plus_repair":
            return "edf_inspired_guard_repair"
        return "edf_inspired_clean"
    if tutor_mode == "bridge_inspired_expert_decision_tutor":
        return "bridge_inspired_expert_decision_clean"
    if tutor_mode == "single_llm_structured":
        if pipeline_mode == "tutor_plus_guard":
            return "single_llm_structured_guard"
        if pipeline_mode == "tutor_plus_guard_plus_repair":
            return "single_llm_structured_guard_repair"
        return "single_llm_structured_clean"
    if tutor_mode == "bridge_contract":
        if pipeline_mode == "deterministic_safe_scaffold":
            return "bridge_contract_safe_scaffold"
        if pipeline_mode == "tutor_plus_guard":
            return "bridge_contract_guard"
        if pipeline_mode == "tutor_plus_guard_plus_repair":
            return "bridge_contract_guard_repair"
        return "bridge_contract_clean"
    if tutor_mode == "bridge_contract_compact":
        if pipeline_mode == "tutor_plus_guard":
            return "bridge_contract_compact_guard"
        if pipeline_mode == "tutor_plus_guard_plus_repair":
            return "bridge_contract_compact_guard_repair"
        return "bridge_contract_compact_clean"
    if tutor_mode == "bridge_contract_minimal":
        if pipeline_mode == "tutor_plus_guard":
            return "bridge_contract_minimal_guard"
        if pipeline_mode == "tutor_plus_guard_plus_repair":
            return "bridge_contract_minimal_guard_repair"
        return "bridge_contract_minimal_clean"
    return f"{tutor_mode}__{pipeline_mode}".strip("_")


def load_review_workbook_rows(path: Path | str) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read .xlsx review workbooks") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    if "盲评表" not in workbook.sheetnames:
        raise ValueError(f"{path}: expected sheet named '盲评表'")
    sheet = workbook["盲评表"]
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[1]]
    output = []
    for raw_row in rows[2:]:
        row = dict(zip(headers, raw_row))
        if row.get("case_id") and row.get("anonymized_response_id"):
            output.append(row)
    return output


def load_key_rows(path: Path | str) -> list[dict]:
    with Path(path).open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def merge_review_and_key_rows(review_rows: list[dict], key_rows: list[dict]) -> list[dict]:
    key_by_response_id = {row["anonymized_response_id"]: row for row in key_rows}
    merged = []
    for review in review_rows:
        response_id = str(review.get("anonymized_response_id") or "")
        key = key_by_response_id.get(response_id)
        if not key:
            raise ValueError(f"missing key row for anonymized_response_id={response_id}")
        scores = {}
        for score_name, column_name in SCORE_COLUMN_MAP.items():
            raw_value = review.get(column_name)
            if score_name == "bridge_oriented_micro_example":
                scores[score_name] = _optional_float(raw_value)
            else:
                scores[score_name] = _float(raw_value)
        condition_id = condition_id_from_key(key)
        merged.append(
            {
                "case_id": str(review.get("case_id") or key.get("case_id") or ""),
                "anonymized_response_id": response_id,
                "problem_ref": str(review.get("problem_ref") or ""),
                "problem_source_platform": str(review.get("problem_source_platform") or ""),
                "problem_source_id": str(review.get("problem_source_id") or ""),
                "problem_source_url": str(review.get("problem_source_url") or ""),
                "problem_statement": str(review.get("problem_statement") or ""),
                "problem_statement_public_summary": str(
                    review.get("problem_statement_public_summary") or ""
                ),
                "problem_statement_rights_note": str(
                    review.get("problem_statement_rights_note") or ""
                ),
                "problem_statement_access_level": str(
                    review.get("problem_statement_access_level") or ""
                ),
                "student_message": str(review.get("student_message") or ""),
                "student_message_length_bucket": str(
                    review.get("student_message_length_bucket") or ""
                ),
                "problem_context": str(review.get("problem_context") or ""),
                "recent_dialogue": str(review.get("recent_dialogue") or ""),
                "context_ai_reply": str(review.get("context_ai_reply") or ""),
                "context_alignment_flag": _choice_key(review.get("context_alignment_flag")),
                "success_criteria": str(review.get("success_criteria") or ""),
                "forbidden_content": str(review.get("forbidden_content") or ""),
                "critical_bridge_boundary": str(review.get("critical_bridge_boundary") or ""),
                "acceptable_reveal": str(review.get("acceptable_reveal") or ""),
                "expected_student_next_action": str(
                    review.get("expected_student_next_action") or ""
                ),
                "response_text": str(review.get("response_text") or ""),
                "condition_id": condition_id,
                "system": condition_id,
                "tutor_mode": str(key.get("tutor_mode") or ""),
                "guard_mode": str(key.get("guard_mode") or ""),
                "pipeline_mode": str(key.get("pipeline_mode") or ""),
                "tutor_model_provider": str(key.get("tutor_model_provider") or ""),
                "chat_thinking_mode": str(key.get("chat_thinking_mode") or ""),
                "final_response_source": str(key.get("final_response_source") or ""),
                "repair_applied": _bool_key(key.get("repair_applied") or "false"),
                "blocked": _bool_key(key.get("blocked") or "false"),
                "scores": scores,
                "micro_example_applicability": _choice_key(
                    review.get("coach_micro_example_applicability")
                ),
                "leakage_label": _choice_key(review.get("coach_leakage_label")),
                "bridge_reveal_justification": _choice_key(
                    review.get("coach_bridge_reveal_justification")
                ),
                "preference_rank": _float(review.get("coach_preference_rank"), default=0.0),
                "overall_quality_score": _float(
                    review.get("coach_overall_quality_score"), default=0.0
                ),
                "would_show_to_student": _choice_key(review.get("coach_would_show_to_student")),
                "student_response_burden": _choice_key(
                    review.get("coach_student_response_burden")
                ),
                "reviewer_confidence": _choice_key(review.get("coach_reviewer_confidence")),
                "needs_discussion": _choice_key(review.get("coach_needs_discussion")),
                "notes": str(review.get("coach_notes") or ""),
                "review_status": _choice_key(review.get("review_status")),
            }
        )
    return merged


def _rows_by_system(rows: list[dict]) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        grouped[str(row.get("condition_id") or row.get("system") or "")].append(row)
    return grouped


def _rows_by_case_and_system(rows: list[dict]) -> dict[str, dict[str, dict]]:
    grouped: dict[str, dict[str, dict]] = defaultdict(dict)
    for row in rows:
        grouped[str(row.get("case_id"))][str(row.get("condition_id") or row.get("system"))] = row
    return grouped


def build_system_summary(rows: list[dict], *, systems: list[str] | None = None) -> dict:
    systems = systems or DEFAULT_SYSTEMS
    grouped = _rows_by_system(rows)
    summary = {}
    for system in systems:
        system_rows = grouped.get(system, [])
        leakage_counts = Counter(row.get("leakage_label") for row in system_rows)
        reveal_counts = Counter(row.get("bridge_reveal_justification") for row in system_rows)
        burden_counts = Counter(row.get("student_response_burden") for row in system_rows)
        show_counts = Counter(row.get("would_show_to_student") for row in system_rows)
        applicable_rows = [row for row in system_rows if _micro_applicable(row)]
        applicable_scored_rows = [
            row for row in applicable_rows if _micro_score_optional(row) is not None
        ]
        summary[system] = {
            "n": len(system_rows),
            "core6_mean": _mean(_mean(_core_scores(row)) for row in system_rows),
            "with_micro7_mean": _mean(_row_with_micro_score(row) for row in system_rows),
            "rubric_eval_score_v1_mean": _mean(
                rubric_eval_score_v1(row) for row in system_rows
            ),
            "bridge_identification_mean": _mean(
                _float((row.get("scores") or {}).get("bridge_identification"))
                for row in system_rows
            ),
            "groundedness_mean": _mean(
                _float((row.get("scores") or {}).get("groundedness"))
                for row in system_rows
            ),
            "scaffold_appropriateness_mean": _mean(
                _float((row.get("scores") or {}).get("scaffold_appropriateness"))
                for row in system_rows
            ),
            "next_step_clarity_mean": _mean(
                _float((row.get("scores") or {}).get("next_step_clarity"))
                for row in system_rows
            ),
            "single_focus_coherence_mean": _mean(
                _float((row.get("scores") or {}).get("single_focus_coherence"))
                for row in system_rows
            ),
            "micro_applicable_mean": _mean(
                float(_micro_score_optional(row)) for row in applicable_scored_rows
            ),
            "micro_applicable_count": len(applicable_rows),
            "micro_scored_count": len(applicable_scored_rows),
            "scaffold_sufficiency_mean": _mean(
                _float((row.get("scores") or {}).get("scaffold_sufficiency"))
                for row in system_rows
                if "scaffold_sufficiency" in (row.get("scores") or {})
            ),
            "overall_quality_mean": _mean(
                _float(row.get("overall_quality_score")) for row in system_rows
            ),
            "rank_mean": _mean(_float(row.get("preference_rank")) for row in system_rows),
            "rank1_count": sum(1 for row in system_rows if _float(row.get("preference_rank")) == 1),
            "student_ready_pass_count": sum(1 for row in system_rows if student_ready_pass(row)),
            "student_ready_safe_pass_count": sum(
                1 for row in system_rows if student_ready_safe_pass(row)
            ),
            "would_show_yes_count": show_counts.get("yes", 0),
            "would_show_borderline_count": show_counts.get("borderline", 0),
            "would_show_no_count": show_counts.get("no", 0),
            "no_leakage_count": leakage_counts.get("no_leakage", 0),
            "minor_bridge_leakage_count": leakage_counts.get("minor_bridge_leakage", 0),
            "major_bridge_leakage_count": leakage_counts.get("major_bridge_leakage", 0),
            "answer_leakage_count": leakage_counts.get("answer_leakage", 0),
            "major_or_answer_leakage_count": leakage_counts.get("major_bridge_leakage", 0)
            + leakage_counts.get("answer_leakage", 0),
            "no_reveal_count": reveal_counts.get("no_reveal", 0),
            "pedagogically_justified_reveal_count": reveal_counts.get("pedagogically_justified", 0),
            "borderline_reveal_count": reveal_counts.get("borderline", 0),
            "unjustified_reveal_count": reveal_counts.get("unjustified", 0),
            "student_response_burden_low_count": burden_counts.get("low", 0),
            "student_response_burden_medium_count": burden_counts.get("medium", 0),
            "student_response_burden_high_count": burden_counts.get("high", 0),
            "repair_applied_count": sum(1 for row in system_rows if row.get("repair_applied") == "true"),
            "blocked_count": sum(1 for row in system_rows if row.get("blocked") == "true"),
            "needs_discussion_count": sum(1 for row in system_rows if row.get("needs_discussion") == "yes"),
            "notes_count": sum(1 for row in system_rows if str(row.get("notes") or "").strip()),
        }
    return summary


def _metric_value(row: dict, metric: str) -> float:
    if metric == "overall_quality_score":
        return _float(row.get("overall_quality_score"))
    if metric == "core6_mean":
        return _mean(_core_scores(row))
    if metric == "with_micro7_mean":
        return _row_with_micro_score(row)
    if metric == "rubric_eval_score_v1":
        return rubric_eval_score_v1(row)
    if metric == "student_ready_pass":
        return 1.0 if student_ready_pass(row) else 0.0
    if metric == "student_ready_safe_pass":
        return 1.0 if student_ready_safe_pass(row) else 0.0
    raise ValueError(f"unsupported metric: {metric}")


def _bootstrap_ci95(diffs: list[float], *, rounds: int = 2000, seed: int = 20260512) -> list[float]:
    if not diffs:
        return [0.0, 0.0]
    rng = random.Random(seed)
    means = []
    n = len(diffs)
    for _ in range(rounds):
        sample = [diffs[rng.randrange(n)] for _ in range(n)]
        means.append(sum(sample) / n)
    means.sort()
    low = means[int(0.025 * (rounds - 1))]
    high = means[int(0.975 * (rounds - 1))]
    return [round(low, 4), round(high, 4)]


def compute_paired_comparisons(
    rows: list[dict],
    *,
    systems: list[str] | None = None,
    metric: str = "overall_quality_score",
    bootstrap_rounds: int = 2000,
) -> dict:
    systems = systems or DEFAULT_SYSTEMS
    grouped = _rows_by_case_and_system(rows)
    comparisons = {}
    for system_a in systems:
        for system_b in systems:
            if system_a == system_b:
                continue
            diffs = []
            wins = ties = losses = 0
            case_diffs = {}
            for case_id, case_rows in grouped.items():
                if system_a not in case_rows or system_b not in case_rows:
                    continue
                diff = round(
                    _metric_value(case_rows[system_a], metric)
                    - _metric_value(case_rows[system_b], metric),
                    4,
                )
                diffs.append(diff)
                case_diffs[case_id] = diff
                if diff > 0:
                    wins += 1
                elif diff < 0:
                    losses += 1
                else:
                    ties += 1
            comparisons[f"{system_a}__vs__{system_b}"] = {
                "system_a": system_a,
                "system_b": system_b,
                "metric": metric,
                "paired_cases": len(diffs),
                "wins": wins,
                "ties": ties,
                "losses": losses,
                "mean_diff": _mean(diffs),
                "bootstrap_ci95": _bootstrap_ci95(diffs, rounds=bootstrap_rounds),
                "case_diffs": case_diffs,
            }
    return comparisons


def _rank1_by_case(rows: list[dict]) -> dict[str, list[str]]:
    grouped = _rows_by_case_and_system(rows)
    output = {}
    for case_id, case_rows in grouped.items():
        output[case_id] = sorted(
            system
            for system, row in case_rows.items()
            if _float(row.get("preference_rank"), default=0) == 1
        )
    return output


def _has_case_specific_rubric(row: dict) -> bool:
    return any(str(row.get(field) or "").strip() for field in CASE_SPECIFIC_RUBRIC_FIELDS)


def build_payload(rows: list[dict], systems: list[str] | None = None) -> dict:
    systems = systems or DEFAULT_SYSTEMS
    row_count_by_system = Counter(str(row.get("condition_id") or row.get("system") or "") for row in rows)
    return {
        "row_count": len(rows),
        "case_count": len({row.get("case_id") for row in rows}),
        "systems": systems,
        "row_count_by_system": dict(sorted(row_count_by_system.items())),
        "overall_leakage_counts": dict(Counter(row.get("leakage_label") for row in rows)),
        "bridge_reveal_justification_counts": dict(
            Counter(row.get("bridge_reveal_justification") for row in rows if row.get("bridge_reveal_justification"))
        ),
        "student_response_burden_counts": dict(
            Counter(row.get("student_response_burden") for row in rows if row.get("student_response_burden"))
        ),
        "overall_show_counts": dict(Counter(row.get("would_show_to_student") for row in rows)),
        "case_specific_rubric_present_count": sum(1 for row in rows if _has_case_specific_rubric(row)),
        "system_summary": build_system_summary(rows, systems=systems),
        "paired_comparisons": compute_paired_comparisons(rows, systems=systems),
        "paired_comparisons_core6": compute_paired_comparisons(
            rows, systems=systems, metric="core6_mean"
        ),
        "paired_comparisons_micro7": compute_paired_comparisons(
            rows, systems=systems, metric="with_micro7_mean"
        ),
        "rank1_by_case": _rank1_by_case(rows),
    }


def _table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    lines.extend("| " + " | ".join(str(value) for value in row) + " |" for row in rows)
    return "\n".join(lines)


def _system_table(payload: dict) -> list[list[object]]:
    rows = []
    for system, item in payload["system_summary"].items():
        rows.append(
            [
                system,
                item["n"],
                item["overall_quality_mean"],
                item["core6_mean"],
                item["scaffold_sufficiency_mean"],
                item["with_micro7_mean"],
                item["rank1_count"],
                item["student_ready_pass_count"],
                item["student_ready_safe_pass_count"],
                f"{item['would_show_yes_count']}/{item['would_show_borderline_count']}/{item['would_show_no_count']}",
                f"{item['no_leakage_count']}/{item['minor_bridge_leakage_count']}/{item['major_or_answer_leakage_count']}",
                f"{item['pedagogically_justified_reveal_count']}/{item['borderline_reveal_count']}/{item['unjustified_reveal_count']}",
                f"{item['student_response_burden_low_count']}/{item['student_response_burden_medium_count']}/{item['student_response_burden_high_count']}",
            ]
        )
    return rows


def _primary_outcome_table(payload: dict) -> list[list[object]]:
    rows = []
    for system, item in payload["system_summary"].items():
        rows.append(
            [
                system,
                item.get("n", 0),
                item.get("overall_quality_mean", 0.0),
                item.get("rubric_eval_score_v1_mean", 0.0),
                item.get("student_ready_pass_count", 0),
                item.get("student_ready_safe_pass_count", 0),
                item.get("major_or_answer_leakage_count", 0),
                item.get("scaffold_sufficiency_mean", 0.0),
                f"{item.get('student_response_burden_low_count', 0)}/{item.get('student_response_burden_medium_count', 0)}/{item.get('student_response_burden_high_count', 0)}",
            ]
        )
    return rows


def _diagnostic_table(payload: dict) -> list[list[object]]:
    rows = []
    for system, item in payload["system_summary"].items():
        rows.append(
            [
                system,
                item.get("n", 0),
                item.get("core6_mean", 0.0),
                item.get("bridge_identification_mean", 0.0),
                item.get("groundedness_mean", 0.0),
                item.get("scaffold_appropriateness_mean", 0.0),
                item.get("next_step_clarity_mean", 0.0),
                item.get("single_focus_coherence_mean", 0.0),
                f"{item.get('micro_applicable_mean', 0.0)} ({item.get('micro_scored_count', 0)}/{item.get('micro_applicable_count', 0)})",
            ]
        )
    return rows


def _important_comparison_rows(payload: dict) -> list[list[object]]:
    overall = payload["paired_comparisons"]
    core6 = payload["paired_comparisons_core6"]
    micro7 = payload["paired_comparisons_micro7"]
    rows = []
    for system_a, system_b in IMPORTANT_COMPARISONS:
        key = f"{system_a}__vs__{system_b}"
        if key not in overall:
            continue
        item = overall[key]
        rows.append(
            [
                f"{system_a} - {system_b}",
                item["paired_cases"],
                item["mean_diff"],
                core6[key]["mean_diff"],
                micro7[key]["mean_diff"],
                f"{item['wins']}/{item['ties']}/{item['losses']}",
                item["bootstrap_ci95"],
            ]
        )
    return rows


def _major_leakage_rows(payload_rows: list[dict]) -> list[list[object]]:
    rows = []
    for row in payload_rows:
        if row.get("leakage_label") not in {"major_bridge_leakage", "answer_leakage"}:
            continue
        notes = str(row.get("notes") or "").replace("\n", " ").strip()
        if len(notes) > 90:
            notes = notes[:87] + "..."
        rows.append(
            [
                row.get("case_id"),
                row.get("condition_id"),
                row.get("overall_quality_score"),
                row.get("would_show_to_student"),
                row.get("bridge_reveal_justification") or "",
                notes,
            ]
        )
    return rows


def _active_system_count(payload: dict) -> int:
    summary = payload.get("system_summary") or {}
    active = [system for system in payload.get("systems", []) if summary.get(system, {}).get("n", 0)]
    return len(active) or len(payload.get("systems", []))


def _is_ai_prelim_review(rows: list[dict]) -> bool:
    return any(
        "ai_prelim" in str(row.get("review_status") or "").lower()
        or "ai 预评" in str(row.get("review_status") or "").lower()
        or "ai self-review" in str(row.get("notes") or "").lower()
        for row in rows
    )


def _safe_scaffold_note_zh(payload: dict) -> str:
    safe_summary = (payload.get("system_summary") or {}).get("bridge_contract_safe_scaffold") or {}
    safe_n = safe_summary.get("n", 0)
    if safe_n:
        return (
            f"- `bridge_contract_safe_scaffold` {safe_summary.get('no_leakage_count', 0)}/{safe_n} 都没有泄露，"
            f"但 overall={safe_summary.get('overall_quality_mean', 0.0)}，"
            f"{safe_summary.get('would_show_no_count', 0)}/{safe_n} 都不建议给学生看；"
            "它应作为兜底 fallback / appendix condition，不适合做常规 tutor baseline。"
        )
    return "- 本轮没有包含 `bridge_contract_safe_scaffold` 常规条件；safe scaffold 只作为 block/fallback 机制和 appendix 候选。"


def _safe_scaffold_note_en(payload: dict) -> str:
    safe_summary = (payload.get("system_summary") or {}).get("bridge_contract_safe_scaffold") or {}
    safe_n = safe_summary.get("n", 0)
    if safe_n:
        return (
            f"- `bridge_contract_safe_scaffold` had {safe_summary.get('no_leakage_count', 0)}/{safe_n} no-leakage labels, "
            f"but overall={safe_summary.get('overall_quality_mean', 0.0)} and "
            f"{safe_summary.get('would_show_no_count', 0)}/{safe_n} no-show labels; "
            "it is a fallback / appendix condition, not a normal tutor baseline."
        )
    return "- This run did not include `bridge_contract_safe_scaffold` as a normal condition; safe scaffold is used only as a block/fallback mechanism and appendix candidate."


def render_report_zh(payload: dict, rows: list[dict], *, review_xlsx: Path, key_csv: Path) -> str:
    ready_leader = max(
        payload["system_summary"].items(),
        key=lambda item: (
            item[1]["student_ready_pass_count"],
            item[1]["overall_quality_mean"],
            -item[1]["major_or_answer_leakage_count"],
        ),
    )
    major_by_system = {
        system: item["major_or_answer_leakage_count"]
        for system, item in payload["system_summary"].items()
        if item["major_or_answer_leakage_count"]
    }
    is_edf_core = payload.get("systems") == EDF_CORE_SYSTEMS
    is_ai_prelim = _is_ai_prelim_review(rows)
    active_condition_count = _active_system_count(payload)
    title_prefix = "EDF-inspired" if is_edf_core else "Dev Ablation"
    review_kind = "AI 初评" if is_ai_prelim else "盲评"
    title = f"# {title_prefix} {payload['case_count']}-case {review_kind}分析"
    scope_sentence = (
        f"本报告分析 {payload['case_count']} 个 dev cases、{active_condition_count} 个匿名系统条件、"
        f"{payload['row_count']} 条 AI 预评回复。它用于 EDF-inspired baseline 开发阶段筛查，"
        "不是教练 gold label，也不作为最终 held-out test 结论。"
        if is_edf_core and is_ai_prelim
        else f"本报告分析 {payload['case_count']} 个 dev cases、{active_condition_count} 个匿名系统条件、"
        f"{payload['row_count']} 条 {'AI 预评' if is_ai_prelim else '盲审'}回复。它用于开发阶段决策和 prompt/regression 修订，"
        "不作为最终 held-out test 结论。"
    )
    return "\n\n".join(
        [
            title,
            scope_sentence,
            f"- 盲审表：`{review_xlsx.name}`（本地填写文件，未提交仓库）\n- 匿名 key：`{key_csv}`",
            (
                f"- 已标注回复：{payload['row_count']} 条\n"
                f"- case 数：{payload['case_count']} 个\n"
                f"- 总体泄露标签：{payload['overall_leakage_counts']}\n"
                f"- 关键桥透露正当性：{payload['bridge_reveal_justification_counts']}\n"
                f"- 学生回复负担：{payload['student_response_burden_counts']}\n"
                f"- 是否愿意给学生看：{payload['overall_show_counts']}\n"
                f"- 含 v3 case-specific rubric 的回复：{payload.get('case_specific_rubric_present_count', 0)} / {payload['row_count']}"
            ),
            "## 主指标",
            "主表只保留少数结果指标；`rubric_eval_score_v1` 仅用于开发阶段筛选，不作为论文唯一结论。",
            _table(
                [
                    "condition",
                    "n",
                    "overall",
                    "rubric_score_dev",
                    "ready",
                    "safe_ready",
                    "major+answer",
                    "sufficiency",
                    "burden low/medium/high",
                ],
                _primary_outcome_table(payload),
            ),
            "## 诊断指标",
            "这些细项用于 error analysis、prompt 修订和 LLM Judge 校准；其中 micro-example 只在适用且已评分时计入。",
            _table(
                [
                    "condition",
                    "n",
                    "core6",
                    "bridge_id",
                    "groundedness",
                    "appropriateness",
                    "next_step",
                    "single_focus",
                    "micro mean (scored/applicable)",
                ],
                _diagnostic_table(payload),
            ),
            "## 关键配对比较",
            _table(
                ["comparison", "cases", "Δ overall", "Δ core6", "Δ micro7", "W/T/L", "CI95"],
                _important_comparison_rows(payload),
            ),
            "## Major/Answer Leakage Case Memo 草稿",
            _table(
                ["case", "condition", "overall", "show", "reveal justification", "coach note"],
                _major_leakage_rows(rows),
            ),
            "## 初步观察",
            "\n".join(
                [
                    f"- `student_ready_pass` 最高的是 `{ready_leader[0]}`：{ready_leader[1]['student_ready_pass_count']} / {ready_leader[1]['n']}。",
                    "- 若 `含 v3 case-specific rubric` 数量为 0，本报告只是 v3 指标重分析；case-specific grader 校准不能把这批旧表当作完整证据。",
                    (
                        "- 本轮是 EDF core 条件集：重点看 EDF 是否接近 enhanced prompt / DBox+Guard，以及 Guard 是否改善 EDF。"
                        if is_edf_core
                        else _safe_scaffold_note_zh(payload)
                    ),
                    "- `major_or_answer_leakage_count` 仍然需要按 case 回看，尤其要检查它是否来自过完整 micro-example、直接补完关键桥，还是 reviewer 对局部代码/状态关系的判定。",
                    "- 新增 `bridge_reveal_justification` 用于区分“有信息量的合理教学透露”和“过早/过完整的无正当性关键桥泄露”，避免把所有关键知识讲解都惩罚为泄露。",
                    "- 新增 `scaffold_sufficiency` 用于识别“低泄露但帮助不足”的过度保留回复；正式比较时应把 leakage 和 sufficiency 放在一起看。",
                    "- `学生回复负担` 只作为交互成本指标，不进入 core score；它用于识别学生需要输入过多而导致在线体验变差的回复。",
                    f"- 出现 major/answer 级泄露的条件：{major_by_system or '无'}。",
                    "- Guard/Repair 的因果效果不能仅凭不同 run 的均值解释；正式判断仍需要 same-candidate before/after repair stress。",
                    (
                        "- 本轮的作用是筛查 EDF-inspired 是否值得进入 appendix/dev 继续观察；按 AI 预评结果，它暂时弱于 DBox+Guard 和 Bridge Contract+Guard+Repair。"
                        if is_edf_core
                        else "- 本轮的作用是筛选 50-case 主实验条件：强 prompt、DBox-inspired、DBox+Guard、Bridge Contract、Safe Scaffold 是否值得保留，要看它们在 paired W/T/L 和 student-ready pass 上是否稳定。"
                    ),
                ]
            ),
            "## 解释边界",
            (
                f"这批 {payload['case_count']}-case 是 dev/regression set。可以据此修 prompt、修 rubric、决定主实验条件；"
                "但不能据此声称某个系统显著优于强 baseline。正式论文 headline 仍需要冻结 prompt/grader 后的 50-case held-out、部分双教练标注和 judge calibration。"
            ),
        ]
    ) + "\n"


def render_report_en(payload: dict, rows: list[dict], *, review_xlsx: Path, key_csv: Path) -> str:
    ready_leader = max(
        payload["system_summary"].items(),
        key=lambda item: (
            item[1]["student_ready_pass_count"],
            item[1]["overall_quality_mean"],
            -item[1]["major_or_answer_leakage_count"],
        ),
    )
    is_edf_core = payload.get("systems") == EDF_CORE_SYSTEMS
    is_ai_prelim = _is_ai_prelim_review(rows)
    active_condition_count = _active_system_count(payload)
    title_prefix = "EDF-inspired" if is_edf_core else "Dev Ablation"
    review_kind = "AI Preliminary Review" if is_ai_prelim else "Blind Review"
    title = f"# {title_prefix} {payload['case_count']}-case {review_kind} Analysis"
    scope_sentence = (
        f"This report analyzes {payload['case_count']} development cases, {active_condition_count} anonymous system conditions, "
        f"and {payload['row_count']} AI-prelim-reviewed responses. It is EDF-inspired baseline development evidence, "
        "not coach gold labels or a final held-out result."
        if is_edf_core and is_ai_prelim
        else f"This report analyzes {payload['case_count']} development cases, {active_condition_count} anonymous system conditions, "
        f"and {payload['row_count']} {'AI-prelim-reviewed' if is_ai_prelim else 'blind-reviewed'} responses. "
        "It is development evidence, not a final held-out result."
    )
    return "\n\n".join(
        [
            title,
            scope_sentence,
            f"- Review workbook: `{review_xlsx.name}` (local filled file, not committed)\n- Anonymous key: `{key_csv}`",
            (
                f"- Reviewed responses: {payload['row_count']}\n"
                f"- Cases: {payload['case_count']}\n"
                f"- Leakage labels: {payload['overall_leakage_counts']}\n"
                f"- Bridge reveal justification labels: {payload['bridge_reveal_justification_counts']}\n"
                f"- Student response burden labels: {payload['student_response_burden_counts']}\n"
                f"- Show-to-student labels: {payload['overall_show_counts']}\n"
                f"- Responses with v3 case-specific rubrics: {payload.get('case_specific_rubric_present_count', 0)} / {payload['row_count']}"
            ),
            "## Primary Outcomes",
            "`rubric_eval_score_v1` is a development-only triage score and must not be used as the sole paper conclusion.",
            _table(
                [
                    "condition",
                    "n",
                    "overall",
                    "rubric_score_dev",
                    "ready",
                    "safe_ready",
                    "major+answer",
                    "sufficiency",
                    "burden low/medium/high",
                ],
                _primary_outcome_table(payload),
            ),
            "## Diagnostic Dimensions",
            "These dimensions support error analysis, prompt revision, and LLM-judge calibration. Micro-example scores are counted only when applicable and filled.",
            _table(
                [
                    "condition",
                    "n",
                    "core6",
                    "bridge_id",
                    "groundedness",
                    "appropriateness",
                    "next_step",
                    "single_focus",
                    "micro mean (scored/applicable)",
                ],
                _diagnostic_table(payload),
            ),
            "## Key Paired Comparisons",
            _table(
                ["comparison", "cases", "Δ overall", "Δ core6", "Δ micro7", "W/T/L", "CI95"],
                _important_comparison_rows(payload),
            ),
            "## Major/Answer Leakage Case Memo Draft",
            _table(
                ["case", "condition", "overall", "show", "reveal justification", "coach note"],
                _major_leakage_rows(rows),
            ),
            "## Development Notes",
            "\n".join(
                [
                    f"- The highest `student_ready_pass` condition is `{ready_leader[0]}`: {ready_leader[1]['student_ready_pass_count']} / {ready_leader[1]['n']}.",
                    "- If the case-specific rubric count is 0, this report is a v3 metric reanalysis only; it should not be used as full evidence for case-specific grader calibration.",
                    (
                        "- This is the EDF core condition set: use it to compare EDF against enhanced prompt / DBox+Guard and to check whether Guard improves EDF."
                        if is_edf_core
                        else _safe_scaffold_note_en(payload)
                    ),
                    "- Major/answer leakage cases should be inspected qualitatively because fully worked micro-examples can leak the critical bridge without giving code.",
                    "- The new bridge-reveal justification field separates useful instructional information from unjustified early completion of the student's current missing bridge.",
                    "- The new scaffold-sufficiency field flags over-withholding: low leakage is not pedagogical success if the response is safe but not useful.",
                    "- Student-response burden is an interaction-cost signal, not part of the core score.",
                    "- Guard/Repair causal effects still require same-candidate before/after repair stress; between-run averages are not enough.",
                    (
                        "- Use this run to decide whether EDF-inspired should remain an appendix/dev baseline; in this AI-prelim review it trails DBox+Guard and Bridge Contract+Guard+Repair."
                        if is_edf_core
                        else "- Use this run to decide which conditions enter the frozen 50-case held-out study, not as a headline result."
                    ),
                ]
            ),
        ]
    ) + "\n"


def write_jsonl(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--review-xlsx", required=True, type=Path)
    parser.add_argument("--key-csv", default=DEFAULT_KEY_CSV, type=Path)
    parser.add_argument("--output-labels-jsonl", default=DEFAULT_OUTPUT_LABELS_JSONL, type=Path)
    parser.add_argument("--output-json", default=DEFAULT_OUTPUT_JSON, type=Path)
    parser.add_argument("--output-md-zh", default=DEFAULT_OUTPUT_MD_ZH, type=Path)
    parser.add_argument("--output-md", default=DEFAULT_OUTPUT_MD, type=Path)
    parser.add_argument(
        "--system-set",
        choices=[
            "default",
            "edf_core",
            "prompt_compression",
            "dbox_bridge_hybrid",
            "dbox_bridge_hybrid_full_fairness",
        ],
        default="default",
        help="Named system list for summary tables and paired comparisons.",
    )
    args = parser.parse_args(argv)

    review_rows = load_review_workbook_rows(args.review_xlsx)
    key_rows = load_key_rows(args.key_csv)
    merged_rows = merge_review_and_key_rows(review_rows, key_rows)
    payload = build_payload(merged_rows, systems=systems_for_set(args.system_set))

    write_jsonl(merged_rows, args.output_labels_jsonl)
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    write_text(
        args.output_md_zh,
        render_report_zh(payload, merged_rows, review_xlsx=args.review_xlsx, key_csv=args.key_csv),
    )
    write_text(
        args.output_md,
        render_report_en(payload, merged_rows, review_xlsx=args.review_xlsx, key_csv=args.key_csv),
    )

    print(f"Wrote labels: {args.output_labels_jsonl}")
    print(f"Wrote summary: {args.output_json}")
    print(f"Wrote zh report: {args.output_md_zh}")
    print(f"Wrote en report: {args.output_md}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
