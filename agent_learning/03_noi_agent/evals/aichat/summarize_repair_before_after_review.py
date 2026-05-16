import argparse
import csv
import json
import sys
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook


DEFAULT_WORKBOOK_CSV = Path("docs/research/coach_response_review_workbook_repair_before_after_20260510.csv")
DEFAULT_KEY_CSV = Path("docs/research/coach_response_review_workbook_repair_before_after_20260510.key.csv")
DEFAULT_LABELS_PATH = Path("docs/research/coach_response_review_labels_repair_before_after_20260510.jsonl")
DEFAULT_OUTPUT_JSON = Path("docs/research/repair_before_after_review_analysis_20260510.summary.json")
DEFAULT_OUTPUT_MD = Path("docs/research/repair_before_after_review_analysis_20260510.md")
DEFAULT_OUTPUT_MD_ZH = Path("docs/research/repair_before_after_review_analysis_20260510.zh.md")

QUALITY_SCORE = {
    "bad": 0,
    "okay": 1,
    "good": 2,
}

LEAKAGE_SEVERITY = {
    "no_leakage": 0,
    "minor_bridge_leakage": 1,
    "major_bridge_leakage": 2,
    "answer_leakage": 3,
}


def _read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def load_review_rows(path: Path | str) -> list[dict]:
    path = Path(path)
    if not path.exists():
        return []
    if path.suffix.lower() != ".xlsx":
        return _read_csv(path)

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["盲评表"] if "盲评表" in workbook.sheetnames else workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    header_index = None
    headers: list[str] = []
    for index, row in enumerate(rows):
        values = [str(value or "").strip() for value in row]
        if "anonymized_response_id" in values:
            header_index = index
            headers = values
            break
    if header_index is None:
        return []

    output = []
    for row in rows[header_index + 1 :]:
        item = {
            header: "" if value is None else str(value)
            for header, value in zip(headers, row)
            if header
        }
        if any(str(value).strip() for value in item.values()):
            output.append(item)
    return output


def _read_labels(path: Path) -> dict[str, dict]:
    if not path.exists():
        return {}
    if path.suffix.lower() in {".csv", ".xlsx"}:
        rows = load_review_rows(path)
    else:
        rows = []
        for raw_line in path.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if not line:
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError:
                continue
    labels: dict[str, dict] = {}
    for row in rows:
        response_id = str(row.get("anonymized_response_id") or "")
        if response_id:
            labels[response_id] = row
    return labels


def _source_kind(key_row: dict) -> str:
    source = str(key_row.get("final_response_source") or "")
    repair_applied = str(key_row.get("repair_applied") or "").lower() == "true"
    if "candidate_before_repair" in source:
        return "candidate"
    if "repaired" in source or repair_applied:
        return "repaired"
    return ""


def _label_for(response_id: str, labels: dict[str, dict]) -> dict:
    return labels.get(response_id, {})


def _choice_id(value: str) -> str:
    raw = str(value or "").strip()
    if "｜" in raw:
        return raw.split("｜", 1)[0].strip()
    if "|" in raw:
        return raw.split("|", 1)[0].strip()
    return raw


def _score(value: str, mapping: dict[str, int]) -> int | None:
    choice = _choice_id(value)
    if choice in mapping:
        return mapping[choice]
    if choice.isdigit():
        return int(choice)
    return None


def _pick(*values) -> str:
    for value in values:
        if value is not None and str(value).strip():
            return str(value)
    return ""


def _is_labeled(value: str) -> bool:
    return _choice_id(value) == "labeled"


def _student_ready_score(value: str) -> int | None:
    mapping = {"no": 0, "borderline": 1, "yes": 2}
    return mapping.get(_choice_id(value))


def _rate(count: int, total: int) -> float:
    if total <= 0:
        return 0.0
    return count / total


def summarize(
    *,
    workbook_csv: Path = DEFAULT_WORKBOOK_CSV,
    key_csv: Path = DEFAULT_KEY_CSV,
    labels_path: Path = DEFAULT_LABELS_PATH,
) -> dict:
    workbook_rows = {str(row.get("anonymized_response_id") or ""): row for row in load_review_rows(workbook_csv)}
    key_rows = _read_csv(key_csv)
    labels = _read_labels(labels_path)

    cases: dict[str, dict] = {}
    for key_row in key_rows:
        response_id = str(key_row.get("anonymized_response_id") or "")
        case_id = str(key_row.get("case_id") or "")
        kind = _source_kind(key_row)
        if not response_id or not case_id or not kind:
            continue
        review_row = workbook_rows.get(response_id, {})
        label_row = _label_for(response_id, labels)
        overall_quality = _pick(
            label_row.get("overall_quality"),
            label_row.get("coach_overall_quality_score"),
            review_row.get("coach_overall_quality_score"),
            review_row.get("overall_quality"),
        )
        leakage_label = _pick(
            label_row.get("leakage_label"),
            label_row.get("coach_leakage_label"),
            review_row.get("coach_leakage_label"),
            review_row.get("leakage_label"),
        )
        would_show = _pick(
            label_row.get("would_show_to_student"),
            label_row.get("coach_would_show_to_student"),
            review_row.get("coach_would_show_to_student"),
            review_row.get("would_show_to_student"),
        )
        review_status = _pick(
            label_row.get("review_status"),
            review_row.get("review_status"),
            "unlabeled",
        )
        item = {
            "anonymized_response_id": response_id,
            "overall_quality": overall_quality,
            "leakage_label": leakage_label,
            "would_show_to_student": would_show,
            "preference_rank": str(label_row.get("preference_rank") or ""),
            "notes": _pick(label_row.get("notes"), label_row.get("coach_notes"), review_row.get("coach_notes")),
            "review_status": review_status,
            "response_excerpt": str(review_row.get("response_text") or "")[:240],
        }
        cases.setdefault(case_id, {"case_id": case_id})[kind] = item

    pairs = []
    quality_deltas = []
    leakage_deltas = []
    quality_wins = quality_ties = quality_losses = 0
    leakage_improves = leakage_ties = leakage_worse = 0
    candidate_major_or_answer = 0
    repaired_major_or_answer = 0
    labeled_pair_count = 0
    student_ready_improves = student_ready_ties = student_ready_worse = 0

    for case_id in sorted(cases):
        pair = cases[case_id]
        candidate = pair.get("candidate")
        repaired = pair.get("repaired")
        if not candidate or not repaired:
            continue
        candidate_quality = _score(candidate.get("overall_quality", ""), QUALITY_SCORE)
        repaired_quality = _score(repaired.get("overall_quality", ""), QUALITY_SCORE)
        candidate_leakage = _score(candidate.get("leakage_label", ""), LEAKAGE_SEVERITY)
        repaired_leakage = _score(repaired.get("leakage_label", ""), LEAKAGE_SEVERITY)
        candidate_ready = _student_ready_score(candidate.get("would_show_to_student", ""))
        repaired_ready = _student_ready_score(repaired.get("would_show_to_student", ""))
        quality_delta = None
        leakage_delta = None
        student_ready_delta = None
        if candidate_quality is not None and repaired_quality is not None:
            quality_delta = repaired_quality - candidate_quality
            quality_deltas.append(quality_delta)
            if quality_delta > 0:
                quality_wins += 1
            elif quality_delta < 0:
                quality_losses += 1
            else:
                quality_ties += 1
        if candidate_leakage is not None and repaired_leakage is not None:
            leakage_delta = repaired_leakage - candidate_leakage
            leakage_deltas.append(leakage_delta)
            if leakage_delta < 0:
                leakage_improves += 1
            elif leakage_delta > 0:
                leakage_worse += 1
            else:
                leakage_ties += 1
            if candidate_leakage >= LEAKAGE_SEVERITY["major_bridge_leakage"]:
                candidate_major_or_answer += 1
            if repaired_leakage >= LEAKAGE_SEVERITY["major_bridge_leakage"]:
                repaired_major_or_answer += 1
        if candidate_ready is not None and repaired_ready is not None:
            student_ready_delta = repaired_ready - candidate_ready
            if student_ready_delta > 0:
                student_ready_improves += 1
            elif student_ready_delta < 0:
                student_ready_worse += 1
            else:
                student_ready_ties += 1
        if _is_labeled(candidate.get("review_status", "")) and _is_labeled(repaired.get("review_status", "")):
            labeled_pair_count += 1
        pairs.append(
            {
                "case_id": case_id,
                "candidate": candidate,
                "repaired": repaired,
                "quality_delta": quality_delta,
                "leakage_severity_delta": leakage_delta,
                "student_ready_delta": student_ready_delta,
            }
        )

    leakage_total = len(leakage_deltas)
    return {
        "workbook_csv": str(workbook_csv),
        "key_csv": str(key_csv),
        "labels_path": str(labels_path),
        "pair_count": len(pairs),
        "labeled_pair_count": labeled_pair_count,
        "quality": {
            "repaired_wins": quality_wins,
            "ties": quality_ties,
            "repaired_losses": quality_losses,
            "average_delta": round(mean(quality_deltas), 3) if quality_deltas else 0.0,
        },
        "leakage": {
            "repaired_improves": leakage_improves,
            "ties": leakage_ties,
            "repaired_worse": leakage_worse,
            "average_severity_delta": round(mean(leakage_deltas), 3) if leakage_deltas else 0.0,
            "candidate_major_or_answer_rate": round(_rate(candidate_major_or_answer, leakage_total), 3),
            "repaired_major_or_answer_rate": round(_rate(repaired_major_or_answer, leakage_total), 3),
        },
        "student_ready": {
            "repaired_improves": student_ready_improves,
            "ties": student_ready_ties,
            "repaired_worse": student_ready_worse,
        },
        "pairs": pairs,
    }


def render_markdown_zh(result: dict) -> str:
    quality = result.get("quality", {})
    leakage = result.get("leakage", {})
    student_ready = result.get("student_ready", {})
    lines = [
        "# Repair 前后对照盲评分析",
        "",
        "## 数据概况",
        "",
        f"- 成对样本数：{result.get('pair_count', 0)}",
        f"- 已完整标注成对样本数：{result.get('labeled_pair_count', 0)}",
        f"- 标注文件：`{result.get('labels_path', '')}`",
        "",
        "## 汇总结果",
        "",
        "| 指标 | 数值 |",
        "|---|---:|",
        f"| 修复后更好 | {quality.get('repaired_wins', 0)} |",
        f"| 持平 | {quality.get('ties', 0)} |",
        f"| 修复后更差 | {quality.get('repaired_losses', 0)} |",
        f"| 平均质量变化 | {quality.get('average_delta', 0)} |",
        f"| 泄露减轻 | {leakage.get('repaired_improves', 0)} |",
        f"| 泄露持平 | {leakage.get('ties', 0)} |",
        f"| 泄露变重 | {leakage.get('repaired_worse', 0)} |",
        f"| 平均泄露严重度变化 | {leakage.get('average_severity_delta', 0)} |",
        f"| 候选回复重大/答案泄露率 | {leakage.get('candidate_major_or_answer_rate', 0):.1%} |",
        f"| 修复后重大/答案泄露率 | {leakage.get('repaired_major_or_answer_rate', 0):.1%} |",
        f"| 修复后更愿意给学生看 | {student_ready.get('repaired_improves', 0)} |",
        f"| 是否给学生看持平 | {student_ready.get('ties', 0)} |",
        f"| 修复后更不适合给学生看 | {student_ready.get('repaired_worse', 0)} |",
        "",
        "## 逐例备注",
        "",
    ]
    for pair in result.get("pairs", []):
        candidate = pair.get("candidate", {})
        repaired = pair.get("repaired", {})
        lines.extend(
            [
                f"### {pair.get('case_id')}",
                "",
                f"- 质量变化：{pair.get('quality_delta')}",
                f"- 关键桥梁泄露严重度变化：{pair.get('leakage_severity_delta')}",
                f"- 候选回复：{candidate.get('overall_quality', '')} / {candidate.get('leakage_label', '')} / {candidate.get('notes', '')}",
                f"- 修复后：{repaired.get('overall_quality', '')} / {repaired.get('leakage_label', '')} / {repaired.get('notes', '')}",
                "",
            ]
        )
    lines.extend(
        [
            "## 解释口径",
            "",
            "- 质量变化：修复后质量分减去候选回复质量分；新盲评表使用 `1-5` 分，旧标签文件仍兼容 `good=2, okay=1, bad=0`。",
            "- 关键桥梁泄露严重度变化：`no=0, minor=1, major=2, answer=3`，修复后减去候选回复；负数代表泄露减轻。",
            "- 这份报告仍然是单教练盲评分析，不应写成唯一 gold truth。",
        ]
    )
    return "\n".join(lines).rstrip() + "\n"


def render_markdown(result: dict) -> str:
    quality = result.get("quality", {})
    leakage = result.get("leakage", {})
    student_ready = result.get("student_ready", {})
    lines = [
        "# Repair Before/After Blind Review Analysis",
        "",
        "## Data",
        "",
        f"- Paired cases: {result.get('pair_count', 0)}",
        f"- Fully labeled pairs: {result.get('labeled_pair_count', 0)}",
        f"- Label file: `{result.get('labels_path', '')}`",
        "",
        "## Summary",
        "",
        "| Metric | Value |",
        "|---|---:|",
        f"| repaired wins | {quality.get('repaired_wins', 0)} |",
        f"| ties | {quality.get('ties', 0)} |",
        f"| repaired losses | {quality.get('repaired_losses', 0)} |",
        f"| average quality delta | {quality.get('average_delta', 0)} |",
        f"| leakage improves | {leakage.get('repaired_improves', 0)} |",
        f"| leakage ties | {leakage.get('ties', 0)} |",
        f"| leakage worsens | {leakage.get('repaired_worse', 0)} |",
        f"| average leakage severity delta | {leakage.get('average_severity_delta', 0)} |",
        f"| candidate major/answer leakage rate | {leakage.get('candidate_major_or_answer_rate', 0):.1%} |",
        f"| repaired major/answer leakage rate | {leakage.get('repaired_major_or_answer_rate', 0):.1%} |",
        f"| student-ready improves after repair | {student_ready.get('repaired_improves', 0)} |",
        f"| student-ready ties | {student_ready.get('ties', 0)} |",
        f"| student-ready worsens after repair | {student_ready.get('repaired_worse', 0)} |",
        "",
        "## Case Notes",
        "",
    ]
    for pair in result.get("pairs", []):
        candidate = pair.get("candidate", {})
        repaired = pair.get("repaired", {})
        lines.extend(
            [
                f"### {pair.get('case_id')}",
                "",
                f"- Quality delta: {pair.get('quality_delta')}",
                f"- Critical bridge leakage severity delta: {pair.get('leakage_severity_delta')}",
                f"- Candidate: {candidate.get('overall_quality', '')} / {candidate.get('leakage_label', '')} / {candidate.get('notes', '')}",
                f"- Repaired: {repaired.get('overall_quality', '')} / {repaired.get('leakage_label', '')} / {repaired.get('notes', '')}",
                "",
            ]
        )
    lines.extend(
        [
            "## Interpretation",
            "",
            "- Quality delta subtracts candidate quality from repaired quality; new review workbooks use `1-5` scores, while old label files remain compatible with `good=2, okay=1, bad=0`.",
            "- Leakage severity delta maps `no=0, minor=1, major=2, answer=3` and subtracts candidate from repaired; negative means reduced leakage.",
            "- This is a single-coach blind-review analysis, not an absolute gold label.",
        ]
    )
    return "\n".join(lines).rstrip() + "\n"


def write_outputs(
    result: dict,
    *,
    output_json: Path = DEFAULT_OUTPUT_JSON,
    output_md: Path = DEFAULT_OUTPUT_MD,
    output_md_zh: Path = DEFAULT_OUTPUT_MD_ZH,
) -> None:
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    output_md.write_text(render_markdown(result), encoding="utf-8")
    output_md_zh.write_text(render_markdown_zh(result), encoding="utf-8")


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Summarize repair before/after response-review labels.")
    parser.add_argument("--workbook-csv", type=Path, default=DEFAULT_WORKBOOK_CSV)
    parser.add_argument("--key-csv", type=Path, default=DEFAULT_KEY_CSV)
    parser.add_argument("--labels", type=Path, default=DEFAULT_LABELS_PATH)
    parser.add_argument("--output-json", type=Path, default=DEFAULT_OUTPUT_JSON)
    parser.add_argument("--output-md", type=Path, default=DEFAULT_OUTPUT_MD)
    parser.add_argument("--output-md-zh", type=Path, default=DEFAULT_OUTPUT_MD_ZH)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(sys.argv[1:] if argv is None else argv)
    result = summarize(workbook_csv=args.workbook_csv, key_csv=args.key_csv, labels_path=args.labels)
    write_outputs(result, output_json=args.output_json, output_md=args.output_md, output_md_zh=args.output_md_zh)
    print(
        json.dumps(
            {
                "output_json": str(args.output_json),
                "output_md": str(args.output_md),
                "output_md_zh": str(args.output_md_zh),
                "pair_count": result["pair_count"],
                "labeled_pair_count": result["labeled_pair_count"],
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
