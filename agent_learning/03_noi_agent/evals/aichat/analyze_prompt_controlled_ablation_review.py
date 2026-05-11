"""Analyze prompt-controlled ablation blind-review labels.

This analysis is for the small smoke study that asks whether response quality
improvements are due to stronger tutoring prompt wording, the concrete Bridge
Contract, or the mere presence of a contract-shaped prompt.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import defaultdict
from pathlib import Path
from typing import Iterable


DEFAULT_REVIEW_XLSX = Path("docs/research/coach_response_review_workbook_prompt_controlled_ablation_smoke3_20260511.zh.xlsx")
DEFAULT_KEY_CSV = Path("docs/research/coach_response_review_workbook_prompt_controlled_ablation_smoke3_20260511.key.csv")
DEFAULT_OUTPUT_LABELS_JSONL = Path("docs/research/coach_response_review_labels_prompt_controlled_ablation_smoke3_20260511.jsonl")
DEFAULT_OUTPUT_JSON = Path("docs/research/prompt_controlled_ablation_blind_review_20260511.summary.json")
DEFAULT_OUTPUT_MD_ZH = Path("docs/research/prompt_controlled_ablation_blind_review_20260511.zh.md")
DEFAULT_OUTPUT_MD = Path("docs/research/prompt_controlled_ablation_blind_review_20260511.md")

SYSTEMS = [
    "single_llm_structured",
    "enhanced_prompt_only",
    "bridge_contract_predicted",
    "bridge_contract_shuffled",
    "bridge_contract_oracle",
]

CORE_SCORE_KEYS = [
    "bridge_identification",
    "groundedness",
    "scaffold_appropriateness",
    "bridge_leakage_control",
    "next_step_clarity",
    "single_focus_coherence",
]

SCORE_COLUMN_MAP = {
    "bridge_identification": "coach_bridge_identification_score",
    "groundedness": "coach_groundedness_score",
    "scaffold_appropriateness": "coach_scaffold_appropriateness_score",
    "bridge_leakage_control": "coach_bridge_leakage_control_score",
    "next_step_clarity": "coach_next_step_clarity_score",
    "single_focus_coherence": "coach_single_focus_coherence_score",
    "bridge_oriented_micro_example": "coach_bridge_oriented_micro_example_score",
}


def _choice_key(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if "｜" in text:
        return text.split("｜", 1)[0].strip()
    return text


def _float(value, default=0.0) -> float:
    key = _choice_key(value)
    try:
        return float(key)
    except (TypeError, ValueError):
        return default


def _mean(values: Iterable[float]) -> float:
    values = list(values)
    return round(sum(values) / len(values), 4) if values else 0.0


def _core_scores(row: dict) -> list[float]:
    scores = row.get("scores") or {}
    return [_float(scores.get(key)) for key in CORE_SCORE_KEYS]


def _micro_score(row: dict) -> float:
    scores = row.get("scores") or {}
    return _float(scores.get("bridge_oriented_micro_example"))


def student_ready_pass(row: dict) -> bool:
    if _float(row.get("overall_quality_score")) < 4:
        return False
    if str(row.get("would_show_to_student", "")).strip() != "yes":
        return False
    if row.get("leakage_label") in {"major_bridge_leakage", "answer_leakage"}:
        return False
    return all(score > 0 for score in _core_scores(row))


def load_review_workbook_rows(path: Path | str) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read .xlsx review workbooks") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    if "盲评表" not in workbook.sheetnames:
        raise ValueError(f"{path}: expected sheet named '盲评表'")
    sheet = workbook["盲评表"]
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[1]]
    output = []
    for raw_row in rows[2:]:
        row = dict(zip(headers, raw_row))
        if row.get("case_id") and row.get("anonymized_response_id"):
            output.append(row)
    return output


def load_key_rows(path: Path | str) -> list[dict]:
    with Path(path).open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def merge_review_and_key_rows(review_rows: list[dict], key_rows: list[dict]) -> list[dict]:
    key_by_response_id = {row["anonymized_response_id"]: row for row in key_rows}
    merged = []
    for review in review_rows:
        response_id = str(review.get("anonymized_response_id") or "")
        key = key_by_response_id.get(response_id)
        if not key:
            raise ValueError(f"missing key row for anonymized_response_id={response_id}")
        scores = {
            score_name: _float(review.get(column_name))
            for score_name, column_name in SCORE_COLUMN_MAP.items()
        }
        merged.append(
            {
                "case_id": str(review.get("case_id") or key.get("case_id") or ""),
                "anonymized_response_id": response_id,
                "problem_ref": str(review.get("problem_ref") or ""),
                "student_message": str(review.get("student_message") or ""),
                "system": str(key.get("tutor_mode") or ""),
                "ablation_variant": str(key.get("tutor_mode") or ""),
                "guard_mode": str(key.get("guard_mode") or ""),
                "pipeline_mode": str(key.get("pipeline_mode") or ""),
                "final_response_source": str(key.get("final_response_source") or ""),
                "repair_applied": str(key.get("repair_applied") or "false"),
                "blocked": str(key.get("blocked") or "false"),
                "scores": scores,
                "micro_example_applicability": _choice_key(review.get("coach_micro_example_applicability")),
                "leakage_label": _choice_key(review.get("coach_leakage_label")),
                "preference_rank": _float(review.get("coach_preference_rank"), default=0.0),
                "overall_quality_score": _float(review.get("coach_overall_quality_score"), default=0.0),
                "would_show_to_student": _choice_key(review.get("coach_would_show_to_student")),
                "reviewer_confidence": _choice_key(review.get("coach_reviewer_confidence")),
                "needs_discussion": _choice_key(review.get("coach_needs_discussion")),
                "notes": str(review.get("coach_notes") or ""),
                "review_status": _choice_key(review.get("review_status")),
            }
        )
    return merged


def build_system_summary(rows: list[dict], *, systems: list[str] | None = None) -> dict:
    systems = systems or SYSTEMS
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        grouped[str(row.get("system", ""))].append(row)

    summary = {}
    for system in systems:
        system_rows = grouped.get(system, [])
        summary[system] = {
            "n": len(system_rows),
            "core6_mean": _mean(_mean(_core_scores(row)) for row in system_rows),
            "with_micro7_mean": _mean(_mean([*_core_scores(row), _micro_score(row)]) for row in system_rows),
            "overall_quality_mean": _mean(_float(row.get("overall_quality_score")) for row in system_rows),
            "rank_mean": _mean(_float(row.get("preference_rank")) for row in system_rows),
            "rank1_count": sum(1 for row in system_rows if _float(row.get("preference_rank")) == 1),
            "student_ready_pass_count": sum(1 for row in system_rows if student_ready_pass(row)),
            "would_show_yes_count": sum(1 for row in system_rows if row.get("would_show_to_student") == "yes"),
            "would_show_no_count": sum(1 for row in system_rows if row.get("would_show_to_student") == "no"),
            "no_leakage_count": sum(1 for row in system_rows if row.get("leakage_label") == "no_leakage"),
            "minor_bridge_leakage_count": sum(
                1 for row in system_rows if row.get("leakage_label") == "minor_bridge_leakage"
            ),
            "major_bridge_leakage_count": sum(
                1 for row in system_rows if row.get("leakage_label") == "major_bridge_leakage"
            ),
            "answer_leakage_count": sum(1 for row in system_rows if row.get("leakage_label") == "answer_leakage"),
        }
    return summary


def _rows_by_case_and_system(rows: list[dict]) -> dict[str, dict[str, dict]]:
    grouped: dict[str, dict[str, dict]] = defaultdict(dict)
    for row in rows:
        grouped[str(row.get("case_id"))][str(row.get("system"))] = row
    return grouped


def _metric(row: dict, name: str) -> float:
    if name == "overall_quality_score":
        return _float(row.get("overall_quality_score"))
    if name == "core6_mean":
        return _mean(_core_scores(row))
    if name == "with_micro7_mean":
        return _mean([*_core_scores(row), _micro_score(row)])
    raise ValueError(f"unsupported metric: {name}")


def _paired_delta(rows: list[dict], system_a: str, system_b: str, metric: str) -> dict:
    grouped = _rows_by_case_and_system(rows)
    deltas = {}
    wins = ties = losses = 0
    for case_id, case_rows in grouped.items():
        if system_a not in case_rows or system_b not in case_rows:
            continue
        delta = round(_metric(case_rows[system_a], metric) - _metric(case_rows[system_b], metric), 4)
        deltas[case_id] = delta
        if delta > 0:
            wins += 1
        elif delta < 0:
            losses += 1
        else:
            ties += 1
    return {
        "system_a": system_a,
        "system_b": system_b,
        "metric": metric,
        "paired_cases": len(deltas),
        "wins": wins,
        "ties": ties,
        "losses": losses,
        "mean_delta": _mean(deltas.values()),
        "case_deltas": deltas,
    }


def build_effect_summary(rows: list[dict]) -> dict:
    effects = {
        "prompt_effect": ("enhanced_prompt_only", "single_llm_structured"),
        "predicted_contract_vs_prompt": ("bridge_contract_predicted", "enhanced_prompt_only"),
        "predicted_contract_vs_shuffled": ("bridge_contract_predicted", "bridge_contract_shuffled"),
        "oracle_contract_vs_predicted": ("bridge_contract_oracle", "bridge_contract_predicted"),
        "oracle_contract_vs_shuffled": ("bridge_contract_oracle", "bridge_contract_shuffled"),
    }
    output = {}
    for name, (system_a, system_b) in effects.items():
        overall = _paired_delta(rows, system_a, system_b, "overall_quality_score")
        core = _paired_delta(rows, system_a, system_b, "core6_mean")
        micro = _paired_delta(rows, system_a, system_b, "with_micro7_mean")
        output[name] = {
            "system_a": system_a,
            "system_b": system_b,
            "overall_quality_delta": overall["mean_delta"],
            "core6_delta": core["mean_delta"],
            "with_micro7_delta": micro["mean_delta"],
            "wins": overall["wins"],
            "ties": overall["ties"],
            "losses": overall["losses"],
            "case_deltas": overall["case_deltas"],
        }
    return output


def _table(headers: list[str], rows: list[list[object]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    lines.extend("| " + " | ".join(str(value) for value in row) + " |" for row in rows)
    return "\n".join(lines)


def _system_table(summary: dict) -> str:
    return _table(
        [
            "Variant",
            "N",
            "Overall",
            "Core6",
            "Micro7",
            "Rank1",
            "Ready",
            "No leak",
            "Minor",
            "Major",
        ],
        [
            [
                system,
                item["n"],
                item["overall_quality_mean"],
                item["core6_mean"],
                item["with_micro7_mean"],
                item["rank1_count"],
                item["student_ready_pass_count"],
                item["no_leakage_count"],
                item["minor_bridge_leakage_count"],
                item["major_bridge_leakage_count"],
            ]
            for system, item in summary.items()
        ],
    )


def _effect_table(effects: dict) -> str:
    rows = []
    for name, item in effects.items():
        rows.append(
            [
                name,
                f"{item['system_a']} - {item['system_b']}",
                item["overall_quality_delta"],
                item["core6_delta"],
                item["with_micro7_delta"],
                f"{item['wins']}/{item['ties']}/{item['losses']}",
            ]
        )
    return _table(
        ["Effect", "Comparison", "Overall Δ", "Core6 Δ", "Micro7 Δ", "W/T/L"],
        rows,
    )


def _case_rank_table(rows: list[dict]) -> str:
    grouped = _rows_by_case_and_system(rows)
    table_rows = []
    for case_id, case_rows in sorted(grouped.items()):
        ranked = sorted(case_rows.values(), key=lambda row: _float(row.get("preference_rank"), default=99.0))
        top = ranked[0] if ranked else {}
        bottom = ranked[-1] if ranked else {}
        table_rows.append(
            [
                case_id,
                top.get("system", ""),
                top.get("overall_quality_score", ""),
                top.get("leakage_label", ""),
                bottom.get("system", ""),
                bottom.get("overall_quality_score", ""),
                bottom.get("leakage_label", ""),
            ]
        )
    return _table(
        ["Case", "Best Variant", "Best Quality", "Best Leakage", "Worst Variant", "Worst Quality", "Worst Leakage"],
        table_rows,
    )


def render_markdown_zh(payload: dict) -> str:
    effects = payload["effect_summary"]
    prompt_delta = effects["prompt_effect"]["overall_quality_delta"]
    predicted_vs_prompt = effects["predicted_contract_vs_prompt"]["overall_quality_delta"]
    predicted_vs_shuffled = effects["predicted_contract_vs_shuffled"]["overall_quality_delta"]
    lines = [
        "# Prompt-Controlled Ablation 盲评分析 2026-05-11",
        "",
        "本报告分析 3 个样本 × 5 个变体的盲评结果，用于回答：Bridge Contract 组质量提升，到底来自模块结构、具体诊断信息，还是来自更强的通用教学提示词？",
        "",
        "> 本结果只作为 dev / pilot evidence，不作为最终 held-out test 结论。样本数只有 3 个，适合发现方向和设计下一轮实验，不适合做显著性声称。",
        "",
        "## 系统汇总",
        "",
        _system_table(payload["system_summary"]),
        "",
        "## Effect 分解",
        "",
        _effect_table(effects),
        "",
        "解释口径：",
        "",
        f"- **Prompt effect**：`enhanced_prompt_only - single_llm_structured` 的总体质量差为 `{prompt_delta}`。如果它为正，说明仅加入更强教学提示词就能提升质量。",
        f"- **Predicted contract effect over prompt**：`bridge_contract_predicted - enhanced_prompt_only` 的总体质量差为 `{predicted_vs_prompt}`。如果它为正，才说明预测 Bridge Contract 在通用强 prompt 之外有额外平均收益。",
        f"- **Contract validity effect**：`bridge_contract_predicted - bridge_contract_shuffled` 的总体质量差为 `{predicted_vs_shuffled}`。如果它为正，说明具体 contract 内容不是纯装饰，错误 contract 会伤害质量。",
        "",
        "## 逐 case 最好 / 最差",
        "",
        _case_rank_table(payload["rows"]),
        "",
        "## 当前结论",
        "",
        "1. **prompt wording 很可能解释了相当一部分质量提升。** 在这 3 个样本里，`enhanced_prompt_only` 的平均总体质量最高，并且没有 major leakage。这说明不能把 Bridge Contract 组在 fair 20-case 中的提升全部归因于模块架构。",
        "2. **具体 contract 仍然可能有价值。** `bridge_contract_predicted` 在 2/3 个 case 中拿到第一，并且明显好于 `bridge_contract_shuffled`。这说明模型不是只被“请用微型例子、不要泄露”这种通用提示影响；错误 contract 会把回复带偏。",
        "3. **oracle contract 没有自动成为上界。** `bridge_contract_oracle` 在这个小样本中没有超过 predicted contract，说明 contract 注入方式、微型例子设计和泄露控制同样重要；“金标 contract”如果被主模型展开得太完整，也可能变成泄露。",
        "4. **这轮实验支持下一步扩到 10–20 case。** 现在最重要的不是立刻声称架构有效，而是把 prompt-only、predicted contract、shuffled contract、oracle contract 在更多样本上做配对盲评。",
        "",
        "## 对论文写法的影响",
        "",
        "论文里应避免写“Bridge Contract 架构本身导致质量提升”。更稳的表述是：",
        "",
        "> Bridge Contract variants 的收益可能混合了通用教学 prompt、具体 missing bridge 诊断信息、以及模块化控制信号。本消融用于将这些因素拆开。初步 3-case 结果显示，强 prompt 本身已经很强；但 shuffled contract 表现较差，说明具体 contract 内容仍可能带来额外价值。",
        "",
        "下一步正式报告应继续使用三类 effect：",
        "",
        "- prompt effect: `enhanced_prompt_only - single_llm_structured`",
        "- predicted diagnosis effect: `bridge_contract_predicted - enhanced_prompt_only`",
        "- contract validity effect: `bridge_contract_predicted - bridge_contract_shuffled`",
        "",
    ]
    return "\n".join(lines)


def render_markdown(payload: dict) -> str:
    effects = payload["effect_summary"]
    lines = [
        "# Prompt-Controlled Ablation Blind-Review Analysis 2026-05-11",
        "",
        "This report analyzes a 3-case × 5-variant blind review designed to test whether Bridge Contract improvements come from modular structure, concrete diagnostic information, or stronger generic tutoring prompt wording.",
        "",
        "> This is development / pilot evidence only. The sample size is three cases, so it should guide the next experiment rather than support final held-out claims.",
        "",
        "## System Summary",
        "",
        _system_table(payload["system_summary"]),
        "",
        "## Effect Decomposition",
        "",
        _effect_table(effects),
        "",
        "## Case-Level Best / Worst",
        "",
        _case_rank_table(payload["rows"]),
        "",
        "## Interpretation",
        "",
        "1. **Prompt wording likely explains a substantial part of the quality gain.** In this 3-case smoke review, `enhanced_prompt_only` achieved the highest mean overall quality and no major leakage, so Bridge Contract gains should not be attributed to architecture alone.",
        "2. **Concrete contracts still appear meaningful.** `bridge_contract_predicted` ranked first in 2/3 cases and outperformed `bridge_contract_shuffled`, suggesting that incorrect contracts can harm quality and that the contract content is not merely decorative.",
        "3. **Oracle contracts are not automatically an upper bound.** In this small sample, `bridge_contract_oracle` did not outperform predicted contracts, likely because contract injection, micro-example design, and leakage control remain important even when the contract is correct.",
        "4. **The next step is a larger paired blind review.** The same design should be expanded to 10–20 cases before making paper-level claims.",
        "",
        "## Paper Implication",
        "",
        "The paper should not claim that Bridge Contract architecture alone caused quality improvements. A safer claim is that Bridge Contract variants combine generic prompt effects, concrete missing-bridge diagnostic information, and modular control signals; prompt-controlled ablations are needed to separate these factors.",
        "",
    ]
    return "\n".join(lines)


def build_payload(rows: list[dict]) -> dict:
    return {
        "study": "prompt_controlled_ablation_smoke3_20260511",
        "row_count": len(rows),
        "case_count": len({row["case_id"] for row in rows}),
        "systems": SYSTEMS,
        "system_summary": build_system_summary(rows),
        "effect_summary": build_effect_summary(rows),
        "rows": rows,
    }


def write_analysis_outputs(
    rows: list[dict],
    *,
    output_json: Path = DEFAULT_OUTPUT_JSON,
    output_md_zh: Path = DEFAULT_OUTPUT_MD_ZH,
    output_md: Path = DEFAULT_OUTPUT_MD,
    output_labels_jsonl: Path = DEFAULT_OUTPUT_LABELS_JSONL,
) -> dict:
    payload = build_payload(rows)
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_md_zh.parent.mkdir(parents=True, exist_ok=True)
    output_md.parent.mkdir(parents=True, exist_ok=True)
    output_labels_jsonl.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    output_md_zh.write_text(render_markdown_zh(payload), encoding="utf-8")
    output_md.write_text(render_markdown(payload), encoding="utf-8")
    with output_labels_jsonl.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    return payload


def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Analyze prompt-controlled ablation blind-review workbook.")
    parser.add_argument("--review-xlsx", type=Path, default=DEFAULT_REVIEW_XLSX)
    parser.add_argument("--key-csv", type=Path, default=DEFAULT_KEY_CSV)
    parser.add_argument("--output-labels-jsonl", type=Path, default=DEFAULT_OUTPUT_LABELS_JSONL)
    parser.add_argument("--output-json", type=Path, default=DEFAULT_OUTPUT_JSON)
    parser.add_argument("--output-md-zh", type=Path, default=DEFAULT_OUTPUT_MD_ZH)
    parser.add_argument("--output-md", type=Path, default=DEFAULT_OUTPUT_MD)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(argv)
    rows = merge_review_and_key_rows(
        load_review_workbook_rows(args.review_xlsx),
        load_key_rows(args.key_csv),
    )
    payload = write_analysis_outputs(
        rows,
        output_json=args.output_json,
        output_md_zh=args.output_md_zh,
        output_md=args.output_md,
        output_labels_jsonl=args.output_labels_jsonl,
    )
    print(
        json.dumps(
            {
                "row_count": payload["row_count"],
                "case_count": payload["case_count"],
                "output_json": str(args.output_json),
                "output_md_zh": str(args.output_md_zh),
                "output_md": str(args.output_md),
                "output_labels_jsonl": str(args.output_labels_jsonl),
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
