"""Summarize same-candidate Repair stress-review labels.

Expected input is the CSV/XLSX template produced by
build_repair_same_candidate_stress_workbook.py after reviewers fill the before
and after label columns. The script is robust to incomplete review sheets and
will report how many pairs are still unlabeled.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional dependency.
    load_workbook = None  # type: ignore[assignment]


DEFAULT_INPUT = Path(
    "evals/aichat/ad_hoc_runs/dialogue_state_v3_main_20260516_merged/human_review/"
    "repair_same_candidate_stress_candidate_list_20260517.csv"
)
DEFAULT_BLIND_KEY = Path(
    "evals/aichat/ad_hoc_runs/dialogue_state_v3_main_20260516_merged/human_review/"
    "repair_same_candidate_stress_blind_review_key_20260517.csv"
)
DEFAULT_OUTPUT_JSON = Path(
    "evals/aichat/ad_hoc_runs/dialogue_state_v3_main_20260516_merged/human_review/"
    "repair_same_candidate_stress_summary_20260517.json"
)
DEFAULT_OUTPUT_MD = Path("docs/research/repair_same_candidate_stress_summary_20260517.md")

LEAKAGE_SEVERITY = {
    "no_leakage": 0,
    "minor_bridge_leakage": 1,
    "major_bridge_leakage": 2,
    "answer_leakage": 3,
}
BURDEN_SCORE = {"low": 0, "medium": 1, "high": 2}


def _choice(value: object) -> str:
    raw = str(value or "").strip()
    if "｜" in raw:
        raw = raw.split("｜", 1)[0].strip()
    if "|" in raw:
        raw = raw.split("|", 1)[0].strip()
    return raw


def _read_csv(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _read_xlsx(path: Path) -> list[dict]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to read xlsx input")
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["盲评表"] if "盲评表" in workbook.sheetnames else workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    output = []
    for row in rows[1:]:
        item = {
            header: "" if value is None else str(value)
            for header, value in zip(headers, row)
            if header
        }
        if any(str(value).strip() for value in item.values()):
            output.append(item)
    return output


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() == ".xlsx":
        return _read_xlsx(path)
    return _read_csv(path)


def read_key(path: Path | None) -> dict[str, dict]:
    if path is None or not path.exists():
        return {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return {str(row.get("stress_id") or ""): row for row in csv.DictReader(handle)}


def normalize_blind_rows(rows: list[dict], key_rows: dict[str, dict]) -> list[dict]:
    if not key_rows:
        return rows
    normalized = []
    for row in rows:
        stress_id = str(row.get("stress_id") or "")
        key = key_rows.get(stress_id)
        if not key:
            normalized.append(row)
            continue

        def pick(prefix: str, field: str) -> str:
            return row.get(f"{prefix}_{field}", "")

        a_source = str(key.get("response_A_source") or "")
        b_source = str(key.get("response_B_source") or "")
        source_to_prefix = {a_source: "A", b_source: "B"}
        before_prefix = source_to_prefix.get("before_repair")
        after_prefix = source_to_prefix.get("after_repair")
        if not before_prefix or not after_prefix:
            normalized.append(row)
            continue
        item = dict(row)
        item.update(
            {
                "source_condition": key.get("source_condition", ""),
                "source_run": key.get("source_run", ""),
                "before_overall_quality_score": pick(before_prefix, "overall_quality_score"),
                "before_would_show_to_student": pick(before_prefix, "would_show_to_student"),
                "before_leakage_label": pick(before_prefix, "leakage_label"),
                "before_student_response_burden": pick(before_prefix, "student_response_burden"),
                "after_overall_quality_score": pick(after_prefix, "overall_quality_score"),
                "after_would_show_to_student": pick(after_prefix, "would_show_to_student"),
                "after_leakage_label": pick(after_prefix, "leakage_label"),
                "after_student_response_burden": pick(after_prefix, "student_response_burden"),
                "after_too_vague": pick(after_prefix, "too_vague"),
            }
        )
        pair_preference = _choice(row.get("pair_preference"))
        if pair_preference in {"A", "B"}:
            preferred_source = key.get(f"response_{pair_preference}_source")
            if preferred_source == "after_repair":
                item["repair_pair_preference"] = "repair_preferred"
            elif preferred_source == "before_repair":
                item["repair_pair_preference"] = "original_preferred"
            else:
                item["repair_pair_preference"] = "unclear"
        elif pair_preference == "tie":
            item["repair_pair_preference"] = "tie"
        elif pair_preference:
            item["repair_pair_preference"] = "unclear"
        normalized.append(item)
    return normalized


def _score_float(value: object) -> float | None:
    text = _choice(value)
    try:
        return float(text)
    except ValueError:
        return None


def _score_map(value: object, mapping: dict[str, int]) -> int | None:
    text = _choice(value)
    return mapping.get(text)


def _is_labeled(row: dict) -> bool:
    required = [
        "before_overall_quality_score",
        "before_leakage_label",
        "before_student_response_burden",
        "after_overall_quality_score",
        "after_leakage_label",
        "after_student_response_burden",
    ]
    return all(_choice(row.get(field)) for field in required)


def summarize(rows: list[dict]) -> dict:
    labeled = [row for row in rows if _is_labeled(row)]
    deltas = []
    wins = ties = losses = 0
    leakage_improved = leakage_same = leakage_worse = 0
    burden_improved = burden_same = burden_worse = 0
    still_leaks = 0
    too_vague = 0
    pair_preference = Counter()
    before_leakage_distribution = Counter()
    after_leakage_distribution = Counter()
    before_burden_distribution = Counter()
    after_burden_distribution = Counter()
    before_show_distribution = Counter()
    after_show_distribution = Counter()
    before_overall_values = []
    after_overall_values = []
    by_family: dict[str, Counter] = {}
    by_source: dict[str, dict] = {}

    for row in labeled:
        before_overall = _score_float(row.get("before_overall_quality_score"))
        after_overall = _score_float(row.get("after_overall_quality_score"))
        before_leak = _score_map(row.get("before_leakage_label"), LEAKAGE_SEVERITY)
        after_leak = _score_map(row.get("after_leakage_label"), LEAKAGE_SEVERITY)
        before_burden = _score_map(row.get("before_student_response_burden"), BURDEN_SCORE)
        after_burden = _score_map(row.get("after_student_response_burden"), BURDEN_SCORE)
        family = str(row.get("bridge_bucket") or "unknown")
        family_counter = by_family.setdefault(family, Counter())
        source_condition = str(row.get("source_condition") or "unknown")
        source_counter = by_source.setdefault(
            source_condition,
            {
                "n": 0,
                "overall_delta_sum": 0.0,
                "overall_delta_count": 0,
                "quality_win": 0,
                "quality_tie": 0,
                "quality_loss": 0,
                "leakage_improved": 0,
                "leakage_same": 0,
                "leakage_worse": 0,
                "burden_improved": 0,
                "burden_same": 0,
                "burden_worse": 0,
                "repair_preferred": 0,
                "original_preferred": 0,
                "tie": 0,
                "unclear": 0,
            },
        )
        source_counter["n"] += 1

        if before_overall is not None and after_overall is not None:
            before_overall_values.append(before_overall)
            after_overall_values.append(after_overall)
            delta = after_overall - before_overall
            deltas.append(delta)
            source_counter["overall_delta_sum"] += delta
            source_counter["overall_delta_count"] += 1
            if delta > 0:
                wins += 1
                family_counter["quality_win"] += 1
                source_counter["quality_win"] += 1
            elif delta < 0:
                losses += 1
                family_counter["quality_loss"] += 1
                source_counter["quality_loss"] += 1
            else:
                ties += 1
                family_counter["quality_tie"] += 1
                source_counter["quality_tie"] += 1
        if before_leak is not None and after_leak is not None:
            before_leakage_distribution[_choice(row.get("before_leakage_label"))] += 1
            after_leakage_distribution[_choice(row.get("after_leakage_label"))] += 1
            if after_leak < before_leak:
                leakage_improved += 1
                family_counter["leakage_improved"] += 1
                source_counter["leakage_improved"] += 1
            elif after_leak > before_leak:
                leakage_worse += 1
                family_counter["leakage_worse"] += 1
                source_counter["leakage_worse"] += 1
            else:
                leakage_same += 1
                family_counter["leakage_same"] += 1
                source_counter["leakage_same"] += 1
            if after_leak >= LEAKAGE_SEVERITY["major_bridge_leakage"]:
                still_leaks += 1
        if before_burden is not None and after_burden is not None:
            before_burden_distribution[_choice(row.get("before_student_response_burden"))] += 1
            after_burden_distribution[_choice(row.get("after_student_response_burden"))] += 1
            if after_burden < before_burden:
                burden_improved += 1
                source_counter["burden_improved"] += 1
            elif after_burden > before_burden:
                burden_worse += 1
                source_counter["burden_worse"] += 1
            else:
                burden_same += 1
                source_counter["burden_same"] += 1
        if _choice(row.get("after_too_vague")) in {"yes", "true", "1"}:
            too_vague += 1
        before_show_distribution[_choice(row.get("before_would_show_to_student"))] += 1
        after_show_distribution[_choice(row.get("after_would_show_to_student"))] += 1
        pref = _choice(row.get("repair_pair_preference"))
        if pref:
            pair_preference[pref] += 1
            if pref in source_counter:
                source_counter[pref] += 1

    by_source_condition = {}
    for source_condition, stats in by_source.items():
        delta_count = stats["overall_delta_count"]
        by_source_condition[source_condition] = {
            "n": stats["n"],
            "mean_overall_delta": stats["overall_delta_sum"] / delta_count if delta_count else None,
            "quality_win_tie_loss": {
                "win": stats["quality_win"],
                "tie": stats["quality_tie"],
                "loss": stats["quality_loss"],
            },
            "leakage_improved_same_worse": {
                "improved": stats["leakage_improved"],
                "same": stats["leakage_same"],
                "worse": stats["leakage_worse"],
            },
            "burden_improved_same_worse": {
                "improved": stats["burden_improved"],
                "same": stats["burden_same"],
                "worse": stats["burden_worse"],
            },
            "repair_pair_preference": {
                "repair_preferred": stats["repair_preferred"],
                "original_preferred": stats["original_preferred"],
                "tie": stats["tie"],
                "unclear": stats["unclear"],
            },
        }

    n = len(labeled)
    return {
        "total_rows": len(rows),
        "labeled_pairs": n,
        "unlabeled_pairs": len(rows) - n,
        "mean_overall_delta": sum(deltas) / len(deltas) if deltas else None,
        "before_overall_mean": sum(before_overall_values) / len(before_overall_values)
        if before_overall_values
        else None,
        "after_overall_mean": sum(after_overall_values) / len(after_overall_values)
        if after_overall_values
        else None,
        "repair_win_tie_loss": {"win": wins, "tie": ties, "loss": losses},
        "repair_pair_preference": dict(pair_preference),
        "before_leakage_distribution": dict(before_leakage_distribution),
        "after_leakage_distribution": dict(after_leakage_distribution),
        "before_burden_distribution": dict(before_burden_distribution),
        "after_burden_distribution": dict(after_burden_distribution),
        "before_show_distribution": dict(before_show_distribution),
        "after_show_distribution": dict(after_show_distribution),
        "leakage_delta": {
            "improved": leakage_improved,
            "same": leakage_same,
            "worse": leakage_worse,
        },
        "burden_delta": {
            "improved": burden_improved,
            "same": burden_same,
            "worse": burden_worse,
        },
        "still_leaks_rate": still_leaks / n if n else None,
        "too_vague_after_repair_rate": too_vague / n if n else None,
        "by_bridge_bucket": {key: dict(value) for key, value in by_family.items()},
        "by_source_condition": by_source_condition,
    }


def _fmt_number(value: object) -> str:
    if value is None:
        return "NA"
    if isinstance(value, float):
        return f"{value:.3f}"
    return str(value)


def write_markdown(summary: dict, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    mean_delta = summary["mean_overall_delta"]
    before_overall = summary["before_overall_mean"]
    after_overall = summary["after_overall_mean"]
    mean_text = "NA" if mean_delta is None else f"{mean_delta:+.3f}"
    still = summary["still_leaks_rate"]
    vague = summary["too_vague_after_repair_rate"]
    still_text = "NA" if still is None else f"{still:.3f}"
    vague_text = "NA" if vague is None else f"{vague:.3f}"
    lines = [
                "# Repair Same-Candidate Stress Summary",
                "",
                "This is an offline summary of filled before/after Repair review labels.",
                "",
                "| metric | value |",
                "| --- | ---: |",
                f"| total rows | {summary['total_rows']} |",
                f"| labeled pairs | {summary['labeled_pairs']} |",
                f"| unlabeled pairs | {summary['unlabeled_pairs']} |",
                f"| mean overall delta | {mean_text} |",
                f"| before overall mean | {_fmt_number(before_overall)} |",
                f"| after overall mean | {_fmt_number(after_overall)} |",
                f"| repair win/tie/loss | {summary['repair_win_tie_loss']['win']}/{summary['repair_win_tie_loss']['tie']}/{summary['repair_win_tie_loss']['loss']} |",
                f"| repair preferred/original preferred/tie | {summary['repair_pair_preference'].get('repair_preferred', 0)}/{summary['repair_pair_preference'].get('original_preferred', 0)}/{summary['repair_pair_preference'].get('tie', 0)} |",
                f"| leakage improved/same/worse | {summary['leakage_delta']['improved']}/{summary['leakage_delta']['same']}/{summary['leakage_delta']['worse']} |",
                f"| burden improved/same/worse | {summary['burden_delta']['improved']}/{summary['burden_delta']['same']}/{summary['burden_delta']['worse']} |",
                f"| still-leaks rate | {still_text} |",
                f"| too-vague-after-repair rate | {vague_text} |",
                "",
                "Interpretation: this supports same-candidate Repair stress evidence, not a main-experiment condition-level causal proof. Report quality and burden trade-offs alongside leakage reduction.",
                "",
            ]
    by_source = summary.get("by_source_condition") or {}
    if by_source:
        lines.extend(
            [
                "## Source Condition Split",
                "",
                "| source condition | n | mean overall delta | quality W/T/L | leakage improved/same/worse | burden improved/same/worse | repair preferred/original/tie |",
                "| --- | ---: | ---: | ---: | ---: | ---: | ---: |",
            ]
        )
        for source_condition, stats in sorted(by_source.items()):
            quality = stats["quality_win_tie_loss"]
            leakage = stats["leakage_improved_same_worse"]
            burden = stats["burden_improved_same_worse"]
            pref = stats["repair_pair_preference"]
            lines.append(
                f"| `{source_condition}` | {stats['n']} | {_fmt_number(stats['mean_overall_delta'])} | "
                f"{quality['win']}/{quality['tie']}/{quality['loss']} | "
                f"{leakage['improved']}/{leakage['same']}/{leakage['worse']} | "
                f"{burden['improved']}/{burden['same']}/{burden['worse']} | "
                f"{pref.get('repair_preferred', 0)}/{pref.get('original_preferred', 0)}/{pref.get('tie', 0)} |"
            )
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument(
        "--key-csv",
        type=Path,
        default=None,
        help="Optional blind A/B key CSV. Use with blinded review workbooks.",
    )
    parser.add_argument("--output-json", type=Path, default=DEFAULT_OUTPUT_JSON)
    parser.add_argument("--output-md", type=Path, default=DEFAULT_OUTPUT_MD)
    args = parser.parse_args(argv)

    rows = normalize_blind_rows(read_rows(args.input), read_key(args.key_csv))
    result = summarize(rows)
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    write_markdown(result, args.output_md)
    print(f"Summarized {result['labeled_pairs']} labeled pairs out of {result['total_rows']} rows")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
