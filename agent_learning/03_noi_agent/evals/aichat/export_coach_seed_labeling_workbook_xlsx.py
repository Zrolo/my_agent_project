import argparse
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


DEFAULT_INPUT_CSV = Path("docs/research/coach_seed_labeling_workbook_v1.csv")
DEFAULT_FOCUS_REGISTRY = Path("docs/research/focus_registry_v1.json")
DEFAULT_OUTPUT_XLSX = Path("docs/research/coach_seed_labeling_workbook_v1.zh.xlsx")

CHINESE_HEADERS = {
    "case_id": "样本编号",
    "problem_ref": "题目编号",
    "student_message": "学生当前问题（先看）",
    "problem_context": "题目/上下文",
    "recent_dialogue": "近期对话",
    "student_code_excerpt": "学生代码片段",
    "coach_problem_solving_state": "学生当前状态（必选）",
    "coach_bridge_family": "缺失桥梁大类（必选）",
    "coach_secondary_bridge_family": "次要桥梁（可空）",
    "coach_bridge_subtype": "桥梁细分（可写中文）",
    "coach_known_focus": "已有知识焦点",
    "coach_bridge_evidence": "标注证据（引用学生话）",
    "coach_missing_bridge_description": "缺的那一步（一句话）",
    "coach_help_seeking_type": "求助类型",
    "coach_allowed_help_level": "本轮最多帮助强度",
    "coach_help_forms": "建议帮助形式",
    "coach_forbidden_content": "本轮不能直接补完什么",
    "coach_needs_new_focus": "是否需要新焦点",
    "coach_confidence": "置信度 1-5",
    "coach_notes": "备注",
    "review_status": "标注状态",
}

FIELD_HELP = {
    "coach_problem_solving_state": "判断学生当前卡在哪个阶段。比如题意看不懂、知道方向但做不出、代码边界有问题。",
    "coach_bridge_family": "判断学生缺的是哪一类推理桥。先选大类，不确定时选 unknown_bridge。",
    "coach_known_focus": "如果 focus_registry 中已有合适焦点就选它；没有就选 unknown，并把是否需要新焦点设为 true。",
    "coach_bridge_evidence": "写你判定的证据。可以直接引用学生的话，例如“学生问 check(mid) true/false”。",
    "coach_missing_bridge_description": "用一句话写学生缺的那一步推理关系。",
    "coach_allowed_help_level": "判断 AI 这一轮最多能帮到什么程度，不是学生想要多少就给多少。",
    "coach_forbidden_content": "写 AI 不能直接说穿的内容，例如完整状态定义、完整 check 条件、完整代码。",
}

STATE_OPTIONS = {
    "text_comprehension_blocked": "连题意或目标都没看懂",
    "problem_representation_unclear": "题意大概懂，但不知道抽象成状态/对象/图",
    "strategy_generation_blocked": "不知道用什么方向或算法",
    "strategy_misconception": "思路错了，但学生以为是对的",
    "strategy_application_gap": "知道方向，但关键步骤做不出来",
    "implementation_execution_gap": "思路基本对，但落到代码做不出来",
    "debugging_verification_gap": "代码、边界、WA/TLE/RE 调试问题",
    "reflection_transfer_gap": "做出来了，但不会总结迁移",
}

BRIDGE_FAMILY_OPTIONS = {
    "representation_bridge": "表示/状态桥：不知道状态、对象、变量表示什么",
    "transition_bridge": "转移/递推桥：不知道从哪些情况转移",
    "predicate_bridge": "判定/check 桥：不知道条件或 true/false 含义",
    "modeling_bridge": "建模桥：不知道把题面抽象成图、状态、约束",
    "selection_bridge": "选择/贪心桥：不知道为什么这样选是安全的",
    "aggregation_bridge": "汇总/贡献桥：不知道怎么合并贡献、前缀、子树等",
    "ordering_bridge": "顺序桥：不知道为什么这样枚举/遍历/处理",
    "mapping_bridge": "映射桥：不知道题目动作对应哪种数据结构或代码操作",
    "boundary_bridge": "边界/实现桥：下标、边界、数据类型、base case 等",
    "complexity_bridge": "复杂度桥：不知道数据范围和算法复杂度是否匹配",
    "unknown_bridge": "信息不足或无法判断",
}

HELP_SEEKING_OPTIONS = {
    "instrumental_help": "想要提示、局部帮助或确认局部理解",
    "executive_help": "想直接要答案、算法名、完整代码或完整步骤",
    "help_avoidance": "回避思考或只想跳过过程",
    "unclear": "信息不足，无法判断",
}

HELP_LEVEL_OPTIONS = {
    "L1": "轻提示：要上下文/当前尝试/一个观察方向，不补关键桥",
    "L2": "中提示：微型例子、反例、半步关系、引导问题",
    "L3": "强提示：学生已有尝试时给局部伪代码、检查清单或局部代码诊断",
}

HELP_FORM_OPTIONS = {
    "guiding_question": "引导问题",
    "micro_example": "微型例子",
    "guiding_question;micro_example": "引导问题 + 微型例子",
    "counterexample": "反例",
    "counterexample;guiding_question": "反例 + 引导问题",
    "constraint_probe": "约束追问",
    "debug_evidence_request": "要求调试证据",
    "debug_evidence_request;local_code_hint": "先要调试证据，再给局部代码提示",
    "local_code_hint": "局部代码提示",
    "checklist": "检查清单",
    "checklist;partial_trace": "检查清单 + 局部手算/跟踪",
    "partial_trace": "局部手算/跟踪",
    "reflection_prompt": "总结迁移问题",
}

BRIDGE_SUBTYPE_OPTIONS = {
    "unknown": "不确定或信息不足",
    "text_goal_decomposition": "题意目标拆解",
    "dp_state_design": "DP 状态设计",
    "transition_design": "递推/转移来源",
    "check_condition": "二分/check 判定条件",
    "modeling_objects_relations": "题面对象和关系建模",
    "method_selection": "算法方向选择",
    "greedy_basis": "贪心依据/交换理由",
    "tree_path_difference": "树上路径差分/贡献标记",
    "prefix_sum_aggregation": "前缀和/区间汇总",
    "difference_array_mapping": "差分数组映射",
    "enumeration_order": "枚举/遍历顺序",
    "data_structure_operation_mapping": "数据结构操作映射",
    "loop_boundary": "循环/下标边界",
    "data_type_overflow": "数据类型/溢出",
    "recursion_base_case": "递归含义/base case",
    "complexity_fit": "复杂度与数据范围匹配",
    "debug_evidence_gap": "调试证据不足",
    "transfer_summary": "迁移总结",
}

BRIDGE_EVIDENCE_OPTIONS = {
    "学生直接问状态/表示含义": "学生问 dp、mask、lazy、next 等到底表示什么",
    "学生直接问转移/来源": "学生问这一格、这一层、这个状态从哪里来",
    "学生直接问 check/条件": "学生问 check true/false、if 条件、判定方向",
    "学生直接问题型/算法名": "学生问是不是某算法或直接要方向",
    "学生给出错误思路": "学生已有想法，但推理方向明显有误",
    "学生有代码但卡实现细节": "学生已写部分代码，卡边界、循环、类型或局部函数",
    "学生只有笼统不会": "学生没有暴露具体卡点",
    "学生在要求完整答案/代码": "学生请求直接给完整题解、算法名或代码",
    "题面上下文显示建模对象不清": "题面对象、关系或限制没有被学生抽象出来",
    "需要讨论": "当前证据不足或多种标签都合理",
}

MISSING_BRIDGE_OPTIONS = {
    "缺少题意目标拆解": "还没把题目目标拆成可操作对象",
    "缺少状态/变量的语义定义": "不知道 dp、mask、数组格子或变量表示什么",
    "缺少转移来源或分类讨论": "不知道当前状态从哪些前置情况来",
    "缺少候选答案到可行性判断的映射": "不知道 check 或 if 条件的语义",
    "缺少题面对象到图/状态/约束的建模": "不知道什么当点、边、状态或限制",
    "缺少局部选择为什么安全的理由": "不知道贪心或选择为什么不亏",
    "缺少贡献如何汇总/还原的关系": "不知道前缀、差分、子树、路径贡献如何合并",
    "缺少正确处理顺序的依赖关系": "不知道为什么先算小区间、前驱、入度为 0 等",
    "缺少题目动作到数据结构操作的映射": "不知道题目里的合并/查询/取最小对应什么操作",
    "缺少边界、下标或类型上限检查": "代码细节与范围、编号或类型上限没有对齐",
    "缺少复杂度和数据范围的匹配判断": "不知道当前做法是否能过",
    "缺少调试证据定位": "只知道错了，但没有定位到最小错误现象",
    "缺少做题后的迁移总结": "做出来但不知道下次如何识别同类题",
    "信息不足，先要上下文/尝试": "当前无法可靠判断 bridge",
}

FORBIDDEN_CONTENT_OPTIONS = {
    "不能直接给完整算法名或题型确认": "避免直接确认是不是 DP/二分/贪心/Trie 等",
    "不能直接给完整状态定义": "避免直接说出 dp/mask/数组格子的完整含义",
    "不能直接给完整转移方程": "避免直接写出完整递推式",
    "不能直接给完整 check 条件和边界更新": "避免直接说 true/false 方向和二分更新",
    "不能直接给完整建模方案": "避免直接给点、边、状态、约束的完整抽象",
    "不能直接给完整贪心准则和证明": "避免直接说按什么排序/选择以及完整证明",
    "不能直接给完整贡献公式": "避免直接给前缀和、差分、树上差分等完整公式",
    "不能直接给完整循环/遍历模板": "避免直接给可粘贴的循环顺序或模板",
    "不能直接给完整数据结构模板": "避免直接给并查集、堆、Trie、线段树等完整实现",
    "不能直接替学生改完整代码": "只做局部诊断，不完整代写",
    "不能给完整题解、完整步骤或完整代码": "答案/代码泄露红线",
    "信息不足时不能猜最终方案": "先要题面、尝试或错误证据",
}

NOTES_OPTIONS = {
    "无备注": "没有特殊情况",
    "需要讨论": "这行边界不清，需要复核",
    "可能多桥混合": "学生同时卡多个 bridge",
    "证据不足": "学生输入太短或上下文不足",
    "可能需要新增 focus": "现有 focus 不贴切",
    "标签可接受但不唯一": "多个标签都说得通",
}

BOOLEAN_OPTIONS = {"false": "不需要", "true": "需要"}
REVIEW_STATUS_OPTIONS = {
    "unlabeled": "未标",
    "labeled": "已标完",
    "needs_discussion": "不确定，需要讨论",
    "review_seed_gold": "内部复核 seed gold",
}

COLUMN_WIDTHS = {
    "A": 15,
    "B": 14,
    "C": 38,
    "D": 42,
    "E": 32,
    "F": 32,
    "G": 28,
    "H": 24,
    "I": 20,
    "J": 24,
    "K": 24,
    "L": 34,
    "M": 38,
    "N": 18,
    "O": 20,
    "P": 24,
    "Q": 38,
    "R": 18,
    "S": 14,
    "T": 30,
    "U": 18,
}


def load_csv_rows(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def load_focus_options(path: Path) -> list[str]:
    if not path.exists():
        return ["未知/不确定（unknown）"]
    data = json.loads(path.read_text(encoding="utf-8"))
    focuses = data.get("focuses", data) if isinstance(data, dict) else data
    labels = {"unknown": "未知/不确定（unknown）"}
    for item in focuses:
        if isinstance(item, str):
            labels[item] = f"{item}（{item}）"
        elif isinstance(item, dict) and item.get("focus_id"):
            focus_id = str(item["focus_id"])
            aliases = item.get("aliases") or []
            if aliases:
                readable = str(aliases[0])
            else:
                readable = str(item.get("description") or focus_id)
            labels[focus_id] = f"{readable}（{focus_id}）"
    return [labels["unknown"], *[labels[key] for key in sorted(labels) if key != "unknown"]]


def _option_label(option_id: str, description: str) -> str:
    if option_id.isascii():
        return f"{description}（{option_id}）"
    return f"{option_id}：{description}"


def _option_labels(options: dict[str, str]) -> list[str]:
    return [_option_label(option_id, description) for option_id, description in options.items()]


def _add_options_sheet(workbook: Workbook, focus_options: list[str]) -> dict[str, str]:
    sheet = workbook.create_sheet("下拉选项")
    lists = {
        "problem_solving_state": _option_labels(STATE_OPTIONS),
        "bridge_family": _option_labels(BRIDGE_FAMILY_OPTIONS),
        "bridge_subtype": _option_labels(BRIDGE_SUBTYPE_OPTIONS),
        "known_focus": focus_options,
        "bridge_evidence": _option_labels(BRIDGE_EVIDENCE_OPTIONS),
        "missing_bridge_description": _option_labels(MISSING_BRIDGE_OPTIONS),
        "help_seeking_type": _option_labels(HELP_SEEKING_OPTIONS),
        "allowed_help_level": _option_labels(HELP_LEVEL_OPTIONS),
        "help_forms": _option_labels(HELP_FORM_OPTIONS),
        "forbidden_content": _option_labels(FORBIDDEN_CONTENT_OPTIONS),
        "needs_new_focus": _option_labels(BOOLEAN_OPTIONS),
        "confidence": [
            "很不确定（1）",
            "较不确定（2）",
            "一般确定（3）",
            "比较确定（4）",
            "非常确定（5）",
        ],
        "notes": _option_labels(NOTES_OPTIONS),
        "review_status": _option_labels(REVIEW_STATUS_OPTIONS),
    }
    ranges = {}
    for col_idx, (name, values) in enumerate(lists.items(), 1):
        sheet.cell(row=1, column=col_idx, value=name)
        sheet.cell(row=1, column=col_idx).font = Font(bold=True)
        for row_idx, value in enumerate(values, 2):
            sheet.cell(row=row_idx, column=col_idx, value=value)
        col_letter = sheet.cell(row=1, column=col_idx).column_letter
        ranges[name] = f"'下拉选项'!${col_letter}$2:${col_letter}${len(values) + 1}"
        sheet.column_dimensions[col_letter].width = 28
    return ranges


def _add_data_validation(sheet, cell_range: str, formula_range: str) -> None:
    validation = DataValidation(type="list", formula1=f"={formula_range}", allow_blank=True)
    validation.error = "请从下拉列表中选择，或留空。"
    validation.errorTitle = "无效选项"
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def _style_labeling_sheet(sheet, headers: list[str]) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    field_fill = PatternFill("solid", fgColor="F2F5F7")
    input_fill = PatternFill("solid", fgColor="FFF7D6")
    thin = Side(style="thin", color="D7DEE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="17324D")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    for cell in sheet[2]:
        cell.font = Font(size=9, color="667085")
        cell.fill = field_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    for row in sheet.iter_rows(min_row=3, max_row=sheet.max_row, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            if cell.column >= 7:
                cell.fill = input_fill
    for col_letter, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[col_letter].width = width
    sheet.row_dimensions[1].height = 36
    sheet.row_dimensions[2].height = 28
    for row_idx in range(3, sheet.max_row + 1):
        sheet.row_dimensions[row_idx].height = 72
    sheet.freeze_panes = "G3"
    sheet.auto_filter.ref = f"A2:{sheet.cell(row=2, column=len(headers)).coordinate}"


def _add_instruction_sheet(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "开始这里"
    rows = [
        ["怎么标注", "只看“标注表”中的学生当前问题和题目/上下文。黄色列是你要填写的教练标注。"],
        ["最小必填", "学生当前状态、缺失桥梁大类、已有知识焦点、标注证据、缺的那一步、本轮最多帮助强度、本轮不能直接补完什么、置信度、标注状态。"],
        ["不确定怎么办", "known_focus 选 unknown；needs_new_focus 选 true；confidence 填 2 或 3；review_status 选 needs_discussion。"],
        ["不要用哪个文件", "不要用 prefilled.csv 做独立标注，它带参考答案。请用这个 Excel 的“标注表”。"],
        ["一句话原则", "判断的是当前这一轮学生缺的推理桥，不是整段聊天，也不是学生整体水平。"],
    ]
    sheet.append(["教练标注说明", ""])
    for row in rows:
        sheet.append(row)
    sheet["A1"].font = Font(bold=True, size=16, color="17324D")
    sheet["A1"].fill = PatternFill("solid", fgColor="D9EAF7")
    for row in sheet.iter_rows(min_row=2, max_col=2):
        row[0].font = Font(bold=True, color="17324D")
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 95


def _add_guide_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("标签说明")
    sections = [
        ("学生当前状态", STATE_OPTIONS),
        ("桥梁大类", BRIDGE_FAMILY_OPTIONS),
        ("桥梁细分", BRIDGE_SUBTYPE_OPTIONS),
        ("标注证据", BRIDGE_EVIDENCE_OPTIONS),
        ("缺失桥梁描述", MISSING_BRIDGE_OPTIONS),
        ("求助类型", HELP_SEEKING_OPTIONS),
        ("帮助强度", HELP_LEVEL_OPTIONS),
        ("帮助形式", HELP_FORM_OPTIONS),
        ("禁止直接补完内容", FORBIDDEN_CONTENT_OPTIONS),
        ("备注", NOTES_OPTIONS),
    ]
    row = 1
    for title, options in sections:
        sheet.cell(row=row, column=1, value=title)
        sheet.cell(row=row, column=1).font = Font(bold=True, size=13, color="17324D")
        row += 1
        sheet.cell(row=row, column=1, value="标签")
        sheet.cell(row=row, column=2, value="解释")
        sheet.cell(row=row, column=1).font = Font(bold=True)
        sheet.cell(row=row, column=2).font = Font(bold=True)
        row += 1
        for option_id, description in options.items():
            sheet.cell(row=row, column=1, value=option_id)
            sheet.cell(row=row, column=2, value=description)
            row += 1
        row += 2
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 88
    for row_cells in sheet.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def export_xlsx(
    *,
    input_csv: Path = DEFAULT_INPUT_CSV,
    focus_registry: Path = DEFAULT_FOCUS_REGISTRY,
    output_xlsx: Path = DEFAULT_OUTPUT_XLSX,
) -> int:
    rows = load_csv_rows(input_csv)
    if not rows:
        raise ValueError(f"No rows found in {input_csv}")
    headers = list(rows[0])
    workbook = Workbook()
    _add_instruction_sheet(workbook)
    ranges = _add_options_sheet(workbook, load_focus_options(focus_registry))
    _add_guide_sheet(workbook)

    sheet = workbook.create_sheet("标注表", 1)
    sheet.append([CHINESE_HEADERS.get(header, header) for header in headers])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    _style_labeling_sheet(sheet, headers)

    field_to_col = {field: idx + 1 for idx, field in enumerate(headers)}
    max_row = sheet.max_row
    validations = {
        "coach_problem_solving_state": ranges["problem_solving_state"],
        "coach_bridge_family": ranges["bridge_family"],
        "coach_secondary_bridge_family": ranges["bridge_family"],
        "coach_bridge_subtype": ranges["bridge_subtype"],
        "coach_known_focus": ranges["known_focus"],
        "coach_bridge_evidence": ranges["bridge_evidence"],
        "coach_missing_bridge_description": ranges["missing_bridge_description"],
        "coach_help_seeking_type": ranges["help_seeking_type"],
        "coach_allowed_help_level": ranges["allowed_help_level"],
        "coach_help_forms": ranges["help_forms"],
        "coach_forbidden_content": ranges["forbidden_content"],
        "coach_needs_new_focus": ranges["needs_new_focus"],
        "coach_confidence": ranges["confidence"],
        "coach_notes": ranges["notes"],
        "review_status": ranges["review_status"],
    }
    for field, formula_range in validations.items():
        col_idx = field_to_col.get(field)
        if col_idx:
            col = sheet.cell(row=1, column=col_idx).column_letter
            _add_data_validation(sheet, f"{col}3:{col}{max_row}", formula_range)
    for field, help_text in FIELD_HELP.items():
        col_idx = field_to_col.get(field)
        if col_idx:
            sheet.cell(row=1, column=col_idx).comment = Comment(help_text, "Codex")

    status_col = sheet.cell(row=1, column=field_to_col["review_status"]).column_letter
    labeled_value = _option_label("labeled", REVIEW_STATUS_OPTIONS["labeled"])
    needs_discussion_value = _option_label("needs_discussion", REVIEW_STATUS_OPTIONS["needs_discussion"])
    sheet.conditional_formatting.add(
        f"A3:U{max_row}",
        FormulaRule(
            formula=[f'${status_col}3="{labeled_value}"'],
            fill=PatternFill("solid", fgColor="E9F7EF"),
        ),
    )
    sheet.conditional_formatting.add(
        f"A3:U{max_row}",
        FormulaRule(
            formula=[f'${status_col}3="{needs_discussion_value}"'],
            fill=PatternFill("solid", fgColor="FFF1CC"),
        ),
    )

    workbook["下拉选项"].sheet_state = "hidden"
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx)
    return len(rows)


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export Chinese coach seed labeling XLSX workbook.")
    parser.add_argument("--input-csv", type=Path, default=DEFAULT_INPUT_CSV)
    parser.add_argument("--focus-registry", type=Path, default=DEFAULT_FOCUS_REGISTRY)
    parser.add_argument("--output-xlsx", type=Path, default=DEFAULT_OUTPUT_XLSX)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(sys.argv[1:] if argv is None else argv)
    row_count = export_xlsx(
        input_csv=args.input_csv,
        focus_registry=args.focus_registry,
        output_xlsx=args.output_xlsx,
    )
    print(json.dumps({"output_xlsx": str(args.output_xlsx), "row_count": row_count}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
