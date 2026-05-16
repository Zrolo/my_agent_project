import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from evals.aichat import apply_heldout_source_completion, export_heldout_source_completion_workbook


def _write_jsonl(path: Path, rows: list[dict]) -> None:
    path.write_text(
        "\n".join(json.dumps(row, ensure_ascii=False, sort_keys=True) for row in rows) + "\n",
        encoding="utf-8",
    )


class HeldoutSourceCompletionTests(unittest.TestCase):
    def test_exports_source_completion_csv_with_context_and_blank_source_fields(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            draft = tmp / "draft.jsonl"
            output_csv = tmp / "source_completion.csv"
            _write_jsonl(
                draft,
                [
                    {
                        "case_id": "heldout_cp_001",
                        "category": "dp_state",
                        "problem_ref": "heldout_dp_state",
                        "student_message": "状态怎么设？",
                        "problem_context": "给若干物品的代价和收益，在总代价限制内最大化收益。",
                        "recent_dialogue": "N/A",
                        "student_code_excerpt": "N/A",
                        "missing_bridge": "缺把下标代表的资源限制和格子值代表的最优目标分开表达。",
                    }
                ],
            )

            row_count = export_heldout_source_completion_workbook.export_source_completion_csv(
                input_jsonl=draft,
                output_csv=output_csv,
            )

            self.assertEqual(1, row_count)
            with output_csv.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual("heldout_cp_001", rows[0]["case_id"])
            self.assertEqual("short", rows[0]["student_message_length_bucket"])
            self.assertEqual("", rows[0]["problem_source_platform"])
            self.assertEqual("", rows[0]["problem_source_url"])
            self.assertIn("资源限制", rows[0]["missing_bridge"])

    def test_exports_source_completion_xlsx_with_chinese_headers(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            draft = tmp / "draft.jsonl"
            output_xlsx = tmp / "source_completion.zh.xlsx"
            _write_jsonl(
                draft,
                [
                    {
                        "case_id": "heldout_cp_001",
                        "category": "dp_state",
                        "problem_ref": "heldout_dp_state",
                        "student_message": "状态怎么设？",
                        "problem_context": "给若干物品的代价和收益，在总代价限制内最大化收益。",
                        "recent_dialogue": "N/A",
                        "student_code_excerpt": "N/A",
                        "missing_bridge": "缺把下标代表的资源限制和格子值代表的最优目标分开表达。",
                    }
                ],
            )

            row_count = export_heldout_source_completion_workbook.export_source_completion_xlsx(
                input_jsonl=draft,
                output_xlsx=output_xlsx,
            )
            workbook = load_workbook(output_xlsx)

        self.assertEqual(1, row_count)
        self.assertIn("题源补全表", workbook.sheetnames)
        sheet = workbook["题源补全表"]
        headers = {
            sheet.cell(row=2, column=column_idx).value: sheet.cell(row=1, column=column_idx).value
            for column_idx in range(1, sheet.max_column + 1)
        }
        self.assertEqual("样本编号", headers["case_id"])
        self.assertEqual("题目来源平台", headers["problem_source_platform"])
        self.assertEqual("原题链接", headers["problem_source_url"])
        self.assertEqual("题面访问级别", headers["problem_statement_access_level"])
        self.assertEqual("heldout_cp_001", sheet["A3"].value)
        self.assertGreaterEqual(len(sheet.data_validations.dataValidation), 1)

    def test_applies_completed_source_metadata_back_to_jsonl(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            draft = tmp / "draft.jsonl"
            completed_csv = tmp / "source_completion_filled.csv"
            output_jsonl = tmp / "draft_with_sources.jsonl"
            _write_jsonl(
                draft,
                [
                    {
                        "case_id": "heldout_cp_001",
                        "category": "dp_state",
                        "problem_ref": "heldout_dp_state",
                        "student_message": "状态怎么设？",
                        "problem_context": "给若干物品的代价和收益，在总代价限制内最大化收益。",
                        "recent_dialogue": "N/A",
                        "student_code_excerpt": "N/A",
                        "missing_bridge": "缺把下标代表的资源限制和格子值代表的最优目标分开表达。",
                    }
                ],
            )
            with completed_csv.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=export_heldout_source_completion_workbook.SOURCE_COMPLETION_COLUMNS,
                    lineterminator="\n",
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "case_id": "heldout_cp_001",
                        "problem_source_platform": "luogu",
                        "problem_source_id": "P1048",
                        "problem_source_url": "https://www.luogu.com.cn/problem/P1048",
                        "problem_statement": "给定若干草药的采摘时间和价值，在总时间内选择若干草药使价值最大。",
                        "problem_statement_public_summary": "洛谷 P1048 采药：01 背包入门题。",
                        "problem_statement_rights_note": "本地教练评审使用必要题面；公开材料仅保留来源链接和改写摘要。",
                        "problem_statement_access_level": "public_summary_only",
                        "source_completion_status": "completed",
                    }
                )

            summary = apply_heldout_source_completion.apply_source_completion(
                input_jsonl=draft,
                source_csv=completed_csv,
                output_jsonl=output_jsonl,
                require_complete=True,
            )

            self.assertEqual({"updated": 1, "missing_source_rows": 0, "incomplete_rows": 0}, summary)
            output_row = json.loads(output_jsonl.read_text(encoding="utf-8").splitlines()[0])
            self.assertEqual("luogu", output_row["problem_source_platform"])
            self.assertEqual("P1048", output_row["problem_source_id"])
            self.assertEqual("public_summary_only", output_row["problem_statement_access_level"])


if __name__ == "__main__":
    unittest.main()
