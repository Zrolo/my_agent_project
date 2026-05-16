import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from evals.aichat import export_coach_adjudication_workbook as exporter


def _write_jsonl(path: Path, rows: list[dict]) -> None:
    path.write_text(
        "\n".join(json.dumps(row, ensure_ascii=False, sort_keys=True) for row in rows) + "\n",
        encoding="utf-8",
    )


class ExportCoachAdjudicationWorkbookTests(unittest.TestCase):
    def test_export_xlsx_contains_disagreement_rows_and_side_by_side_labels(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            draft_jsonl = tmp / "draft.jsonl"
            annotator_a_jsonl = tmp / "a.jsonl"
            annotator_b_jsonl = tmp / "b.jsonl"
            agreement_json = tmp / "agreement.json"
            output_xlsx = tmp / "adjudication.xlsx"

            _write_jsonl(
                draft_jsonl,
                [
                    {
                        "case_id": "case_1",
                        "student_message": "check(mid) 怎么判断？",
                        "problem_context": "二分答案。",
                        "recent_dialogue": "N/A",
                        "student_code_excerpt": "N/A",
                    }
                ],
            )
            _write_jsonl(
                annotator_a_jsonl,
                [
                    {
                        "case_id": "case_1",
                        "student_problem_solving_state": "method_application_gap",
                        "primary_bridge_family": "predicate_condition_bridge",
                        "primary_bridge_subtype_id": "predicate.feasibility_truth_direction",
                        "registered_focus_id": "check_condition",
                        "max_scaffold_level": "L2",
                        "leakage_risk": "high",
                        "bridge_specific_forbidden_content": ["no_exact_check_condition"],
                        "evidence_quote": "check(mid) 怎么判断",
                    }
                ],
            )
            _write_jsonl(
                annotator_b_jsonl,
                [
                    {
                        "case_id": "case_1",
                        "student_problem_solving_state": "implementation_translation_gap",
                        "primary_bridge_family": "implementation_boundary_bridge",
                        "primary_bridge_subtype_id": "implementation.loop_boundary",
                        "registered_focus_id": "binary_search_bound_direction",
                        "max_scaffold_level": "L3",
                        "leakage_risk": "medium",
                        "bridge_specific_forbidden_content": ["no_full_boundary_update"],
                        "evidence_quote": "判断",
                    }
                ],
            )
            agreement_json.write_text(
                json.dumps({"needs_adjudication_case_ids": ["case_1"]}, ensure_ascii=False),
                encoding="utf-8",
            )

            row_count = exporter.export_xlsx(
                agreement_json=agreement_json,
                annotator_a_jsonl=annotator_a_jsonl,
                annotator_b_jsonl=annotator_b_jsonl,
                draft_jsonl=draft_jsonl,
                output_xlsx=output_xlsx,
            )
            workbook = load_workbook(output_xlsx)
            sheet = workbook["裁决表"]
            headers = [sheet.cell(row=2, column=col).value for col in range(1, sheet.max_column + 1)]

        self.assertEqual(1, row_count)
        self.assertEqual("case_id", headers[0])
        self.assertIn("a_primary_bridge_family", headers)
        self.assertIn("b_primary_bridge_family", headers)
        self.assertIn("adjudicated_primary_bridge_family", headers)
        self.assertEqual("case_1", sheet["A3"].value)
        self.assertEqual("predicate_condition_bridge", sheet["I3"].value)
        self.assertEqual("implementation_boundary_bridge", sheet["J3"].value)
        self.assertIsNone(sheet["K3"].value)

    def test_export_uses_all_paired_rows_when_no_needs_adjudication_filter(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            draft_jsonl = tmp / "draft.jsonl"
            annotator_a_jsonl = tmp / "a.jsonl"
            annotator_b_jsonl = tmp / "b.jsonl"
            agreement_json = tmp / "agreement.json"
            output_xlsx = tmp / "adjudication.xlsx"
            _write_jsonl(draft_jsonl, [{"case_id": "case_1", "student_message": "不会", "problem_context": "context"}])
            _write_jsonl(annotator_a_jsonl, [{"case_id": "case_1", "primary_bridge_family": "method_selection_bridge"}])
            _write_jsonl(annotator_b_jsonl, [{"case_id": "case_1", "primary_bridge_family": "method_selection_bridge"}])
            agreement_json.write_text(json.dumps({"needs_adjudication_case_ids": []}), encoding="utf-8")

            row_count = exporter.export_xlsx(
                agreement_json=agreement_json,
                annotator_a_jsonl=annotator_a_jsonl,
                annotator_b_jsonl=annotator_b_jsonl,
                draft_jsonl=draft_jsonl,
                output_xlsx=output_xlsx,
                include_all_paired=True,
            )
            workbook = load_workbook(output_xlsx)
            sheet = workbook["裁决表"]

        self.assertEqual(1, row_count)
        self.assertEqual("case_1", sheet["A3"].value)


if __name__ == "__main__":
    unittest.main()
