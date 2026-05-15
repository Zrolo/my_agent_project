import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from evals.aichat import export_dialogue_state_v3_calibration_workbook
from evals.aichat import generate_dialogue_state_v3_50
from evals.aichat import summarize_dialogue_state_case_review
from test_dialogue_state_v3_generation_unit import _v2_case


def _write_jsonl(path: Path, rows: list[dict]) -> None:
    path.write_text(
        "\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n",
        encoding="utf-8",
    )


class DialogueStateV3CaseReviewWorkflowTests(unittest.TestCase):
    def test_select_calibration_cases_covers_review_categories(self):
        cases = generate_dialogue_state_v3_50.build_dialogue_state_cases([_v2_case(i) for i in range(1, 51)])

        selected = export_dialogue_state_v3_calibration_workbook.select_calibration_cases(cases)

        self.assertEqual(6, len(selected))
        self.assertEqual(
            {
                "initial_question",
                "followup_after_correct_short_answer",
                "followup_after_partial_answer",
                "followup_after_wrong_answer",
                "followup_after_prerequisite_gap",
                "policy_direct_answer_special",
            },
            {row["context_type"] for row in selected},
        )
        self.assertEqual({"NA", "F1", "F2", "F3", "F4"}, {row["student_scaffold_followability"] for row in selected})

    def test_calibration_export_main_writes_bilingual_workbooks(self):
        cases = generate_dialogue_state_v3_50.build_dialogue_state_cases([_v2_case(i) for i in range(1, 51)])

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            source = tmp / "dialogue_v3.jsonl"
            zh = tmp / "calibration.zh.xlsx"
            en = tmp / "calibration.en.xlsx"
            report_json = tmp / "calibration.json"
            report_zh = tmp / "calibration.zh.md"
            report_en = tmp / "calibration.md"
            _write_jsonl(source, cases)

            exit_code = export_dialogue_state_v3_calibration_workbook.main(
                [
                    "--source-jsonl",
                    str(source),
                    "--output-zh-xlsx",
                    str(zh),
                    "--output-en-xlsx",
                    str(en),
                    "--report-json",
                    str(report_json),
                    "--report-zh",
                    str(report_zh),
                    "--report-en",
                    str(report_en),
                ]
            )

            self.assertEqual(0, exit_code)
            zh_book = load_workbook(zh)
            en_book = load_workbook(en)

        self.assertEqual(8, zh_book["总表"].max_row)
        self.assertEqual(8, en_book["dialogue_state_review"].max_row)
        self.assertIn("评审说明", zh_book.sheetnames)
        self.assertIn("Instructions", en_book.sheetnames)

    def test_case_review_summary_counts_structured_review_fields(self):
        cases = generate_dialogue_state_v3_50.build_dialogue_state_cases([_v2_case(i) for i in range(1, 51)])

        with tempfile.TemporaryDirectory() as tmpdir:
            workbook_path = Path(tmpdir) / "review.zh.xlsx"
            generate_dialogue_state_v3_50.export_dialogue_state_review_workbook(cases, workbook_path)
            workbook = load_workbook(workbook_path)
            sheet = workbook["总表"]
            machine_headers = [sheet.cell(row=2, column=col).value for col in range(1, sheet.max_column + 1)]
            decision_col = machine_headers.index("case_decision") + 1
            issue_col = machine_headers.index("issue_type") + 1
            confidence_col = machine_headers.index("reviewer_confidence") + 1
            reviewer_col = machine_headers.index("reviewer_id") + 1
            round_col = machine_headers.index("review_round") + 1
            sheet.cell(row=3, column=decision_col).value = "accept"
            sheet.cell(row=3, column=reviewer_col).value = "coach_A"
            sheet.cell(row=3, column=round_col).value = "calibration"
            sheet.cell(row=4, column=decision_col).value = "revise"
            sheet.cell(row=4, column=issue_col).value = "context_mismatch"
            sheet.cell(row=4, column=reviewer_col).value = "coach_A"
            sheet.cell(row=4, column=round_col).value = "calibration"
            sheet.cell(row=5, column=decision_col).value = "drop"
            sheet.cell(row=5, column=confidence_col).value = "low"
            sheet.cell(row=5, column=reviewer_col).value = "coach_B"
            sheet.cell(row=5, column=round_col).value = "round1"
            workbook.save(workbook_path)

            summary = summarize_dialogue_state_case_review.summarize_workbook(workbook_path, sheet_name="总表")

        self.assertEqual(50, summary["row_count"])
        self.assertEqual({"accept": 1, "revise": 1, "drop": 1, "blank": 47}, summary["case_decision_counts"])
        self.assertEqual(1, summary["issue_type_counts"]["context_mismatch"])
        self.assertEqual(1, summary["reviewer_confidence_counts"]["low"])
        self.assertEqual({"coach_A": 2, "coach_B": 1, "blank": 47}, summary["reviewer_id_counts"])
        self.assertEqual({"calibration": 2, "round1": 1, "blank": 47}, summary["review_round_counts"])
        self.assertEqual(2, summary["needs_followup_count"])

    def test_case_review_summary_normalizes_chinese_review_choices(self):
        cases = generate_dialogue_state_v3_50.build_dialogue_state_cases([_v2_case(i) for i in range(1, 51)])

        with tempfile.TemporaryDirectory() as tmpdir:
            workbook_path = Path(tmpdir) / "review.zh.xlsx"
            generate_dialogue_state_v3_50.export_dialogue_state_review_workbook(cases, workbook_path)
            workbook = load_workbook(workbook_path)
            sheet = workbook["总表"]
            machine_headers = [sheet.cell(row=2, column=col).value for col in range(1, sheet.max_column + 1)]
            decision_col = machine_headers.index("case_decision") + 1
            issue_col = machine_headers.index("issue_type") + 1
            confidence_col = machine_headers.index("reviewer_confidence") + 1
            sheet.cell(row=3, column=decision_col).value = "接受"
            sheet.cell(row=4, column=decision_col).value = "修改"
            sheet.cell(row=4, column=issue_col).value = "上下文"
            sheet.cell(row=4, column=confidence_col).value = "低"
            workbook.save(workbook_path)

            summary = summarize_dialogue_state_case_review.summarize_workbook(workbook_path, sheet_name="总表")

            self.assertEqual({"accept": 1, "revise": 1, "blank": 48}, summary["case_decision_counts"])
            self.assertEqual(1, summary["issue_type_counts"]["context_mismatch"])
            self.assertEqual(1, summary["reviewer_confidence_counts"]["low"])

            workbook = load_workbook(workbook_path)
            sheet = workbook["总表"]
            machine_headers = [sheet.cell(row=2, column=col).value for col in range(1, sheet.max_column + 1)]
            issue_col = machine_headers.index("issue_type") + 1
            sheet.cell(row=4, column=issue_col).value = "跟随状态"
            workbook.save(workbook_path)

            summary = summarize_dialogue_state_case_review.summarize_workbook(workbook_path, sheet_name="总表")
            self.assertEqual(1, summary["issue_type_counts"]["followability_issue"])

    def test_case_review_summary_prefers_reviewed_bucket_sheets(self):
        cases = generate_dialogue_state_v3_50.build_dialogue_state_cases([_v2_case(i) for i in range(1, 51)])

        with tempfile.TemporaryDirectory() as tmpdir:
            workbook_path = Path(tmpdir) / "review.zh.xlsx"
            generate_dialogue_state_v3_50.export_dialogue_state_review_workbook(cases, workbook_path)
            workbook = load_workbook(workbook_path)
            bucket_sheet = workbook["判定条件"]
            machine_headers = [bucket_sheet.cell(row=2, column=col).value for col in range(1, bucket_sheet.max_column + 1)]
            decision_col = machine_headers.index("case_decision") + 1
            issue_col = machine_headers.index("issue_type") + 1
            bucket_sheet.cell(row=3, column=decision_col).value = "接受"
            bucket_sheet.cell(row=4, column=decision_col).value = "修改"
            bucket_sheet.cell(row=4, column=issue_col).value = "上下文不一致"
            workbook.save(workbook_path)

            summary = summarize_dialogue_state_case_review.summarize_workbook(workbook_path)

        self.assertEqual("bucket_sheets", summary["source_mode_used"])
        self.assertEqual(50, summary["row_count"])
        self.assertEqual({"accept": 1, "revise": 1, "blank": 48}, summary["case_decision_counts"])
        self.assertEqual(1, summary["issue_type_counts"]["context_mismatch"])
        self.assertEqual(1, summary["needs_followup_count"])


if __name__ == "__main__":
    unittest.main()
