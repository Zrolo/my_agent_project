import tempfile
import unittest
import json
from pathlib import Path

from openpyxl import load_workbook

from evals.aichat import export_coach_seed_labeling_workbook_v2 as workbook_v2


def _option_values(workbook, option_group: str) -> list[str]:
    sheet = workbook["下拉选项"]
    for column in range(1, sheet.max_column + 1):
        if sheet.cell(row=1, column=column).value == option_group:
            return [
                sheet.cell(row=row, column=column).value
                for row in range(2, sheet.max_row + 1)
                if sheet.cell(row=row, column=column).value
            ]
    raise AssertionError(f"Missing option group: {option_group}")


class CoachSeedLabelingV2WorkbookTests(unittest.TestCase):
    def test_export_xlsx_has_v2_columns_and_no_placeholder_ones(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "coach_v2.xlsx"

            row_count = workbook_v2.export_xlsx(
                seed_jsonl=Path("docs/research/bridgebench_cp_seed_v1.jsonl"),
                focus_registry=Path("docs/research/focus_registry_v1.json"),
                output_xlsx=output_path,
            )

            workbook = load_workbook(output_path)

        self.assertEqual(50, row_count)
        self.assertIn("标注表", workbook.sheetnames)
        self.assertIn("标签说明", workbook.sheetnames)
        sheet = workbook["标注表"]
        machine_headers = [cell.value for cell in sheet[2]]

        self.assertIn("turn_type", machine_headers)
        self.assertIn("problem_source_url", machine_headers)
        self.assertIn("problem_statement", machine_headers)
        self.assertIn("policy_risk_type", machine_headers)
        self.assertIn("registered_focus_id", machine_headers)
        self.assertIn("student_already_stated_bridge", machine_headers)
        self.assertIn("bridge_specific_forbidden_content", machine_headers)
        self.assertIn("coach_note_tags", machine_headers)
        self.assertIn("coach_free_notes", machine_headers)
        self.assertNotIn("coach_known_focus", machine_headers)
        self.assertNotIn("coach_forbidden_content", machine_headers)
        self.assertNotIn("coach_notes", machine_headers)

        recent_dialogue_col = machine_headers.index("recent_dialogue") + 1
        code_col = machine_headers.index("student_code_excerpt") + 1
        self.assertEqual("N/A", sheet.cell(row=3, column=recent_dialogue_col).value)
        self.assertEqual("N/A", sheet.cell(row=3, column=code_col).value)

        review_status_col = machine_headers.index("review_status") + 1
        self.assertEqual("未标：还没有完成这一行（unlabeled）", sheet.cell(row=3, column=review_status_col).value)

    def test_export_xlsx_includes_real_problem_source_fields_when_present(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            seed_jsonl = tmp / "luogu_seed.jsonl"
            output_path = tmp / "coach_v2.xlsx"
            seed_jsonl.write_text(
                json.dumps(
                    {
                        "case_id": "heldout_v2_luogu_001",
                        "problem_ref": "luogu_p1002_state",
                        "problem_source_platform": "luogu",
                        "problem_source_id": "P1002",
                        "problem_source_url": "https://www.luogu.com.cn/problem/P1002",
                        "problem_statement": "题目：过河卒\n题意摘录：棋盘上卒从 A 到 B，部分点被马控制。",
                        "problem_statement_public_summary": "洛谷 P1002 过河卒：棋盘路径计数题。",
                        "problem_statement_rights_note": "本地教练复核使用必要题面；公开材料保留链接和摘要。",
                        "problem_statement_access_level": "local_review_only",
                        "student_message": "状态到底记录什么？",
                        "problem_context": "洛谷 P1002 过河卒。",
                        "recent_dialogue": "N/A",
                        "student_code_excerpt": "N/A",
                    },
                    ensure_ascii=False,
                )
                + "\n",
                encoding="utf-8",
            )

            row_count = workbook_v2.export_xlsx(
                seed_jsonl=seed_jsonl,
                focus_registry=Path("docs/research/focus_registry_v1.json"),
                output_xlsx=output_path,
            )
            workbook = load_workbook(output_path)

        sheet = workbook["标注表"]
        headers = [cell.value for cell in sheet[2]]
        self.assertEqual(1, row_count)
        self.assertEqual(
            "https://www.luogu.com.cn/problem/P1002",
            sheet.cell(row=3, column=headers.index("problem_source_url") + 1).value,
        )
        self.assertIn(
            "过河卒",
            sheet.cell(row=3, column=headers.index("problem_statement") + 1).value,
        )

    def test_export_xlsx_carries_optional_dialogue_state_fields(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            seed_jsonl = tmp / "dialogue_v3_seed.jsonl"
            output_path = tmp / "coach_v2.xlsx"
            seed_jsonl.write_text(
                json.dumps(
                    {
                        "case_id": "dialogue_v3_001_predicate_check_semantics",
                        "source_case_id": "heldout_v2_luogu_001",
                        "problem_ref": "luogu_p1001_check",
                        "problem_source_platform": "luogu",
                        "problem_source_id": "P1001",
                        "problem_source_url": "https://www.luogu.com.cn/problem/P1001",
                        "problem_statement": "题目：样例题\n题意摘录：判断候选值是否满足限制。",
                        "problem_statement_public_summary": "洛谷 P1001 样例题：判断候选值。",
                        "problem_statement_rights_note": "local review only",
                        "problem_statement_access_level": "local_review_only",
                        "student_message": "true 就是能满足限制。",
                        "problem_context": "洛谷 P1001。",
                        "recent_dialogue": "学生：check 该判什么？\nAI：你先说 true 的含义。",
                        "student_code_excerpt": "N/A",
                        "turn_position": "followup",
                        "context_type": "followup_after_correct_short_answer",
                        "student_scaffold_followability": "F1",
                        "followability_label_confidence": "high",
                        "followability_evidence_quote": "能满足限制",
                        "prior_ai_scaffold": "你先说 true 的含义。",
                        "student_reply_to_prior_scaffold": "true 就是能满足限制。",
                        "expected_tutor_move": "advance",
                        "fixed_recent_dialogue_source": "synthetic_dialogue_state_v3",
                    },
                    ensure_ascii=False,
                )
                + "\n",
                encoding="utf-8",
            )

            workbook_v2.export_xlsx(
                seed_jsonl=seed_jsonl,
                focus_registry=Path("docs/research/focus_registry_v1.json"),
                output_xlsx=output_path,
            )
            workbook = load_workbook(output_path)

        sheet = workbook["标注表"]
        headers = [cell.value for cell in sheet[2]]
        self.assertIn("turn_position", headers)
        self.assertIn("student_scaffold_followability", headers)
        self.assertIn("prior_ai_scaffold", headers)
        self.assertIn("student_reply_to_prior_scaffold", headers)
        self.assertEqual("F1", sheet.cell(row=3, column=headers.index("student_scaffold_followability") + 1).value)
        self.assertEqual("advance", sheet.cell(row=3, column=headers.index("expected_tutor_move") + 1).value)

    def test_export_xlsx_uses_chinese_dropdown_labels_and_v21_options(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "coach_v2.xlsx"
            workbook_v2.export_xlsx(
                seed_jsonl=Path("docs/research/bridgebench_cp_seed_v1.jsonl"),
                focus_registry=Path("docs/research/focus_registry_v1.json"),
                output_xlsx=output_path,
            )

            workbook = load_workbook(output_path)

        turn_type_options = _option_values(workbook, "turn_type")
        scaffold_options = _option_values(workbook, "max_scaffold_level")
        policy_risk_options = _option_values(workbook, "policy_risk_type")
        self.assertIn("可诊断学习轮：学生给出具体卡点或尝试，可判断下一步（diagnosable_learning_turn）", turn_type_options)
        self.assertIn("完整代码请求：要求可提交代码或大段代写（complete_code_request）", turn_type_options)
        self.assertIn("关键桥请求：直接索要状态、转移、check、公式、标记规则（critical_bridge_request）", turn_type_options)
        self.assertFalse(any("direct_answer_request" in value for value in turn_type_options))
        self.assertIn("澄清/要证据：只要题面、尝试、代码、错误现象，不给实质解题提示（L0）", scaffold_options)
        self.assertFalse(any("L4_forbidden" in value for value in scaffold_options))
        self.assertIn("完整代码风险：可能代写可提交代码（complete_code_risk）", policy_risk_options)
        self.assertGreaterEqual(len(workbook["标注表"].data_validations.dataValidation), 18)

        state_options = _option_values(workbook, "student_problem_solving_state")
        self.assertIn(
            "方法应用/变式卡住：知道方向、算法或知识点，但不会落到当前题的关键步骤（method_application_gap）",
            state_options,
        )

    def test_export_xlsx_uses_abstract_subtype_and_forbidden_options(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "coach_v2.xlsx"
            workbook_v2.export_xlsx(
                seed_jsonl=Path("docs/research/bridgebench_cp_seed_v1.jsonl"),
                focus_registry=Path("docs/research/focus_registry_v1.json"),
                output_xlsx=output_path,
            )

            workbook = load_workbook(output_path)

        subtype_options = _option_values(workbook, "bridge_subtype")
        forbidden_options = _option_values(workbook, "bridge_specific_forbidden_content")

        self.assertTrue(any("过期候选判断" in value for value in subtype_options))
        self.assertTrue(any("单次使用更新顺序" in value for value in subtype_options))
        self.assertTrue(any("单调候选答案搜索信号" in value for value in subtype_options))
        self.assertTrue(any("表格/记忆化格子语义" in value for value in subtype_options))
        self.assertTrue(any("延迟更新语义" in value for value in subtype_options))
        self.assertTrue(any("边界推进方向" in value for value in subtype_options))
        self.assertTrue(any("边界差量更新" in value for value in subtype_options))
        self.assertTrue(any("优先级候选操作映射" in value for value in subtype_options))
        self.assertTrue(any("连通分量合并/查询映射" in value for value in subtype_options))
        self.assertTrue(any("被支配候选栈映射" in value for value in subtype_options))
        self.assertFalse(any("Dijkstra 旧项判断" in value for value in subtype_options))
        self.assertFalse(any("KMP 前缀函数语义" in value for value in subtype_options))
        for algorithm_specific_label in [
            "二分答案信号",
            "DP 状态候选信号",
            "DP 状态语义",
            "lazy 标记语义",
            "check 真假方向",
            "二分边界更新",
            "最短路 relax 条件",
            "区间 DP 顺序",
            "一维前缀和",
            "二维前缀和",
            "差分数组区间更新",
            "堆操作映射",
            "并查集操作映射",
            "单调栈操作映射",
            "贪心交换论证",
        ]:
            self.assertFalse(any(algorithm_specific_label in value for value in subtype_options))

        self.assertTrue(any("完整守卫/跳过条件" in value for value in forbidden_options))
        self.assertTrue(any("完整边界更新规则" in value for value in forbidden_options))
        self.assertTrue(any("完整推演轨迹" in value for value in forbidden_options))
        self.assertFalse(any("Dijkstra" in value or "SPFA" in value or "Floyd" in value for value in forbidden_options))

    def test_export_xlsx_can_limit_trial_workbook_rows(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "coach_v2_20.xlsx"

            row_count = workbook_v2.export_xlsx(
                seed_jsonl=Path("docs/research/bridgebench_cp_seed_v1.jsonl"),
                focus_registry=Path("docs/research/focus_registry_v1.json"),
                output_xlsx=output_path,
                limit=20,
            )

            workbook = load_workbook(output_path)

        sheet = workbook["标注表"]
        self.assertEqual(20, row_count)
        self.assertEqual(22, sheet.max_row)
        self.assertEqual("cp_bridge_001", sheet["A3"].value)
        self.assertEqual("cp_bridge_020", sheet["A22"].value)


if __name__ == "__main__":
    unittest.main()
