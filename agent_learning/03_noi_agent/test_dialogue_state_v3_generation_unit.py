import json
import tempfile
import unittest
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from evals.aichat import generate_dialogue_state_v3_50
from evals.aichat import validate_heldout_50_dataset


def _v2_case(idx: int, bucket: str = "predicate_check_semantics") -> dict:
    return {
        "case_id": f"heldout_v2_luogu_{idx:03d}",
        "id": f"heldout_v2_luogu_{idx:03d}",
        "category": bucket,
        "bridge_bucket": bucket,
        "bridge_bucket_zh": "判定条件/check",
        "problem_ref": f"luogu_p{1000 + idx}_{bucket}",
        "problem_source_platform": "luogu",
        "problem_source_id": f"P{1000 + idx}",
        "problem_source_url": f"https://www.luogu.com.cn/problem/P{1000 + idx}",
        "problem_statement": f"题目：样例题 {idx}\n题意摘录：给定若干数据，要求判断候选答案是否满足限制。",
        "problem_statement_public_summary": f"洛谷 P{1000 + idx}《样例题 {idx}》：判断候选答案是否满足限制。",
        "problem_statement_rights_note": "local review only",
        "problem_statement_access_level": "local_review_only",
        "problem_title": f"样例题 {idx}",
        "problem_tags": ["二分"],
        "student_message": "check 该判什么？",
        "student_message_length_bucket": "short",
        "problem_context": f"洛谷 P{1000 + idx}《样例题 {idx}》：判断候选答案是否满足限制。",
        "recent_dialogue": "N/A",
        "recent_dialogue_bucket": "none",
        "student_code_excerpt": "N/A",
        "student_known_state": "学生知道可能要判断候选值。",
        "missing_bridge": "缺少把候选值、限制和可行性方向对应起来的关系。",
        "allowed_help_level": "L2",
        "forbidden_content": ["no_complete_check_condition"],
        "success_criteria": ["回复让学生判断 true/false 的含义。"],
        "review_notes_for_coach": "draft",
        "reference_label_status": "draft_needs_coach_review",
        "primary_bridge_family": "predicate_check_bridge",
        "primary_bridge_subtype_id": "predicate.feasibility_truth_direction",
        "registered_focus_id": "check_truth_direction",
        "algorithm_topic_l1": "binary_search",
    }


def _length_bucket(value: str) -> str:
    length = len("".join(str(value or "").split()))
    if length <= 30:
        return "short"
    if length <= 70:
        return "medium_short"
    if length <= 140:
        return "medium_long"
    return "long"


class DialogueStateV3GenerationTests(unittest.TestCase):
    def test_build_dialogue_state_cases_adds_followup_fields(self):
        v2_rows = [_v2_case(i) for i in range(1, 51)]

        cases = generate_dialogue_state_v3_50.build_dialogue_state_cases(v2_rows)

        self.assertEqual(50, len(cases))
        self.assertEqual(50, len({row["case_id"] for row in cases}))
        self.assertTrue(all(row["case_id"].startswith("dialogue_v3_") for row in cases))
        self.assertTrue(any(row["turn_position"] == "followup" for row in cases))
        for row in cases:
            self.assertIn(row["turn_position"], {"initial", "followup"})
            self.assertIn("context_type", row)
            self.assertEqual("synthetic_dialogue_state_v3", row["fixed_recent_dialogue_source"])
            if row["turn_position"] == "followup":
                self.assertIn(row["student_scaffold_followability"], {"F1", "F2", "F3", "F4"})
                self.assertIn(
                    row["expected_tutor_move"],
                    {"advance", "clarify", "micro_step", "prerequisite_repair", "safe_redirect"},
                )
                self.assertIn(row["followability_label_confidence"], {"high", "medium", "low"})
                self.assertTrue(row["followability_evidence_quote"].strip())
                self.assertIn(row["followability_evidence_quote"], row["student_reply_to_prior_scaffold"])
                self.assertIn("AI：", row["recent_dialogue"])
                self.assertIn("学生：", row["recent_dialogue"])
                self.assertEqual(row["prior_ai_scaffold"], row["context_ai_reply"])
            else:
                self.assertEqual("NA", row["student_scaffold_followability"])
                self.assertEqual("", row["context_ai_reply"])

    def test_default_distribution_is_followup_majority(self):
        v2_rows = [_v2_case(i) for i in range(1, 51)]

        cases = generate_dialogue_state_v3_50.build_dialogue_state_cases(v2_rows)
        context_counts = Counter(row["context_type"] for row in cases)
        followability_counts = Counter(row.get("student_scaffold_followability", "NA") for row in cases)

        self.assertEqual(10, context_counts["initial_question"])
        self.assertEqual(40, sum(count for context, count in context_counts.items() if context != "initial_question"))
        self.assertEqual(7, followability_counts["F1"])
        self.assertEqual(18, followability_counts["F2"])
        self.assertEqual(10, followability_counts["F3"])
        self.assertEqual(5, followability_counts["F4"])
        self.assertEqual(10, followability_counts["NA"])

    def test_generated_cases_preserve_heldout_length_and_context_quotas(self):
        target_buckets = ["short"] * 20 + ["medium_short"] * 15 + ["medium_long"] * 10 + ["long"] * 5
        target_recent = ["none"] * 10 + ["short"] * 25 + ["long"] * 15
        v2_rows = []
        for index in range(1, 51):
            row = _v2_case(index)
            row["student_message_length_bucket"] = target_buckets[index - 1]
            row["recent_dialogue_bucket"] = target_recent[index - 1]
            v2_rows.append(row)

        cases = generate_dialogue_state_v3_50.build_dialogue_state_cases(v2_rows)
        actual_length = Counter(_length_bucket(row["student_message"]) for row in cases)
        actual_recent = Counter(row["recent_dialogue_bucket"] for row in cases)

        self.assertEqual(
            {"short": 20, "medium_short": 15, "medium_long": 10, "long": 5},
            dict(actual_length),
        )
        self.assertEqual({"none": 10, "short": 25, "long": 15}, dict(actual_recent))
        self.assertTrue(all(row["context_ai_reply"] for row in cases if row["turn_position"] == "followup"))

    def test_followup_replies_match_bridge_bucket(self):
        v2_rows = [_v2_case(i) for i in range(1, 51)]
        overrides = {
            11: ("transition_recurrence_source", "转移/递推来源"),
            18: ("boundary_update_order", "边界更新/循环方向"),
            26: ("modeling_object_relation", "建模/对象关系"),
            43: ("implementation_boundary", "实现边界/初始化/类型"),
        }
        for index, (bucket, bucket_zh) in overrides.items():
            v2_rows[index - 1]["bridge_bucket"] = bucket
            v2_rows[index - 1]["category"] = bucket
            v2_rows[index - 1]["bridge_bucket_zh"] = bucket_zh
        cases = generate_dialogue_state_v3_50.build_dialogue_state_cases(v2_rows)
        by_source = {row["source_case_id"]: row for row in cases}

        transition = by_source["heldout_v2_luogu_011"]
        self.assertIn("来源", transition["student_message"])

        boundary = by_source["heldout_v2_luogu_018"]
        self.assertNotIn("分支怎么拆", boundary["student_message"])
        self.assertTrue(any(token in boundary["student_message"] for token in ["旧值", "覆盖", "顺序"]))

        modeling = by_source["heldout_v2_luogu_026"]
        self.assertNotIn("true", modeling["student_message"])
        self.assertNotIn("边界", modeling["student_message"])
        self.assertTrue(any(token in modeling["student_message"] for token in ["对象", "关系", "覆盖"]))

        implementation = by_source["heldout_v2_luogu_043"]
        self.assertNotIn("哪一边还可能包含答案", implementation["prior_ai_scaffold"])
        self.assertNotIn("可行性/状态语义", implementation["student_message"])
        self.assertTrue(any(token in implementation["student_message"] for token in ["下标", "字符", "输入"]))

    def test_coach_a_calibration_revisions_are_reflected_in_generated_cases(self):
        v2_rows = [_v2_case(i) for i in range(1, 51)]
        overrides = {
            11: {
                "bridge_bucket": "transition_recurrence_source",
                "category": "transition_recurrence_source",
                "bridge_bucket_zh": "转移/递推来源",
            },
            18: {
                "bridge_bucket": "boundary_update_order",
                "category": "boundary_update_order",
                "bridge_bucket_zh": "边界更新/循环方向",
                "problem_source_id": "P1060",
                "problem_title": "[NOIP 2006 普及组] 开心的金明",
                "problem_tags": ["动态规划 DP", "背包 DP"],
            },
            43: {
                "bridge_bucket": "implementation_boundary",
                "category": "implementation_boundary",
                "bridge_bucket_zh": "实现边界/初始化/类型",
                "problem_source_id": "P1765",
                "problem_title": "手机",
                "problem_tags": ["模拟", "字符串"],
                "student_message_length_bucket": "long",
            },
        }
        for index, patch in overrides.items():
            v2_rows[index - 1].update(patch)

        cases = generate_dialogue_state_v3_50.build_dialogue_state_cases(v2_rows)
        by_source = {row["source_case_id"]: row for row in cases}

        transition = by_source["heldout_v2_luogu_011"]
        self.assertEqual("followup_after_partial_answer", transition["context_type"])
        self.assertEqual("F2", transition["student_scaffold_followability"])
        self.assertEqual("clarify", transition["expected_tutor_move"])
        self.assertIn("没对应到", transition["student_message"])
        self.assertEqual("short", _length_bucket(transition["student_message"]))
        self.assertTrue(any("前驱来源" in item for item in transition["success_criteria"]))

        knapsack = by_source["heldout_v2_luogu_018"]
        self.assertNotIn("哪一边还可能包含答案", knapsack["prior_ai_scaffold"])
        self.assertIn("旧值", knapsack["prior_ai_scaffold"])
        self.assertTrue(any(token in knapsack["prior_ai_scaffold"] for token in ["覆盖", "复用"]))
        self.assertTrue(all("哪一段仍可能包含答案" not in item for item in knapsack["success_criteria"]))
        self.assertTrue(any("本轮新值" in item or "上一轮" in item for item in knapsack["success_criteria"]))

        other_knapsack_tag = by_source["heldout_v2_luogu_017"]
        self.assertNotIn("本轮刚算出的新值覆盖或复用", other_knapsack_tag["prior_ai_scaffold"])

        phone = by_source["heldout_v2_luogu_043"]
        self.assertIn("字符", phone["prior_ai_scaffold"])
        self.assertIn("按键", phone["prior_ai_scaffold"])
        self.assertNotIn("输入范围", phone["student_message"])
        self.assertNotIn("下标或输入边界", phone["student_message"])
        self.assertNotIn("如果可以的话", phone["student_message"])
        self.assertIn("按键", phone["student_message"])
        self.assertEqual("medium_long", _length_bucket(phone["student_message"]))
        self.assertIn("单个字符", phone["missing_bridge"])
        self.assertIn("no_full_keypress_mapping_table", phone["forbidden_content"])
        self.assertNotIn("no_direct_current_bridge", phone["forbidden_content"])
        self.assertTrue(any("按键次数" in item for item in phone["success_criteria"]))

    def test_default_luogu_source_keeps_review_gate_distributions_after_calibration(self):
        source_path = Path("docs/research/bridgebench_cp_heldout_v2_50_draft.jsonl")
        source_rows = [
            json.loads(line)
            for line in source_path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]

        cases = generate_dialogue_state_v3_50.build_dialogue_state_cases(source_rows)
        by_id = {row["case_id"]: row for row in cases}
        length_counts = Counter(_length_bucket(row["student_message"]) for row in cases)
        followability_counts = Counter(row["student_scaffold_followability"] for row in cases)

        self.assertEqual({"short": 20, "medium_short": 15, "medium_long": 10, "long": 5}, dict(length_counts))
        self.assertEqual({"NA": 10, "F1": 6, "F2": 19, "F3": 10, "F4": 5}, dict(followability_counts))
        self.assertEqual("F2", by_id["dialogue_v3_011_transition_recurrence_source"]["student_scaffold_followability"])
        self.assertEqual("followup_after_partial_answer", by_id["dialogue_v3_011_transition_recurrence_source"]["context_type"])
        self.assertEqual("short", _length_bucket(by_id["dialogue_v3_011_transition_recurrence_source"]["student_message"]))
        self.assertNotIn(
            "本轮刚算出的新值覆盖或复用",
            by_id["dialogue_v3_017_boundary_update_order"]["prior_ai_scaffold"],
        )

    def test_generated_cases_pass_heldout_dataset_validator(self):
        target_buckets = ["short"] * 20 + ["medium_short"] * 15 + ["medium_long"] * 10 + ["long"] * 5
        target_recent = ["none"] * 10 + ["short"] * 25 + ["long"] * 15
        v2_rows = []
        for index in range(1, 51):
            row = _v2_case(index)
            row["student_message_length_bucket"] = target_buckets[index - 1]
            row["recent_dialogue_bucket"] = target_recent[index - 1]
            v2_rows.append(row)

        with tempfile.TemporaryDirectory() as tmpdir:
            dataset = Path(tmpdir) / "dialogue_v3.jsonl"
            dataset.write_text(
                "\n".join(
                    json.dumps(row, ensure_ascii=False)
                    for row in generate_dialogue_state_v3_50.build_dialogue_state_cases(v2_rows)
                )
                + "\n",
                encoding="utf-8",
            )
            report = validate_heldout_50_dataset.validate_dataset(
                dataset,
                expected_count=50,
                dev_seed_paths=[],
                max_no_recent_dialogue=10,
                min_long_recent_dialogue=15,
                min_code_excerpt=10,
                min_student_message_length_buckets={
                    "short": 20,
                    "medium_short": 15,
                    "medium_long": 10,
                    "long": 5,
                },
            )

        self.assertTrue(report["ok"], report["errors"])

    def test_export_review_workbook_has_dialogue_state_columns(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "dialogue_review.xlsx"
            cases = generate_dialogue_state_v3_50.build_dialogue_state_cases([_v2_case(i) for i in range(1, 51)])

            generate_dialogue_state_v3_50.export_dialogue_state_review_workbook(cases, output)
            workbook = load_workbook(output)

        self.assertIn("总表", workbook.sheetnames)
        sheet = workbook["总表"]
        headers = [sheet.cell(row=1, column=col).value for col in range(1, sheet.max_column + 1)]
        self.assertIn("上下文类型", headers)
        self.assertIn("上下文 AI 回复", headers)
        self.assertIn("学生跟随状态", headers)
        self.assertIn("上一轮 AI 脚手架", headers)
        self.assertIn("学生对脚手架的回答", headers)
        self.assertIn("跟随状态证据", headers)
        self.assertIn("题源是否可用", headers)
        self.assertIn("上下文是否连贯", headers)
        self.assertIn("学生话术是否真实", headers)
        self.assertIn("样本处理决定", headers)
        self.assertIn("教练修改建议", headers)

    def test_export_review_workbook_has_structured_review_validations(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "dialogue_review.xlsx"
            cases = generate_dialogue_state_v3_50.build_dialogue_state_cases([_v2_case(i) for i in range(1, 51)])

            generate_dialogue_state_v3_50.export_dialogue_state_review_workbook(cases, output)
            workbook = load_workbook(output)

        sheet = workbook["总表"]
        headers = [sheet.cell(row=1, column=col).value for col in range(1, sheet.max_column + 1)]
        decision_col = headers.index("样本处理决定") + 1
        confidence_col = headers.index("审核置信度") + 1
        self.assertIn(sheet.cell(row=3, column=decision_col).value, {"", None})
        self.assertIn(sheet.cell(row=3, column=confidence_col).value, {"", None})
        validation_formulas = {validation.formula1 for validation in sheet.data_validations.dataValidation}
        self.assertIn('"接受,修改,丢弃,讨论"', validation_formulas)
        self.assertIn('"高,中,低"', validation_formulas)

    def test_chinese_review_workbook_uses_chinese_display_values(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "dialogue_review.xlsx"
            cases = generate_dialogue_state_v3_50.build_dialogue_state_cases([_v2_case(i) for i in range(1, 51)])

            generate_dialogue_state_v3_50.export_dialogue_state_review_workbook(cases, output)
            workbook = load_workbook(output)

        sheet = workbook["总表"]
        headers = [sheet.cell(row=1, column=col).value for col in range(1, sheet.max_column + 1)]
        row3 = {
            header: sheet.cell(row=3, column=index + 1).value
            for index, header in enumerate(headers)
        }
        self.assertEqual("初始提问", row3["轮次位置"])
        self.assertEqual("初始提问", row3["上下文类型"])
        self.assertEqual("不适用", row3["学生跟随状态"])
        self.assertEqual("拆成更小一步", row3["期望下一步教学动作"])
        self.assertNotIn("initial_question", row3.values())
        self.assertNotIn("micro_step", row3.values())

    def test_export_review_workbook_has_instruction_and_bridge_bucket_sheets(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "dialogue_review.xlsx"
            v2_rows = [_v2_case(i) for i in range(1, 51)]
            v2_rows[0]["bridge_bucket"] = "state_representation_semantics"
            v2_rows[0]["category"] = "state_representation_semantics"
            v2_rows[0]["bridge_bucket_zh"] = "状态/表示语义"
            cases = generate_dialogue_state_v3_50.build_dialogue_state_cases(v2_rows)

            generate_dialogue_state_v3_50.export_dialogue_state_review_workbook(cases, output)
            workbook = load_workbook(output)

        self.assertIn("评审说明", workbook.sheetnames)
        self.assertIn("总表", workbook.sheetnames)
        self.assertIn("状态表示", workbook.sheetnames)
        self.assertIn("判定条件", workbook.sheetnames)
        self.assertNotIn("dialogue_state_review", workbook.sheetnames)
        self.assertTrue(all(not name.startswith("bucket_") for name in workbook.sheetnames))
        instruction_values = [
            workbook["评审说明"].cell(row=row, column=1).value
            for row in range(1, workbook["评审说明"].max_row + 1)
        ]
        self.assertTrue(any("先看原题题面" in str(value or "") for value in instruction_values))
        self.assertTrue(any("优先在各桥梁桶 sheet 中填写审核列" in str(value or "") for value in instruction_values))
        self.assertTrue(any("总表用于全局查看" in str(value or "") for value in instruction_values))
        self.assertTrue(any("固定中文选项" in str(value or "") for value in instruction_values))
        self.assertTrue(any("接受=accept" in str(value or "") for value in instruction_values))
        self.assertTrue(any("reviewer_id 建议填写 coach_A" in str(value or "") for value in instruction_values))
        self.assertTrue(any("review_round 建议填写 calibration" in str(value or "") for value in instruction_values))
        self.assertTrue(any("context_ai_reply" in str(value or "") for value in instruction_values))
        self.assertTrue(any("不用于评价该 AI 回复本身好坏" in str(value or "") for value in instruction_values))
        bucket_headers = [
            workbook["状态表示"].cell(row=1, column=col).value
            for col in range(1, workbook["状态表示"].max_column + 1)
        ]
        self.assertIn("目标缺失桥梁", bucket_headers)
        self.assertIn("审核人", bucket_headers)
        self.assertIn("审核轮次", bucket_headers)

    def test_export_review_workbook_can_emit_english_coach_version(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "dialogue_review.en.xlsx"
            v2_rows = [_v2_case(i) for i in range(1, 51)]
            v2_rows[0]["bridge_bucket"] = "state_representation_semantics"
            v2_rows[0]["category"] = "state_representation_semantics"
            cases = generate_dialogue_state_v3_50.build_dialogue_state_cases(v2_rows)

            generate_dialogue_state_v3_50.export_dialogue_state_review_workbook(cases, output, language="en")
            workbook = load_workbook(output)

        self.assertIn("Instructions", workbook.sheetnames)
        self.assertIn("dialogue_state_review", workbook.sheetnames)
        self.assertIn("bucket_state_representation", workbook.sheetnames)
        instruction_values = [
            workbook["Instructions"].cell(row=row, column=1).value
            for row in range(1, workbook["Instructions"].max_row + 1)
        ]
        self.assertTrue(any("Read the problem statement first" in str(value or "") for value in instruction_values))
        self.assertTrue(any("Prefer filling the bucket sheets" in str(value or "") for value in instruction_values))
        self.assertTrue(any("reviewer_id should be coach_A" in str(value or "") for value in instruction_values))
        self.assertTrue(any("review_round should be calibration" in str(value or "") for value in instruction_values))
        headers = [
            workbook["dialogue_state_review"].cell(row=1, column=col).value
            for col in range(1, workbook["dialogue_state_review"].max_column + 1)
        ]
        self.assertIn("Current Student Message / Reply", headers)
        self.assertIn("Target Missing Bridge", headers)
        self.assertIn("Forbidden Content", headers)
        self.assertIn("Case Decision", headers)
        self.assertIn("Coach Fix Suggestion", headers)
        self.assertIn("Reviewer ID", headers)
        self.assertIn("Review Round", headers)

    def test_main_does_not_overwrite_v2_source(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            source = tmp / "v2.jsonl"
            output = tmp / "v3.jsonl"
            workbook = tmp / "review.xlsx"
            workbook_en = tmp / "review.en.xlsx"
            report_json = tmp / "report.json"
            report_zh = tmp / "report.zh.md"
            report_en = tmp / "report.md"
            source.write_text(
                "\n".join(json.dumps(_v2_case(i), ensure_ascii=False) for i in range(1, 51)) + "\n",
                encoding="utf-8",
            )
            before = source.read_text(encoding="utf-8")

            exit_code = generate_dialogue_state_v3_50.main(
                [
                    "--source-jsonl",
                    str(source),
                    "--output-jsonl",
                    str(output),
                    "--review-xlsx",
                    str(workbook),
                    "--review-xlsx-en",
                    str(workbook_en),
                    "--report-json",
                    str(report_json),
                    "--report-zh",
                    str(report_zh),
                    "--report-en",
                    str(report_en),
                ]
            )

            self.assertEqual(0, exit_code)
            self.assertEqual(before, source.read_text(encoding="utf-8"))
            self.assertTrue(output.exists())
            self.assertTrue(workbook.exists())
            self.assertTrue(workbook_en.exists())
            self.assertTrue(report_json.exists())


if __name__ == "__main__":
    unittest.main()
