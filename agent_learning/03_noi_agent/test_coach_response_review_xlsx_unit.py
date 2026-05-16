import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from evals.aichat import export_coach_response_review_workbook_xlsx as xlsx_export


class CoachResponseReviewXlsxTests(unittest.TestCase):
    def test_export_xlsx_uses_chinese_headers_and_dropdown_options(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            input_csv = Path(tmpdir) / "review.csv"
            output_xlsx = Path(tmpdir) / "review.zh.xlsx"
            with input_csv.open("w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=xlsx_export.REVIEW_COLUMNS, lineterminator="\n")
                writer.writeheader()
                writer.writerow(
                    {
                        "case_id": "case_1",
                        "anonymized_response_id": "resp_001",
                        "problem_ref": "P1048",
                        "problem_source_platform": "luogu",
                        "problem_source_id": "P1048",
                        "problem_source_url": "https://www.luogu.com.cn/problem/P1048",
                        "problem_statement": "给定 n 个草药，每个草药有时间和价值，在总时间 M 内选择若干草药使价值最大。",
                        "problem_statement_public_summary": "洛谷 P1048 采药：在总时间限制内选择若干物品最大化价值。",
                        "problem_statement_rights_note": "本地教练评审使用必要题面；公开材料仅保留来源链接和改写摘要。",
                        "problem_statement_access_level": "public_summary_only",
                        "student_message": "我不知道状态怎么设。",
                        "student_message_length_bucket": "short",
                        "problem_context": "采药。",
                        "recent_dialogue": "",
                        "context_ai_reply": "上一轮 AI 让学生先观察容量变化。",
                        "context_alignment_flag": "no_recent_dialogue",
                        "response_text": "先想一格 dp 应该存什么。",
                        "review_status": "unlabeled",
                    }
                )

            row_count = xlsx_export.export_xlsx(input_csv=input_csv, output_xlsx=output_xlsx)
            workbook = load_workbook(output_xlsx)

        self.assertEqual(1, row_count)
        self.assertIn("盲评表", workbook.sheetnames)
        self.assertIn("评分流程", workbook.sheetnames)
        self.assertIn("评分说明", workbook.sheetnames)
        self.assertIn("下拉选项", workbook.sheetnames)
        sheet = workbook["盲评表"]
        columns_by_field = {
            sheet.cell(row=2, column=column_idx).value: column_idx
            for column_idx in range(1, sheet.max_column + 1)
        }

        def chinese_header(field: str) -> str:
            return sheet.cell(row=1, column=columns_by_field[field]).value

        def value(field: str):
            return sheet.cell(row=3, column=columns_by_field[field]).value

        self.assertEqual("题目来源平台", chinese_header("problem_source_platform"))
        self.assertEqual("平台题号", chinese_header("problem_source_id"))
        self.assertEqual("原题链接", chinese_header("problem_source_url"))
        self.assertEqual("原题题面/必要题面", chinese_header("problem_statement"))
        self.assertEqual("公开题面摘要", chinese_header("problem_statement_public_summary"))
        self.assertEqual("题面版权/使用说明", chinese_header("problem_statement_rights_note"))
        self.assertEqual("题面访问级别", chinese_header("problem_statement_access_level"))
        self.assertEqual("学生当前问题", chinese_header("student_message"))
        self.assertEqual("学生问题长度类型", chinese_header("student_message_length_bucket"))
        self.assertEqual("近期对话", chinese_header("recent_dialogue"))
        self.assertEqual("上下文 AI 回复", chinese_header("context_ai_reply"))
        self.assertEqual("上下文对齐状态", chinese_header("context_alignment_flag"))
        self.assertEqual("本 case 成功标准", chinese_header("success_criteria"))
        self.assertEqual("本轮禁止补完", chinese_header("forbidden_content"))
        self.assertEqual("关键桥泄露边界", chinese_header("critical_bridge_boundary"))
        self.assertEqual("允许透露/合理解释", chinese_header("acceptable_reveal"))
        self.assertEqual("期望学生下一步", chinese_header("expected_student_next_action"))
        self.assertEqual("AI 回复（要评分）", chinese_header("response_text"))
        self.assertEqual("是否抓住卡点 0-2", chinese_header("coach_bridge_identification_score"))
        self.assertEqual("case_1", sheet["A3"].value)
        self.assertEqual("luogu", value("problem_source_platform"))
        self.assertEqual("P1048", value("problem_source_id"))
        self.assertEqual("https://www.luogu.com.cn/problem/P1048", value("problem_source_url"))
        self.assertEqual("给定 n 个草药，每个草药有时间和价值，在总时间 M 内选择若干草药使价值最大。", value("problem_statement"))
        self.assertEqual("洛谷 P1048 采药：在总时间限制内选择若干物品最大化价值。", value("problem_statement_public_summary"))
        self.assertEqual("short", value("student_message_length_bucket"))
        self.assertEqual("帮助是否足够 0-2", chinese_header("coach_scaffold_sufficiency_score"))
        self.assertEqual("桥梁导向微型例子 0-2", chinese_header("coach_bridge_oriented_micro_example_score"))
        self.assertEqual("微型例子是否适用", chinese_header("coach_micro_example_applicability"))
        self.assertEqual("泄露标签", chinese_header("coach_leakage_label"))
        self.assertEqual("关键桥透露是否有教学理由", chinese_header("coach_bridge_reveal_justification"))
        self.assertEqual("总体质量 1-5", chinese_header("coach_overall_quality_score"))
        self.assertEqual("是否愿意给学生看", chinese_header("coach_would_show_to_student"))
        self.assertEqual("学生回复负担", chinese_header("coach_student_response_burden"))
        self.assertEqual("评分置信度", chinese_header("coach_reviewer_confidence"))
        self.assertEqual("是否需要讨论", chinese_header("coach_needs_discussion"))
        self.assertEqual("unlabeled", value("review_status"))
        self.assertTrue(sheet.freeze_panes)
        self.assertGreaterEqual(len(sheet.data_validations.dataValidation), 16)
        self.assertGreaterEqual(len(sheet.conditional_formatting), 2)
        options = workbook["下拉选项"]
        self.assertEqual("2｜好：直接针对学生当前缺失的桥梁/卡点", options["A2"].value)
        self.assertEqual("2｜好：信息足够推进，既不泄露也不空泛", options["D2"].value)
        self.assertEqual("2｜好：例子能引导学生抽象可迁移的桥梁关系", options["H2"].value)
        self.assertEqual("applicable｜适用：这条回复使用或应该使用微型例子", options["I2"].value)
        self.assertEqual("no_leakage｜无泄露：没有说穿当前关键桥", options["J2"].value)
        self.assertEqual("no_reveal｜未透露：没有实质透露当前关键桥", options["K2"].value)
        self.assertEqual("5｜优秀：非常愿意给学生看", options["M2"].value)
        self.assertEqual("yes｜愿意：可以直接给学生看", options["N2"].value)
        self.assertEqual(
            "low｜低：一两个关键词、局部判断或一句短句；选择题仅限不夹答案的低风险判断",
            options["O2"].value,
        )
        guide = workbook["评分说明"]
        guide_values = "\n".join(str(cell.value or "") for row in guide.iter_rows() for cell in row)
        self.assertIn("0-2", guide_values)
        self.assertIn("例子", guide_values)
        self.assertIn("状态怎么设", guide_values)
        self.assertIn("总体质量", guide_values)
        self.assertIn("评分置信度", guide_values)
        self.assertIn("关键桥透露是否有教学理由", guide_values)
        self.assertIn("过早、过完整", guide_values)
        self.assertIn("最低足够学生努力", guide_values)
        self.assertIn("优先短生成", guide_values)
        self.assertIn("慎用选择题", guide_values)
        self.assertIn("一两个关键词", guide_values)
        self.assertIn("学生回复负担", guide_values)
        self.assertIn("完整表格", guide_values)
        self.assertIn("帮助是否足够", guide_values)
        self.assertIn("过度保留", guide_values)
        self.assertIn("安全但没帮助", guide_values)
        self.assertIn("校准与双标", guide_values)
        self.assertIn("agreement / adjudication", guide_values)
        self.assertIn("重大泄露", guide_values)
        self.assertIn("推荐格式", guide_values)

        workflow = workbook["评分流程"]
        workflow_values = "\n".join(str(cell.value or "") for row in workflow.iter_rows() for cell in row)
        self.assertIn("教练评分顺序", workflow_values)
        self.assertIn("校准轮", workflow_values)
        self.assertIn("calibration samples", workflow_values)
        self.assertIn("先看题目/题面摘要", workflow_values)
        self.assertIn("再看近期对话", workflow_values)
        self.assertIn("上下文 AI 回复", workflow_values)
        self.assertIn("学生当前问题", workflow_values)
        self.assertIn("最后只评价 AI 回复（要评分）", workflow_values)
        self.assertIn("派生显示字段", workflow_values)
        self.assertIn("不是额外模型输入", workflow_values)
        self.assertIn("近期对话和学生当前问题明显不接", workflow_values)
        self.assertIn("补强制备注", workflow_values)
        self.assertIn("低置信", workflow_values)

    def test_export_xlsx_can_show_dialogue_state_context_columns(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            input_csv = Path(tmpdir) / "review.csv"
            output_xlsx = Path(tmpdir) / "review.zh.xlsx"
            fieldnames = [
                *xlsx_export.REVIEW_COLUMNS,
            ]
            with input_csv.open("w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=fieldnames, lineterminator="\n")
                writer.writeheader()
                writer.writerow(
                    {
                        "case_id": "dialogue_v3_001_predicate_check_semantics",
                        "anonymized_response_id": "resp_001",
                        "problem_ref": "P1001",
                        "student_message": "true 就是能满足限制。",
                        "student_message_length_bucket": "short",
                        "problem_context": "洛谷 P1001。",
                        "recent_dialogue": "学生：check 该判什么？\nAI：你先说 true 的含义。",
                        "context_ai_reply": "你先说 true 的含义。",
                        "context_alignment_flag": "aligned_prior_context_ends_with_assistant",
                        "turn_position": "followup",
                        "context_type": "followup_after_correct_short_answer",
                        "student_scaffold_followability": "F1",
                        "expected_tutor_move": "advance",
                        "prior_ai_scaffold": "你先说 true 的含义。",
                        "student_reply_to_prior_scaffold": "true 就是能满足限制。",
                        "response_text": "对，这说明下一步可以看 true 后保留哪一边。",
                        "review_status": "unlabeled",
                    }
                )

            xlsx_export.export_xlsx(input_csv=input_csv, output_xlsx=output_xlsx)
            workbook = load_workbook(output_xlsx)

        sheet = workbook["盲评表"]
        headers = {sheet.cell(row=2, column=column_idx).value: column_idx for column_idx in range(1, sheet.max_column + 1)}
        self.assertIn("turn_position", headers)
        self.assertIn("student_scaffold_followability", headers)
        self.assertIn("prior_ai_scaffold", headers)
        self.assertIn("student_reply_to_prior_scaffold", headers)
        self.assertEqual("轮次位置", sheet.cell(row=1, column=headers["turn_position"]).value)
        self.assertEqual("F1", sheet.cell(row=3, column=headers["student_scaffold_followability"]).value)
        self.assertEqual("你先说 true 的含义。", sheet.cell(row=3, column=headers["prior_ai_scaffold"]).value)

    def test_export_xlsx_can_add_by_case_sheets_for_same_case_comparison(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            input_csv = Path(tmpdir) / "review.csv"
            output_xlsx = Path(tmpdir) / "review.by_case.zh.xlsx"
            with input_csv.open("w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=xlsx_export.REVIEW_COLUMNS, lineterminator="\n")
                writer.writeheader()
                for case_id, problem_ref, student_message in [
                    ("dialogue_v3_001_state_representation_semantics", "P1048", "状态怎么设？"),
                    ("dialogue_v3_002_state_representation_semantics", "P1005", "这一格要记录什么？"),
                ]:
                    for response_id in ["resp_a", "resp_b"]:
                        writer.writerow(
                            {
                                "case_id": case_id,
                                "anonymized_response_id": response_id,
                                "problem_ref": problem_ref,
                                "problem_source_platform": "luogu",
                                "problem_source_id": problem_ref,
                                "problem_source_url": f"https://www.luogu.com.cn/problem/{problem_ref}",
                                "problem_statement": "必要题面。",
                                "student_message": student_message,
                                "success_criteria": "能引导学生说出当前一步。",
                                "forbidden_content": "不能直接补完关键桥。",
                                "response_text": f"{response_id} 的回复",
                                "review_status": "unlabeled",
                            }
                        )

            row_count = xlsx_export.export_xlsx(
                input_csv=input_csv,
                output_xlsx=output_xlsx,
                by_case_sheets=True,
            )
            workbook = load_workbook(output_xlsx)

        self.assertEqual(4, row_count)
        self.assertIn("盲评表", workbook.sheetnames)
        self.assertIn("按题索引", workbook.sheetnames)
        self.assertIn("题001_P1048", workbook.sheetnames)
        self.assertIn("题002_P1005", workbook.sheetnames)
        index = workbook["按题索引"]
        self.assertEqual("sheet_name", index["A1"].value)
        self.assertEqual("题001_P1048", index["A2"].value)
        self.assertEqual("dialogue_v3_001_state_representation_semantics", index["B2"].value)
        case_sheet = workbook["题001_P1048"]
        field_headers = {
            case_sheet.cell(row=2, column=column_idx).value: column_idx
            for column_idx in range(1, case_sheet.max_column + 1)
        }
        self.assertEqual(4, case_sheet.max_row)
        self.assertEqual("AI 回复（要评分）", case_sheet.cell(row=1, column=field_headers["response_text"]).value)
        self.assertEqual("resp_a 的回复", case_sheet.cell(row=3, column=field_headers["response_text"]).value)
        self.assertEqual("resp_b 的回复", case_sheet.cell(row=4, column=field_headers["response_text"]).value)


if __name__ == "__main__":
    unittest.main()
